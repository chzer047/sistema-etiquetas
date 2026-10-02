# ==========================================
# C XML BR Engine
# COMPLETO + REV + HISTÓRICO + BANCO + REGISTROS
# ==========================================

import streamlit as st
import pdfplumber
import fitz
import re
import pandas as pd
from pathlib import Path
import zipfile
import tempfile
from ftfy import fix_text
import sqlite3
from datetime import datetime
from io import BytesIO
from zoneinfo import ZoneInfo
from openpyxl import load_workbook

# ==========================================
# CONFIG
# ==========================================

st.set_page_config(
    page_title="C XML BR Engine",
    layout="wide",
    initial_sidebar_state="expanded"
)

st.markdown("""<style>
@import url('https://fonts.googleapis.com/css2?family=Inter:wght@400;500;600;700&display=swap');

/* ===== RESET GERAL ===== */
html, body, [class*="css"] { font-family: 'Inter', sans-serif !important; }

/* ===== SIDEBAR ESCURA ===== */
[data-testid="stSidebar"] {
    background: #16213E !important;
    border-right: 1px solid rgba(255,255,255,0.05) !important;
    padding-top: 0 !important;
}
[data-testid="stSidebar"] > div:first-child {
    padding-top: 1rem !important;
}
[data-testid="stSidebar"] * { color: rgba(255,255,255,0.75) !important; }
[data-testid="stSidebar"] h1,
[data-testid="stSidebar"] h2,
[data-testid="stSidebar"] h3 {
    color: rgba(255,255,255,0.35) !important;
    font-size: 10px !important;
    font-weight: 700 !important;
    letter-spacing: 0.1em !important;
    text-transform: uppercase !important;
    margin-bottom: 4px !important;
    margin-top: 16px !important;
}
[data-testid="stSidebar"] p,
[data-testid="stSidebar"] span,
[data-testid="stSidebar"] label {
    font-size: 13px !important;
    color: rgba(255,255,255,0.6) !important;
}

/* Botão principal sidebar */
[data-testid="stSidebar"] .stButton > button {
    background: transparent !important;
    border: none !important;
    color: rgba(255,255,255,0.6) !important;
    border-radius: 8px !important;
    font-weight: 400 !important;
    font-size: 13px !important;
    width: 100% !important;
    text-align: left !important;
    padding: 8px 12px !important;
    transition: all 0.15s !important;
    letter-spacing: 0 !important;
}
[data-testid="stSidebar"] .stButton > button:hover {
    background: rgba(255,255,255,0.07) !important;
    color: rgba(255,255,255,0.9) !important;
}
[data-testid="stSidebar"] .stButton > button:contains("▶") {
    background: rgba(233,69,96,0.15) !important;
    color: #E94560 !important;
    font-weight: 600 !important;
}
[data-testid="stSidebar"] .stDownloadButton > button {
    background: rgba(16,185,129,0.1) !important;
    border: 1px solid rgba(16,185,129,0.25) !important;
    color: #10B981 !important;
    border-radius: 8px !important;
    width: 100% !important;
    font-size: 12px !important;
}
[data-testid="stSidebar"] hr {
    border-color: rgba(255,255,255,0.07) !important;
    margin: 12px 0 !important;
}
[data-testid="stSidebar"] .stFileUploader {
    background: rgba(255,255,255,0.03) !important;
    border: 1px dashed rgba(255,255,255,0.12) !important;
    border-radius: 8px !important;
}
[data-testid="stSidebar"] .stFileUploader label {
    color: rgba(255,255,255,0.5) !important;
    font-size: 11px !important;
}

/* ===== FILE UPLOADER — lista de arquivos compacta (evita captura de scroll) ===== */
[data-testid="stFileUploaderFileList"] {
    max-height: 96px !important;
    overflow-y: auto !important;
}

/* ===== CONTEÚDO PRINCIPAL ===== */
.main { background: #F8FAFC !important; }
.main .block-container {
    padding-top: 1.75rem !important;
    padding-left: 2.5rem !important;
    padding-right: 2.5rem !important;
    max-width: 1440px !important;
    background: #F8FAFC !important;
}

/* ===== TÍTULOS ===== */
h1 { font-size: 26px !important; font-weight: 700 !important; color: #0F172A !important; letter-spacing: -0.3px !important; }
h2 { font-size: 18px !important; font-weight: 600 !important; color: #1E293B !important; }
h3 { font-size: 15px !important; font-weight: 600 !important; color: #334155 !important; }

/* ===== TABS ===== */
[data-testid="stTabs"] [data-baseweb="tab-list"] {
    background: transparent !important;
    border-bottom: 1px solid #E2E8F0 !important;
    gap: 2px !important;
    padding-bottom: 0 !important;
}
[data-testid="stTabs"] [data-baseweb="tab"] {
    background: transparent !important;
    border-radius: 6px 6px 0 0 !important;
    font-size: 13px !important;
    font-weight: 500 !important;
    color: #94A3B8 !important;
    padding: 10px 14px !important;
    border-bottom: 2px solid transparent !important;
    transition: all 0.15s !important;
}
[data-testid="stTabs"] [data-baseweb="tab"]:hover {
    color: #475569 !important;
    background: rgba(233,69,96,0.04) !important;
}
[data-testid="stTabs"] [aria-selected="true"] {
    color: #E94560 !important;
    border-bottom: 2px solid #E94560 !important;
    background: rgba(233,69,96,0.05) !important;
    font-weight: 600 !important;
}

/* ===== BOTÕES ===== */
.stButton > button {
    border-radius: 8px !important;
    font-weight: 500 !important;
    font-size: 13px !important;
    transition: all 0.15s !important;
    border: 1px solid #E2E8F0 !important;
}
.stButton > button[kind="primary"] {
    background: #E94560 !important;
    border: none !important;
    color: #fff !important;
    box-shadow: 0 2px 8px rgba(233,69,96,0.3) !important;
}
.stButton > button[kind="primary"]:hover {
    background: #C73652 !important;
    box-shadow: 0 4px 14px rgba(233,69,96,0.4) !important;
    transform: translateY(-1px) !important;
}
.stButton > button:hover {
    border-color: #CBD5E1 !important;
    background: #F1F5F9 !important;
}
.stDownloadButton > button {
    border-radius: 8px !important;
    font-weight: 500 !important;
    font-size: 13px !important;
}

/* ===== MÉTRICAS ===== */
[data-testid="metric-container"] {
    background: #FFFFFF !important;
    border: 1px solid #E2E8F0 !important;
    border-radius: 12px !important;
    padding: 18px 22px !important;
    box-shadow: 0 1px 4px rgba(0,0,0,0.04) !important;
    transition: box-shadow 0.15s !important;
}
[data-testid="metric-container"]:hover {
    box-shadow: 0 4px 12px rgba(0,0,0,0.08) !important;
}
[data-testid="metric-container"] label {
    font-size: 12px !important;
    font-weight: 500 !important;
    color: #94A3B8 !important;
    text-transform: uppercase !important;
    letter-spacing: 0.05em !important;
}
[data-testid="metric-container"] [data-testid="stMetricValue"] {
    font-size: 30px !important;
    font-weight: 700 !important;
    color: #0F172A !important;
    letter-spacing: -0.5px !important;
}
[data-testid="metric-container"] [data-testid="stMetricDelta"] {
    font-size: 12px !important;
}

/* ===== EXPANDERS ===== */
[data-testid="stExpander"] {
    background: #FFFFFF !important;
    border: 1px solid #E2E8F0 !important;
    border-radius: 10px !important;
    overflow: hidden !important;
    margin-bottom: 8px !important;
    box-shadow: 0 1px 3px rgba(0,0,0,0.03) !important;
}
[data-testid="stExpander"] summary {
    font-weight: 500 !important;
    font-size: 13px !important;
    padding: 12px 16px !important;
    background: #FAFAFA !important;
    color: #334155 !important;
    border-bottom: 1px solid #F1F5F9 !important;
}
[data-testid="stExpander"] summary:hover {
    background: #F1F5F9 !important;
}

/* ===== DATAFRAMES ===== */
[data-testid="stDataFrame"] {
    border: 1px solid #E2E8F0 !important;
    border-radius: 10px !important;
    overflow: hidden !important;
    box-shadow: 0 1px 3px rgba(0,0,0,0.03) !important;
}

/* ===== INPUTS ===== */
[data-testid="stTextInput"] input,
[data-testid="stTextArea"] textarea,
[data-testid="stSelectbox"] > div > div {
    border-radius: 8px !important;
    border: 1px solid #E2E8F0 !important;
    font-size: 13px !important;
    background: #FFFFFF !important;
    transition: all 0.15s !important;
}
[data-testid="stTextInput"] input:focus,
[data-testid="stTextArea"] textarea:focus {
    border-color: #E94560 !important;
    box-shadow: 0 0 0 3px rgba(233,69,96,0.1) !important;
    outline: none !important;
}

/* ===== ALERTAS ===== */
[data-testid="stAlert"] {
    border-radius: 10px !important;
    border-left-width: 4px !important;
    font-size: 13px !important;
}

/* ===== FILE UPLOADER ===== */
[data-testid="stFileUploader"] {
    border: 2px dashed #E2E8F0 !important;
    border-radius: 10px !important;
    background: #FAFAFA !important;
    transition: all 0.15s !important;
}
[data-testid="stFileUploader"]:hover {
    border-color: #E94560 !important;
    background: rgba(233,69,96,0.02) !important;
}

/* ===== SELECTBOX ===== */
[data-testid="stSelectbox"] label {
    font-size: 13px !important;
    font-weight: 500 !important;
    color: #475569 !important;
}

/* ===== CHECKBOX / TOGGLE ===== */
[data-testid="stCheckbox"] label {
    font-size: 13px !important;
    color: #475569 !important;
}
[data-testid="stToggle"] span[data-checked="true"] {
    background: #E94560 !important;
}

/* ===== PROGRESS ===== */
[data-testid="stProgressBar"] > div > div {
    background: linear-gradient(90deg, #E94560, #C73652) !important;
    border-radius: 4px !important;
}

/* ===== DIVIDER ===== */
hr {
    border: none !important;
    border-top: 1px solid #E2E8F0 !important;
    margin: 1.5rem 0 !important;
}

/* ===== SUCCESS / WARNING / ERROR ===== */
.stSuccess { border-left-color: #10B981 !important; }
.stWarning { border-left-color: #F59E0B !important; }
.stError   { border-left-color: #E94560 !important; }
.stInfo    { border-left-color: #3B82F6 !important; }

/* ===== CÓDIGO ===== */
code, pre {
    border-radius: 6px !important;
    font-size: 12px !important;
}

/* ===== CAPTION ===== */
.stCaptionContainer p {
    font-size: 12px !important;
    color: #94A3B8 !important;
}

/* ===== RADIO ===== */
[data-testid="stRadio"] label {
    font-size: 13px !important;
    color: #475569 !important;
}
</style>""", unsafe_allow_html=True)



DB_PATH = "certificados.db"

# ==========================================
# CONEXÃO
# ==========================================

if "db_version" not in st.session_state:
    st.session_state["db_version"] = 0

def _abrir_conexao(db_path):
    c = sqlite3.connect(db_path, check_same_thread=False, timeout=30)
    c.execute("PRAGMA journal_mode=WAL")
    c.execute("PRAGMA synchronous=NORMAL")
    c.execute("PRAGMA cache_size=-65536")       # 64MB de cache (era 32MB)
    c.execute("PRAGMA temp_store=MEMORY")
    c.execute("PRAGMA mmap_size=268435456")     # 256MB memory-mapped I/O
    c.execute("PRAGMA busy_timeout=10000")
    c.commit()
    return c


def commit_seguro(tentativas=5, espera=0.5):
    """Faz commit com retry em caso de lock do SQLite."""
    import time
    for i in range(tentativas):
        try:
            conn.commit()
            return
        except sqlite3.OperationalError as e:
            if "locked" in str(e).lower() and i < tentativas - 1:
                time.sleep(espera)
            else:
                raise

def _conexao_valida(c):
    try:
        c.execute("SELECT 1")
        return True
    except Exception:
        return False

@st.cache_resource(show_spinner=False)
def get_connection(db_path, _version):
    return _abrir_conexao(db_path)

conn = get_connection(DB_PATH, st.session_state["db_version"])

# Fallback: se a conexão cacheada estiver fechada, abre diretamente
if not _conexao_valida(conn):
    st.cache_resource.clear()
    conn = _abrir_conexao(DB_PATH)

cursor = conn.cursor()

# ==========================================
# TABELAS
# ==========================================


@st.cache_resource(show_spinner=False)
def _preparar_schema(_db_version):
    """Cria tabelas e roda migrações UMA VEZ por sessão (via cache_resource),
    em vez de a cada re-execução do script — isso elimina lentidão por clique."""
    cursor.execute("""
    CREATE TABLE IF NOT EXISTS certificados (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        ip_bri TEXT UNIQUE,
        produto TEXT,
        ce_bri TEXT,
        familia TEXT,
        rev INTEGER,
        data_emissao TEXT,
        arquivo_pdf TEXT,
        data_cadastro TEXT,
        data_atualizacao TEXT
    )
    """)

    cursor.execute("""
    CREATE TABLE IF NOT EXISTS itens (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        certificado_id INTEGER,
        ordem INTEGER,
        marca TEXT,
        modelo TEXT,
        nome TEXT,
        codigo TEXT,
        modelo_ref_key TEXT
    )
    """)

    cursor.execute("""
    CREATE TABLE IF NOT EXISTS historico_alteracoes (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        ip_bri TEXT,
        rev_antiga INTEGER,
        rev_nova INTEGER,
        modelo TEXT,
        codigo TEXT,
        tipo_alteracao TEXT,
        campo_alterado TEXT,
        valor_antigo TEXT,
        valor_novo TEXT,
        arquivo_pdf TEXT,
        data_hora TEXT
    )
    """)

    cursor.execute("""
    CREATE TABLE IF NOT EXISTS registros (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        cliente_base TEXT,
        fabrica TEXT,
        ce_bri TEXT,
        familia TEXT,
        registro TEXT,
        endereco_fabrica TEXT,
        arquivo_excel TEXT,
        data_cadastro TEXT
    )
    """)


    cursor.execute("""
    CREATE TABLE IF NOT EXISTS sistema5_clientes (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        categoria TEXT,
        cliente_base TEXT,
        data_cadastro TEXT,
        UNIQUE(categoria, cliente_base)
    )
    """)

    cursor.execute("""
    CREATE TABLE IF NOT EXISTS sistema5_fabricas (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        cliente_id INTEGER,
        fabrica TEXT,
        ce_bri TEXT,
        endereco_fabrica TEXT,
        data_cadastro TEXT,
        UNIQUE(cliente_id, fabrica)
    )
    """)

    cursor.execute("""
    CREATE TABLE IF NOT EXISTS sistema5_arquivos (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        cliente_id INTEGER,
        fabrica_id INTEGER,
        tipo_processo TEXT,
        ip_processo TEXT,
        data_processo TEXT,
        arquivo_nome TEXT,
        data_upload TEXT
    )
    """)

    cursor.execute("""
    CREATE TABLE IF NOT EXISTS sistema5_itens (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        arquivo_id INTEGER,
        cliente_base TEXT,
        categoria TEXT,
        fabrica TEXT,
        ce_bri TEXT,
        endereco_fabrica TEXT,
        tipo_processo TEXT,
        ip_processo TEXT,
        data_processo TEXT,
        familia TEXT,
        item TEXT,
        marca TEXT,
        modelo TEXT,
        nome TEXT,
        codigo TEXT,
        modelo_ref_key TEXT,
        arquivo_nome TEXT,
        data_upload TEXT
    )
    """)


    cursor.execute("""
    CREATE TABLE IF NOT EXISTS ip_bri_familias (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        ce_bri TEXT,
        familia TEXT,
        ip_bri TEXT,
        observacao TEXT,
        data_cadastro TEXT,
        data_atualizacao TEXT,
        UNIQUE(ce_bri, familia)
    )
    """)

    # ---- Biblioteca de selos "Segurança/REGISTRO" por Cliente+Fábrica+Família ----
    # Base pro futuro módulo de Geração de Etiquetas.
    # Cada combinação Cliente+Fábrica+Família pode ter até 4 variantes visuais do
    # selo (amarelo, P&B, compacto colorido, compacto P&B), todas com o MESMO
    # número de registro (lido por OCR ao subir a primeira imagem).
    cursor.execute("""
    CREATE TABLE IF NOT EXISTS selos_registro (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        cliente_base TEXT,
        fabrica TEXT,
        familia TEXT,
        registro TEXT,
        selo_amarelo BLOB,
        selo_pb BLOB,
        selo_compacto_cor BLOB,
        selo_compacto_pb BLOB,
        data_cadastro TEXT,
        data_atualizacao TEXT,
        UNIQUE(cliente_base, fabrica, familia)
    )
    """)

    # Migração: se a tabela selos_registro veio de uma versão antiga (coluna 'imagem'
    # única em vez das 4 variantes), adiciona as colunas novas sem perder o que existe.
    def _garantir_coluna_selos(coluna, tipo):
        try:
            info = cursor.execute("PRAGMA table_info(selos_registro)").fetchall()
            existe = any(row[1] == coluna for row in info)
            if not existe:
                cursor.execute(f"ALTER TABLE selos_registro ADD COLUMN {coluna} {tipo}")
        except Exception:
            pass

    for _col, _tipo in [
        ("registro", "TEXT"),
        ("selo_amarelo", "BLOB"),
        ("selo_pb", "BLOB"),
        ("selo_compacto_cor", "BLOB"),
        ("selo_compacto_pb", "BLOB"),
    ]:
        _garantir_coluna_selos(_col, _tipo)

    # ---- Imagens genéricas reaproveitáveis em qualquer etiqueta (chorão, pilha/bateria) ----
    cursor.execute("""
    CREATE TABLE IF NOT EXISTS assets_etiqueta_genericos (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        tipo TEXT UNIQUE,
        imagem BLOB,
        nome_arquivo TEXT,
        largura INTEGER,
        altura INTEGER,
        data_cadastro TEXT,
        data_atualizacao TEXT
    )
    """)

    # ---- Textos padrão que entram nas etiquetas (advertências, pilha, restritivos etc) ----
    cursor.execute("""
    CREATE TABLE IF NOT EXISTS textos_etiqueta (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        categoria TEXT,
        titulo TEXT,
        conteudo TEXT,
        data_cadastro TEXT,
        data_atualizacao TEXT
    )
    """)

    # 4 modelos-Word fixos de selo (amarelo, pb, compacto_cor, compacto_pb).
    # O número de registro dentro deles é texto editável -> na geração, o sistema
    # substitui pelo registro correto vindo da tabela Registros. Não usa mais OCR.
    cursor.execute("""
    CREATE TABLE IF NOT EXISTS modelos_selo (
        variante TEXT PRIMARY KEY,
        arquivo_docx BLOB,
        nome_arquivo TEXT,
        registro_exemplo TEXT,
        data_atualizacao TEXT
    )
    """)

    # Cadastro de clientes (importadores) pra geração de etiquetas.
    # Preenchido manualmente pra bater com o cartão CNPJ. Editável.
    cursor.execute("""
    CREATE TABLE IF NOT EXISTS clientes_etiqueta (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        apelido TEXT UNIQUE,
        razao_social TEXT,
        cnpj TEXT,
        endereco TEXT,
        sac TEXT,
        origem TEXT,
        data_cadastro TEXT,
        data_atualizacao TEXT
    )
    """)

    commit_seguro()

_preparar_schema(st.session_state["db_version"])



def coluna_existe(tabela, coluna):
    """Verifica se uma coluna já existe na tabela antes de tentar ALTER TABLE."""
    try:
        info = cursor.execute(f"PRAGMA table_info({tabela})").fetchall()
        return any(row[1] == coluna for row in info)
    except Exception:
        return True  # assume que existe para não tentar adicionar

# Garante colunas novas antes de criar índices em bancos antigos
if not coluna_existe("sistema5_itens", "modelo_ref_key"):
    try:
        cursor.execute("ALTER TABLE sistema5_itens ADD COLUMN modelo_ref_key TEXT")
        commit_seguro()
    except Exception:
        pass

if not coluna_existe("itens", "modelo_ref_key"):
    try:
        cursor.execute("ALTER TABLE itens ADD COLUMN modelo_ref_key TEXT")
        commit_seguro()
    except Exception:
        pass

# Índices para performance em buscas por campo mais usados
cursor.execute("CREATE INDEX IF NOT EXISTS idx_itens_codigo ON itens(codigo)")
cursor.execute("CREATE INDEX IF NOT EXISTS idx_itens_modelo ON itens(modelo)")
cursor.execute("CREATE INDEX IF NOT EXISTS idx_certificados_ce_bri ON certificados(ce_bri)")
cursor.execute("CREATE INDEX IF NOT EXISTS idx_certificados_ip_bri ON certificados(ip_bri)")
cursor.execute("CREATE INDEX IF NOT EXISTS idx_s5_itens_modelo ON sistema5_itens(modelo)")
cursor.execute("CREATE INDEX IF NOT EXISTS idx_s5_itens_arquivo_id ON sistema5_itens(arquivo_id)")
cursor.execute("CREATE INDEX IF NOT EXISTS idx_s5_itens_modelo_ref_key ON sistema5_itens(modelo_ref_key)")
cursor.execute("CREATE INDEX IF NOT EXISTS idx_itens_modelo_ref_key ON itens(modelo_ref_key)")
cursor.execute("CREATE INDEX IF NOT EXISTS idx_ip_bri_familias_ce_fam ON ip_bri_familias(ce_bri, familia)")
cursor.execute("CREATE INDEX IF NOT EXISTS idx_s5_itens_cliente_base ON sistema5_itens(cliente_base)")
cursor.execute("CREATE INDEX IF NOT EXISTS idx_s5_itens_fabrica ON sistema5_itens(fabrica)")
cursor.execute("CREATE INDEX IF NOT EXISTS idx_s5_itens_ip_processo ON sistema5_itens(ip_processo)")
cursor.execute("CREATE INDEX IF NOT EXISTS idx_s5_itens_codigo ON sistema5_itens(codigo)")
cursor.execute("CREATE INDEX IF NOT EXISTS idx_s5_arquivos_ip_processo ON sistema5_arquivos(ip_processo)")
cursor.execute("CREATE INDEX IF NOT EXISTS idx_s5_arquivos_cliente_id ON sistema5_arquivos(cliente_id)")
cursor.execute("CREATE INDEX IF NOT EXISTS idx_registros_ce_bri ON registros(ce_bri)")
cursor.execute("CREATE INDEX IF NOT EXISTS idx_registros_cliente_base ON registros(cliente_base)")
cursor.execute("CREATE INDEX IF NOT EXISTS idx_registros_fabrica ON registros(fabrica)")
commit_seguro()

# ==========================================
# FUNÇÕES GERAIS
# ==========================================


def clean(x):
    if x is None:
        return ""
    s = str(x)
    # Normaliza caracteres tipográficos que estão fora do ISO-8859-1 (evita '?' no XML)
    s = (s
         .replace('–', '-')   # en dash  –
         .replace('—', '-')   # em dash  —
         .replace('―', '-')   # horizontal bar
         .replace('‘', "'")   # ' (aspas simples esquerda)
         .replace('’', "'")   # ' (aspas simples direita)
         .replace('“', '"')   # " (aspas duplas esquerda)
         .replace('”', '"')   # " (aspas duplas direita)
         .replace('…', '...')  # … (reticências)
         .replace(' ', ' ')   # non-breaking space
         )
    return re.sub(r"\s+", " ", s.replace("\n", " ")).strip()




def normalizar_hifens_ref(valor):
    """
    Padroniza hífens unicode, remove espaços invisíveis e normaliza espaços.
    NÃO remove hífens comuns. SK-126 continua diferente de SK126.
    """
    texto = clean(valor).upper()

    for h in ["‐", "‑", "‒", "–", "—", "−", "﹣", "－"]:
        texto = texto.replace(h, "-")

    texto = texto.replace("\u200b", "")
    texto = texto.replace("\u200c", "")
    texto = texto.replace("\u200d", "")
    texto = texto.replace("\ufeff", "")
    texto = texto.replace("\xa0", " ")

    texto = re.sub(r"\s+", " ", texto).strip()

    return texto


def modelo_corresponde_referencia_exata(modelo_banco, ref_busca, filtro_rigoroso=False):
    """
    Verifica se um modelo do banco corresponde à referência buscada.
    
    filtro_rigoroso=False (padrão): aceita qualquer modelo que comece com a referência
    filtro_rigoroso=True: rejeita variantes numéricas (ex: ref=0122 não aceita 0122-18)
    """
    modelo = normalizar_hifens_ref(modelo_banco)
    ref    = normalizar_hifens_ref(ref_busca)

    if not modelo or not ref:
        return False
    if modelo == ref:
        return True
    if not modelo.startswith(ref):
        return False

    resto = modelo[len(ref):]
    if not resto:
        return True

    if not filtro_rigoroso:
        # Rejeita se o sufixo começa com dígito — seria outra referência (DT-1039 ≠ DT-103)
        return not re.match(r'^\d', resto)

    # Modo rigoroso: só aceita hífen seguido de letra (descrição), nunca número
    if re.match(r"^-[A-Za-zÀ-ÿ].*", resto):
        return True

    return False


def filtrar_resultado_por_referencia_exata(df, ref_busca):
    if df is None or df.empty:
        return pd.DataFrame()

    col_modelo = None

    if "MODELO" in df.columns:
        col_modelo = "MODELO"
    elif "modelo" in df.columns:
        col_modelo = "modelo"

    if not col_modelo:
        return pd.DataFrame()

    return df[
        df[col_modelo].apply(
            lambda modelo: modelo_corresponde_referencia_exata(
                modelo, ref_busca,
                filtro_rigoroso=st.session_state.get("filtro_referencia_rigoroso", False)
            )
        )
    ].copy()



def codigo_barras_para_texto(valor, number_format=None):
    """
    Converte código de barras para texto preservando zeros à esquerda quando possível.

    Problema:
    Excel pode armazenar 0789123456789 como número 789123456789.
    Se a célula tiver formato 0000000000000, reconstruímos o zero perdido com zfill.
    """

    if valor is None:
        return None

    texto = str(valor).strip()

    if texto.lower() in ["", "nan", "none", "null"]:
        return None

    # Remove .0 quando veio de número inteiro lido como float.
    if re.fullmatch(r"\d+\.0", texto):
        texto = texto[:-2]

    # Primeiro trecho quando vier "codigo/codigo"
    texto = re.split(r"[\/\n]", texto)[0]

    digitos = re.sub(r"\D", "", texto)

    if not digitos:
        return None

    # Se a célula Excel tem formatação do tipo 0000000000000,
    # usa a quantidade de zeros do formato para restaurar zeros à esquerda.
    if number_format:
        fmt = str(number_format)

        # Pega o maior bloco contínuo de zeros do formato.
        blocos_zeros = re.findall(r"0{8,14}", fmt)

        if blocos_zeros:
            tamanho = max(len(b) for b in blocos_zeros)

            if len(digitos) < tamanho:
                digitos = digitos.zfill(tamanho)

    if len(digitos) < 8 or len(digitos) > 14:
        return None

    return digitos



def ler_valor_celula_preservando_zeros(cell):
    """
    Lê o valor de uma célula preservando zeros à esquerda.
    Se a célula tem valor texto, usa diretamente.
    Se tem valor numérico, converte para inteiro (remove .0) sem alterar dígitos.
    Não tenta "adivinhar" zeros perdidos — deixa a busca ser flexível.
    """
    valor = cell.value
    if valor is None:
        return ""

    # Se já é string, usa direto (preserva "0156" como está)
    if isinstance(valor, str):
        texto = valor.strip()
        if texto.lower() in ["nan", "none", "null", ""]:
            return ""
        return texto

    # Se é número, converte sem .0
    texto = str(valor).strip()
    if re.fullmatch(r"\d+\.0", texto):
        texto = texto[:-2]

    if not texto or texto.lower() in ["nan", "none", "null"]:
        return ""

    return texto


def referencia_chave_exata(valor):
    texto = normalizar_hifens_ref(valor) if "normalizar_hifens_ref" in globals() else clean(valor).upper()
    return texto.strip()


def referencia_inicial_chave_modelo(modelo):
    texto = referencia_chave_exata(modelo)
    partes = re.split(r"\s+-\s+", texto, maxsplit=1)
    if partes:
        return partes[0].strip()
    return texto.strip()


def atualizar_chaves_referencia_banco():
    try:
        pendentes_s5 = cursor.execute("""
        SELECT id, modelo
        FROM sistema5_itens
        WHERE modelo IS NOT NULL
        AND modelo != ''
        AND (modelo_ref_key IS NULL OR modelo_ref_key = '')
        LIMIT 10000
        """).fetchall()

        for item_id, modelo in pendentes_s5:
            cursor.execute(
                "UPDATE sistema5_itens SET modelo_ref_key = ? WHERE id = ?",
                (referencia_inicial_chave_modelo(modelo), item_id)
            )

        pendentes_cert = cursor.execute("""
        SELECT id, modelo
        FROM itens
        WHERE modelo IS NOT NULL
        AND modelo != ''
        AND (modelo_ref_key IS NULL OR modelo_ref_key = '')
        LIMIT 10000
        """).fetchall()

        for item_id, modelo in pendentes_cert:
            cursor.execute(
                "UPDATE itens SET modelo_ref_key = ? WHERE id = ?",
                (referencia_inicial_chave_modelo(modelo), item_id)
            )

        commit_seguro()
    except Exception:
        pass


def corrigir_texto(texto):
    if texto is None:
        return ""

    return fix_text(str(texto))



def extrair_codigo_unico(codigo_raw):
    return codigo_barras_para_texto(codigo_raw)


def limitar_nome(nome):
    nome = clean(nome)
    if len(nome) <= 200:
        return nome

    # Aplicadas em ordem, parando assim que ≤200 — mesmo padrão do XML Rápido.
    # O início da descrição é sempre preservado: o corte final remove do final.
    # Tuplas: (padrão, abreviação, é_regex)
    _abrevs = [
        ("MECANISMO SIMPLES",                  "MEC.SIMPL.", False),
        ("BASICAMENTE",                         "BASIC.",     False),
        ("DENOMINADO",                          "DENOM.",     False),
        ("UMEDECENDO",                          "UMED.",      False),
        ("INVISIVEIS",                          "INVIS.",     False),
        ("INVISIVEL",                           "INVIS.",     False),
        ("VARIADOS",                            "VAR.",       False),
        ("VISIVEL",                             "VIS.",       False),
        ("PRODUZIDO",                           "PROD.",      False),
        ("INDICATIVO",                          "IND.",       False),
        ("RESTRITIVO",                          "REST.",      False),
        ("RESTRIÇÃO",                           "REST.",      False),
        (r'(?<![A-Za-z])ANOS\b',               "A",          True),
        (r'(?<![A-Za-z])MESES\b',              "MES.",       True),
        (r'\bINJE[ÇC]ÃO\b|\bINJECAO\b',        "INJ.",       True),
        (r'\bM[ÁA]XIMA\b',                     "MAX.",       True),
        (r'\bPL[ÁA]STICO\b',                   "PLAST.",     True),
        ("CONTROLE",                            "CONT.",      False),
        ("REMOTO",                              "REM.",       False),
        ("VELOCIDADE",                          "VEL.",       False),
        (r'\bMEDIDAS?\b',                       "MED.",       True),
        # --- abreviações secundárias (aplicadas só se ainda >200) ---
        ("ILUSTRAÇÕES",                         "ILUSTR.",    False),
        ("INFANTIL",                            "INF.",       False),
        (r'\bCONJUNTOS?\b',                     "CONJ.",      True),
        (r'\bDESENHOS?\b',                      "DES.",       True),
        ("CANETA",                              "CAN.",       False),
        ("SORTIDO",                             "SORT.",      False),
        (r'\bLIVROS?\b',                        "LIV.",       True),
        ("D'AGUA",                              "D'AG.",      False),
    ]
    for _old, _new, _regex in _abrevs:
        if len(nome) <= 200:
            break
        if _regex:
            nome = re.sub(_old, _new, nome, flags=re.IGNORECASE)
        else:
            nome = nome.replace(_old, _new)

    return nome[:200]



def escape_xml(texto):
    return str(texto)\
        .replace("&", "&amp;")\
        .replace("<", "&lt;")\
        .replace(">", "&gt;")



def verificar_codigos_duplicados(df):
    duplicados = df[
        df.duplicated(
            subset=["CODIGO"],
            keep=False
        )
    ].copy()

    if not duplicados.empty:
        duplicados = duplicados.sort_values(by="CODIGO")

    return duplicados


# ==========================================
# REV
# ==========================================


def extrair_rev(texto):
    texto = corrigir_texto(texto)

    padroes = [
        r"\bREV\.?\s*:?\s*(\d+)",
        r"\bREVISÃO\s*:?\s*(\d+)",
        r"\bREVISAO\s*:?\s*(\d+)"
    ]

    for padrao in padroes:
        match = re.search(
            padrao,
            texto,
            flags=re.IGNORECASE
        )

        if match:
            return int(match.group(1))

    linhas = [clean(l) for l in texto.split("\n") if clean(l)]

    for i, linha in enumerate(linhas):
        if linha.upper() in ["REV", "REV.", "REVISÃO", "REVISAO"] and i + 1 < len(linhas):
            numero = re.sub(r"\D", "", linhas[i + 1])
            if numero:
                return int(numero)

    return 0


# ==========================================
# EXTRAÇÃO CERTIFICADO
# ==========================================


def extrair_familia_ip_bri(ip_bri):
    if not ip_bri:
        return None

    match = re.search(r"-([0-9]+)$", ip_bri)

    if not match:
        return None

    return str(int(match.group(1)))



def extrair_dados_certificado(pdf_path):
    doc = fitz.open(str(pdf_path))

    texto = ""

    for page in doc:
        texto += page.get_text() + "\n"

    texto = corrigir_texto(texto)

    ip_bri = None
    produto = None
    ce_bri = None
    data_emissao = None

    ip_match = re.search(
        r"IP-BRI-\d+\/\d+-\d+",
        texto
    )

    if ip_match:
        ip_bri = ip_match.group(0)

    ce_match = re.search(
        r"CE-BRI-[A-Z0-9\-]+",
        texto,
        flags=re.IGNORECASE
    )

    if ce_match:
        ce_bri = ce_match.group(0).upper()

    produto_match = re.search(
        r"Produto:\s*(.+)",
        texto
    )

    if produto_match:
        produto = clean(produto_match.group(1))

    emissao_match = re.search(
        r"Data de Emissão:\s*(\d{2}\/\d{2}\/\d{4})",
        texto,
        flags=re.IGNORECASE
    )

    if emissao_match:
        data_emissao = emissao_match.group(1)

    rev = extrair_rev(texto)
    familia = extrair_familia_ip_bri(ip_bri)

    return {
        "ip_bri": ip_bri,
        "produto": produto,
        "ce_bri": ce_bri,
        "rev": rev,
        "familia": familia,
        "data_emissao": data_emissao
    }


# ==========================================
# PARSE PDF
# ==========================================


def parse_pdf(pdf_path):
    rows = []

    with pdfplumber.open(str(pdf_path)) as pdf:
        for page in pdf.pages:
            for table in page.extract_tables() or []:
                for row in table:
                    if not row or len(row) < 6:
                        continue

                    ordem = clean(row[1])

                    if not re.fullmatch(r"\d{3,4}", ordem):
                        continue

                    ordem = int(ordem)

                    marca = clean(corrigir_texto(row[2]))
                    modelo = clean(corrigir_texto(row[3]))
                    nome = clean(corrigir_texto(row[4]))

                    codigo_raw = clean(corrigir_texto(row[5]))
                    codigo = extrair_codigo_unico(codigo_raw)

                    if not codigo:
                        continue

                    rows.append([
                        ordem,
                        marca,
                        modelo,
                        nome,
                        codigo
                    ])

    rows.sort(key=lambda x: x[0])

    return rows


# ==========================================
# PARSE PDF INNAC (sem salvar no banco)
# ==========================================

def _parse_pdf_innac_fitz(pdf_path):
    """Extração via fitz (PyMuPDF). Funciona bem quando o PDF repete cabeçalhos em todas as páginas."""
    _rows = []
    _doc = fitz.open(str(pdf_path))
    for _page in _doc:
        _tabs = _page.find_tables()
        for _tab in _tabs.tables:
            for _row in _tab.extract():
                if not _row or len(_row) < 6:
                    continue
                _ordem_raw = str(_row[1] or '').strip()
                if re.fullmatch(r'\d{1,4}', _ordem_raw):
                    _ordem = int(_ordem_raw)
                    _marca       = clean(corrigir_texto(str(_row[2] or '')))
                    _modelo_full = clean(corrigir_texto(str(_row[3] or ''))).strip('"').strip("'")
                    _desc        = clean(corrigir_texto(str(_row[4] or '').replace('\n', ' ')))
                    _codigo_raw  = str(_row[5] or '').strip()
                    _tokens = [t for t in re.split(r'[\s\-]', _modelo_full.strip()) if t]
                    _ref = _tokens[0] if _tokens else ''
                    _codigo = re.sub(r'[^0-9]', '', _codigo_raw)
                    if len(_codigo) < 8 or not _ref:
                        continue
                    _rows.append([_ordem, _marca, _modelo_full, _desc, _codigo])
                else:
                    _c2 = str(_row[2] or '').strip()
                    _c3 = str(_row[3] or '').strip()
                    if not _c2 and not _c3 and _rows:
                        _cont = clean(corrigir_texto(str(_row[4] or '').replace('\n', ' ')))
                        if _cont:
                            _rows[-1][3] = (_rows[-1][3].rstrip(' ') + ' ' + _cont).strip()
    _seen = set()
    _unique = []
    for _r in sorted(_rows, key=lambda x: x[0]):
        if _r[4] not in _seen:
            _seen.add(_r[4])
            _unique.append(_r)
    return _unique, len(_doc)


def _parse_pdf_innac_pdfplumber(pdf_path):
    """
    Fallback via pdfplumber. Usado quando fitz falha em certificados cujo cabeçalho
    aparece apenas na primeira página da tabela (sem repetição nas demais páginas).
    pdfplumber usa análise posicional de palavras e linhas, não depende de cabeçalhos.
    """
    import pdfplumber
    _rows = []
    with pdfplumber.open(str(pdf_path)) as _pdf:
        for _page in _pdf.pages:
            _table = _page.extract_table()
            if not _table:
                continue
            for _row in _table:
                if not _row or len(_row) < 6:
                    continue
                _ordem_raw = str(_row[1] or '').strip()
                if not re.fullmatch(r'\d{1,4}', _ordem_raw):
                    continue
                _ordem = int(_ordem_raw)
                _marca       = clean(corrigir_texto(str(_row[2] or '')))
                _modelo_full = clean(corrigir_texto(str(_row[3] or ''))).strip('"').strip("'")
                _desc        = clean(corrigir_texto(str(_row[4] or '').replace('\n', ' ')))
                _codigo_raw  = str(_row[5] or '').strip()
                _tokens = [t for t in re.split(r'[\s\-]', _modelo_full.strip()) if t]
                _ref = _tokens[0] if _tokens else ''
                _codigo = re.sub(r'[^0-9]', '', _codigo_raw)
                if len(_codigo) < 8 or not _ref:
                    continue
                _rows.append([_ordem, _marca, _modelo_full, _desc, _codigo])
    _seen = set()
    _unique = []
    for _r in sorted(_rows, key=lambda x: x[0]):
        if _r[4] not in _seen:
            _seen.add(_r[4])
            _unique.append(_r)
    return _unique


def _parse_pdf_innac(pdf_path):
    """
    Parser para certificados INNAC (Instituto Nacional de Avaliação da Conformidade).
    Tabela com 7 colunas: [* | Ordem | Marca | Modelo | Descrição | Código de Barras | Pai]

    Tenta fitz primeiro. Se o rendimento for suspeitamente baixo para o tamanho do PDF
    (certificado com cabeçalho ausente na maioria das páginas), aciona fallback pdfplumber.
    """
    _result, _n_pages = _parse_pdf_innac_fitz(pdf_path)

    # Heurística: PDFs grandes (>10 páginas) devem render pelo menos 1 item a cada 4 páginas.
    # Se ficou abaixo disso, o fitz provavelmente perdeu o alinhamento de colunas.
    _limiar = max(3, _n_pages // 4)
    if _n_pages > 10 and len(_result) < _limiar:
        try:
            _result_pb = _parse_pdf_innac_pdfplumber(pdf_path)
            if len(_result_pb) > len(_result):
                return _result_pb
        except Exception:
            pass  # pdfplumber falhou — devolve o que fitz conseguiu

    return _result


def _limitar_nome_innac(nome):
    """
    Aplica abreviações padrão PT-BR para descrições do certificado INNAC.
    As abreviações são diferentes do padrão interno pois o INNAC usa terminologia
    própria (REVESTIMENTO EXTERNO, FAIXA ETARIA, COSTURA INVISÍVEL, etc.).
    Garante máximo de 200 caracteres.
    """
    nome = clean(nome)
    if len(nome) <= 200:
        return nome
    _abrevs_innac = [
        # Expressões compostas — aplicar ANTES das palavras individuais
        ('REVESTIMENTO EXTERNO',       'REV.EXT.'),
        ('DETALHES EM TECIDO BORDADO', 'DET.BORD.'),
        ('FAIXA ETARIA INDICATIVA',    'F.IND.'),
        ('FAIXA ETARIA RESTRITIVA',    'F.REST.'),
        ('COSTURA INDUSTRIAL',         'CST.IND.'),
        ('COSTURA INVISÍVEL',          'CST.INV.'),
        ('COSTURA INVISIVEL',          'CST.INV.'),
        ('FIXAÇÃO DE COMPONENTES',     'FIX.COMP.'),
        ('FIXACAO DE COMPONENTES',     'FIX.COMP.'),
        ('PONTO ESCADA',               'PT.ESC.'),
        ('TECIDO E METAL',             'TEC.MET.'),
        # Palavras individuais — apenas se ainda ultrapassar 200
        ('ENCHIMENTO',                 'ENCH.'),
        ('REVESTIMENTO',               'REVEST.'),
        ('FAIXA ETARIA',               'F.ET.'),
        ('INDICATIVA',                 'INDIC.'),
        ('RESTRITIVA',                 'RESTR.'),
        ('INDUSTRIAL',                 'IND.'),
        ('INVISÍVEL',                  'INVIS.'),
        ('INVISIVEL',                  'INVIS.'),
        ('COMPONENTES',                'COMP.'),
        ('EXTERNO',                    'EXT.'),
        ('BORDADO',                    'BORD.'),
        ('COSTURA',                    'COST.'),
        ('FIXAÇÃO',                    'FIX.'),
        ('FIXACAO',                    'FIX.'),
        ('TECIDO',                     'TEC.'),
    ]
    _d = nome
    for _old, _new in _abrevs_innac:
        if len(_d) <= 200:
            break
        _d = _d.replace(_old, _new)
    return _d[:200].rstrip(',; ') if len(_d) > 200 else _d


def _gerar_xml_innac(df, sufixo):
    """Gera XML INNAC usando _limitar_nome_innac (abreviações específicas do certificado INNAC)."""
    linhas = [
        '<?xml version="1.0" encoding="ISO-8859-1"?>',
        '<ArrayOfItemSolicitacao>'
    ]
    for _, r in df.iterrows():
        modelo = r["MODELO"].rstrip(".,") + sufixo
        linhas.append(f"""
<ItemSolicitacao>
<Marca>{escape_xml(r['MARCA'])}</Marca>
<Modelo>{escape_xml(modelo)}</Modelo>
<Nome>{escape_xml(_limitar_nome_innac(r['NOME']))}</Nome>
<CodigosBarras>
<Codigo>{str(r['CODIGO'])}</Codigo>
</CodigosBarras>
</ItemSolicitacao>
""")
    linhas.append("</ArrayOfItemSolicitacao>")
    return "\n".join(linhas)


def _extrair_header_innac(pdf_path):
    """Extrai cabeçalho do certificado INNAC: IP-BRI, CE-BRI, data de emissão, REV."""
    _doc = fitz.open(str(pdf_path))
    _texto = _doc[0].get_text()  # cabeçalho está na página 1
    _linhas = _texto.split('\n')

    # IP-BRI  ex: IP-BRI-2281/2025-05
    _ip = re.search(r'IP-BRI-\d+\/\d+-\d+', _texto)
    _ip_bri = _ip.group(0) if _ip else None

    # CE-BRI  ex: CE-BRI/INNAC-02724-01A
    _ce = re.search(r'CE-BRI[\/\-][A-Z0-9\/\-]+', _texto)
    _ce_bri = _ce.group(0).rstrip('-').strip() if _ce else None

    # Data de Emissão  ex: 29/12/2025
    _em = re.search(r'Data de Emiss[aã]o:\s*(\d{2}\/\d{2}\/\d{4})', _texto, re.IGNORECASE)
    _data_emissao = _em.group(1) if _em else None

    # REV: número isolado na linha imediatamente antes do label 'Revisão'
    _rev = None
    for _i, _l in enumerate(_linhas):
        if _l.strip() == 'Revisão':
            for _j in range(max(0, _i - 7), _i):
                _m = re.fullmatch(r'\d+', _linhas[_j].strip())
                if _m:
                    _rev = int(_m.group(0))
                    break
            break

    # Quantidade de produtos declarada no certificado
    _qtd = re.search(r'Quantidade de produtos:\s*(\d+)', _texto)
    _qtd_prod = int(_qtd.group(1)) if _qtd else None

    return {
        'ip_bri':       _ip_bri,
        'ce_bri':       _ce_bri,
        'data_emissao': _data_emissao,
        'rev':          _rev,
        'qtd_produtos': _qtd_prod,
    }


# ==========================================
# XML
# ==========================================


def gerar_xml(df, sufixo):
    linhas = [
        '<?xml version="1.0" encoding="ISO-8859-1"?>',
        '<ArrayOfItemSolicitacao>'
    ]

    for _, r in df.iterrows():
        modelo = r["MODELO"].rstrip(".,") + sufixo

        linhas.append(f"""
<ItemSolicitacao>
<Marca>{escape_xml(r['MARCA'])}</Marca>
<Modelo>{escape_xml(modelo)}</Modelo>
<Nome>{escape_xml(limitar_nome(r['NOME']))}</Nome>
<CodigosBarras>
<Codigo>{str(r['CODIGO'])}</Codigo>
</CodigosBarras>
</ItemSolicitacao>
""")

    linhas.append("</ArrayOfItemSolicitacao>")

    return "\n".join(linhas)


# ==========================================
# HISTÓRICO
# ==========================================


def registrar_historico(
    ip_bri,
    rev_antiga,
    rev_nova,
    modelo,
    codigo,
    tipo,
    campo,
    antigo,
    novo,
    arquivo
):
    cursor.execute("""
    INSERT INTO historico_alteracoes (
        ip_bri,
        rev_antiga,
        rev_nova,
        modelo,
        codigo,
        tipo_alteracao,
        campo_alterado,
        valor_antigo,
        valor_novo,
        arquivo_pdf,
        data_hora
    )
    VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
    """, (
        ip_bri,
        rev_antiga,
        rev_nova,
        modelo,
        codigo,
        tipo,
        campo,
        antigo,
        novo,
        arquivo,
        datetime.now().strftime("%d/%m/%Y %H:%M:%S")
    ))

    commit_seguro()



def comparar_itens(
    ip_bri,
    rev_antiga,
    rev_nova,
    antigos,
    novos,
    arquivo
):
    antigos_dict = {
        r[2]: r for r in antigos
    }

    novos_dict = {
        r[2]: r for r in novos
    }

    for modelo in antigos_dict:
        if modelo not in novos_dict:
            antigo = antigos_dict[modelo]

            registrar_historico(
                ip_bri,
                rev_antiga,
                rev_nova,
                antigo[2],
                antigo[4],
                "ITEM_REMOVIDO",
                "",
                f"MARCA: {antigo[1]} | MODELO: {antigo[2]} | NOME: {antigo[3]} | CÓDIGO: {antigo[4]}",
                "",
                arquivo
            )

    for modelo in novos_dict:
        if modelo not in antigos_dict:
            novo = novos_dict[modelo]

            registrar_historico(
                ip_bri,
                rev_antiga,
                rev_nova,
                novo[2],
                novo[4],
                "ITEM_NOVO",
                "",
                "",
                f"MARCA: {novo[1]} | MODELO: {novo[2]} | NOME: {novo[3]} | CÓDIGO: {novo[4]}",
                arquivo
            )

    for modelo in novos_dict:
        if modelo in antigos_dict:
            antigo = antigos_dict[modelo]
            novo = novos_dict[modelo]

            campos = {
                "MARCA": (antigo[1], novo[1]),
                "NOME": (antigo[3], novo[3]),
                "CODIGO": (antigo[4], novo[4])
            }

            for campo in campos:
                antigo_valor, novo_valor = campos[campo]

                if str(antigo_valor) != str(novo_valor):
                    registrar_historico(
                        ip_bri,
                        rev_antiga,
                        rev_nova,
                        modelo,
                        novo[4],
                        "CAMPO_ALTERADO",
                        campo,
                        antigo_valor,
                        novo_valor,
                        arquivo
                    )


# ==========================================
# SALVAR / ATUALIZAR
# ==========================================


def garantir_estrutura_banco():
    """
    Fonte única de verdade para estrutura do banco.
    Chamada UMA VEZ na inicialização via atualizar_familias_certificados().
    NÃO chamar por transação individual — custo desnecessário.
    """
    cursor.execute("""
    CREATE TABLE IF NOT EXISTS certificados (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        ip_bri TEXT UNIQUE,
        produto TEXT,
        ce_bri TEXT,
        rev INTEGER,
        data_emissao TEXT,
        arquivo_pdf TEXT,
        data_cadastro TEXT,
        data_atualizacao TEXT
    )
    """)

    cursor.execute("""
    CREATE TABLE IF NOT EXISTS itens (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        certificado_id INTEGER,
        ordem INTEGER,
        marca TEXT,
        modelo TEXT,
        nome TEXT,
        codigo TEXT
    )
    """)

    cursor.execute("""
    CREATE TABLE IF NOT EXISTS registros (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        cliente_base TEXT,
        fabrica TEXT,
        ce_bri TEXT,
        familia TEXT,
        registro TEXT,
        arquivo_excel TEXT,
        data_cadastro TEXT
    )
    """)

    commit_seguro()

    # Ajustes para bancos antigos (colunas que podem não existir)
    ajustes = [
        ("itens", "certificado_id", "INTEGER"),
        ("itens", "ordem", "INTEGER"),
        ("itens", "marca", "TEXT"),
        ("itens", "modelo", "TEXT"),
        ("itens", "nome", "TEXT"),
        ("itens", "codigo", "TEXT"),
        ("certificados", "produto", "TEXT"),
        ("certificados", "ce_bri", "TEXT"),
        ("certificados", "familia", "TEXT"),
        ("certificados", "rev", "INTEGER"),
        ("certificados", "data_emissao", "TEXT"),
        ("certificados", "arquivo_pdf", "TEXT"),
        ("certificados", "data_cadastro", "TEXT"),
        ("certificados", "data_atualizacao", "TEXT"),
        ("registros", "cliente_base", "TEXT"),
        ("registros", "ce_bri", "TEXT"),
        ("registros", "familia", "TEXT"),
        ("registros", "registro", "TEXT"),
        ("registros", "endereco_fabrica", "TEXT"),
        ("registros", "fabrica", "TEXT"),
        ("registros", "arquivo_excel", "TEXT"),
        ("registros", "data_cadastro", "TEXT"),
        ("sistema5_itens", "familia", "TEXT"),
        ("sistema5_itens", "item", "TEXT"),
        ("sistema5_itens", "modelo_ref_key", "TEXT"),
        ("itens", "modelo_ref_key", "TEXT"),
        ("ip_bri_familias", "ce_bri", "TEXT"),
        ("ip_bri_familias", "familia", "TEXT"),
        ("ip_bri_familias", "ip_bri", "TEXT"),
        ("ip_bri_familias", "observacao", "TEXT"),
        ("ip_bri_familias", "data_cadastro", "TEXT"),
        ("ip_bri_familias", "data_atualizacao", "TEXT"),
        ("textos_etiqueta", "tipo", "TEXT DEFAULT 'padrao'"),
    ]

    for tabela, coluna, tipo in ajustes:
        try:
            cursor.execute(f"ALTER TABLE {tabela} ADD COLUMN {coluna} {tipo}")
            commit_seguro()
        except Exception:
            pass

    try:
        cursor.execute("UPDATE registros SET cliente_base = 'BOLSA' WHERE cliente_base IS NULL OR cliente_base = ''")
        commit_seguro()
    except Exception:
        pass


def reparar_tabela_itens_se_necessario():
    """
    Verifica e repara a estrutura da tabela itens se necessário.
    Chamada UMA VEZ na inicialização. Nunca chamar por item inserido.
    """
    obrigatorias = {
        "id",
        "certificado_id",
        "ordem",
        "marca",
        "modelo",
        "nome",
        "codigo"
    }

    try:
        info = cursor.execute("PRAGMA table_info(itens)").fetchall()
        existentes = {linha[1] for linha in info}
    except Exception:
        existentes = set()

    if obrigatorias.issubset(existentes):
        return

    # Cria uma tabela nova correta e tenta preservar o que for possível
    cursor.execute("""
    CREATE TABLE IF NOT EXISTS itens_corrigida (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        certificado_id INTEGER,
        ordem INTEGER,
        marca TEXT,
        modelo TEXT,
        nome TEXT,
        codigo TEXT
    )
    """)

    colunas_comuns = [c for c in ["id", "certificado_id", "ordem", "marca", "modelo", "nome", "codigo"] if c in existentes]

    if colunas_comuns:
        cols = ", ".join(colunas_comuns)
        try:
            cursor.execute(f"INSERT OR IGNORE INTO itens_corrigida ({cols}) SELECT {cols} FROM itens")
        except Exception:
            pass

    try:
        cursor.execute("DROP TABLE itens")
    except Exception:
        pass

    cursor.execute("ALTER TABLE itens_corrigida RENAME TO itens")
    commit_seguro()


def inserir_item_seguro(certificado_id, r):
    cursor.execute("""
    INSERT INTO itens (
        certificado_id,
        ordem,
        marca,
        modelo,
        nome,
        codigo
    )
    VALUES (?, ?, ?, ?, ?, ?)
    """, (
        certificado_id,
        r[0],
        r[1],
        r[2],
        r[3],
        r[4]
    ))


def salvar_ou_atualizar_certificado(
    dados_certificado,
    rows,
    nome_arquivo
):
    ip_bri = dados_certificado["ip_bri"]
    rev_nova = dados_certificado["rev"]

    if not ip_bri:
        return "erro", "IP-BRI não encontrado"

    existente = cursor.execute("""
    SELECT id, rev
    FROM certificados
    WHERE ip_bri = ?
    """, (ip_bri,)).fetchone()

    if not existente:
        cursor.execute("""
        INSERT INTO certificados (
            ip_bri,
            produto,
            ce_bri,
            familia,
            rev,
            data_emissao,
            arquivo_pdf,
            data_cadastro,
            data_atualizacao
        )
        VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)
        """, (
            ip_bri,
            dados_certificado["produto"],
            dados_certificado["ce_bri"],
            dados_certificado.get("familia"),
            rev_nova,
            dados_certificado["data_emissao"],
            nome_arquivo,
            datetime.now().strftime("%d/%m/%Y %H:%M:%S"),
            datetime.now().strftime("%d/%m/%Y %H:%M:%S")
        ))

        commit_seguro()

        certificado_id = cursor.lastrowid

        for r in rows:
            cursor.execute("""
            INSERT INTO itens (
                certificado_id,
                ordem,
                marca,
                modelo,
                nome,
                codigo
            )
            VALUES (?, ?, ?, ?, ?, ?)
            """, (
                certificado_id,
                r[0],
                r[1],
                r[2],
                r[3],
                r[4]
            ))

        commit_seguro()

        registrar_historico(
            ip_bri,
            None,
            rev_nova,
            "",
            "",
            "CERTIFICADO_NOVO",
            "",
            "",
            f"Novo certificado com {len(rows)} itens",
            nome_arquivo
        )

        return "novo", "Novo certificado salvo"

    certificado_id = existente[0]
    rev_antiga = int(existente[1] or 0)

    # Verifica se o certificado existe, mas ficou sem itens no banco.
    # Nesse caso, mesmo com REV igual ou menor, o sistema preenche os itens sem apagar dados úteis.
    qtd_itens_existentes = cursor.execute("""
    SELECT COUNT(*)
    FROM itens
    WHERE certificado_id = ?
    """, (certificado_id,)).fetchone()[0]

    if rev_nova <= rev_antiga and qtd_itens_existentes > 0:
        registrar_historico(
            ip_bri,
            rev_antiga,
            rev_nova,
            "",
            "",
            "REV_IGNORADA",
            "",
            f"REV banco: {rev_antiga}",
            f"REV enviada: {rev_nova}",
            nome_arquivo
        )

        return "ignorado", "REV menor ou igual. Banco não alterado."

    if rev_nova <= rev_antiga and qtd_itens_existentes == 0:
        cursor.execute("""
        UPDATE certificados
        SET
            produto = ?,
            ce_bri = ?,
            familia = ?,
            rev = ?,
            data_emissao = ?,
            arquivo_pdf = ?,
            data_atualizacao = ?
        WHERE id = ?
        """, (
            dados_certificado["produto"],
            dados_certificado["ce_bri"],
            dados_certificado.get("familia"),
            rev_nova,
            dados_certificado["data_emissao"],
            nome_arquivo,
            datetime.now().strftime("%d/%m/%Y %H:%M:%S"),
            certificado_id
        ))

        for r in rows:
            inserir_item_seguro(certificado_id, r)

        commit_seguro()

        registrar_historico(
            ip_bri,
            rev_antiga,
            rev_nova,
            "",
            "",
            "ITENS_REINSERIDOS",
            "",
            "Certificado existia sem itens vinculados",
            f"{len(rows)} itens inseridos na REV {rev_nova}",
            nome_arquivo
        )

        return "corrigido", f"Certificado já existia, mas estava sem itens. {len(rows)} itens foram inseridos."

    antigos = cursor.execute("""
    SELECT
        ordem,
        marca,
        modelo,
        nome,
        codigo
    FROM itens
    WHERE certificado_id = ?
    """, (certificado_id,)).fetchall()

    comparar_itens(
        ip_bri,
        rev_antiga,
        rev_nova,
        antigos,
        rows,
        nome_arquivo
    )

    cursor.execute("""
    DELETE FROM itens
    WHERE certificado_id = ?
    """, (certificado_id,))

    cursor.execute("""
    UPDATE certificados
    SET
        produto = ?,
        ce_bri = ?,
        familia = ?,
        rev = ?,
        data_emissao = ?,
        arquivo_pdf = ?,
        data_atualizacao = ?
    WHERE id = ?
    """, (
        dados_certificado["produto"],
        dados_certificado["ce_bri"],
        dados_certificado.get("familia"),
        rev_nova,
        dados_certificado["data_emissao"],
        nome_arquivo,
        datetime.now().strftime("%d/%m/%Y %H:%M:%S"),
        certificado_id
    ))

    for r in rows:
        inserir_item_seguro(certificado_id, r)

    commit_seguro()

    registrar_historico(
        ip_bri,
        rev_antiga,
        rev_nova,
        "",
        "",
        "CERTIFICADO_ATUALIZADO",
        "",
        f"REV {rev_antiga}",
        f"REV {rev_nova}",
        nome_arquivo
    )

    return "atualizado", f"Atualizado REV {rev_antiga} → {rev_nova}"


# ==========================================
# EXPORTAR
# ==========================================


def exportar_todos_itens():
    return pd.read_sql_query("""
    SELECT
        i.marca AS marca,
        i.modelo AS modelo,
        i.nome AS nome,
        i.codigo AS codigo,
        c.ip_bri AS ip_bri,
        c.ce_bri AS ce_bri,
        r.registro AS registro,
        c.familia AS fam,
        r.fabrica AS fabrica,
        c.rev AS rev,
        c.produto AS produto,
        c.data_emissao AS data_emissao,
        i.ordem AS ordem
    FROM itens i
    INNER JOIN certificados c
    ON c.id = i.certificado_id
    LEFT JOIN registros r
    ON UPPER(r.ce_bri) = UPPER(c.ce_bri)
    AND r.familia = c.familia
    ORDER BY i.marca, i.modelo, c.ip_bri, i.ordem
    """, conn)


# ==========================================
# REGISTROS EXCEL
# ==========================================


def normalizar_familia(valor):
    if pd.isna(valor):
        return ""

    texto = str(valor).strip()

    if texto.endswith(".0"):
        texto = texto[:-2]

    numeros = re.findall(r"[0-9]+", texto)

    if numeros:
        return str(int(numeros[0]))

    return ""



def normalizar_registro(valor):
    if pd.isna(valor):
        return ""

    texto = str(valor).strip()

    if texto.endswith(".0"):
        texto = texto[:-2]

    return texto



def parece_registro(valor):
    texto = normalizar_registro(valor)

    if not texto:
        return False

    return bool(re.search(r"[0-9]+[ ]*/[ ]*[0-9]+", texto))



def extrair_ce_bri_da_aba(aba):
    match = re.search(
        r"CE-BRI-[A-Z0-9\-]+",
        str(aba),
        flags=re.IGNORECASE
    )

    if match:
        return match.group(0).upper()

    return None



def ler_registros_aba(excel_file, aba):
    bruto = pd.read_excel(
        excel_file,
        sheet_name=aba,
        header=None,
        dtype=object
    )

    registros = []

    for i, row in bruto.iterrows():
        valores = [str(v).strip().upper() for v in row.values]

        cols_familia = [
            idx for idx, valor in enumerate(valores)
            if "FAMILIA" in valor or "FAMÍLIA" in valor
        ]

        cols_registro = [
            idx for idx, valor in enumerate(valores)
            if "REGISTRO" in valor
        ]

        if not cols_familia or not cols_registro:
            continue

        for col_familia in cols_familia:
            registros_direita = [c for c in cols_registro if c > col_familia]

            if registros_direita:
                col_registro = registros_direita[0]
            else:
                col_registro = min(cols_registro, key=lambda c: abs(c - col_familia))

            for j in range(i + 1, len(bruto)):
                linha_atual = [str(v).strip().upper() for v in bruto.iloc[j].values]
                linha_texto = " ".join([v for v in linha_atual if v and v.lower() != "nan"])

                if ("FAMILIA" in linha_texto or "FAMÍLIA" in linha_texto) and "REGISTRO" in linha_texto:
                    break

                familia_raw = bruto.iat[j, col_familia] if col_familia < bruto.shape[1] else None
                registro_raw = bruto.iat[j, col_registro] if col_registro < bruto.shape[1] else None

                familia = normalizar_familia(familia_raw)
                registro = normalizar_registro(registro_raw)

                if familia and registro and parece_registro(registro):
                    registros.append({
                        "FAMILIA": familia,
                        "REGISTRO": registro,
                        "FABRICA": aba,
                        "CE_BRI": extrair_ce_bri_da_aba(aba)
                    })

    if not registros:
        for i, row in bruto.iterrows():
            for col in range(bruto.shape[1] - 1):
                familia = normalizar_familia(bruto.iat[i, col])
                registro = normalizar_registro(bruto.iat[i, col + 1])

                if familia and registro and parece_registro(registro):
                    registros.append({
                        "FAMILIA": familia,
                        "REGISTRO": registro,
                        "FABRICA": aba,
                        "CE_BRI": extrair_ce_bri_da_aba(aba)
                    })

    if not registros:
        return None

    dados = pd.DataFrame(registros)

    dados = dados.drop_duplicates(
        subset=["FAMILIA", "REGISTRO", "FABRICA", "CE_BRI"]
    )

    dados["FAMILIA_NUM"] = pd.to_numeric(dados["FAMILIA"], errors="coerce")
    dados = dados.sort_values(
        by=["FAMILIA_NUM", "REGISTRO"]
    ).drop(columns=["FAMILIA_NUM"])

    return dados



def salvar_registros_no_banco(banco_registros, arquivo_nome, cliente_base, enderecos_fabricas=None):
    if banco_registros is None or banco_registros.empty:
        return 0

    cliente_base = clean(cliente_base).upper()
    enderecos_fabricas = enderecos_fabricas or {}

    # Apaga somente os registros da base atual.
    # Exemplo: salvar BOLSA não apaga MOHNISH.
    cursor.execute("DELETE FROM registros WHERE UPPER(cliente_base) = UPPER(?)", (cliente_base,))

    total = 0

    for _, r in banco_registros.iterrows():
        fabrica = r.get("FABRICA", "")
        endereco = enderecos_fabricas.get(str(fabrica), "")

        cursor.execute("""
        INSERT INTO registros (
            cliente_base,
            fabrica,
            ce_bri,
            familia,
            registro,
            endereco_fabrica,
            arquivo_excel,
            data_cadastro
        )
        VALUES (?, ?, ?, ?, ?, ?, ?, ?)
        """, (
            cliente_base,
            fabrica,
            r.get("CE_BRI", ""),
            str(r.get("FAMILIA", "")),
            r.get("REGISTRO", ""),
            endereco,
            arquivo_nome,
            datetime.now().strftime("%d/%m/%Y %H:%M:%S")
        ))
        total += 1

    commit_seguro()
    return total


def salvar_registros_mesclado(banco_registros, arquivo_nome, cliente_base, enderecos_fabricas_novas=None):
    """Mescla os registros do Excel com o que já existe no banco, SEM apagar
    nada (ao contrário de salvar_registros_no_banco, que substitui tudo da
    base). Fábricas que ainda não existem pra esse cliente são criadas
    normalmente, com o endereço informado no upload. Fábricas que JÁ
    existem mantêm o endereço que já está no sistema (o do formulário de
    upload é ignorado pra elas) e só recebem os registros (família/CE-BRI)
    que ainda não estavam cadastrados — os que já existem ficam como estão.
    Retorna um dict-resumo do que foi feito."""
    if banco_registros is None or banco_registros.empty:
        return {"fabricas_novas": [], "fabricas_existentes": [], "novos_registros": 0, "mantidos": 0}

    cliente_base = clean(cliente_base).upper()
    enderecos_fabricas_novas = enderecos_fabricas_novas or {}

    # Snapshot de quais fábricas já existiam ANTES deste upload -- fixo,
    # pra cada fábrica cair em só uma categoria no resumo mesmo que ganhe
    # registros novos ao longo do laço abaixo.
    _fabricas_originais = {
        r[0] for r in cursor.execute(
            "SELECT DISTINCT fabrica FROM registros WHERE UPPER(cliente_base) = UPPER(?)",
            (cliente_base,)
        ).fetchall()
    }

    _resumo = {"fabricas_novas": set(), "fabricas_existentes": set(), "novos_registros": 0, "mantidos": 0}

    for _, r in banco_registros.iterrows():
        fabrica = str(r.get("FABRICA", ""))
        familia = str(r.get("FAMILIA", ""))
        ce_bri = str(r.get("CE_BRI", "") or "")
        registro_valor = r.get("REGISTRO", "")

        if fabrica in _fabricas_originais:
            _end_row = cursor.execute(
                """SELECT endereco_fabrica FROM registros
                   WHERE UPPER(cliente_base) = UPPER(?) AND fabrica = ?
                   AND endereco_fabrica IS NOT NULL AND endereco_fabrica != '' LIMIT 1""",
                (cliente_base, fabrica)
            ).fetchone()
            endereco = _end_row[0] if _end_row else ""
            _resumo["fabricas_existentes"].add(fabrica)
        else:
            endereco = enderecos_fabricas_novas.get(fabrica, "")
            _resumo["fabricas_novas"].add(fabrica)

        _ja_tem = cursor.execute(
            """SELECT id FROM registros
               WHERE UPPER(cliente_base) = UPPER(?) AND fabrica = ?
               AND familia = ? AND IFNULL(ce_bri, '') = ?""",
            (cliente_base, fabrica, familia, ce_bri)
        ).fetchone()

        if _ja_tem:
            _resumo["mantidos"] += 1
            continue

        cursor.execute("""
        INSERT INTO registros (
            cliente_base, fabrica, ce_bri, familia, registro,
            endereco_fabrica, arquivo_excel, data_cadastro
        ) VALUES (?, ?, ?, ?, ?, ?, ?, ?)
        """, (
            cliente_base, fabrica, ce_bri, familia, registro_valor,
            endereco, arquivo_nome, datetime.now().strftime("%d/%m/%Y %H:%M:%S")
        ))
        _resumo["novos_registros"] += 1

    commit_seguro()
    _resumo["fabricas_novas"] = sorted(_resumo["fabricas_novas"])
    _resumo["fabricas_existentes"] = sorted(_resumo["fabricas_existentes"])
    return _resumo


def exportar_registros_banco():
    return pd.read_sql_query("""
    SELECT
        cliente_base,
        fabrica,
        ce_bri,
        familia,
        registro,
        endereco_fabrica,
        arquivo_excel,
        data_cadastro
    FROM registros
    ORDER BY cliente_base, fabrica, CAST(familia AS INTEGER), registro
    """, conn)



@st.cache_data(ttl=60, show_spinner=False)
def buscar_registro_por_certificado(ce_bri, familia):
    if not ce_bri or not familia:
        return pd.DataFrame()

    return pd.read_sql_query("""
    SELECT
        fabrica,
        ce_bri,
        familia,
        registro,
        endereco_fabrica,
        arquivo_excel,
        data_cadastro
    FROM registros
    WHERE UPPER(ce_bri) = UPPER(?)
    AND familia = ?
    ORDER BY fabrica, familia, registro
    """, conn, params=(ce_bri, str(int(familia))))


def gerar_backup_geral_zip():
    buffer = BytesIO()

    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as zipf:
        data_hora = datetime.now().strftime("%Y-%m-%d_%H-%M-%S")

        if Path(DB_PATH).exists():
            zipf.write(DB_PATH, f"backup_db/certificados_{data_hora}.db")

        try:
            banco_completo = exportar_todos_itens()
            zipf.writestr(
                f"csv/banco_completo_{data_hora}.csv",
                banco_completo.to_csv(index=False, sep=";").encode("ISO-8859-1", errors="replace")
            )
        except Exception:
            pass

        try:
            registros = exportar_registros_banco()
            zipf.writestr(
                f"csv/registros_{data_hora}.csv",
                registros.to_csv(index=False, sep=";").encode("ISO-8859-1", errors="replace")
            )
        except Exception:
            pass

        try:
            certificados = pd.read_sql_query("SELECT * FROM certificados ORDER BY ip_bri", conn)
            zipf.writestr(
                f"csv/certificados_{data_hora}.csv",
                certificados.to_csv(index=False, sep=";").encode("ISO-8859-1", errors="replace")
            )
        except Exception:
            pass

        try:
            itens = pd.read_sql_query("SELECT * FROM itens ORDER BY certificado_id, ordem", conn)
            zipf.writestr(
                f"csv/itens_puros_{data_hora}.csv",
                itens.to_csv(index=False, sep=";").encode("ISO-8859-1", errors="replace")
            )
        except Exception:
            pass

        try:
            historico = pd.read_sql_query("SELECT * FROM historico_alteracoes ORDER BY id DESC", conn)
            zipf.writestr(
                f"csv/historico_{data_hora}.csv",
                historico.to_csv(index=False, sep=";").encode("ISO-8859-1", errors="replace")
            )
        except Exception:
            pass

        try:
            resumo = f"""BACKUP GERAL C XML BR ENGINE
Gerado em: {datetime.now().strftime('%d/%m/%Y %H:%M:%S')}

Arquivos incluídos:
- certificados.db
- banco_completo.csv
- registros.csv
- certificados.csv
- itens_puros.csv
- historico.csv
"""
            zipf.writestr("LEIA-ME.txt", resumo.encode("ISO-8859-1", errors="replace"))
        except Exception:
            pass

    buffer.seek(0)
    return buffer



def aviso_backup_diario():
    try:
        hora_sp = datetime.now(ZoneInfo("America/Sao_Paulo")).hour
        if hora_sp >= 18:
            st.sidebar.error("⚠️ Faça o backup geral antes de sair. O Streamlit pode reiniciar e apagar os dados.")
    except Exception:
        pass



@st.cache_data(ttl=10, show_spinner=False)
def metricas_banco():
    """Métricas gerais do banco — cached 10s para evitar 12 queries COUNT redundantes."""
    try:
        total_certs   = cursor.execute("SELECT COUNT(*) FROM certificados").fetchone()[0]
        total_itens   = cursor.execute("SELECT COUNT(*) FROM itens").fetchone()[0]
        total_marcas  = cursor.execute("SELECT COUNT(DISTINCT marca) FROM itens WHERE marca IS NOT NULL AND marca != ''").fetchone()[0]
        total_regs    = cursor.execute("SELECT COUNT(*) FROM registros").fetchone()[0]
        total_s5      = cursor.execute("SELECT COUNT(*) FROM sistema5_itens").fetchone()[0]
        ultimo_cert   = cursor.execute("SELECT ip_bri FROM certificados ORDER BY id DESC LIMIT 1").fetchone()
        return {
            "total_certs":  total_certs,
            "total_itens":  total_itens,
            "total_marcas": total_marcas,
            "total_regs":   total_regs,
            "total_s5":     total_s5,
            "ultimo_cert":  ultimo_cert[0] if ultimo_cert else "—",
        }
    except Exception:
        return {}


def atualizar_familias_certificados():
    """
    Executada UMA VEZ na inicialização.
    Centraliza: estrutura do banco, reparo da tabela itens, famílias faltantes
    e normalização de IPs de processo.
    """
    try:
        garantir_estrutura_banco()
    except Exception:
        pass

    try:
        reparar_tabela_itens_se_necessario()
    except Exception:
        pass

    try:
        certificados = cursor.execute("SELECT id, ip_bri FROM certificados").fetchall()
        for cert_id, ip_bri_atual in certificados:
            fam = extrair_familia_ip_bri(ip_bri_atual)
            if fam:
                cursor.execute("UPDATE certificados SET familia = ? WHERE id = ?", (fam, cert_id))
        commit_seguro()
    except Exception:
        pass

    # Normaliza IPs de processo salvos com espaço (IP 0852-26 → IP-0852-26)
    try:
        ips_antigos = cursor.execute(
            "SELECT id, ip_processo FROM sistema5_arquivos WHERE ip_processo LIKE 'IP %'"
        ).fetchall()
        for arq_id, ip_old in ips_antigos:
            ip_new = normalizar_ip_processo(ip_old)
            if ip_new != ip_old:
                cursor.execute(
                    "UPDATE sistema5_arquivos SET ip_processo = ? WHERE id = ?",
                    (ip_new, arq_id)
                )
        if ips_antigos:
            commit_seguro()
    except Exception:
        pass


atualizar_familias_certificados()

# ==========================================
# PREENCHIMENTO DE CONFIRMAÇÃO
# ==========================================


def cor_rgb(cell):
    try:
        fill = cell.fill
        if fill is None or fill.fill_type is None:
            return None

        color = fill.fgColor

        if color is None:
            return None

        rgb = color.rgb

        if not rgb:
            return None

        rgb = str(rgb).replace("#", "")

        if len(rgb) == 8:
            rgb = rgb[2:]

        if len(rgb) != 6:
            return None

        return tuple(int(rgb[i:i+2], 16) for i in (0, 2, 4))
    except Exception:
        return None



def eh_amarelo(cell):
    rgb = cor_rgb(cell)
    if not rgb:
        return False

    r, g, b = rgb

    # Amarelo comum/forte: FFEB00, FFFF00, FFE699, etc.
    if r >= 170 and g >= 140 and b <= 170:
        return True

    # Amarelo claro de templates
    if r >= 220 and g >= 200 and b <= 190:
        return True

    return False


def eh_verde(cell):
    rgb = cor_rgb(cell)
    if not rgb:
        return False

    r, g, b = rgb

    # Verde forte/claro: 92D050, 70AD47, A9D18E, etc.
    if g >= 110 and g >= r and g >= b and r <= 210:
        return True

    # Verde bem claro usado em alguns templates
    if g >= 170 and r <= 220 and b <= 220 and g > r - 10:
        return True

    return False


def eh_azul(cell):
    rgb = cor_rgb(cell)
    if not rgb:
        return False

    r, g, b = rgb
    return b >= 120 and b > r and b >= g



def normalizar_cabecalho(valor):
    texto = clean(valor).upper()
    texto = texto.replace("CÓDIGO", "CODIGO")
    texto = texto.replace("COD.", "CODIGO")
    texto = texto.replace("CÓD.", "CODIGO")
    texto = texto.replace("IP DO PROCESSO", "IP_PROCESSO")
    texto = texto.replace("IP_PROCESSO", "IP_PROCESSO")
    texto = texto.replace("IP-PROCESSO", "IP_PROCESSO")
    texto = texto.replace("IP PROCESSO", "IP_PROCESSO")
    texto = texto.replace("IP-BRI", "IP_BRI")
    texto = texto.replace("CE-BRI", "CE_BRI")
    texto = texto.replace("ENDEREÇO", "ENDERECO")
    texto = texto.replace("END.", "ENDERECO")
    texto = texto.replace("ENDERECO_DA_FABRICA", "ENDERECO")
    texto = texto.replace("ENDEREÇO_DA_FABRICA", "ENDERECO")
    texto = texto.replace("ENDERECO_FABRICA", "ENDERECO")
    texto = texto.replace("ENDEREÇO_FABRICA", "ENDERECO")
    texto = texto.replace(" ", "_")
    texto = texto.replace("ENDERECO_DA_FABRICA", "ENDERECO")
    texto = texto.replace("ENDERECO_FABRICA", "ENDERECO")
    return texto



@st.cache_data(ttl=60, show_spinner=False)
def buscar_candidatos_confirmacao(valor_ref):
    """
    Retorna TODOS os candidatos únicos para uma referência.
    Se houver mais de 1 modelo distinto, o usuário precisa escolher.
    """
    ref = clean(valor_ref)
    if not ref:
        return []

    todos_candidatos = []

    # ---- Busca no Sistema 5 via função completa (já enriquece IP_BRI e REGISTRO) ----
    try:
        ref_key = referencia_chave_exata(ref)
        ref_sem_zero = ref.lstrip("0") or ref

        s5_result = pd.read_sql_query("""
        SELECT DISTINCT
            marca AS MARCA, modelo AS MODELO, nome AS NOME,
            codigo AS CODIGO, ip_processo AS IP_PROCESSO,
            ce_bri AS CE_BRI, NULL AS REGISTRO,
            endereco_fabrica AS ENDERECO,
            fabrica AS FABRICA, tipo_processo AS TIPO_PROCESSO,
            data_processo AS DATA_PROCESSO,
            familia AS FAMILIA, arquivo_nome AS ARQUIVO_ORIGEM,
            id AS ID_SISTEMA5
        FROM sistema5_itens
        WHERE (modelo_ref_key = ?
           OR UPPER(modelo) LIKE UPPER(?)
           OR UPPER(modelo) LIKE UPPER(?))
          AND UPPER(COALESCE(familia,'')) NOT LIKE '%DESCONSIDER%'
        ORDER BY COALESCE(
            data_processo,
            substr(data_upload,7,4) || '-' || substr(data_upload,4,2) || '-' || substr(data_upload,1,2),
            '1900-01-01'
        ) DESC, id DESC
        LIMIT 50
        """, conn, params=(ref_key, f"{ref}%", f"{ref_sem_zero}%"))

        if not s5_result.empty:
            s5_result = filtrar_resultado_por_referencia_exata(s5_result, ref)
            s5_result = s5_result.drop_duplicates(subset=["MODELO", "MARCA"])

            # Enriquecer cada candidato do S5 com IP_BRI e REGISTRO corretos
            for _, row in s5_result.iterrows():
                item = row.to_dict()
                ce_bri = clean(item.get("CE_BRI"))
                familia = clean(item.get("FAMILIA"))

                # IP_BRI: busca certificado oficial, senão usa família
                if ce_bri and familia:
                    oficial = cursor.execute("""
                    SELECT ip_bri FROM certificados
                    WHERE UPPER(ce_bri) = UPPER(?) AND familia = ?
                    ORDER BY rev DESC LIMIT 1
                    """, (ce_bri, familia)).fetchone()
                    if oficial:
                        item["IP_BRI"] = oficial[0]
                    else:
                        ip_manual = buscar_ip_bri_manual_familia(ce_bri, familia)
                        item["IP_BRI"] = ip_manual if ip_manual else familia

                # REGISTRO: busca na tabela registros pelo CE-BRI + família
                if ce_bri and familia:
                    reg = cursor.execute("""
                    SELECT registro FROM registros
                    WHERE UPPER(ce_bri) = UPPER(?) AND familia = ?
                    ORDER BY id DESC LIMIT 1
                    """, (ce_bri, familia)).fetchone()
                    if reg:
                        item["REGISTRO"] = reg[0]

                todos_candidatos.append(item)
    except Exception:
        pass

    # ---- Busca nos certificados ----
    try:
        ref_key = referencia_chave_exata(ref)

        cert_result = pd.read_sql_query("""
        SELECT
            i.marca AS MARCA, i.modelo AS MODELO, i.nome AS NOME,
            i.codigo AS CODIGO, c.ip_bri AS IP_BRI,
            c.ce_bri AS CE_BRI, r.registro AS REGISTRO,
            r.endereco_fabrica AS ENDERECO
        FROM itens i
        INNER JOIN certificados c ON c.id = i.certificado_id
        LEFT JOIN registros r
            ON UPPER(r.ce_bri) = UPPER(c.ce_bri)
            AND r.familia = c.familia
        WHERE i.modelo_ref_key = ?
           OR UPPER(i.modelo) LIKE UPPER(?)
        ORDER BY c.rev DESC, c.ip_bri DESC
        LIMIT 50
        """, conn, params=(ref_key, f"{ref}%"))

        if not cert_result.empty:
            cert_result = filtrar_resultado_por_referencia_exata(cert_result, ref)
            cert_result = cert_result.drop_duplicates(subset=["MODELO", "MARCA"])
            todos_candidatos.extend(cert_result.to_dict("records"))
    except Exception:
        pass

    # Deduplica por (modelo, marca) — mesma marca + mesmo modelo é o mesmo item
    vistos = set()
    unicos = []
    for c in todos_candidatos:
        modelo = str(c.get("MODELO", "")).strip().upper()
        marca  = str(c.get("MARCA",  "")).strip().upper()
        chave  = (modelo, marca)
        if modelo and chave not in vistos:
            vistos.add(chave)
            unicos.append(c)

    # Só há conflito real se os modelos COMPLETOS são diferentes
    if len(unicos) <= 1:
        return unicos

    modelos_distintos = set(str(c.get("MODELO","")).strip().upper() for c in unicos)
    if len(modelos_distintos) <= 1:
        return [unicos[0]]

    # ---- Colapsa variantes textuais do mesmo item ----
    # O que define o item é a "palavra-núcleo": a primeira palavra depois de
    # "BRINQUEDO" que não seja um termo genérico de embalagem/quantidade
    # (CONJUNTO, KIT, etc). Adjetivos ao redor (DE PLÁSTICO, DE METAL, C/ PILHA)
    # não contam. Se o núcleo é igual, é o mesmo item -> mantém só o primeiro
    # (a lista já vem ordenada pela prioridade padrão: Sistema 5 mais recente
    # primeiro, depois certificados por revisão).
    _PALAVRAS_GENERICAS = {
        "BRINQUEDO", "CONJUNTO", "KIT", "SET", "PACOTE", "MINI", "MODELO",
        "GRANDE", "PEQUENO", "MEDIO", "SORTIDO", "SORTIDOS",
    }

    def _nucleo_modelo(modelo_txt):
        palavras = re.findall(r'[A-ZÀ-Ü0-9]+', modelo_txt.upper())
        for p in palavras:
            if p in _PALAVRAS_GENERICAS or p.isdigit():
                continue
            return p
        return None

    # Colapsa variantes do mesmo item só quando marca E núcleo são iguais.
    # Marcas diferentes nunca são colapsadas — cada uma precisa ser apresentada ao usuário.
    _colapsados = []
    for cand in unicos:
        _nucleo_atual = _nucleo_modelo(str(cand.get("MODELO", "")))
        _marca_atual  = str(cand.get("MARCA", "")).strip().upper()
        _eh_variante  = False
        if _nucleo_atual:
            for existente in _colapsados:
                _nucleo_existente = _nucleo_modelo(str(existente.get("MODELO", "")))
                _marca_existente  = str(existente.get("MARCA", "")).strip().upper()
                if (_nucleo_existente and _nucleo_atual == _nucleo_existente
                        and _marca_atual == _marca_existente):
                    _eh_variante = True
                    break
        if not _eh_variante:
            _colapsados.append(cand)

    return _colapsados


def buscar_item_confirmacao(valor_ref, _debug_ref=None):
    ref = clean(valor_ref)
    if not ref:
        return None

    item_s5 = buscar_item_sistema5_confirmacao(ref, _debug_ref=_debug_ref)
    if item_s5:
        return item_s5

    ref_key = referencia_chave_exata(ref)

    resultado = pd.read_sql_query("""
    SELECT
        i.marca AS MARCA,
        i.modelo AS MODELO,
        i.nome AS NOME,
        i.codigo AS CODIGO,
        c.ip_bri AS IP_BRI,
        c.ce_bri AS CE_BRI,
        r.registro AS REGISTRO,
        r.endereco_fabrica AS ENDERECO
    FROM itens i
    INNER JOIN certificados c
    ON c.id = i.certificado_id
    LEFT JOIN registros r
    ON UPPER(r.ce_bri) = UPPER(c.ce_bri)
    AND r.familia = c.familia
    WHERE i.modelo_ref_key = ?
    ORDER BY c.rev DESC, c.ip_bri DESC
    LIMIT 10
    """, conn, params=(ref_key,))

    # Valida que o modelo encontrado realmente corresponde à referência
    if not resultado.empty:
        resultado = filtrar_resultado_por_referencia_exata(resultado, ref)

    if resultado.empty:
        candidatos = pd.read_sql_query("""
        SELECT
            i.marca AS MARCA,
            i.modelo AS MODELO,
            i.nome AS NOME,
            i.codigo AS CODIGO,
            c.ip_bri AS IP_BRI,
            c.ce_bri AS CE_BRI,
            r.registro AS REGISTRO,
            r.endereco_fabrica AS ENDERECO
        FROM itens i
        INNER JOIN certificados c
        ON c.id = i.certificado_id
        LEFT JOIN registros r
        ON UPPER(r.ce_bri) = UPPER(c.ce_bri)
        AND r.familia = c.familia
        WHERE UPPER(i.modelo) LIKE UPPER(?)
        ORDER BY c.rev DESC, c.ip_bri DESC
        LIMIT 100
        """, conn, params=(f"{ref}%",))
        resultado = filtrar_resultado_por_referencia_exata(candidatos, ref)

    if resultado.empty:
        return None

    return resultado.iloc[0].to_dict()


def descobrir_cabecalho_coluna(ws, row_idx, col_idx):
    # Primeiro procura cabeçalho azul acima da célula amarela
    for r in range(row_idx - 1, 0, -1):
        cell = ws.cell(r, col_idx)
        valor = clean(cell.value)

        if valor and (eh_azul(cell) or normalizar_cabecalho(valor) in ["MARCA", "MODELO", "NOME", "CODIGO", "IP_BRI", "CE_BRI", "REGISTRO", "ENDERECO", "FAMILIA", "ITEM", "TIPO_PROCESSO", "DATA_PROCESSO", "ARQUIVO_ORIGEM", "IP_PROCESSO"]):
            return normalizar_cabecalho(valor)

    # Plano B: procura qualquer texto de cabeçalho acima
    for r in range(row_idx - 1, 0, -1):
        valor = clean(ws.cell(r, col_idx).value)
        cab = normalizar_cabecalho(valor)

        if cab in ["MARCA", "MODELO", "NOME", "CODIGO", "IP_BRI", "CE_BRI", "REGISTRO", "ENDERECO", "FAMILIA", "ITEM", "TIPO_PROCESSO", "DATA_PROCESSO", "ARQUIVO_ORIGEM", "IP_PROCESSO"]:
            return cab

    return ""



def _limpar_excel_imagens(uploaded_file):
    """
    Remove arquivos de imagem E drawings problemáticos do Excel (ZIP).
    Também limpa referências a esses arquivos nos XMLs de relacionamento.
    Cobre:
      - imagens com extensão problemática (.mpo, .emf, .wmf)
      - drawings com XML inválido para o openpyxl (ex: valor fora de range no fill)
    """
    try:
        uploaded_file.seek(0)
    except Exception:
        pass

    data = uploaded_file.read()
    import zipfile as _zf
    import re as _re

    buf_limpo = BytesIO()
    extensoes_imagem = {".mpo", ".emf", ".wmf", ".jpeg", ".jpg", ".png", ".gif", ".bmp", ".tiff"}
    arquivos_removidos = set()

    with _zf.ZipFile(BytesIO(data), "r") as zin:
        nomes = {i.filename for i in zin.infolist()}

        # 1ª passagem: identifica quais arquivos remover
        for item in zin.infolist():
            nome_lower = item.filename.lower()
            ext = "." + nome_lower.rsplit(".", 1)[-1] if "." in nome_lower else ""

            # Remove imagens e drawings (drawings têm XML inválido com frequência)
            if ext in extensoes_imagem:
                arquivos_removidos.add(item.filename)
            elif "xl/drawings" in nome_lower and nome_lower.endswith(".xml"):
                arquivos_removidos.add(item.filename)

        # 2ª passagem: grava ZIP sem os arquivos removidos, corrigindo .rels
        with _zf.ZipFile(buf_limpo, "w", _zf.ZIP_DEFLATED) as zout:
            for item in zin.infolist():
                if item.filename in arquivos_removidos:
                    continue

                conteudo = zin.read(item.filename)
                nome_lower = item.filename.lower()

                # Nos arquivos de relacionamento, remove <Relationship> que apontam
                # para arquivos removidos (evita que o openpyxl tente resolver a ref)
                if nome_lower.endswith(".rels"):
                    try:
                        texto = conteudo.decode("utf-8", errors="replace")
                        for arq_rem in arquivos_removidos:
                            nome_base = arq_rem.split("/")[-1]
                            # Remove a linha <Relationship ... Target="...nomebase..." .../>
                            texto = _re.sub(
                                rf'<Relationship\s[^>]*["\'/]{_re.escape(nome_base)}["\'][^>]*/?>',
                                '',
                                texto,
                                flags=_re.IGNORECASE
                            )
                        conteudo = texto.encode("utf-8")
                    except Exception:
                        pass

                # No [Content_Types].xml remove Default/Override das extensões/arquivos
                # removidos — sem isso o openpyxl levanta KeyError('<ext>') ao abrir o ZIP
                elif nome_lower == "[content_types].xml":
                    try:
                        texto = conteudo.decode("utf-8", errors="replace")
                        extensoes_removidas = {
                            ("." + f.rsplit(".", 1)[-1]).lower()
                            for f in arquivos_removidos
                            if "." in f.rsplit("/", 1)[-1]
                        }
                        for ext_rem in extensoes_removidas:
                            ext_sem_ponto = _re.escape(ext_rem.lstrip("."))
                            texto = _re.sub(
                                rf'<Default\b[^>]*\bExtension=["\']?{ext_sem_ponto}["\']?[^>]*/?>',
                                '',
                                texto,
                                flags=_re.IGNORECASE
                            )
                        for arq_rem in arquivos_removidos:
                            nome_base = _re.escape(arq_rem.split("/")[-1])
                            texto = _re.sub(
                                rf'<Override\b[^>]*\bPartName=["\'][^"\']*{nome_base}["\'][^>]*/?>',
                                '',
                                texto,
                                flags=_re.IGNORECASE
                            )
                        conteudo = texto.encode("utf-8")
                    except Exception:
                        pass

                zout.writestr(item, conteudo)

    buf_limpo.seek(0)
    return buf_limpo


def preencher_excel_confirmacao(uploaded_file, escolhas_pre=None):
    """
    escolhas_pre: dict {ref -> item_dict} com escolhas manuais para referências ambíguas.
    """
    if escolhas_pre is None:
        escolhas_pre = {}

    # Tenta abrir normalmente primeiro
    try:
        uploaded_file.seek(0)
    except Exception:
        pass

    wb = None
    _exc_final = None

    # Tentativa 1: abertura normal
    try:
        uploaded_file.seek(0)
        wb = load_workbook(uploaded_file)
    except Exception as _e1:
        _exc_final = _e1

    # Tentativa 2: limpa imagens problemáticas do ZIP (mpo/emf/wmf)
    if wb is None:
        try:
            wb = load_workbook(_limpar_excel_imagens(uploaded_file))
        except Exception as _e2:
            _exc_final = _e2

    # Tentativa 3: read_only=True (parser alternativo, mais tolerante a XML inválido)
    if wb is None:
        try:
            uploaded_file.seek(0)
            wb = load_workbook(uploaded_file, read_only=True, data_only=True)
        except Exception as _e3:
            _exc_final = _e3

    # Tentativa 4: repassa pelo LibreOffice (re-salva como XLSX limpo) e abre novamente
    if wb is None:
        try:
            import tempfile as _tmp_conf, subprocess as _sub_conf, glob as _glob_conf, os as _os_conf
            _soff = _localizar_soffice()
            _tdir = _tmp_conf.mkdtemp()
            _src = _os_conf.path.join(_tdir, "entrada.xlsx")
            uploaded_file.seek(0)
            with open(_src, "wb") as _f:
                _f.write(uploaded_file.read())
            _sub_conf.run(
                [_soff, "--headless", "--convert-to", "xlsx", "--outdir", _tdir, _src,
                 f"-env:UserInstallation=file://{_tdir}/lo"],
                capture_output=True, timeout=120
            )
            _convertidos = _glob_conf.glob(_os_conf.path.join(_tdir, "*.xlsx"))
            _convertido = next((x for x in _convertidos if "entrada" in x), None) or (_convertidos[0] if _convertidos else None)
            if _convertido:
                wb = load_workbook(_convertido)
        except Exception as _e4:
            _exc_final = _e4

    if wb is None:
        raise Exception(f"Não foi possível abrir o Excel após 4 tentativas. Último erro: {_exc_final}")

    preenchidos = 0
    nao_encontrados = []

    diagnostico = {
        "referencias_verdes_lidas": 0,
        "linhas_com_referencia": 0,
        "itens_encontrados": 0,
        "celulas_amarelas_detectadas": 0,
        "campos_reconhecidos": 0
    }

    campos_validos = ["MARCA", "MODELO", "NOME", "CODIGO", "IP_BRI", "CE_BRI", "REGISTRO", "ENDERECO", "FAMILIA", "ITEM", "TIPO_PROCESSO", "DATA_PROCESSO", "ARQUIVO_ORIGEM", "IP_PROCESSO"]

    for ws in wb.worksheets:
        for row in ws.iter_rows():
            refs = []

            # REGRA OFICIAL:
            # Qualquer célula verde da linha é considerada referência.
            for cell in row:
                if eh_verde(cell) and clean(cell.value):
                    val = ler_valor_celula_preservando_zeros(cell)
                    if val:
                        refs.append(val)

            if not refs:
                continue

            diagnostico["linhas_com_referencia"] += 1
            diagnostico["referencias_verdes_lidas"] += len(refs)

            item = None

            for ref in refs:
                # Usa escolha manual se o usuário desambiguou
                if ref in escolhas_pre:
                    item = escolhas_pre[ref]
                else:
                    item = buscar_item_confirmacao(ref)
                if item:
                    break

            if not item:
                nao_encontrados.extend(refs)
                continue

            diagnostico["itens_encontrados"] += 1

            for cell in row:
                if not eh_amarelo(cell):
                    continue

                diagnostico["celulas_amarelas_detectadas"] += 1

                campo = descobrir_cabecalho_coluna(ws, cell.row, cell.column)

                if campo not in campos_validos:
                    continue

                diagnostico["campos_reconhecidos"] += 1

                valor = item.get(campo)

                if valor is None:
                    continue

                valor = str(valor).strip()

                if valor == "" or valor.lower() == "nan":
                    continue

                if campo == "CODIGO":
                    cell.number_format = "@"
                    cell.value = str(valor)
                elif campo == "NOME":
                    cell.value = limitar_nome(valor)
                else:
                    cell.value = valor

                preenchidos += 1

    output = BytesIO()
    wb.save(output)
    output.seek(0)

    st.session_state["diagnostico_confirmacao"] = diagnostico

    return output, preenchidos, nao_encontrados

def atualizar_endereco_fabrica_existente(cliente_base, fabrica, novo_endereco):
    cliente_base = clean(cliente_base).upper()
    fabrica = clean(fabrica)
    novo_endereco = clean(novo_endereco)

    cursor.execute("""
    UPDATE registros
    SET endereco_fabrica = ?
    WHERE UPPER(cliente_base) = UPPER(?)
    AND fabrica = ?
    """, (
        novo_endereco,
        cliente_base,
        fabrica
    ))

    commit_seguro()
    return cursor.rowcount


# ==========================================
# BIBLIOTECA DE SELOS "SEGURANÇA/REGISTRO" (base pro módulo de Geração de Etiquetas)
# ==========================================

_MAPA_VARIANTE_SELO = {
    "amarelo": "selo_amarelo",
    "pb": "selo_pb",
    "compacto_cor": "selo_compacto_cor",
    "compacto_pb": "selo_compacto_pb",
}


def salvar_variante_selo(cliente_base, fabrica, familia, variante, imagem_bytes, registro=None):
    """Salva UMA variante de selo (amarelo/pb/compacto_cor/compacto_pb) para a
    combinação Cliente+Fábrica+Família. Se a linha ainda não existe, cria; se já
    existe, atualiza só a variante enviada (preserva as outras). Se um registro
    for informado (lido por OCR), grava também."""
    cliente_base = clean(cliente_base).upper()
    fabrica = clean(fabrica)
    familia = clean(familia)

    coluna = _MAPA_VARIANTE_SELO.get(variante)
    if not coluna:
        raise ValueError(f"Variante inválida: {variante}")

    agora = datetime.now().strftime("%d/%m/%Y %H:%M:%S")

    # Garante que a linha existe
    cursor.execute("""
    INSERT INTO selos_registro (cliente_base, fabrica, familia, data_cadastro, data_atualizacao)
    VALUES (?, ?, ?, ?, ?)
    ON CONFLICT(cliente_base, fabrica, familia) DO NOTHING
    """, (cliente_base, fabrica, familia, agora, agora))

    # Atualiza a variante enviada
    cursor.execute(f"""
    UPDATE selos_registro
    SET {coluna} = ?, data_atualizacao = ?
    WHERE cliente_base = ? AND fabrica = ? AND familia = ?
    """, (imagem_bytes, agora, cliente_base, fabrica, familia))

    # Grava o registro só se veio um valor e ainda não há um salvo
    if registro:
        cursor.execute("""
        UPDATE selos_registro
        SET registro = ?
        WHERE cliente_base = ? AND fabrica = ? AND familia = ?
          AND (registro IS NULL OR registro = '')
        """, (clean(registro), cliente_base, fabrica, familia))

    commit_seguro()


def atualizar_registro_selo(cliente_base, fabrica, familia, registro):
    """Atualiza manualmente o número de registro (caso o OCR erre)."""
    cliente_base = clean(cliente_base).upper()
    cursor.execute("""
    UPDATE selos_registro SET registro = ?
    WHERE cliente_base = ? AND fabrica = ? AND familia = ?
    """, (clean(registro), cliente_base, clean(fabrica), clean(familia)))
    commit_seguro()


def listar_selos_registro():
    return pd.read_sql_query("""
    SELECT id, cliente_base, fabrica, familia, registro,
           (selo_amarelo IS NOT NULL) AS tem_amarelo,
           (selo_pb IS NOT NULL) AS tem_pb,
           (selo_compacto_cor IS NOT NULL) AS tem_compacto_cor,
           (selo_compacto_pb IS NOT NULL) AS tem_compacto_pb,
           data_cadastro, data_atualizacao
    FROM selos_registro
    ORDER BY cliente_base, fabrica, CAST(familia AS INTEGER)
    """, conn)


def buscar_variante_selo(cliente_base, fabrica, familia, variante):
    coluna = _MAPA_VARIANTE_SELO.get(variante)
    if not coluna:
        return None
    row = cursor.execute(f"""
    SELECT {coluna} FROM selos_registro
    WHERE UPPER(cliente_base) = UPPER(?) AND fabrica = ? AND familia = ?
    """, (clean(cliente_base).upper(), clean(fabrica), clean(familia))).fetchone()
    return row[0] if row and row[0] else None


def buscar_dados_selo(cliente_base, fabrica, familia):
    """Retorna o registro + quais variantes existem para essa combinação."""
    row = cursor.execute("""
    SELECT registro, selo_amarelo, selo_pb, selo_compacto_cor, selo_compacto_pb
    FROM selos_registro
    WHERE UPPER(cliente_base) = UPPER(?) AND fabrica = ? AND familia = ?
    """, (clean(cliente_base).upper(), clean(fabrica), clean(familia))).fetchone()
    if not row:
        return None
    return {
        "registro": row[0],
        "amarelo": row[1],
        "pb": row[2],
        "compacto_cor": row[3],
        "compacto_pb": row[4],
    }


def excluir_selo_registro(id_selo):
    cursor.execute("DELETE FROM selos_registro WHERE id = ?", (id_selo,))
    commit_seguro()
    return cursor.rowcount


# ---- Textos padrão de etiqueta (advertências, pilha, restritivos etc) ----

_NUM_EXTENSO = {
    0: "ZERO", 1: "UM", 2: "DOIS", 3: "TRÊS", 4: "QUATRO", 5: "CINCO",
    6: "SEIS", 7: "SETE", 8: "OITO", 9: "NOVE", 10: "DEZ", 11: "ONZE",
    12: "DOZE", 13: "TREZE", 14: "QUATORZE", 15: "QUINZE", 16: "DEZESSEIS",
    17: "DEZESSETE", 18: "DEZOITO",
}


def gerar_codigo_barras_ean13(numero_13_digitos):
    """Gera a imagem PNG de um código de barras EAN-13 a partir dos 13 dígitos
    completos (como vêm no Excel). Retorna (png_bytes, aviso):
    - png_bytes: imagem do código de barras, ou None se falhar
    - aviso: mensagem de alerta se o dígito verificador não bater (possível erro
      de digitação no Excel), ou None se tudo certo."""
    import io as _io_bc
    numero = re.sub(r'\D', '', str(numero_13_digitos or ''))

    if len(numero) != 13:
        return None, f"Código de barras não tem 13 dígitos (tem {len(numero)}): {numero!r}"

    try:
        from barcode import EAN13
        from barcode.writer import ImageWriter, mm2px, pt2mm
        from PIL import Image as _PILImage, ImageDraw as _PILDraw, ImageFont as _PILFont
    except ImportError:
        return None, "Biblioteca de código de barras (python-barcode) não instalada."

    # Parâmetros do desenho — recalibrados a partir do arquivo codigo_de_barras.docx
    # (mesmo componente SmartCode, mas exibido bem maior: 260,25 x 108,95pt =
    # proporção 2,39:1 — mais "alto e estreito" do que a calibração anterior
    # de 2,07:1). Módulo mais alto pra bater essa proporção sem espremer.
    MODULE_WIDTH = 0.3   # mm por módulo
    MODULE_HEIGHT = 7.36 # mm de altura da barra "normal"
    QUIET_ZONE = 2.0     # mm de margem em branco de cada lado
    FONT_SIZE = 8         # pt
    TEXT_DISTANCE = 0.3  # mm entre a base da barra e o topo do texto (bem colado, como no SmartCode)
    GUARD_HEIGHT_FACTOR = 1.3  # barra de guarda nitidamente mais alta, ajustado a partir do print comparativo
    DPI = 300

    def _localizar_fonte_sem_zero_riscado():
        """A fonte padrão da lib (DejaVuSansMono) desenha o '0' com um risco
        no meio pra não confundir com 'O' — não é o padrão usado nas etiquetas.
        Usa a Arial do Windows (mesma fonte do resto da etiqueta), que não tem
        esse risco; se não achar, cai numa equivalente portátil."""
        import os as _osF
        candidatos = [
            r"C:\Windows\Fonts\arial.ttf",
            r"C:\Windows\Fonts\Arial.ttf",
            "/usr/share/fonts/truetype/liberation/LiberationSans-Regular.ttf",
        ]
        for c in candidatos:
            if _osF.path.exists(c):
                return c
        return None

    try:
        # EAN13 recebe 12 dígitos e calcula o verificador. Passamos os 12 primeiros
        # e conferimos se o 13º calculado bate com o que veio no Excel.
        # guardbar=True ativa as barras de guarda mais altas (esquerda/centro/direita),
        # no padrão visual de código de barras de verdade.
        writer = ImageWriter(mode="RGBA")
        ean = EAN13(numero[:12], writer=writer, guardbar=True)
        codigo_calculado = ean.ean  # dígitos puros (sem formatação)

        aviso = None
        if codigo_calculado != numero:
            aviso = (f"⚠️ Dígito verificador não confere: Excel='{numero}' | "
                     f"correto seria '{codigo_calculado}'. Verifique o número no Excel.")

        # 1) Gera só as barras (sem texto — o texto é desenhado manualmente
        # abaixo, dígito por dígito, pra cada um ficar sob seu grupo de 7
        # módulos, igual ao padrão real, em vez de um bloco só de texto colado).
        buffer_barras = _io_bc.BytesIO()
        ean.write(buffer_barras, options={
            'module_width': MODULE_WIDTH,
            'module_height': MODULE_HEIGHT,
            'quiet_zone': QUIET_ZONE,
            'font_size': FONT_SIZE,
            'text_distance': TEXT_DISTANCE,
            'write_text': False,
            'background': (255, 255, 255, 0),
            'guard_height_factor': GUARD_HEIGHT_FACTOR,
        })
        buffer_barras.seek(0)
        imagem = _PILImage.open(buffer_barras).convert("RGBA")

        # 2) Abre espaço embaixo pra caber o texto (a lib só reserva essa
        # altura sozinha quando write_text=True, que não usamos aqui) e o
        # tanto a mais que a barra de guarda ocupa além da margem padrão (1mm)
        altura_texto_mm = pt2mm(FONT_SIZE) + TEXT_DISTANCE
        altura_guarda_extra_mm = max(0.0, MODULE_HEIGHT * (GUARD_HEIGHT_FACTOR - 1) - 1.0)
        altura_extra_px = int(mm2px(altura_texto_mm + altura_guarda_extra_mm, DPI))
        tela = _PILImage.new("RGBA", (imagem.width, imagem.height + altura_extra_px), (255, 255, 255, 0))
        tela.paste(imagem, (0, 0))

        # 3) Desenha cada dígito centralizado no seu próprio grupo de 7 módulos
        # (estrutura padrão do EAN-13: guarda(3 módulos) + 6 dígitos×7 módulos
        # + guarda central(5) + 6 dígitos×7 módulos + guarda(3) = 95 módulos).
        desenho = _PILDraw.Draw(tela)
        caminho_fonte = _localizar_fonte_sem_zero_riscado() or writer.font_path
        fonte = _PILFont.truetype(caminho_fonte, int(mm2px(pt2mm(FONT_SIZE), DPI)))
        # anchor="ma": o ponto dado é o TOPO do texto (não a base) — assim o
        # texto sempre começa exatamente 'TEXT_DISTANCE' abaixo da barra, sem
        # depender de métrica de fonte (evita o texto subir e encostar na barra).
        topo_texto_mm = 1 + MODULE_HEIGHT + TEXT_DISTANCE  # 1mm = margem superior padrão da lib
        inicio_barras_mm = QUIET_ZONE

        def _desenhar_digito(caractere, centro_x_mm):
            desenho.text(
                (mm2px(centro_x_mm, DPI), mm2px(topo_texto_mm, DPI)),
                caractere, font=fonte, fill=(0, 0, 0, 255), anchor="ma"
            )

        _desenhar_digito(codigo_calculado[0], inicio_barras_mm - 4 * MODULE_WIDTH)
        x_grupo_esquerdo = inicio_barras_mm + 3 * MODULE_WIDTH
        for i in range(6):
            _desenhar_digito(codigo_calculado[1 + i], x_grupo_esquerdo + (i + 0.5) * 7 * MODULE_WIDTH)
        x_grupo_direito = inicio_barras_mm + (3 + 42 + 5) * MODULE_WIDTH
        for i in range(6):
            _desenhar_digito(codigo_calculado[7 + i], x_grupo_direito + (i + 0.5) * 7 * MODULE_WIDTH)

        buffer_final = _io_bc.BytesIO()
        tela.save(buffer_final, format="PNG")
        return buffer_final.getvalue(), aviso
    except Exception as e:
        return None, f"Erro ao gerar código de barras: {e}"


def formatar_quantidade(qtd_pecas):
    """Monta o texto de quantidade no padrão da etiqueta.
    - 1 peça  -> '1 Unidade'
    - >1 peça -> '1 Conjunto c/ 03 Peças' (peças com zero à esquerda)
    Retorna None se a quantidade for inválida."""
    try:
        n = int(qtd_pecas)
    except (ValueError, TypeError):
        return None
    if n <= 0:
        return None
    if n == 1:
        return "1 Unidade"
    return f"1 Conjunto c/ {n:02d} Peças"


def formatar_idade(numero, unidade):
    """Formata a idade no padrão da etiqueta.
    - Anos: número com zero à esquerda + extenso maiúsculo entre parênteses -> '03 (TRÊS) ANOS'
    - Meses: número com zero à esquerda, sem extenso -> '06 MESES'
    Retorna None se não conseguir montar."""
    try:
        n = int(numero)
    except (ValueError, TypeError):
        return None

    unidade = (unidade or "").strip().upper()
    num_fmt = f"{n:02d}"  # zero à esquerda

    if "MES" in unidade or "MÊS" in unidade:
        plural = "MESES" if n != 1 else "MÊS"
        return f"{num_fmt} {plural}"
    else:  # anos (default)
        extenso = _NUM_EXTENSO.get(n, str(n))
        plural = "ANOS" if n != 1 else "ANO"
        return f"{num_fmt} ({extenso}) {plural}"


def preencher_idade_no_texto(texto, idade_formatada):
    """Substitui o marcador {IDADE} pela idade formatada. Se não houver idade
    disponível, mantém o marcador visível pra você perceber que faltou preencher."""
    if not texto:
        return texto
    if idade_formatada:
        return texto.replace("{IDADE}", idade_formatada)
    return texto


def detectar_idade_do_nome(nome):
    """Detecta a idade (indicativo/restritivo) a partir do NOME do produto no Excel.
    Retorna dict: {'numero': int, 'unidade': 'ANOS'/'MESES', 'tipo': 'indicativo'/'restritivo'}
    ou None se não achar. Prioriza o INDICATIVO (que é o que vai no texto de idade)."""
    if not nome:
        return None
    nome_upper = str(nome).upper()

    ind = re.search(r'INDICATIVO\s*\+?\s*(\d+)\s*(ANOS?|MESES?|M[ÊE]S)', nome_upper)
    if not ind:
        ind = re.search(r'\+\s*(\d+)\s*(ANOS?|MESES?|M[ÊE]S)', nome_upper)
    if ind:
        unidade = "MESES" if "MES" in ind.group(2) or "MÊS" in ind.group(2) else "ANOS"
        return {"numero": int(ind.group(1)), "unidade": unidade, "tipo": "indicativo"}

    return None


def detectar_tem_pilha_do_nome(nome):
    """Detecta se o produto tem pilha/bateria pelo NOME do Excel."""
    if not nome:
        return False
    return bool(re.search(r'PILHA|BATERIA|À PILHA|A PILHA', str(nome).upper()))


def separar_referencia_nome(texto_referencia):
    """Separa 'KK-1827 - BRINQUEDO MÓBILE DE PLÁSTICO' em:
    ('KK-1827', 'BRINQUEDO MÓBILE DE PLÁSTICO'). Limpa apóstrofo inicial."""
    if not texto_referencia:
        return "", ""
    txt = str(texto_referencia).strip().lstrip("'\"`´").strip()
    m = re.match(r'^([A-Z0-9][\w\-]*)\s*-\s*(.+)$', txt, re.IGNORECASE)
    if m:
        return m.group(1).strip(), m.group(2).strip()
    return txt, ""


def aplicar_markdown_negrito_no_paragrafo(paragrafo, texto_com_marcacao):
    """Recebe um parágrafo do python-docx e um texto com marcação **negrito**.
    Divide o texto em trechos, aplicando negrito real onde estiver entre ** **,
    e removendo os asteriscos. Ex: '**ATENÇÃO:** não recomendável' vira
    'ATENÇÃO:' em negrito + ' não recomendável' normal."""
    partes = re.split(r'\*\*(.+?)\*\*', texto_com_marcacao)
    for i, parte in enumerate(partes):
        if not parte:
            continue
        run = paragrafo.add_run(parte)
        if i % 2 == 1:  # trecho que estava entre ** **
            run.bold = True
    return paragrafo


def texto_para_preview_html(texto_com_marcacao):
    """Converte a marcação **negrito** em <b> pra pré-visualizar na tela."""
    import html as _html
    escapado = _html.escape(texto_com_marcacao or "")
    return re.sub(r'\*\*(.+?)\*\*', r'<b>\1</b>', escapado)


def salvar_texto_etiqueta(categoria, titulo, conteudo, tipo="padrao", id_existente=None):
    agora = datetime.now().strftime("%d/%m/%Y %H:%M:%S")
    if id_existente:
        cursor.execute("""
        UPDATE textos_etiqueta SET categoria = ?, titulo = ?, conteudo = ?, tipo = ?, data_atualizacao = ?
        WHERE id = ?
        """, (clean(categoria), clean(titulo), conteudo, tipo, agora, id_existente))
    else:
        cursor.execute("""
        INSERT INTO textos_etiqueta (categoria, titulo, conteudo, tipo, data_cadastro, data_atualizacao)
        VALUES (?, ?, ?, ?, ?, ?)
        """, (clean(categoria), clean(titulo), conteudo, tipo, agora, agora))
    commit_seguro()


def listar_textos_etiqueta(tipo=None, categoria=None):
    """Lista textos cadastrados. Filtra por tipo e/ou categoria quando informado."""
    query = "SELECT id, tipo, categoria, titulo, conteudo, data_atualizacao FROM textos_etiqueta"
    condicoes = []
    params = []
    if tipo:
        condicoes.append("tipo = ?")
        params.append(tipo)
    if categoria:
        condicoes.append("categoria = ?")
        params.append(categoria)
    if condicoes:
        query += " WHERE " + " AND ".join(condicoes)
    query += " ORDER BY tipo, categoria, titulo"
    return pd.read_sql_query(query, conn, params=params)


def excluir_texto_etiqueta(id_texto):
    cursor.execute("DELETE FROM textos_etiqueta WHERE id = ?", (id_texto,))
    commit_seguro()
    return cursor.rowcount


# ---- Textos padrão extraídos dos moldes reais (MODELO_ETIQUETA.docx, 22 exemplos) ----
# Cada entrada é um texto distinto de ATENÇÃO/INDICAÇÃO/ADVERTÊNCIA/CUIDADOS DE USO/
# COMPOSIÇÃO encontrado nos exemplos reais da empresa, com a idade generalizada pra
# {IDADE}. Importável uma vez via botão na aba "Textos padrão da etiqueta".
_TEXTOS_PADRAO_SEED = [
    {'tipo': 'padrao', 'categoria': 'INDICAÇÃO', 'titulo': 'Padrão — indicação de idade (genérica, usada em quase tudo)', 'conteudo': '**INDICAÇÃO: **ESTE PRODUTO É INDICADO PARA CRIANÇAS A PARTIR DE {IDADE}.'},
    {'tipo': 'padrao', 'categoria': 'ATENÇÃO', 'titulo': 'Padrão — genérico (partes pequenas)', 'conteudo': '**ATENÇÃO: **NÃO RECOMENDÁVEL PARA CRIANÇAS MENORES DE {IDADE} POR CONTER PARTE (S) PEQUENA (S) QUE PODE(M) SER ENGOLIDA (S) OU ASPIRADA (S).'},
    {'tipo': 'padrao', 'categoria': 'ATENÇÃO', 'titulo': 'Padrão — genérico + montagem por um adulto', 'conteudo': '**ATENÇÃO: **NÃO RECOMENDÁVEL PARA CRIANÇAS MENORES DE {IDADE} POR CONTER PARTE (S) PEQUENA (S) QUE PODE(M) SER ENGOLIDA (S) OU ASPIRADA (S). ESTE PRODUTO DEVE SER MONTADO POR UM ADULTO ANTES DE SER ENTREGUE À CRIANÇA.'},
    {'tipo': 'padrao', 'categoria': 'ATENÇÃO', 'titulo': 'Padrão — genérico + montagem por um adulto (variante de digitação)', 'conteudo': '**ATENÇÃO: **NÃO RECOMENDÁVEL PARA CRIANÇAS MENORES DE {IDADE} POR CONTER PARTE (S) PEQUENA (S) QUE PODE (M) SER ENGOLIDA (S) OU ASPIRADA (S). ESTE PRODUTO DEVERÁ SER MONTADO POR UM ADULTO ANTES DE SER ENTREGUE A CRIANÇA'},
    {'tipo': 'padrao', 'categoria': 'ATENÇÃO', 'titulo': 'Padrão — genérico + risco de asfixia', 'conteudo': '**ATENÇÃO: **NÃO RECOMENDÁVEL PARA CRIANÇAS MENORES DE {IDADE} POR CONTER PARTE (S) PEQUENA (S) QUE PODE (M) SER ENGOLIDA (S) OU ASPIRADA (S). PARA EVITAR O PERIGO DE ASFIXIA, MANTER ESTA EMBALAGEM LONGE DO ALCANCE DAS CRIANÇAS.'},
    {'tipo': 'padrao', 'categoria': 'ATENÇÃO', 'titulo': 'Padrão — brinquedo tipo boneca de sentar (não é assento de verdade)', 'conteudo': '**ATENÇÃO: **NÃO RECOMENDÁVEL PARA CRIANÇAS MENORES DE {IDADE} POR CONTER PARTE (S) PEQUENA (S) QUE PODE(M) SER ENGOLIDA (S) OU ASPIRADA (S). ESTE BRINQUEDO É DESTINADO PARA USO COM BONECAS, A CRIANÇA NÃO DEVE SENTAR NO ASSENTO.'},
    {'tipo': 'padrao', 'categoria': 'ATENÇÃO', 'titulo': 'Padrão — item para berço/cama/carrinho de bebê', 'conteudo': '**ATENÇÃO:**ESTE PRODUTO FOI PROJETADO PARA SER INSTALADO EM BERÇOS, CAMAS OU CARRINHOS DE BEBÊ. DEVE SER INSTALADO DE ACORDO COM AS INSTRUÇÕES. NÃO DEVE SER ENTREGUE SOLTO Á CRIANÇA. PARA EVITAR QUE A CRIANÇA POSSA PRENDER-SE E FERIR-SE, RETIRAR O BRINQUEDO QUANDO A CRIANÇA COMEÇAR A SE LEVANTAR SOBRE AS MÃOS E OS JOELHOS.'},
    {'tipo': 'padrao', 'categoria': 'ATENÇÃO', 'titulo': 'Padrão — mordedor (não pode ir ao congelador/freezer + higienização)', 'conteudo': '**ATENÇÃO!** ESTE PRODUTO NÃO PODE SER COLOCADO EM CONGELADOR E FREEZER. ANTES DO USO RECOMENDA-SE COLOCAR EM ÁGUA FERVENTE OU IMERGIR DURANTE 15 MINUTOS EM SOLUÇÃO DE HIPOCLORITO DE SÓDIO COM 0,5% DE CLORO ATIVO, QUE PODE SER PREPARADA UTILIZANDO FRUTAS E VERDURAS AO FERVER O PRODUTO, ESPERE ESFRIAR TOTALMENTE ANTES DE FORNECER A CRIANÇA.'},
    {'tipo': 'padrao', 'categoria': 'ATENÇÃO', 'titulo': 'Padrão — odor característico (ventilar antes de usar)', 'conteudo': '**ATENÇÃO: **MANTER O PRODUTO FORA DA EMBALAGEM EM AMBIENTE BEM VENTILADO ANTES DE COLOCÁ-LO EM USO, PARA REDUÇÃO DO SEU ODOR CARACTERÍSTICO.'},
    {'tipo': 'padrao', 'categoria': 'ATENÇÃO', 'titulo': 'Padrão — pipa/similar (fios elétricos e tempestade)', 'conteudo': '**ATENÇÃO: **NÃO RECOMENDÁVEL PARA CRIANÇAS MENORES DE {IDADE} POR CONTER PARTE (S) PEQUENA (S) QUE PODE(M) SER ENGOLIDA (S) OU ASPIRADA (S). NÃO DEVE SER UTILIZADO PERTO DE FIOS ELÉTRICOS OU DURANTE UMA TEMPESTADE.'},
    {'tipo': 'padrao', 'categoria': 'ATENÇÃO', 'titulo': 'Padrão — boia/flutuador (não é equipamento salva-vidas)', 'conteudo': '**ATENÇÃO: **NÃO RECOMENDÁVEL PARA CRIANÇAS MENORES DE {IDADE} POR CONTER PARTE (S) PEQUENA (S) QUE PODE (M) SER ENGOLIDA (S) OU ASPIRADA(S). ESTE BRINQUEDO NÃO É UM EQUIPAMENTO SALVA-VIDAS.'},
    {'tipo': 'padrao', 'categoria': 'ADVERTÊNCIA', 'titulo': 'Padrão — patinete/bike (equipamento de proteção completo)', 'conteudo': '**ADVERTÊNCIA!** UTILIZAR COM EQUIPAMENTO DE PROTEÇÃO. TAIS COMO CAPACETE MUNHEQUEIRAS, JOELHEIRAS E COTOVELEIRAS.'},
    {'tipo': 'pilha', 'categoria': 'INDICAÇÃO', 'titulo': 'C/Pilha — indicação de idade (genérica)', 'conteudo': '**INDICAÇÃO: **ESTE PRODUTO É INDICADO PARA CRIANÇAS A PARTIR DE {IDADE}.'},
    {'tipo': 'pilha', 'categoria': 'ATENÇÃO', 'titulo': 'C/Pilha — genérico (partes pequenas) + montagem por adulto', 'conteudo': '**ATENÇÃO: **NÃO RECOMENDÁVEL PARA CRIANÇAS MENORES DE {IDADE} POR CONTER PARTE (S) PEQUENA (S) QUE PODE(M) SER ENGOLIDA (S) OU ASPIRADA (S). ESTE PRODUTO DEVE SER MONTADO POR UM ADULTO ANTES DE SER ENTREGUE À CRIANÇA.'},
    {'tipo': 'pilha', 'categoria': 'ATENÇÃO', 'titulo': 'C/Pilha — só montagem por adulto (sem menção a partes pequenas)', 'conteudo': '**ATENÇÃO: **ESTE PRODUTO DEVE SER MONTADO POR UM ADULTO ANTES DE SER ENTREGUE À CRIANÇA.'},
    {'tipo': 'pilha', 'categoria': 'ADVERTÊNCIA', 'titulo': 'C/Pilha — instruções de pilhas e baterias (texto fixo obrigatório)', 'conteudo': '**ADVERTÊNCIA:** COMO RETIRAR E COMO COLOCAR AS PILHAS E AS BATERIAS SUBSTITUIVEIS: AS PILHAS NÃO DEVEM SER RECARREGADAS; AS BATERIAS DEVEM SER RETIRADAS DO BRINQUEDO ANTES DE SEREM RECARREGADAS; AS BATERIAS SOMENTE DEVEM SER RECARREGADAS SOB A SUPERVISÃO DE UM ADULTO; DIFERENTES TIPOS DE PILHAS E BATERIAS NOVAS E USADAS NÃO DEVEM SER MISTURADAS; SÓ DEVEM SER USADAS PILHAS E BATERIAS DO TIPO RECOMENDADO OU UM SIMILAR; AS PILHAS E BATERIAS DEVEM SER COLOCADAS RESPEITANDO A POLARIDADE; AS PILHAS E BATERIAS DESCARREGADAS DEVEM SER RETIRADAS DO BRINQUEDO; OS TERMINAIS DE UMA PILHA OU BATERIA NÃO DEVEM SER COLOCADOS EM CURTO-CIRCUITO.'},
    {'tipo': 'maquiagem', 'categoria': 'INDICAÇÃO', 'titulo': 'Maquiagem — indicação de idade', 'conteudo': '**INDICAÇÃO: **ESTE PRODUTO É INDICADO PARA CRIANÇAS A PARTIR DE {IDADE}.'},
    {'tipo': 'maquiagem', 'categoria': 'ATENÇÃO', 'titulo': 'Maquiagem — partes pequenas', 'conteudo': '**ATENÇÃO: **NÃO RECOMENDÁVEL PARA CRIANÇAS MENORES DE {IDADE} POR CONTER PARTE (S) PEQUENA (S) QUE PODE(M) SER ENGOLIDA (S) OU ASPIRADA (S)'},
    {'tipo': 'maquiagem', 'categoria': 'ADVERTÊNCIA', 'titulo': 'Maquiagem — advertência de uso', 'conteudo': '**ADVERTÊNCIA!** NÃO PODE SER UTILIZADO EM CRIANÇAS'},
    {'tipo': 'massa', 'categoria': 'INDICAÇÃO', 'titulo': 'Massa de modelar — indicação de idade', 'conteudo': '**INDICAÇÃO: **ESTE PRODUTO É INDICADO PARA CRIANÇAS A PARTIR DE {IDADE}.'},
    {'tipo': 'massa', 'categoria': 'ATENÇÃO', 'titulo': 'Massa de modelar — partes pequenas + montagem por adulto', 'conteudo': '**ATENÇÃO: **NÃO RECOMENDÁVEL PARA CRIANÇAS MENORES DE {IDADE} POR CONTER PARTE (S) PEQUENA (S) QUE PODE(M) SER ENGOLIDA (S) OU ASPIRADA (S). ESTE PRODUTO DEVE SER MONTADO POR UM ADULTO ANTES DE SER ENTREGUE À CRIANÇA.'},
    {'tipo': 'massa', 'categoria': 'CUIDADOS DE USO', 'titulo': 'Massa de modelar — cuidados de uso', 'conteudo': '**CUIDADOS DE USO:** USAR SOB SUPERFÍCIES LIMPAS. O PRODUTO PODE GRUDAR EM CABELOS, TAPETES E OUTROS TECIDOS. NÃO RECOMENDÁVEL PARA PESSOAS HIPERSENSÍVEIS A QUALQUER UM DOS COMPONENTES. MANTER A TAMPA FECHADA APÓS O USO. PRODUTO NÃO TÓXICO. NÃO INGERIR.'},
    {'tipo': 'massa', 'categoria': 'COMPOSIÇÃO', 'titulo': 'Massa de modelar — composição', 'conteudo': '**COMPOSIÇÃO: **ÁGUA, GLICERINA, ESPESSANTES, DISPERSANTES, AROMATIZANTES E PIGMENTOS ARTIFICIAIS, CONSERVANTES.'},
]


def importar_textos_padrao_seed():
    """Insere os textos de _TEXTOS_PADRAO_SEED que ainda não existem no banco
    (checa por tipo+categoria+titulo, pra poder rodar de novo sem duplicar)."""
    existentes = pd.read_sql_query("SELECT tipo, categoria, titulo FROM textos_etiqueta", conn)
    chaves_existentes = {(r['tipo'], r['categoria'], r['titulo']) for _, r in existentes.iterrows()}
    inseridos = 0
    for item in _TEXTOS_PADRAO_SEED:
        chave = (item['tipo'], item['categoria'], item['titulo'])
        if chave in chaves_existentes:
            continue
        salvar_texto_etiqueta(item['categoria'], item['titulo'], item['conteudo'], tipo=item['tipo'])
        inseridos += 1
    return inseridos, len(_TEXTOS_PADRAO_SEED) - inseridos


# ---- Modelos-Word de selo (4 variantes fixas, número de registro editável) ----

_VARIANTES_SELO_VALIDAS = {
    "amarelo": "🟨 Amarelo",
    "pb": "⬛ Preto e Branco",
    "compacto_cor": "🎨 Compacto Colorido",
    "compacto_pb": "◽ Compacto Preto e Branco",
}


def _detectar_registro_no_docx(docx_bytes):
    """Lê o número de registro (formato NNN NNN/AAAA) de dentro de um selo-modelo Word."""
    import zipfile as _zf
    from io import BytesIO as _Bio
    try:
        with _zf.ZipFile(_Bio(docx_bytes)) as z:
            xml = z.read('word/document.xml').decode('utf-8', 'ignore')
        texto = ''.join(re.findall(r'<w:t[^>]*>(.*?)</w:t>', xml))
        m = re.search(r'\d{3}\s*\d{3}\s*/\s*\d{4}', texto)
        return m.group(0) if m else None
    except Exception:
        return None


def salvar_modelo_selo(variante, docx_bytes, nome_arquivo):
    if variante not in _VARIANTES_SELO_VALIDAS:
        raise ValueError(f"Variante inválida: {variante}")
    registro_exemplo = _detectar_registro_no_docx(docx_bytes)
    agora = datetime.now().strftime("%d/%m/%Y %H:%M:%S")
    cursor.execute("""
    INSERT INTO modelos_selo (variante, arquivo_docx, nome_arquivo, registro_exemplo, data_atualizacao)
    VALUES (?, ?, ?, ?, ?)
    ON CONFLICT(variante) DO UPDATE SET
        arquivo_docx = excluded.arquivo_docx,
        nome_arquivo = excluded.nome_arquivo,
        registro_exemplo = excluded.registro_exemplo,
        data_atualizacao = excluded.data_atualizacao
    """, (variante, docx_bytes, nome_arquivo, registro_exemplo, agora))
    commit_seguro()
    return registro_exemplo


def buscar_modelo_selo(variante):
    row = cursor.execute(
        "SELECT arquivo_docx, nome_arquivo, registro_exemplo FROM modelos_selo WHERE variante = ?",
        (variante,)
    ).fetchone()
    if not row:
        return None
    return {"docx": row[0], "nome": row[1], "registro_exemplo": row[2]}


def listar_modelos_selo():
    return pd.read_sql_query("""
    SELECT variante, nome_arquivo, registro_exemplo, data_atualizacao
    FROM modelos_selo ORDER BY variante
    """, conn)


def excluir_modelo_selo(variante):
    cursor.execute("DELETE FROM modelos_selo WHERE variante = ?", (variante,))
    commit_seguro()
    return cursor.rowcount


def gerar_selo_com_registro(variante, registro_sem_espaco):
    """Pega o modelo-Word da variante, substitui o número de registro pelo
    registro informado (formatando NNNNNN/AAAA -> NNN NNN/AAAA) e devolve os
    bytes do .docx pronto. Retorna None se o modelo não estiver cadastrado."""
    import zipfile as _zf
    from io import BytesIO as _Bio

    modelo = buscar_modelo_selo(variante)
    if not modelo or not modelo["docx"]:
        return None

    # Formata: 003565/2023 -> 003 565/2023
    reg_limpo = re.sub(r'\s', '', registro_sem_espaco or '')
    m = re.match(r'(\d{3})(\d{3})/(\d{4})', reg_limpo)
    novo_fmt = f"{m.group(1)} {m.group(2)}/{m.group(3)}" if m else registro_sem_espaco

    with _zf.ZipFile(_Bio(modelo["docx"])) as z:
        xml = z.read('word/document.xml').decode('utf-8')

    def processar_paragrafo(mp):
        bloco = mp.group(0)
        ts = re.findall(r'<w:t[^>]*>(.*?)</w:t>', bloco)
        concat = ''.join(ts)
        if not re.search(r'\d{3}\s*\d{3}\s*/\s*\d{4}', concat):
            return bloco
        novo_concat = re.sub(r'\d{3}\s*\d{3}\s*/\s*\d{4}', novo_fmt, concat)
        contador = {'i': 0}
        def repl_t(mt):
            abertura = re.match(r'(<w:t[^>]*>)', mt.group(0)).group(1)
            if contador['i'] == 0:
                contador['i'] += 1
                if 'xml:space' not in abertura:
                    abertura = abertura[:-1] + ' xml:space="preserve">'
                return abertura + novo_concat + '</w:t>'
            contador['i'] += 1
            return abertura + '</w:t>'
        return re.sub(r'<w:t[^>]*>.*?</w:t>', repl_t, bloco)

    novo_xml = re.sub(r'<w:p[ >].*?</w:p>', processar_paragrafo, xml, flags=re.DOTALL)

    saida = _Bio()
    with _zf.ZipFile(_Bio(modelo["docx"])) as zin:
        with _zf.ZipFile(saida, 'w', _zf.ZIP_DEFLATED) as zout:
            for item in zin.namelist():
                if item == 'word/document.xml':
                    zout.writestr(item, novo_xml)
                else:
                    zout.writestr(item, zin.read(item))
    return saida.getvalue()


def _localizar_soffice():
    """Localiza o executável do LibreOffice (versão global, usada por qualquer aba).
    Prioriza soffice.bin/soffice.com sobre soffice.exe, porque no LibreOffice Portable
    o soffice.exe é um wrapper que trava pedindo 'Press Enter to continue' em modo
    automático — o .bin/.com roda direto, sem travar."""
    import shutil as _sh0
    import os as _os0
    _userprofile = _os0.environ.get("USERPROFILE", "")
    try:
        _pasta_app = _os0.path.dirname(_os0.path.abspath(__file__))
    except Exception:
        _pasta_app = ""

    # Pastas onde o LibreOffice pode estar (o "program")
    _pastas_program = [
        _os0.path.join(_pasta_app, "LibreOfficePortable", "App", "libreoffice", "program"),
        r"C:\LibreOfficePortable\App\libreoffice\program",
        r"C:\Program Files\LibreOffice\program",
        r"C:\Program Files (x86)\LibreOffice\program",
        _os0.path.join(_userprofile, "Desktop", "LibreOfficePortable", "App", "libreoffice", "program"),
        _os0.path.join(_userprofile, "Downloads", "LibreOfficePortable", "App", "libreoffice", "program"),
        r"\\SOPTWVFLS01\perfil$\Carlos.patricio\Downloads\LibreOfficePortable\App\libreoffice\program",
    ]

    candidatos = []
    # Para cada pasta, tenta primeiro .bin, depois .com, depois .exe
    for _pasta in _pastas_program:
        candidatos.append(_os0.path.join(_pasta, "soffice.bin"))
        candidatos.append(_os0.path.join(_pasta, "soffice.com"))
        candidatos.append(_os0.path.join(_pasta, "soffice.exe"))
    # PATH do sistema (último recurso)
    candidatos.append(_sh0.which("soffice.bin"))
    candidatos.append(_sh0.which("soffice"))
    candidatos.append(_sh0.which("soffice.exe"))

    for c in candidatos:
        if c and _os0.path.exists(c):
            return c
    return "soffice"


def _matar_processos_libreoffice():
    """Mata qualquer processo LibreOffice pendurado (soffice.bin/.exe/_safe).
    O LibreOffice tem um comportamento conhecido no Windows: a primeira conversão
    headless pode deixar um processo 'de fundo' vivo (pra acelerar conversões
    seguintes), e esse processo pendurado causa o erro código 81 ('já existe uma
    instância aberta') nas próximas tentativas, mesmo usando perfis isolados.
    Chamada ANTES de iniciar uma tarefa (garante início limpo) e DEPOIS de
    terminar (não deixa nada pra trás pra próxima vez). Não gera erro se não
    houver nada rodando — é seguro chamar sempre."""
    import subprocess as _sp0, os as _os0b
    if _os0b.name != "nt":
        try:
            _sp0.run(["pkill", "-9", "-f", "soffice"], capture_output=True, timeout=10)
        except Exception:
            pass
        return
    for _nome_proc in ("soffice.bin", "soffice.exe", "soffice_safe.exe"):
        try:
            _sp0.run(["taskkill", "/F", "/IM", _nome_proc, "/T"], capture_output=True, timeout=10)
        except Exception:
            pass


class _TarefaLibreOfficeLimpa:
    """Context manager: garante que não há LibreOffice pendurado antes de começar
    a tarefa, e limpa de novo ao final (sucesso ou erro), pra próxima tarefa
    começar sempre do zero. Uso: with _TarefaLibreOfficeLimpa(): ...código...

    A limpeza-antes evita que uma instância deixada por uma execução anterior
    (às vezes a própria página do Streamlit rodando outra tarefa em paralelo,
    ou um resquício de conversão anterior) atrapalhe esta nova tarefa.

    NOTA: usada apenas como MÉTODO DE RESERVA (fallback), quando o listener
    persistente (ver abaixo) não estiver disponível. Quando o listener funciona,
    não faz sentido matar processos, pois isso mataria o próprio listener."""
    def __enter__(self):
        _matar_processos_libreoffice()
        return self

    def __exit__(self, exc_type, exc_val, exc_tb):
        _matar_processos_libreoffice()
        return False  # não engole exceções


# ==========================================================================
# LISTENER PERSISTENTE DO LIBREOFFICE (solução definitiva pro código 81)
# ==========================================================================
# Em vez de abrir e fechar uma instância do LibreOffice a cada conversão (o que
# no Windows pode deixar recursos internos "presos", causando o erro código 81
# em chamadas seguintes), mantemos UMA ÚNICA instância do LibreOffice sempre
# aberta, rodando como um "servidor" (--accept=socket). Todas as conversões
# conversam com essa mesma instância via a API UNO do LibreOffice, sem nunca
# precisar abrir uma instância nova. Se o suporte a isso não estiver disponível
# nesta instalação específica do LibreOffice Portable, o sistema cai automati-
# camente no método antigo (abrir uma instância por conversão).

_LO_LISTENER_PORTA = "2002"


def _pasta_trabalho_uno():
    """Pasta fixa (fora do TEMP do Windows) pra guardar o script auxiliar e o
    perfil do listener. Usa uma pasta ao lado do app.py quando possível."""
    import os as _osw
    try:
        _pasta_app = _osw.path.dirname(_osw.path.abspath(__file__))
    except Exception:
        _pasta_app = r"C:\CXMLBREngine"
    _pasta = _osw.path.join(_pasta_app, "_lo_listener")
    try:
        _osw.makedirs(_pasta, exist_ok=True)
    except Exception:
        import tempfile as _tmpw
        _pasta = _osw.path.join(_tmpw.gettempdir(), "cxml_lo_listener")
        _osw.makedirs(_pasta, exist_ok=True)
    return _pasta


def _python_bundled_libreoffice(soffice_path):
    """A partir do caminho do soffice.bin/exe, deduz o caminho do python.exe
    que vem EMBUTIDO no próprio LibreOffice (é esse Python que tem o módulo
    'uno' funcionando corretamente — o Python do Windows normal não tem)."""
    import os as _osw
    if not soffice_path or soffice_path == "soffice":
        return None
    _pasta_program = _osw.path.dirname(soffice_path)
    for _nome in ("python.exe", "python"):
        _cand = _osw.path.join(_pasta_program, _nome)
        if _osw.path.exists(_cand):
            return _cand
    return None


_SCRIPT_CONVERSOR_UNO = '''import sys
import uno
from com.sun.star.beans import PropertyValue


def make_prop(name, value):
    p = PropertyValue()
    p.Name = name
    p.Value = value
    return p


def connect(host="localhost", port="PORTA_AQUI"):
    local_context = uno.getComponentContext()
    resolver = local_context.ServiceManager.createInstanceWithContext(
        "com.sun.star.bridge.UnoUrlResolver", local_context
    )
    ctx = resolver.resolve(
        "uno:socket,host=" + host + ",port=" + port + ";urp;StarOffice.ComponentContext"
    )
    smgr = ctx.ServiceManager
    desktop = smgr.createInstanceWithContext("com.sun.star.frame.Desktop", ctx)
    return desktop


def convert(desktop, input_path, output_path, filter_name):
    input_url = uno.systemPathToFileUrl(input_path)
    output_url = uno.systemPathToFileUrl(output_path)
    in_props = (make_prop("Hidden", True),)
    doc = desktop.loadComponentFromURL(input_url, "_blank", 0, in_props)
    try:
        out_props = (make_prop("FilterName", filter_name),)
        doc.storeToURL(output_url, out_props)
    finally:
        doc.close(False)


if __name__ == "__main__":
    input_path, output_path, filter_name = sys.argv[1], sys.argv[2], sys.argv[3]
    desktop = connect()
    convert(desktop, input_path, output_path, filter_name)
    print("OK")
'''


def _garantir_script_conversor_uno():
    """Escreve (uma vez) o script auxiliar de conversão via UNO em disco."""
    import os as _osw
    _pasta = _pasta_trabalho_uno()
    _caminho = _osw.path.join(_pasta, "conversor_uno.py")
    _conteudo = _SCRIPT_CONVERSOR_UNO.replace("PORTA_AQUI", _LO_LISTENER_PORTA)
    try:
        _precisa_escrever = True
        if _osw.path.exists(_caminho):
            with open(_caminho, "r", encoding="utf-8") as _f:
                if _f.read() == _conteudo:
                    _precisa_escrever = False
        if _precisa_escrever:
            with open(_caminho, "w", encoding="utf-8") as _f:
                _f.write(_conteudo)
    except Exception:
        return None
    return _caminho


_SCRIPT_CONVERSOR_UNO_LOTE = '''import sys
import uno
from com.sun.star.beans import PropertyValue


def make_prop(name, value):
    p = PropertyValue()
    p.Name = name
    p.Value = value
    return p


def connect(host="localhost", port="PORTA_AQUI"):
    local_context = uno.getComponentContext()
    resolver = local_context.ServiceManager.createInstanceWithContext(
        "com.sun.star.bridge.UnoUrlResolver", local_context
    )
    ctx = resolver.resolve(
        "uno:socket,host=" + host + ",port=" + port + ";urp;StarOffice.ComponentContext"
    )
    smgr = ctx.ServiceManager
    desktop = smgr.createInstanceWithContext("com.sun.star.frame.Desktop", ctx)
    return desktop


def convert(desktop, input_path, output_path, filter_name):
    input_url = uno.systemPathToFileUrl(input_path)
    output_url = uno.systemPathToFileUrl(output_path)
    in_props = (make_prop("Hidden", True),)
    doc = desktop.loadComponentFromURL(input_url, "_blank", 0, in_props)
    try:
        out_props = (make_prop("FilterName", filter_name),)
        doc.storeToURL(output_url, out_props)
    finally:
        doc.close(False)


if __name__ == "__main__":
    manifesto_path = sys.argv[1]
    with open(manifesto_path, "r", encoding="utf-8") as f:
        linhas = [l.rstrip("\\n") for l in f if l.strip()]
    desktop = connect()
    for linha in linhas:
        input_path, output_path, filter_name = linha.split("\\t")
        try:
            convert(desktop, input_path, output_path, filter_name)
            print("OK\\t" + input_path)
        except Exception as e:
            print("ERR\\t" + input_path + "\\t" + str(e).replace("\\n", " ").replace("\\t", " "))
        sys.stdout.flush()
'''


def _garantir_script_conversor_uno_lote():
    """Escreve (uma vez) o script auxiliar de conversão EM LOTE via UNO em
    disco. Diferença pro script de arquivo único: conecta na ponte UNO
    (connect()) UMA VEZ só e reaproveita pra converter todos os arquivos do
    manifesto, em vez de reconectar a cada arquivo — é essa reconexão
    repetida (processo novo + handshake UNO novo por arquivo) que domina o
    tempo quando o lote tem muitos códigos de barras (EMF)."""
    import os as _osw
    _pasta = _pasta_trabalho_uno()
    _caminho = _osw.path.join(_pasta, "conversor_uno_lote.py")
    _conteudo = _SCRIPT_CONVERSOR_UNO_LOTE.replace("PORTA_AQUI", _LO_LISTENER_PORTA)
    try:
        _precisa_escrever = True
        if _osw.path.exists(_caminho):
            with open(_caminho, "r", encoding="utf-8") as _f:
                if _f.read() == _conteudo:
                    _precisa_escrever = False
        if _precisa_escrever:
            with open(_caminho, "w", encoding="utf-8") as _f:
                _f.write(_conteudo)
    except Exception:
        return None
    return _caminho


def _converter_lote_via_uno_listener(jobs, soffice_path, _log=None):
    """Converte VÁRIOS arquivos numa ÚNICA chamada do Python embutido do
    LibreOffice, reaproveitando a MESMA ponte UNO pra todos. Antes, cada
    arquivo de um lote (ex: os ~200 códigos de barras de uma lista grande)
    pagava o custo de abrir um processo novo + reconectar a ponte UNO do
    zero; isso é o maior custo de tempo da conversão em lotes grandes.
    'jobs': lista de tuplas (input_path, output_path, filter_name).
    Retorna dict {input_path: (ok, erro)}, ou None se o lote inteiro falhou
    (quem chamou deve cair pro método de reserva arquivo-a-arquivo)."""
    import subprocess as _spwL, tempfile as _tmpL2, os as _osL2, datetime as _dtL

    def _registrar(msg):
        if _log is not None:
            _log.append(f"[{_dtL.datetime.now().strftime('%H:%M:%S')}] {msg}")

    if not jobs:
        return {}

    _suporte = _verificar_suporte_uno(soffice_path)
    if not _suporte["suportado"]:
        _registrar(f"UNO não suportado (lote): {_suporte['motivo']}")
        return None

    _script = _garantir_script_conversor_uno_lote()
    if not _script:
        _registrar("Falha ao escrever o script auxiliar de conversão em lote")
        return None

    if not _iniciar_listener_libreoffice(soffice_path):
        _registrar("Não consegui iniciar/conectar ao listener (lote)")
        return None

    _manifesto_dir = _tmpL2.mkdtemp(prefix="uno_lote_manifesto_")
    try:
        _manifesto_path = _osL2.path.join(_manifesto_dir, "manifesto.txt")
        with open(_manifesto_path, "w", encoding="utf-8") as _f:
            for _inp, _outp, _filt in jobs:
                _f.write(f"{_inp}\t{_outp}\t{_filt}\n")

        _registrar(f"Convertendo lote de {len(jobs)} arquivo(s) numa única ponte UNO...")
        try:
            _res = _spwL.run(
                [_suporte["python_lo"], _script, _manifesto_path],
                capture_output=True, timeout=max(60, 5 * len(jobs)), text=True
            )
        except Exception as _e:
            _registrar(f"Exceção ao chamar conversor em lote: {type(_e).__name__}: {_e}")
            return None

        _resultados = {}
        for _linha in (_res.stdout or "").splitlines():
            _partes = _linha.split("\t")
            if _partes[0] == "OK" and len(_partes) >= 2:
                _resultados[_partes[1]] = (True, None)
            elif _partes[0] == "ERR" and len(_partes) >= 3:
                _resultados[_partes[1]] = (False, _partes[2])

        _registrar(f"Lote finalizado: código={_res.returncode} | "
                   f"convertidos={sum(1 for v in _resultados.values() if v[0])} de {len(jobs)} | "
                   f"stderr='{(_res.stderr or '').strip()[:200]}'")

        if not _resultados and _res.returncode != 0:
            # Nada foi processado (provável falha ao conectar/rodar o script
            # inteiro) -> sinaliza falha total pra cair no método de reserva.
            return None
        return _resultados
    finally:
        import shutil as _shL2
        _shL2.rmtree(_manifesto_dir, ignore_errors=True)


def _listener_uno_respondendo(porta=_LO_LISTENER_PORTA, timeout=1.0):
    """Testa (rápido, via socket puro, sem precisar do módulo uno) se já existe
    algo escutando na porta do listener."""
    import socket as _sockw
    try:
        with _sockw.create_connection(("localhost", int(porta)), timeout=timeout):
            return True
    except Exception:
        return False


@st.cache_resource(show_spinner=False)
def _verificar_suporte_uno(_soffice_path):
    """Testa se o Python embutido no LibreOffice consegue importar o módulo
    'uno'. Se não conseguir, o listener persistente não é viável nesta
    instalação e o sistema deve usar o método antigo (fallback).

    Faz 2 tentativas antes de "desistir" e guardar um resultado negativo —
    evita que um problema passageiro (ex: sistema ocupado num instante) vire
    uma crença fixa e errada de 'não suportado' pelo resto da sessão."""
    import subprocess as _spw, time as _timeuno
    _python_lo = _python_bundled_libreoffice(_soffice_path)
    if not _python_lo:
        return {"suportado": False, "motivo": "python.exe do LibreOffice não encontrado", "python_lo": None}

    _ultimo_erro = None
    for _tentativa in range(2):
        try:
            _res = _spw.run([_python_lo, "-c", "import uno; print('OK')"],
                             capture_output=True, timeout=20, text=True)
            if _res.returncode == 0 and "OK" in (_res.stdout or ""):
                return {"suportado": True, "motivo": None, "python_lo": _python_lo}
            _ultimo_erro = f"módulo uno não importou (código {_res.returncode}): {_res.stderr[:200]}"
        except Exception as _e:
            _ultimo_erro = f"erro ao testar: {_e}"
        if _tentativa == 0:
            _timeuno.sleep(1.5)  # espera antes de tentar de novo

    return {"suportado": False, "motivo": _ultimo_erro, "python_lo": _python_lo}


def _iniciar_listener_libreoffice(soffice_path, porta=_LO_LISTENER_PORTA, timeout_inicio=15):
    """Garante que o listener persistente do LibreOffice está rodando. Se já
    estiver (de uma execução anterior desta mesma sessão do Windows), não faz
    nada. Se não estiver, inicia ele em segundo plano (independente do processo
    do Streamlit) e espera até responder ou até o tempo limite."""
    import subprocess as _spw, os as _osw, time as _timew

    if _listener_uno_respondendo(porta, timeout=0.5):
        return True

    _perfil = _osw.path.join(_pasta_trabalho_uno(), "listener_profile")
    _perfil_uri = _montar_uri_perfil(_perfil)

    _flags = 0
    if _osw.name == "nt":
        _flags = _spw.CREATE_NEW_PROCESS_GROUP | getattr(_spw, "DETACHED_PROCESS", 0x00000008)

    try:
        _spw.Popen(
            [soffice_path, "--headless", "--invisible", "--nologo",
             "--nofirststartwizard", "--nodefault", "--norestore",
             f"--accept=socket,host=localhost,port={porta};urp;",
             f"-env:UserInstallation={_perfil_uri}"],
            creationflags=_flags,
            stdout=_spw.DEVNULL, stderr=_spw.DEVNULL, stdin=_spw.DEVNULL,
            close_fds=True,
        )
    except Exception:
        return False

    _inicio = _timew.time()
    while _timew.time() - _inicio < timeout_inicio:
        if _listener_uno_respondendo(porta, timeout=0.5):
            return True
        _timew.sleep(0.5)
    return False


def _converter_via_uno_listener(entrada_path, saida_path, filtro_nome, soffice_path, _log=None):
    """Converte um documento usando o listener persistente (via UNO). Retorna
    (sucesso: bool, erro: str|None). Não abre nenhuma instância nova do
    LibreOffice — só conversa com a que já está rodando.

    Detecta e se recupera de 'listener zumbi': às vezes a porta responde
    (a conexão de rede funciona) mas o LibreOffice ali dentro está travado/
    corrompido internamente e a conversão falha mesmo assim. Nesse caso, mata
    o processo e inicia um novo, tentando de novo uma vez.

    _log: se for uma lista, recebe cada passo com hora, pra dar visibilidade
    total do que está acontecendo (pedido explícito: poder auditar o log)."""
    import subprocess as _spw, time as _timeu, datetime as _dtu

    def _registrar(msg):
        if _log is not None:
            _log.append(f"[{_dtu.datetime.now().strftime('%H:%M:%S')}] {msg}")

    _suporte = _verificar_suporte_uno(soffice_path)
    if not _suporte["suportado"]:
        _registrar(f"UNO não suportado: {_suporte['motivo']}")
        return False, f"UNO não suportado: {_suporte['motivo']}"

    _script = _garantir_script_conversor_uno()
    if not _script:
        _registrar("Falha ao escrever o script auxiliar de conversão")
        return False, "não consegui escrever o script auxiliar de conversão"

    def _tentar_converter_uma_vez():
        _ja_respondia = _listener_uno_respondendo(_LO_LISTENER_PORTA, timeout=0.5)
        _registrar(f"Porta {_LO_LISTENER_PORTA} já respondia antes de iniciar? {_ja_respondia}")

        if not _iniciar_listener_libreoffice(soffice_path):
            _registrar("Não consegui iniciar/conectar ao listener")
            return False, "não consegui iniciar/conectar ao listener do LibreOffice"

        try:
            _res = _spw.run(
                [_suporte["python_lo"], _script, entrada_path, saida_path, filtro_nome],
                capture_output=True, timeout=60, text=True
            )
            _registrar(f"Conversão via listener: código={_res.returncode} | "
                       f"stdout='{(_res.stdout or '').strip()[:100]}' | "
                       f"stderr='{(_res.stderr or '').strip()[:200]}'")
            if _res.returncode == 0 and "OK" in (_res.stdout or ""):
                return True, None
            return False, f"conversão via listener falhou: {_res.stderr[:300]}"
        except Exception as _e:
            _registrar(f"Exceção ao chamar conversor: {type(_e).__name__}: {_e}")
            return False, f"erro ao chamar conversor via listener: {_e}"

    _ok, _erro = _tentar_converter_uma_vez()
    if _ok:
        return True, None

    # LISTENER ZUMBI: a porta pode estar "respondendo" (TCP) mas o LibreOffice
    # internamente travado/corrompido, fazendo a conversão real falhar. Mata o
    # processo (libera a porta) e tenta de novo com uma instância nova.
    _registrar(f"Primeira tentativa falhou ({_erro}). Suspeita de listener zumbi — "
               f"matando processo e tentando de novo com instância nova.")
    _matar_processos_libreoffice()
    _timeu.sleep(1.5)
    _ok2, _erro2 = _tentar_converter_uma_vez()
    if _ok2:
        _registrar("✅ Recuperação automática funcionou (segunda tentativa, listener novo).")
        return True, None

    _registrar(f"Segunda tentativa também falhou: {_erro2}")
    return False, f"{_erro} | após reiniciar listener: {_erro2}"


def _localizar_tesseract():
    """Localiza o executável do Tesseract (versão global, usada por qualquer aba)."""
    import shutil as _sh1
    import os as _os1
    try:
        _pasta_app = _os1.path.dirname(_os1.path.abspath(__file__))
    except Exception:
        _pasta_app = ""
    candidatos = [
        _sh1.which("tesseract"),
        _sh1.which("tesseract.exe"),
        _os1.path.join(_pasta_app, "Tesseract-OCR", "tesseract.exe"),
        _os1.path.expandvars(r"%LOCALAPPDATA%\Programs\Tesseract-OCR\tesseract.exe"),
        r"C:\Program Files\Tesseract-OCR\tesseract.exe",
        r"C:\Program Files (x86)\Tesseract-OCR\tesseract.exe",
    ]
    for c in candidatos:
        if c and _os1.path.exists(c):
            return c
    return None


def _montar_uri_perfil(caminho_pasta):
    """Monta um file:// URI válido pro -env:UserInstallation do LibreOffice,
    funcionando tanto no Windows (C:\\... -> file:///C:/...) quanto no Linux."""
    import os as _osu
    _abs = _osu.path.abspath(caminho_pasta)
    _abs = _abs.replace("\\", "/")  # barras invertidas -> normais (Windows)
    if not _abs.startswith("/"):
        # Windows: "C:/..." -> "file:///C:/..."
        return "file:///" + _abs
    # Linux/Mac: "/tmp/..." -> "file:///tmp/..."
    return "file://" + _abs


def _selo_docx_para_png_metodo_antigo(selo_docx_bytes, soffice_path, coletar_erro=None):
    """MÉTODO DE RESERVA (fallback): abre uma instância nova do LibreOffice pra
    cada conversão. Usado apenas se o listener persistente (método principal,
    ver selo_docx_para_png) não estiver disponível nesta instalação."""
    import tempfile as _tmp, os as _os, subprocess as _sub, glob as _glob, time as _time
    import uuid as _uuid, shutil as _shutil
    from PIL import Image as _ImgS, ImageOps as _IOS
    from io import BytesIO as _BioS

    def _log(msg):
        if coletar_erro is not None:
            coletar_erro.append(msg)

    def _uma_tentativa(tag):
        # Pasta de trabalho única + perfil único (PID + uuid) para NUNCA colidir
        # com outra instância do LibreOffice — evita o erro código 81.
        _base = _os.path.join(_tmp.gettempdir(), f"cxml_selo_{_os.getpid()}_{_uuid.uuid4().hex}_{tag}")
        _os.makedirs(_base, exist_ok=True)
        _docx_path = _os.path.join(_base, "selo.docx")
        with open(_docx_path, "wb") as f:
            f.write(selo_docx_bytes)
        _profile_dir = _os.path.join(_base, "loprofile")
        _perfil = _montar_uri_perfil(_profile_dir)

        # Ambiente isolado: HOME/TMP próprios evitam o lock global do LibreOffice
        _env = dict(_os.environ)
        _env["HOME"] = _base
        _env["TMPDIR"] = _base
        _env["TMP"] = _base
        _env["TEMP"] = _base

        _res = _sub.run(
            [soffice_path, "--headless", "--norestore", "--invisible", "--nologo",
             "--nofirststartwizard", "--nodefault", "--nolockcheck",
             f"-env:UserInstallation={_perfil}",
             "--convert-to", "png", "--outdir", _base, _docx_path],
            capture_output=True, timeout=120, env=_env
        )
        _pngs = _glob.glob(_os.path.join(_base, "*.png"))
        return _res, _pngs, _base

    if not _os.path.exists(soffice_path) and soffice_path not in ("soffice",):
        _log(f"Executável do LibreOffice não encontrado: {soffice_path}")
        return None

    try:
        _res, _pngs, _base = _uma_tentativa("t1")

        # Se falhou, tenta de novo com perfil totalmente novo (não mata processo,
        # só usa outro perfil isolado — mais confiável no Windows).
        if not _pngs:
            _time.sleep(1)
            _res, _pngs, _base = _uma_tentativa("t2")

        if not _pngs:
            _stderr = _res.stderr.decode('utf-8', 'ignore')[:300] if _res.stderr else ""
            _stdout = _res.stdout.decode('utf-8', 'ignore')[:300] if _res.stdout else ""
            _dica = ""
            if _res.returncode == 81:
                _dica = (" | Código 81 persistente mesmo com perfil isolado. "
                         "Pode ser antivírus bloqueando a pasta temporária, ou o soffice.bin "
                         "sem permissão de escrita no Temp.")
            _log(f"LibreOffice não gerou PNG. código de saída={_res.returncode} | "
                 f"stderr='{_stderr}' | stdout='{_stdout}'{_dica}")
            return None

        _img = _ImgS.open(_pngs[0])
        _inv = _IOS.invert(_img.convert("RGB"))
        _bbox = _inv.getbbox()
        if _bbox:
            _img = _img.crop(_bbox)
        _out = _BioS()
        _img.save(_out, format="PNG")
        return _out.getvalue()
    except _sub.TimeoutExpired:
        _log(f"LibreOffice travou (timeout). soffice='{soffice_path}'.")
        return None
    except Exception as _e:
        _log(f"Erro na conversão do selo: {type(_e).__name__}: {_e}")
        return None


def selo_docx_para_png(selo_docx_bytes, soffice_path, coletar_erro=None):
    """Converte o selo (Word) em PNG recortado (sem espaço branco em volta),
    pra ser inserido como imagem na etiqueta. Retorna bytes PNG ou None.

    MÉTODO PRINCIPAL: usa o listener persistente do LibreOffice (uma única
    instância sempre aberta, evita o erro código 81 causado por abrir/fechar
    repetidamente). Se não for possível (instalação sem suporte a UNO, ou
    listener não conseguiu iniciar), cai automaticamente no método antigo."""
    import tempfile as _tmp2, os as _os2, uuid as _uuid2
    from PIL import Image as _ImgS2, ImageOps as _IOS2
    from io import BytesIO as _BioS2

    def _log(msg):
        if coletar_erro is not None:
            coletar_erro.append(msg)

    _base = _os2.path.join(_tmp2.gettempdir(), f"cxml_selo_uno_{_uuid2.uuid4().hex}")
    _os2.makedirs(_base, exist_ok=True)
    _docx_path = _os2.path.join(_base, "selo.docx")
    _png_path = _os2.path.join(_base, "selo.png")
    with open(_docx_path, "wb") as f:
        f.write(selo_docx_bytes)

    _ok, _erro = _converter_via_uno_listener(_docx_path, _png_path, "writer_png_Export", soffice_path)

    if _ok and _os2.path.exists(_png_path):
        try:
            _img = _ImgS2.open(_png_path)
            _inv = _IOS2.invert(_img.convert("RGB"))
            _bbox = _inv.getbbox()
            if _bbox:
                _img = _img.crop(_bbox)
            _out = _BioS2()
            _img.save(_out, format="PNG")
            return _out.getvalue()
        except Exception as _e:
            _log(f"Selo convertido via listener, mas erro ao processar imagem: {_e}")
            # cai pro método antigo como última tentativa
        finally:
            import shutil as _shutil2
            _shutil2.rmtree(_base, ignore_errors=True)
    else:
        _log(f"Listener não disponível/falhou ({_erro}) — usando método antigo (mais lento).")
        import shutil as _shutil3
        _shutil3.rmtree(_base, ignore_errors=True)

    # Fallback: método antigo (abre uma instância por conversão)
    return _selo_docx_para_png_metodo_antigo(selo_docx_bytes, soffice_path, coletar_erro=coletar_erro)


def _montar_etiqueta_no_doc(doc, dados, selo_png, barcode_png, chorao_png=None, pilha_png=None):
    """Monta uma etiqueta (tabela 2 colunas) dentro do documento `doc`.
    Medidas extraídas do padrão real da empresa:
    - Fonte Arial 7pt (corpo), 8,5pt (prefixo ATENÇÃO/INDICAÇÃO/ADVERTÊNCIA e rodapé), 8pt (referência/nome)
    - Etiqueta 11.48cm de largura total (2 colunas de ~5.74cm)
    - Selo 5.54cm × 2.5cm, código de barras 5.34cm de largura"""
    from docx.shared import Pt as _Pt, Cm as _Cm
    from docx.enum.text import WD_ALIGN_PARAGRAPH as _ALIGN
    from docx.enum.table import WD_ALIGN_VERTICAL as _VALIGN
    from docx.oxml.ns import qn as _qn
    from docx.oxml import OxmlElement as _OxmlElement
    from io import BytesIO as _Bio

    FONTE = "Arial"
    TAM_CORPO = 7      # tamanho principal (padrão real)
    TAM_PREFIXO = 8.5  # prefixo "ATENÇÃO:" / "INDICAÇÃO:" / "ADVERTÊNCIA:" e rodapé
    TAM_REF = 8        # referência/nome
    LARG_COL = 5.74    # cada coluna (total 11.48cm)
    LARG_SELO = 5.54   # largura do selo (padrão real: 5,54cm)
    ALT_SELO = 2.5     # altura do selo (padrão real: 2,5cm)
    LARG_BARRAS = 5.34 # largura do código de barras
    ALT_BARRAS = 2.24  # altura — proporção 2,39:1 (medida do codigo_de_barras.docx)

    def _fmt(run, tam=TAM_CORPO, bold=False):
        run.font.name = FONTE
        run.font.size = _Pt(tam)
        run.bold = bold

    def _tornar_imagem_flutuante_canto_superior_esquerdo(run, largura_coluna_cm, doc_pr_id=100):
        """Converte a imagem (recém-inserida em `run`) de inline pra ANCORADA
        (flutuante) no canto superior ESQUERDO do parágrafo, com o texto do
        próprio parágrafo contornando pela direita — igual ao que o Word
        gera quando você mesmo seleciona a imagem → Opções de Layout →
        'Com Disposição do Texto' → 'Quadrado' (wrapSquare), alinhada à
        esquerda da coluna. Precisa ser o PRIMEIRO run do parágrafo (antes do
        texto) pra âncora calcular a partir do início certo."""
        drawing = run._element.find(_qn('w:drawing'))
        inline = drawing.find(_qn('wp:inline'))
        filhos = {child.tag.split('}')[-1]: child for child in inline}
        extent = filhos['extent']
        docPr = filhos['docPr']
        cNvGraphicFramePr = filhos.get('cNvGraphicFramePr')
        graphic = filhos['graphic']

        anchor = _OxmlElement('wp:anchor')
        for _attr, _val in [('distT', '0'), ('distB', '0'), ('distL', '114300'), ('distR', '0'),
                             ('simplePos', '0'), ('relativeHeight', str(251659264 + doc_pr_id)),
                             ('behindDoc', '0'), ('locked', '0'), ('layoutInCell', '1'), ('allowOverlap', '1')]:
            anchor.set(_attr, _val)

        simplePos = _OxmlElement('wp:simplePos')
        simplePos.set('x', '0'); simplePos.set('y', '0')
        anchor.append(simplePos)

        positionH = _OxmlElement('wp:positionH')
        positionH.set('relativeFrom', 'column')
        alignH = _OxmlElement('wp:align')
        alignH.text = 'left'
        positionH.append(alignH)
        anchor.append(positionH)

        positionV = _OxmlElement('wp:positionV')
        positionV.set('relativeFrom', 'paragraph')
        alignV = _OxmlElement('wp:align')
        alignV.text = 'top'
        positionV.append(alignV)
        anchor.append(positionV)

        anchor.append(extent)

        effectExtent = _OxmlElement('wp:effectExtent')
        for _attr in ('l', 't', 'r', 'b'):
            effectExtent.set(_attr, '0')
        anchor.append(effectExtent)

        # wrapSquare ("Quadrado" no Word) — confirmado por você como o certo.
        # O "Através" seguia as partes transparentes do ícone redondo e
        # bagunçava o espaçamento do texto (espaços estranhos entre palavras).
        wrapSquare = _OxmlElement('wp:wrapSquare')
        wrapSquare.set('wrapText', 'right')
        anchor.append(wrapSquare)

        anchor.append(docPr)
        if cNvGraphicFramePr is not None:
            anchor.append(cNvGraphicFramePr)
        anchor.append(graphic)

        drawing.remove(inline)
        drawing.append(anchor)

    def _adicionar_borda_superior(paragrafo):
        """Linha divisória preta acima do parágrafo, igual à do molde real
        (<w:pBdr><w:top w:val="single" w:sz="4" w:space="1" w:color="auto"/></w:pBdr>)."""
        pPr = paragrafo._p.get_or_add_pPr()
        pBdr = pPr.find(_qn('w:pBdr'))
        if pBdr is None:
            pBdr = _OxmlElement('w:pBdr')
            pPr.append(pBdr)
        top = _OxmlElement('w:top')
        top.set(_qn('w:val'), 'single')
        top.set(_qn('w:sz'), '4')
        top.set(_qn('w:space'), '1')
        top.set(_qn('w:color'), 'auto')
        pBdr.append(top)

    tabela = doc.add_table(rows=2, cols=2)
    tabela.style = 'Table Grid'
    tabela.autofit = False
    for row in tabela.rows:
        row.cells[0].width = _Cm(LARG_COL)
        row.cells[1].width = _Cm(LARG_COL)

    # ---- ESQUERDA ----
    cell_esq = tabela.cell(0, 0)
    cell_esq.vertical_alignment = _VALIGN.TOP
    p0 = cell_esq.paragraphs[0]
    p0.paragraph_format.space_after = _Pt(0)
    if selo_png:
        p0.alignment = _ALIGN.LEFT
        p0.add_run().add_picture(_Bio(selo_png), width=_Cm(LARG_SELO), height=_Cm(ALT_SELO))

    def add_linha(cell, texto, tam=TAM_CORPO, bold_prefix=None, center=False, bold_all=False):
        pp = cell.add_paragraph()
        pp.paragraph_format.space_after = _Pt(0)
        pp.paragraph_format.space_before = _Pt(0)
        if center:
            pp.alignment = _ALIGN.CENTER
        if bold_prefix:
            r1 = pp.add_run(bold_prefix); _fmt(r1, tam, bold=True)
            r2 = pp.add_run(texto); _fmt(r2, tam)
        else:
            r = pp.add_run(texto); _fmt(r, tam, bold=bold_all)
        return pp

    if dados.get('solicitante_cnpj'):
        add_linha(cell_esq, "Solicitante da certificação:")
        add_linha(cell_esq, f"CNPJ: {dados['solicitante_cnpj']}")
    add_linha(cell_esq, dados['importador_razao'], bold_prefix="Importador: ")
    add_linha(cell_esq, dados['importador_endereco'])
    add_linha(cell_esq, f"CNPJ: {dados['importador_cnpj']} / Origem: {dados['origem']}")
    p_qtd = add_linha(cell_esq, f"Quantidade: {dados['quantidade']}")
    _adicionar_borda_superior(p_qtd)
    p_dtfab = add_linha(cell_esq, f"Data de Fabricação: {dados['data_fabricacao']} / Lote: {dados['lote']}")
    _adicionar_borda_superior(p_dtfab)
    add_linha(cell_esq, "Data de validade: Indeterminado")
    add_linha(cell_esq, f"SAC: {dados['sac']}")

    # ---- DIREITA ----
    cell_dir = tabela.cell(0, 1)
    cell_dir.vertical_alignment = _VALIGN.TOP

    for _idx_bloco, bloco in enumerate(dados['textos_direita']):
        if _idx_bloco == 0:
            pb = cell_dir.paragraphs[0]
        else:
            pb = cell_dir.add_paragraph()
        pb.paragraph_format.space_after = _Pt(1)

        # Chorão (ícone -3 anos): imagem FLUTUANTE ancorada no canto superior
        # ESQUERDO do primeiro parágrafo (ATENÇÃO), com contorno "Quadrado"
        # (wrapSquare) — precisa ser o PRIMEIRO run do parágrafo, antes do
        # texto, senão a âncora calcula a partir do lugar errado.
        if _idx_bloco == 0 and chorao_png:
            _run_chorao = pb.add_run()
            _run_chorao.add_picture(_Bio(chorao_png), width=_Cm(0.9))
            _tornar_imagem_flutuante_canto_superior_esquerdo(_run_chorao, LARG_COL)

        # Prefixos conhecidos: renderizados em TAM_PREFIXO (8,5pt) bold,
        # o restante do texto em TAM_CORPO (7pt) via markdown negrito normal.
        # Suporta texto armazenado como "ATENÇÃO: ..." ou "**ATENÇÃO:** ...".
        _PREFIXOS = ('ATENÇÃO:', 'INDICAÇÃO:', 'ADVERTÊNCIA:',
                     'CUIDADOS DE USO:', 'COMPOSIÇÃO:')
        _bloco_restante = bloco
        _achou_pref = False
        for _pref in _PREFIXOS:
            # Verifica forma direta ("ATENÇÃO: ...") e forma markdown ("**ATENÇÃO:** ...")
            _pref_up = _pref.upper()
            _bloco_up = _bloco_restante.upper().lstrip()
            _offset = len(_bloco_restante) - len(_bloco_restante.lstrip())
            if _bloco_up.startswith(_pref_up):
                _run_pref = pb.add_run(_bloco_restante[_offset:_offset + len(_pref)] + ' ')
                _run_pref.font.name = FONTE
                _run_pref.font.size = _Pt(TAM_PREFIXO)
                _run_pref.bold = True
                _bloco_restante = _bloco_restante[_offset + len(_pref):].lstrip()
                _achou_pref = True
                break
            elif _bloco_up.startswith('**' + _pref_up):
                # Formato markdown: extrai texto real (sem os **)
                _run_pref = pb.add_run(_pref + ' ')
                _run_pref.font.name = FONTE
                _run_pref.font.size = _Pt(TAM_PREFIXO)
                _run_pref.bold = True
                # Remove o padrão **PREFIXO:** do início do bloco
                _bloco_restante = re.sub(r'^\*\*' + re.escape(_pref) + r'\*\*\s*', '',
                                         _bloco_restante, flags=re.IGNORECASE).lstrip()
                _achou_pref = True
                break
        aplicar_markdown_negrito_no_paragrafo(pb, _bloco_restante)
        for run in pb.runs:
            if run.font.size is None:
                run.font.name = FONTE
                run.font.size = _Pt(TAM_CORPO)

    # Referência + nome + marca + código de barras: normalmente ficam na
    # coluna DIREITA (embaixo dos textos de ATENÇÃO/INDICAÇÃO/ADVERTÊNCIA),
    # centralizados. Na etiqueta C/Pilha esse bloco inteiro muda de COLUNA:
    # vai para a ESQUERDA, junto do selo/importador (não é só alinhamento).
    _bloco_na_esquerda = dados.get('tipo') == 'pilha'
    _cell_ref_barras = cell_esq if _bloco_na_esquerda else cell_dir
    _alinhamento_ref_barras = _ALIGN.LEFT if _bloco_na_esquerda else _ALIGN.CENTER

    pref = _cell_ref_barras.add_paragraph()
    pref.alignment = _alinhamento_ref_barras
    pref.paragraph_format.space_before = _Pt(3)
    pref.paragraph_format.space_after = _Pt(0)
    rref = pref.add_run(f"{dados['referencia']} - {dados['nome_curto']}")
    _fmt(rref, TAM_REF, bold=True)

    pmarca = _cell_ref_barras.add_paragraph()
    pmarca.alignment = _alinhamento_ref_barras
    pmarca.paragraph_format.space_after = _Pt(2)
    rmarca = pmarca.add_run(f"MARCA: {dados['marca']}")
    _fmt(rmarca, TAM_CORPO)

    if barcode_png:
        pbar = _cell_ref_barras.add_paragraph()
        pbar.alignment = _alinhamento_ref_barras
        pbar.paragraph_format.space_after = _Pt(0)
        pbar.add_run().add_picture(_Bio(barcode_png), width=_Cm(LARG_BARRAS), height=_Cm(ALT_BARRAS))

    # ---- RODAPÉ ----
    cell_rod = tabela.cell(1, 0).merge(tabela.cell(1, 1))
    prod = cell_rod.paragraphs[0]
    prod.alignment = _ALIGN.CENTER
    prod.paragraph_format.space_after = _Pt(0)
    prod.paragraph_format.space_before = _Pt(0)
    rrod = prod.add_run('"GUARDAR A EMBALAGEM POR CONTER INFORMAÇÕES IMPORTANTES"')
    _fmt(rrod, TAM_PREFIXO, bold=True)
    return tabela


# Quais categorias de texto se aplicam a cada tipo de etiqueta, e em que ordem
# aparecem no lado direito (mesma ordem observada nos modelos reais).
CATEGORIAS_POR_TIPO_ETIQUETA = {
    'padrao': ['ATENÇÃO', 'INDICAÇÃO', 'ADVERTÊNCIA'],
    'pilha': ['ATENÇÃO', 'INDICAÇÃO', 'ADVERTÊNCIA'],
    'maquiagem': ['ATENÇÃO', 'INDICAÇÃO', 'ADVERTÊNCIA'],
    'massa': ['ATENÇÃO', 'INDICAÇÃO', 'CUIDADOS DE USO', 'COMPOSIÇÃO'],
}


def _montar_textos_direita(tipo, idade_formatada, titulos_escolhidos, textos_por_chave):
    """Monta a lista de blocos de texto do lado direito conforme o tipo de etiqueta.
    titulos_escolhidos: dict {categoria: titulo} com o texto escolhido pelo usuário
    pra esse item, em cada categoria aplicável ao tipo (None/'' pra pular a categoria).
    textos_por_chave: dict {(tipo, categoria, titulo): conteudo}, pré-montado a partir
    de listar_textos_etiqueta(). Preenche {IDADE} automaticamente."""
    blocos = []
    for categoria in CATEGORIAS_POR_TIPO_ETIQUETA.get(tipo, []):
        titulo = (titulos_escolhidos or {}).get(categoria)
        if not titulo:
            continue
        conteudo = textos_por_chave.get((tipo, categoria, titulo))
        if conteudo:
            blocos.append(preencher_idade_no_texto(conteudo, idade_formatada))

    # fallback: se não houver textos escolhidos/cadastrados, avisa no próprio bloco
    if not blocos:
        blocos.append("**ATENÇÃO:** (cadastre e escolha os textos padrão na aba 'Textos padrão da etiqueta')")
    return blocos


def _gerar_lote_etiquetas(itens_config, cliente, origem, solicitante_cnpj, data_fab, lote,
                          cliente_base_registro, chorao_png_bytes=None, pilha_png_bytes=None):
    """Gera o Word com todas as etiquetas. O registro vem da BASE (tabela registros),
    buscado por fábrica+família, avisando se não bater com o Excel.
    cliente_base_registro: qual base usar pra buscar o registro (ex: 'BOLSA' ou o cliente escolhido)."""
    from docx import Document as _Doc
    from docx.shared import Cm as _Cm
    from io import BytesIO as _Bio

    avisos = []
    soffice = _localizar_soffice()
    import os as _os_diag
    _soffice_existe = (soffice == "soffice") or _os_diag.path.exists(soffice)
    if not _soffice_existe:
        avisos.append(f"⚠️ LibreOffice NÃO encontrado em '{soffice}'. Os selos podem não ser gerados. Instale o LibreOffice e reinicie.")

    _df_txt = listar_textos_etiqueta()
    textos_por_chave = {}
    if not _df_txt.empty:
        for _, _row in _df_txt.iterrows():
            textos_por_chave[(_row['tipo'], _row['categoria'], _row['titulo'])] = _row['conteudo']

    _cache_selo = {}

    doc = _Doc()
    for section in doc.sections:
        section.top_margin = _Cm(2.5)
        section.bottom_margin = _Cm(2.5)
        section.left_margin = _Cm(3.0)
        section.right_margin = _Cm(3.0)

    # NÃO mata/reinicia o LibreOffice aqui — o listener persistente é
    # gerenciado internamente por selo_docx_para_png() (inicia se precisar,
    # reaproveita se já estiver rodando). Matar processos aqui destruiria o
    # listener a cada geração, obrigando um cold-start lento toda vez.
    for idx, cfg in enumerate(itens_config):
        item = cfg['item']
        ref_curta, nome_curto = separar_referencia_nome(item['referencia_full'])

        idade_fmt = formatar_idade(cfg['idade_num'], cfg['idade_uni']) if cfg.get('idade_num') else None
        qtd_fmt = formatar_quantidade(cfg.get('pecas', 1)) or "1 Unidade"

        # --- REGISTRO: busca na BASE por fábrica+família ---
        registro_excel = re.sub(r'\s', '', item.get('registro', '') or '')
        registro_base = buscar_registro_por_fabrica_familia(
            cliente_base_registro, item.get('fabrica', ''), item.get('familia', '')
        )
        registro_base = re.sub(r'\s', '', registro_base or '')

        if registro_base:
            registro = registro_base
            # Avisa se o registro da base não bate com o do Excel
            if registro_excel and registro_base != registro_excel:
                avisos.append(
                    f"⚠️ Item {ref_curta} (Fáb {item.get('fabrica','?')}/Fam {item.get('familia','?')}): "
                    f"registro da base ({registro_base}) ≠ Excel ({registro_excel}). Usei o da base."
                )
        else:
            # Não achou na base: usa o do Excel como fallback, mas avisa
            registro = registro_excel
            avisos.append(
                f"⚠️ Item {ref_curta}: não achei registro na base para Fábrica '{item.get('fabrica','?')}' "
                f"+ Família '{item.get('familia','?')}' (cliente base '{cliente_base_registro}'). "
                f"Usei o do Excel ({registro_excel or 'vazio'})."
            )

        # Selo
        _chave_selo = (cfg['variante'], registro)
        selo_png = _cache_selo.get(_chave_selo)
        if selo_png is None:
            _modelo_existe = buscar_modelo_selo(cfg['variante'])
            if not _modelo_existe:
                avisos.append(f"Item {ref_curta}: modelo de selo '{cfg['variante']}' NÃO está cadastrado "
                              f"na aba 'Modelos de selo'. Cadastre-o primeiro.")
            elif not registro:
                avisos.append(f"Item {ref_curta}: sem número de registro pra gravar no selo.")
            else:
                _selo_docx = gerar_selo_com_registro(cfg['variante'], registro)
                if _selo_docx:
                    _erros_conv = []
                    selo_png = selo_docx_para_png(_selo_docx, soffice, coletar_erro=_erros_conv)
                    if not selo_png and _erros_conv:
                        for _e in _erros_conv:
                            avisos.append(f"Item {ref_curta} — falha ao converter selo: {_e}")
                else:
                    avisos.append(f"Item {ref_curta}: não consegui gravar o registro no modelo de selo.")
            _cache_selo[_chave_selo] = selo_png
        if not selo_png:
            avisos.append(f"Item {ref_curta}: selo '{cfg['variante']}' não gerado (veja o motivo detalhado acima).")

        # Código de barras
        barcode_png, aviso_bc = gerar_codigo_barras_ean13(item.get('codigo_barras', ''))
        if aviso_bc:
            avisos.append(f"Item {ref_curta}: {aviso_bc}")

        # Chorão: entra se for produto restritivo (-3 anos / idade em anos <= 3 tipicamente)
        # Regra: usa o chorão quando a idade detectada for em ANOS (indicativo +3 = restritivo -3)
        usar_chorao = (cfg.get('idade_uni') == "ANOS")
        chorao_para_esse = chorao_png_bytes if usar_chorao else None

        textos_dir = _montar_textos_direita(cfg['tipo'], idade_fmt, cfg.get('titulos_texto', {}), textos_por_chave)

        dados = {
            'solicitante_cnpj': solicitante_cnpj or None,
            'importador_razao': cliente['razao_social'],
            'importador_endereco': cliente['endereco'],
            'importador_cnpj': cliente['cnpj'],
            'origem': origem,
            'quantidade': qtd_fmt,
            'data_fabricacao': data_fab,
            'lote': lote,
            'sac': cliente['sac'],
            'referencia': ref_curta,
            'nome_curto': nome_curto,
            'marca': item.get('marca', ''),
            'textos_direita': textos_dir,
            'tipo': cfg['tipo'],
        }

        _montar_etiqueta_no_doc(doc, dados, selo_png, barcode_png,
                                chorao_png=chorao_para_esse, pilha_png=pilha_png_bytes if cfg['tipo']=="pilha" else None)

        if idx < len(itens_config) - 1:
            doc.add_page_break()

    _buf_docx = _Bio()
    doc.save(_buf_docx)
    docx_bytes = _buf_docx.getvalue()

    return {'docx': docx_bytes, 'avisos': avisos}


# ---- Clientes (importadores) para geração de etiquetas ----

def salvar_cliente_etiqueta(apelido, razao_social, cnpj, endereco, sac, origem, id_existente=None):
    agora = datetime.now().strftime("%d/%m/%Y %H:%M:%S")
    if id_existente:
        cursor.execute("""
        UPDATE clientes_etiqueta
        SET apelido=?, razao_social=?, cnpj=?, endereco=?, sac=?, origem=?, data_atualizacao=?
        WHERE id=?
        """, (clean(apelido), clean(razao_social), clean(cnpj), clean(endereco),
              clean(sac), clean(origem), agora, id_existente))
    else:
        cursor.execute("""
        INSERT INTO clientes_etiqueta (apelido, razao_social, cnpj, endereco, sac, origem, data_cadastro, data_atualizacao)
        VALUES (?, ?, ?, ?, ?, ?, ?, ?)
        ON CONFLICT(apelido) DO UPDATE SET
            razao_social=excluded.razao_social, cnpj=excluded.cnpj,
            endereco=excluded.endereco, sac=excluded.sac, origem=excluded.origem,
            data_atualizacao=excluded.data_atualizacao
        """, (clean(apelido), clean(razao_social), clean(cnpj), clean(endereco),
              clean(sac), clean(origem), agora, agora))
    commit_seguro()


def listar_clientes_etiqueta():
    return pd.read_sql_query("""
    SELECT id, apelido, razao_social, cnpj, endereco, sac, origem, data_atualizacao
    FROM clientes_etiqueta ORDER BY apelido
    """, conn)


def buscar_registro_por_fabrica_familia(cliente_base, fabrica, familia):
    """Busca o registro na tabela `registros` pela combinação cliente + fábrica + família.
    Compara família por número (FAM 20 == 20). Retorna o registro (str) ou None.
    Como fábrica+família é único, retorna o primeiro match."""
    fam_num = re.sub(r'\D', '', str(familia or ''))
    fab_limpa = clean(fabrica).upper()

    rows = cursor.execute("""
    SELECT fabrica, familia, registro FROM registros
    WHERE UPPER(cliente_base) = UPPER(?)
    """, (clean(cliente_base),)).fetchall()

    for _fab, _fam, _reg in rows:
        _fab_r = clean(_fab).upper()
        _fam_r = re.sub(r'\D', '', str(_fam or ''))
        _fab_ok = (fab_limpa == _fab_r) or (fab_limpa and fab_limpa in _fab_r) or (_fab_r and _fab_r in fab_limpa)
        _fam_ok = (fam_num and fam_num == _fam_r)
        if _fab_ok and _fam_ok and _reg:
            return _reg
    return None


def listar_clientes_base_registros():
    """Lista os clientes distintos cadastrados na tabela registros."""
    rows = cursor.execute("""
    SELECT DISTINCT cliente_base FROM registros
    WHERE cliente_base IS NOT NULL AND cliente_base != ''
    ORDER BY cliente_base
    """).fetchall()
    return [r[0] for r in rows]


def buscar_cliente_etiqueta(apelido):
    row = cursor.execute("""
    SELECT id, apelido, razao_social, cnpj, endereco, sac, origem
    FROM clientes_etiqueta WHERE apelido = ?
    """, (clean(apelido),)).fetchone()
    if not row:
        return None
    return {
        "id": row[0], "apelido": row[1], "razao_social": row[2],
        "cnpj": row[3], "endereco": row[4], "sac": row[5], "origem": row[6],
    }


def excluir_cliente_etiqueta(id_cliente):
    cursor.execute("DELETE FROM clientes_etiqueta WHERE id = ?", (id_cliente,))
    commit_seguro()
    return cursor.rowcount


def ler_registro_via_ocr(imagem_bytes):
    """Lê o número de registro de uma imagem de selo via OCR (Tesseract).
    Reutiliza a mesma lógica da conferência de etiquetas: recorte central,
    upscale, variantes de contraste/binarização/nitidez, votação por maioria.
    Retorna a string do registro (ex: '004871/2025') ou None se não conseguir."""
    try:
        import os as _osR
        import shutil as _shR
        from PIL import Image as _ImgR, ImageEnhance as _IER, ImageFilter as _IFR
        from io import BytesIO as _BioR
        from collections import Counter as _CounterR
    except Exception:
        return None

    # Localizar Tesseract
    _tess = None
    for c in [
        _shR.which("tesseract"), _shR.which("tesseract.exe"),
        _osR.path.expandvars(r"%LOCALAPPDATA%\Programs\Tesseract-OCR\tesseract.exe"),
        r"C:\Program Files\Tesseract-OCR\tesseract.exe",
        r"C:\Program Files (x86)\Tesseract-OCR\tesseract.exe",
    ]:
        if c and _osR.path.exists(c):
            _tess = c
            break
    if not _tess:
        return None

    try:
        import pytesseract as _pt
        _pt.pytesseract.tesseract_cmd = _tess
    except ImportError:
        return None

    try:
        _img = _ImgR.open(_BioR(imagem_bytes))
    except Exception:
        return None

    w, h = _img.size

    def _variantes(base_img, regiao):
        cinza = base_img.convert("L")
        bw, bh = cinza.size
        cinza = cinza.resize((bw * 4, bh * 4), _ImgR.LANCZOS)
        cinza = _IER.Contrast(cinza).enhance(2.0)
        bin_ = cinza.point(lambda p: 255 if p > 140 else 0)
        bin_inv = cinza.point(lambda p: 0 if p > 140 else 255)
        nitida = cinza.filter(_IFR.UnsharpMask(radius=2, percent=200, threshold=2))
        return [(cinza, f'{regiao}/contraste'), (bin_, f'{regiao}/binário'),
                (bin_inv, f'{regiao}/binário-inv'), (nitida, f'{regiao}/nítida')]

    def _tentar(variantes):
        candidatos = []
        for img_t, _rot in variantes:
            for psm in ('6', '11'):
                try:
                    txt = _pt.image_to_string(img_t, lang='por', config=f'--psm {psm} --oem 3')
                except Exception:
                    try:
                        txt = _pt.image_to_string(img_t, config=f'--psm {psm} --oem 3')
                    except Exception:
                        txt = ""
                m = re.search(r'(?<!\d)(\d{3})\s?(\d{3})(?!\d)\s*/\s*(\d{4})', txt)
                if m:
                    candidatos.append(f"{m.group(1)}{m.group(2)}/{m.group(3)}")
        for img_t, _rot in variantes:
            for psm in ('7', '8'):
                for oem in ('3', '0'):
                    try:
                        txt = _pt.image_to_string(img_t, config=f'--psm {psm} --oem {oem} -c tessedit_char_whitelist=0123456789/')
                    except Exception:
                        txt = ""
                    m = re.search(r'(?<!\d)(\d{6})(?!\d)\s*/\s*(\d{4})', txt)
                    if m:
                        candidatos.append(f"{m.group(1)}/{m.group(2)}")
        return candidatos

    crop_box = (int(w * 0.10), int(h * 0.25), int(w * 0.90), int(h * 0.75))
    candidatos = _tentar(_variantes(_img.crop(crop_box), 'recorte'))
    if not candidatos:
        candidatos = _tentar(_variantes(_img, 'completa'))

    if not candidatos:
        return None

    escolhido, _n = _CounterR(candidatos).most_common(1)[0]
    return escolhido


# ---- Assets genéricos reaproveitáveis (chorão, pilha/bateria) ----

def salvar_asset_generico(tipo, imagem_bytes, nome_arquivo):
    tipo = clean(tipo).lower()

    try:
        from PIL import Image as _ImgAsset
        from io import BytesIO as _BioAsset
        _w, _h = _ImgAsset.open(_BioAsset(imagem_bytes)).size
    except Exception:
        _w, _h = None, None

    agora = datetime.now().strftime("%d/%m/%Y %H:%M:%S")

    cursor.execute("""
    INSERT INTO assets_etiqueta_genericos (tipo, imagem, nome_arquivo, largura, altura, data_cadastro, data_atualizacao)
    VALUES (?, ?, ?, ?, ?, ?, ?)
    ON CONFLICT(tipo) DO UPDATE SET
        imagem = excluded.imagem,
        nome_arquivo = excluded.nome_arquivo,
        largura = excluded.largura,
        altura = excluded.altura,
        data_atualizacao = excluded.data_atualizacao
    """, (tipo, imagem_bytes, nome_arquivo, _w, _h, agora, agora))

    commit_seguro()


def buscar_asset_generico(tipo):
    tipo = clean(tipo).lower()
    row = cursor.execute("""
    SELECT imagem, nome_arquivo FROM assets_etiqueta_genericos WHERE tipo = ?
    """, (tipo,)).fetchone()
    if not row:
        return None, None
    return row[0], row[1]


def listar_assets_genericos():
    return pd.read_sql_query("""
    SELECT tipo, nome_arquivo, largura, altura, data_cadastro, data_atualizacao
    FROM assets_etiqueta_genericos
    ORDER BY tipo
    """, conn)






# ==========================================
# IP-BRI MANUAL POR FAMÍLIA
# ==========================================

def normalizar_ip_bri_manual(valor):
    texto = clean(valor).upper()

    if not texto:
        return ""

    if texto.startswith("IP-BRI-"):
        return texto

    if texto.startswith("IP-"):
        numero = texto.replace("IP-", "")
        return f"IP-BRI-{numero}"

    return f"IP-BRI-{texto}"


def salvar_ip_bri_familia(ce_bri, familia, ip_bri, observacao=""):
    ce_bri = clean(ce_bri).upper()
    familia = normalizar_familia(familia)
    ip_bri = normalizar_ip_bri_manual(ip_bri)
    observacao = clean(observacao)

    if not ce_bri or not familia or not ip_bri:
        return False, "Informe CE-BRI, família e IP-BRI."

    # Bloqueia o mesmo IP-BRI em outra família/CE-BRI.
    duplicado = cursor.execute("""
    SELECT ce_bri, familia
    FROM ip_bri_familias
    WHERE UPPER(ip_bri) = UPPER(?)
    AND NOT (
        UPPER(ce_bri) = UPPER(?)
        AND familia = ?
    )
    LIMIT 1
    """, (ip_bri, ce_bri, familia)).fetchone()

    if duplicado:
        return False, (
            f"Este IP-BRI já está cadastrado em outro vínculo: "
            f"{duplicado[0]} / família {duplicado[1]}. "
            f"Edite ou exclua o cadastro antigo antes de salvar."
        )

    existente = cursor.execute("""
    SELECT id
    FROM ip_bri_familias
    WHERE UPPER(ce_bri) = UPPER(?)
    AND familia = ?
    """, (ce_bri, familia)).fetchone()

    agora_txt = datetime.now().strftime("%d/%m/%Y %H:%M:%S")

    if existente:
        cursor.execute("""
        UPDATE ip_bri_familias
        SET ip_bri = ?, observacao = ?, data_atualizacao = ?
        WHERE id = ?
        """, (ip_bri, observacao, agora_txt, existente[0]))
    else:
        cursor.execute("""
        INSERT INTO ip_bri_familias (
            ce_bri, familia, ip_bri, observacao, data_cadastro, data_atualizacao
        )
        VALUES (?, ?, ?, ?, ?, ?)
        """, (ce_bri, familia, ip_bri, observacao, agora_txt, agora_txt))

    commit_seguro()
    return True, f"IP-BRI {ip_bri} salvo para {ce_bri} / família {familia}."

@st.cache_data(ttl=60, show_spinner=False)
def buscar_ip_bri_manual_familia(ce_bri, familia):
    ce_bri = clean(ce_bri).upper()
    familia = normalizar_familia(familia)

    if not ce_bri or not familia:
        return ""

    resultado = pd.read_sql_query("""
    SELECT ip_bri
    FROM ip_bri_familias
    WHERE UPPER(ce_bri) = UPPER(?)
    AND familia = ?
    ORDER BY id DESC
    LIMIT 1
    """, conn, params=(ce_bri, familia))

    if resultado.empty:
        return ""

    return clean(resultado.iloc[0]["ip_bri"])


@st.cache_data(ttl=20, show_spinner=False)
def listar_ip_bri_pendentes_sistema5():
    return pd.read_sql_query("""
    SELECT
        s.cliente_base,
        s.categoria,
        s.fabrica,
        s.ce_bri,
        s.familia,
        MAX(r.registro) AS registro,
        GROUP_CONCAT(DISTINCT s.tipo_processo) AS tipos_processo,
        GROUP_CONCAT(DISTINCT s.ip_processo) AS ips_processos,
        MAX(s.data_processo) AS ultima_data_processo,
        COUNT(DISTINCT s.ip_processo) AS qtd_processos,
        COUNT(*) AS qtd_itens
    FROM sistema5_itens s
    LEFT JOIN certificados c
    ON UPPER(c.ce_bri) = UPPER(s.ce_bri)
    AND c.familia = s.familia
    LEFT JOIN ip_bri_familias m
    ON UPPER(m.ce_bri) = UPPER(s.ce_bri)
    AND m.familia = s.familia
    LEFT JOIN registros r
    ON UPPER(r.ce_bri) = UPPER(s.ce_bri)
    AND r.familia = s.familia
    WHERE s.ce_bri IS NOT NULL
    AND s.ce_bri != ''
    AND s.familia IS NOT NULL
    AND s.familia != ''
    AND c.ip_bri IS NULL
    AND m.ip_bri IS NULL
    GROUP BY
        s.cliente_base,
        s.categoria,
        s.fabrica,
        s.ce_bri,
        s.familia
    ORDER BY
        s.cliente_base,
        s.fabrica,
        CAST(s.familia AS INTEGER)
    """, conn)

@st.cache_data(ttl=20, show_spinner=False)
def listar_ip_bri_manuais():
    return pd.read_sql_query("""
    SELECT
        id,
        ce_bri,
        familia,
        ip_bri,
        observacao,
        data_cadastro,
        data_atualizacao
    FROM ip_bri_familias
    ORDER BY ce_bri, CAST(familia AS INTEGER)
    """, conn)

# ==========================================
# PESQUISA AVANÇADA
# ==========================================

@st.cache_data(ttl=30, show_spinner=False)
def pesquisar_por_ip_bri(ip_bri_busca):
    termo = clean(ip_bri_busca).upper()

    if not termo:
        return pd.DataFrame()

    termo_sem_prefixo = termo.replace("IP-BRI-", "")
    termo_com_prefixo = termo if termo.startswith("IP-BRI-") else f"IP-BRI-{termo}"

    return pd.read_sql_query("""
    SELECT DISTINCT
        c.ip_bri AS ip_bri,
        c.ce_bri AS ce_bri,
        c.familia AS fam,
        r.registro AS registro,
        r.fabrica AS fabrica,
        r.endereco_fabrica AS endereco,
        c.rev AS rev,
        c.produto AS produto,
        c.data_emissao AS data_emissao
    FROM certificados c
    LEFT JOIN registros r
    ON UPPER(r.ce_bri) = UPPER(c.ce_bri)
    AND r.familia = c.familia
    WHERE UPPER(c.ip_bri) LIKE ?
       OR REPLACE(UPPER(c.ip_bri), 'IP-BRI-', '') LIKE ?
    ORDER BY c.ip_bri, r.fabrica, r.registro
    """, conn, params=(f"{termo_com_prefixo}%", f"{termo_sem_prefixo}%"))


@st.cache_data(ttl=30, show_spinner=False)
def pesquisar_por_registro(registro_busca):
    termo = clean(registro_busca)

    if not termo:
        return pd.DataFrame()

    return pd.read_sql_query("""
    SELECT
        r.registro AS registro,
        c.ip_bri AS ip_bri,
        c.ce_bri AS ce_bri,
        c.familia AS fam,
        r.fabrica AS fabrica,
        r.endereco_fabrica AS endereco,
        marcas.fabricante AS fabricante,
        c.rev AS rev,
        c.produto AS produto,
        c.data_emissao AS data_emissao
    FROM registros r
    LEFT JOIN certificados c
    ON UPPER(c.ce_bri) = UPPER(r.ce_bri)
    AND c.familia = r.familia
    LEFT JOIN (
        SELECT
            certificado_id,
            MIN(marca) AS fabricante
        FROM itens
        GROUP BY certificado_id
    ) marcas
    ON marcas.certificado_id = c.id
    WHERE r.registro LIKE ?
    ORDER BY r.registro, c.ip_bri, marcas.fabricante
    """, conn, params=(f"{termo}%",))


# ==========================================
# SISTEMA 5 - INCLUSÕES / MANUTENÇÕES
# ==========================================

def normalizar_categoria_sistema5(valor):
    valor = clean(valor).upper()
    if "NOVO" in valor:
        return "SISTEMA 5 NOVO PROJETO"
    if "PROPR" in valor:
        return "SISTEMA 5 PROPRIOS"
    if "FOCUS" in valor:
        return "SISTEMA 5 FOCUS"
    return valor


def normalizar_df_para_exibicao(df):
    """
    Converte colunas com tipos mistos (int/str) para string
    antes de exibir no st.dataframe, evitando ArrowTypeError.
    """
    if df is None or df.empty:
        return df
    df = df.copy()
    for col in df.columns:
        if df[col].dtype == object:
            try:
                df[col] = df[col].astype(str).replace("None", "").replace("nan", "")
            except Exception:
                pass
    return df


def normalizar_ip_processo(ip):
    """
    Normaliza IP de processo para formato padrão: IP-XXXX-XX
    Aceita: IP-0852-26, IP 0852-26, IP0852-26, ip-0852-26
    Sempre retorna no formato IP-XXXX-XX em maiúsculas.
    """
    if not ip:
        return ""
    # Remove espaços extras, uppercase
    ip = clean(ip).upper()
    # Garante que entre IP e os números sempre tem hífen
    ip = re.sub(r"\bIP[\s\-]*(\d)", r"IP-\1", ip)
    return ip


def extrair_ip_processo(nome):
    texto = clean(nome).upper()
    match = re.search(r"\bIP[- ]?\d+[-/]\d+\b", texto)
    if match:
        return normalizar_ip_processo(match.group(0))
    return None


def normalizar_data_processo_manual(texto):
    """Normaliza uma data digitada manualmente (nos campos 'Data' do upload do
    Sistema 5) para o formato ISO 'AAAA-MM-DD', que é o único formato que
    ordena corretamente como texto no banco. Sem essa normalização, digitar
    '09/04/2026' (formato brasileiro natural) em vez de '2026-04-09' quebra
    silenciosamente a ordenação por 'processo mais recente', pois a comparação
    de texto trata '09/04/2026' como sendo 'menor' que '2025-10-27' (compara
    caractere a caractere, '0' < '2'). Aceita AAAA-MM-DD, DD-MM-AAAA, DD/MM/AAAA,
    DD.MM.AAAA. Se não reconhecer o formato, devolve o texto original (para não
    apagar silenciosamente o que a pessoa digitou) e um aviso."""
    texto = clean(texto)
    if not texto:
        return texto, None

    # Já está em ISO? (AAAA-MM-DD)
    m_iso = re.fullmatch(r"(\d{4})-(\d{2})-(\d{2})", texto)
    if m_iso:
        ano, mes, dia = m_iso.groups()
        try:
            if 1 <= int(mes) <= 12 and 1 <= int(dia) <= 31:
                return texto, None
        except ValueError:
            pass

    # Formato brasileiro DD-MM-AAAA / DD/MM/AAAA / DD.MM.AAAA
    m_br = re.fullmatch(r"(\d{2})[-/.](\d{2})[-/.](\d{4})", texto)
    if m_br:
        dia, mes, ano = m_br.groups()
        try:
            if 1 <= int(dia) <= 31 and 1 <= int(mes) <= 12:
                return f"{ano}-{mes}-{dia}", None
        except ValueError:
            pass

    # Não reconheceu o formato — devolve como veio, mas avisa
    return texto, (f"⚠️ Data '{texto}' não está em um formato reconhecido (esperado AAAA-MM-DD "
                   f"ou DD-MM-AAAA). Foi salva como digitada, mas pode prejudicar a ordenação "
                   f"por 'processo mais recente'. Corrija para o formato AAAA-MM-DD se possível.")


def extrair_data_processo(nome):
    """Extrai a data (AAAA-MM-DD) do nome do arquivo do Sistema 5.
    CORRIGIDO (bug encontrado): a versão anterior usava \\b (fronteira de palavra),
    que NÃO reconhece fronteira entre '_' e dígito (ex: '..._09_04_2026' falhava
    silenciosamente, fazendo o processo virar '1900-01-01' no COALESCE e perder
    pra processos mais antigos). Também podia capturar fragmentos de códigos como
    'ZS02-26-07-03' como se fosse uma data (2003-07-26), o que é errado.
    Agora: usa (?<!\\d)/(?!\\d) em vez de \\b, e valida que dia/mês/ano formam uma
    data real e plausível (ano entre 2015 e 2035) antes de aceitar o candidato."""
    texto = clean(nome)
    candidatos = re.finditer(r"(?<!\d)(\d{2})[-_.](\d{2})[-_.](\d{2,4})(?!\d)", texto)
    melhor = None
    for m in candidatos:
        dia_s, mes_s, ano_s = m.groups()
        try:
            dia, mes = int(dia_s), int(mes_s)
        except ValueError:
            continue
        if not (1 <= dia <= 31 and 1 <= mes <= 12):
            continue  # não forma uma data real -> provavelmente fragmento de código
        ano = ano_s if len(ano_s) == 4 else "20" + ano_s
        try:
            ano_num = int(ano)
        except ValueError:
            continue
        if not (2015 <= ano_num <= 2035):
            continue  # ano fora de faixa plausível pro sistema
        candidato_data = f"{ano}-{mes_s.zfill(2)}-{dia_s.zfill(2)}"
        if melhor is None:
            melhor = candidato_data
        if len(ano_s) == 4:
            melhor = candidato_data
            break  # ano de 4 dígitos explícito é o mais confiável, já pode parar
    return melhor


def corrigir_datas_processo_existentes():
    """Varre TODOS os registros já salvos em sistema5_itens e sistema5_arquivos
    procurando data_processo que não esteja em formato ISO (AAAA-MM-DD) — geralmente
    porque foi digitada manualmente no formato brasileiro (DD/MM/AAAA) num upload
    anterior, antes da normalização automática existir. Corrige em massa.
    Retorna um relatório: {corrigidos: int, ja_estavam_ok: int, nao_reconhecidos: [...]}"""
    relatorio = {"corrigidos": 0, "ja_estavam_ok": 0, "nao_reconhecidos": []}

    for tabela in ["sistema5_itens", "sistema5_arquivos"]:
        linhas = cursor.execute(
            f"SELECT id, data_processo FROM {tabela} WHERE data_processo IS NOT NULL AND data_processo != ''"
        ).fetchall()

        for row_id, data_atual in linhas:
            data_normalizada, aviso = normalizar_data_processo_manual(data_atual)
            if aviso:
                relatorio["nao_reconhecidos"].append({
                    "tabela": tabela, "id": row_id, "valor_atual": data_atual
                })
            elif data_normalizada != data_atual:
                cursor.execute(
                    f"UPDATE {tabela} SET data_processo = ? WHERE id = ?",
                    (data_normalizada, row_id)
                )
                relatorio["corrigidos"] += 1
            else:
                relatorio["ja_estavam_ok"] += 1

    commit_seguro()
    return relatorio


def detectar_tipo_processo(nome):
    texto = clean(nome).upper()
    if "MANUT" in texto:
        return "MANUTENCAO"
    if "RECERT" in texto:
        return "RECERTIFICACAO"
    if "INICIAL" in texto:
        return "INICIAL"
    if "INCLUS" in texto:
        return "INCLUSAO"
    return "OUTROS"


def get_or_create_cliente_sistema5(categoria, cliente_base):
    categoria = normalizar_categoria_sistema5(categoria)
    cliente_base = clean(cliente_base).upper()

    cursor.execute("""
    INSERT OR IGNORE INTO sistema5_clientes (categoria, cliente_base, data_cadastro)
    VALUES (?, ?, ?)
    """, (categoria, cliente_base, datetime.now().strftime("%d/%m/%Y %H:%M:%S")))
    commit_seguro()

    return cursor.execute("""
    SELECT id FROM sistema5_clientes
    WHERE categoria = ? AND cliente_base = ?
    """, (categoria, cliente_base)).fetchone()[0]


def get_or_create_fabrica_sistema5(cliente_id, fabrica, ce_bri="", endereco_fabrica=""):
    fabrica = clean(fabrica).upper()
    ce_bri = clean(ce_bri).upper()
    endereco_fabrica = clean(endereco_fabrica)

    cursor.execute("""
    INSERT OR IGNORE INTO sistema5_fabricas (
        cliente_id, fabrica, ce_bri, endereco_fabrica, data_cadastro
    )
    VALUES (?, ?, ?, ?, ?)
    """, (
        cliente_id,
        fabrica,
        ce_bri,
        endereco_fabrica,
        datetime.now().strftime("%d/%m/%Y %H:%M:%S")
    ))

    cursor.execute("""
    UPDATE sistema5_fabricas
    SET
        ce_bri = COALESCE(NULLIF(?, ''), ce_bri),
        endereco_fabrica = COALESCE(NULLIF(?, ''), endereco_fabrica)
    WHERE cliente_id = ? AND fabrica = ?
    """, (ce_bri, endereco_fabrica, cliente_id, fabrica))

    commit_seguro()

    return cursor.execute("""
    SELECT id FROM sistema5_fabricas
    WHERE cliente_id = ? AND fabrica = ?
    """, (cliente_id, fabrica)).fetchone()[0]


def normalizar_coluna_excel_s5(valor):
    txt = clean(valor).upper()

    # Remove acentos para padronizar
    troca = {
        "Á": "A", "À": "A", "Â": "A", "Ã": "A",
        "É": "E", "Ê": "E",
        "Í": "I",
        "Ó": "O", "Ô": "O", "Õ": "O",
        "Ú": "U",
        "Ç": "C"
    }

    for antigo, novo in troca.items():
        txt = txt.replace(antigo, novo)

    txt = re.sub(r"\s+", " ", txt).strip()

    # Ignora colunas auxiliares que citam "modelo", mas não são referência.
    if "APLICAVEL SOMENTE" in txt or "QUANTIDADE" in txt or "NCM" in txt:
        return txt

    # Item
    if txt == "ITEM" or txt.startswith("ITEM "):
        return "ITEM"

    # Códigos
    if (
        "CODIGO DE BARRAS" in txt
        or "COD DE BARRAS" in txt
        or "COD. DE BARRAS" in txt
        or txt in ["CODIGO", "COD.", "COD"]
        or "BARRAS" in txt
    ):
        return "CODIGO"

    # Modelo / Referência
    # IMPORTANTE:
    # Não pode reconhecer qualquer texto com "MODELO",
    # porque existe coluna "Quantidade aplicável somente para o modelo 1b".
    if (
        txt.startswith("MODELO")
        or txt.startswith("REFERENCIA")
        or "MODELO / REFERENCIA" in txt
        or "MODELO/REFERENCIA" in txt
        or "DESIGNACAO COMERCIAL" in txt
    ):
        return "MODELO"

    # Marca
    if (
        txt == "MARCA"
        or txt.startswith("MARCA ")
        or "MARCA COMERCIALIZADA" in txt
        or "FABRICANTE" in txt
    ):
        return "MARCA"

    # Descrição técnica
    if (
        "DESCRICAO TECNICA" in txt
        or "DESCRICAO TECNICA DO MODELO" in txt
        or txt.startswith("DESCRICAO")
        or txt == "NOME"
        or "PROCESSO PRODUTIVO" in txt
    ):
        return "NOME"

    return txt



def valor_valido_item_s5(valor):
    texto = clean(valor)

    if not texto:
        return False

    if texto.lower() in ["nan", "none", "null", "-", "--"]:
        return False

    return True


def codigo_barras_valido_s5(valor, number_format=None):
    return codigo_barras_para_texto(valor, number_format)


def familia_desconsiderada_s5(nome_aba):
    texto = clean(nome_aba).upper()
    texto = texto.replace("Á", "A").replace("À", "A").replace("Â", "A").replace("Ã", "A")
    texto = texto.replace("É", "E").replace("Ê", "E")
    texto = texto.replace("Í", "I")
    texto = texto.replace("Ó", "O").replace("Ô", "O").replace("Õ", "O")
    texto = texto.replace("Ú", "U")
    texto = texto.replace("Ç", "C")

    return "DESCONSIDERADO" in texto or "DESCONSIDERADA" in texto



def limpar_itens_desconsiderados_sistema5():
    """
    Remove itens antigos que já foram cadastrados no Sistema 5
    em famílias/abas marcadas como DESCONSIDERADO ou DESCONSIDERADA.
    """

    try:
        qtd_antes = cursor.execute("""
        SELECT COUNT(*)
        FROM sistema5_itens
        WHERE UPPER(familia) LIKE '%DESCONSIDERADO%'
        OR UPPER(familia) LIKE '%DESCONSIDERADA%'
        """).fetchone()[0]

        cursor.execute("""
        DELETE FROM sistema5_itens
        WHERE UPPER(familia) LIKE '%DESCONSIDERADO%'
        OR UPPER(familia) LIKE '%DESCONSIDERADA%'
        """)

        commit_seguro()

        return qtd_antes

    except Exception:
        return 0



def parece_codigo_barras_s5(valor):
    codigo = re.sub(r"\D", "", clean(valor))
    return bool(codigo and len(codigo) >= 8 and len(codigo) <= 14)


def parece_marca_s5(valor):
    texto = clean(valor)
    if not texto or texto.lower() in ["nan", "none", "null"]:
        return False
    # Marca normalmente tem letras; código de barras normalmente só números.
    return bool(re.search(r"[A-Za-zÀ-ÿ]", texto))


def corrigir_marca_codigo_invertidos_s5(marca, codigo_raw):
    """
    Alguns memoriais vêm com o cabeçalho:
    Marca Comercializada | Código de Barras
    mas os dados vêm invertidos:
    código na coluna da marca e marca na coluna do código.

    Se detectar isso, troca automaticamente.
    """
    marca_limpa = clean(marca)
    codigo_limpo = clean(codigo_raw)

    if parece_codigo_barras_s5(marca_limpa) and parece_marca_s5(codigo_limpo):
        return codigo_limpo, marca_limpa

    return marca_limpa, codigo_limpo


def ler_excel_inclusao_sistema5(uploaded_file):
    """
    Leitor específico do Sistema 5 com preservação de código de barras.

    Principais cuidados:
    - Lê todas as abas.
    - Cada aba representa uma família.
    - Ignora abas DESCONSIDERADO/DESCONSIDERADA.
    - Preserva códigos de barras que começam com 0 usando o number_format do Excel.
    - Corrige casos em que MARCA e CÓDIGO vêm invertidos.
    - Ignora imagens corrompidas (.mpo e outros formatos não suportados).
    """

    try:
        uploaded_file.seek(0)
    except Exception:
        pass

    try:
        wb = load_workbook(uploaded_file, data_only=True)
    except Exception as e:
        erro_str = str(e).lower()
        if any(x in erro_str for x in ["mpo", "image", "picture", "drawing", "wmf", "emf"]):
            try:
                wb = load_workbook(_limpar_excel_imagens(uploaded_file), data_only=True)
            except Exception:
                return pd.DataFrame()
        else:
            return pd.DataFrame()

    todos_itens = []

    for ws in wb.worksheets:
        aba = ws.title

        if familia_desconsiderada_s5(aba):
            continue

        header_row_idx = None
        colunas = {}

        for row in ws.iter_rows():
            temp = {}

            for cell in row:
                nome = normalizar_coluna_excel_s5(cell.value)

                if nome == "ITEM":
                    temp["ITEM"] = cell.column
                elif nome == "MODELO":
                    temp["MODELO"] = cell.column
                elif nome == "MARCA":
                    temp["MARCA"] = cell.column
                elif nome == "CODIGO":
                    temp["CODIGO"] = cell.column
                elif nome == "NOME":
                    temp["NOME"] = cell.column

            if all(campo in temp for campo in ["MODELO", "MARCA", "CODIGO", "NOME"]):
                header_row_idx = row[0].row
                colunas = temp
                break

        if header_row_idx is None:
            continue

        for row_idx in range(header_row_idx + 1, ws.max_row + 1):
            def cell_col(nome_coluna):
                col = colunas.get(nome_coluna)
                if not col:
                    return None
                return ws.cell(row_idx, col)

            def valor_coluna(nome_coluna):
                cell = cell_col(nome_coluna)
                if cell is None:
                    return ""
                return clean(cell.value)

            item = valor_coluna("ITEM")
            modelo = valor_coluna("MODELO")
            marca = valor_coluna("MARCA")
            nome = valor_coluna("NOME")

            marca_cell = cell_col("MARCA")
            codigo_cell = cell_col("CODIGO")

            marca_raw = clean(marca_cell.value) if marca_cell else ""
            codigo_raw = clean(codigo_cell.value) if codigo_cell else ""

            # Caso normal: código está na coluna de código.
            codigo = codigo_barras_para_texto(
                codigo_cell.value if codigo_cell else codigo_raw,
                codigo_cell.number_format if codigo_cell else None
            )

            # Caso especial: código e marca vieram invertidos.
            if parece_codigo_barras_s5(marca_raw) and parece_marca_s5(codigo_raw):
                marca = codigo_raw
                codigo = codigo_barras_para_texto(
                    marca_cell.value if marca_cell else marca_raw,
                    marca_cell.number_format if marca_cell else None
                )
            else:
                marca = marca_raw

            linha_texto = " ".join([modelo, marca, codigo_raw, nome]).upper()

            palavras_ignorar = [
                "OBS",
                "OBSERVAÇÃO",
                "OBSERVACAO",
                "LOCAL E DATA",
                "PLACE AND DATE",
                "ASSINATURA",
                "SIGNATURE",
                "RESPONSÁVEL",
                "RESPONSAVEL",
                "APROVAÇÃO",
                "APROVACAO"
            ]

            if any(p in linha_texto for p in palavras_ignorar):
                continue

            if not valor_valido_item_s5(modelo):
                continue

            if not valor_valido_item_s5(marca):
                continue

            if not valor_valido_item_s5(nome):
                continue

            if not codigo:
                continue

            if marca.upper().strip() in ["DIRETA", "INDIRETA"]:
                continue

            if modelo.upper().strip() in ["DIRETA", "INDIRETA"]:
                continue

            if nome.upper().strip() in ["DIRETA", "INDIRETA"]:
                continue

            if normalizar_coluna_excel_s5(modelo) == "MODELO":
                continue

            if normalizar_coluna_excel_s5(marca) == "MARCA":
                continue

            todos_itens.append({
                "FAMILIA": normalizar_familia_aba_s5(aba),
                "ITEM": item,
                "MARCA": marca,
                "MODELO": modelo,
                "NOME": nome,
                "CODIGO": codigo
            })

    if not todos_itens:
        return pd.DataFrame()

    df_final = pd.DataFrame(todos_itens)

    df_final = df_final.drop_duplicates(
        subset=["FAMILIA", "MODELO", "CODIGO"],
        keep="first"
    )

    return df_final


@st.cache_data(ttl=15, show_spinner=False)
def listar_clientes_sistema5_banco(categoria=""):
    """
    Retorna clientes já salvos na tabela sistema5_clientes.
    Usado para popular dropdown em SISTEMA 5 PROPRIOS.
    """
    categoria = normalizar_categoria_sistema5(categoria) if categoria else ""

    if categoria:
        return pd.read_sql_query("""
        SELECT DISTINCT cliente_base
        FROM sistema5_clientes
        WHERE categoria = ?
        AND cliente_base IS NOT NULL AND cliente_base != ''
        ORDER BY cliente_base
        """, conn, params=(categoria,))

    return pd.read_sql_query("""
    SELECT DISTINCT cliente_base
    FROM sistema5_clientes
    WHERE cliente_base IS NOT NULL AND cliente_base != ''
    ORDER BY cliente_base
    """, conn)


@st.cache_data(ttl=15, show_spinner=False)
def listar_fabricas_sistema5_banco(cliente_base="", categoria=""):
    """
    Retorna fabricas de duas fontes: sistema5_fabricas + registros.
    Unifica para o selectbox de fabrica.
    Retorna vazio se cliente_base não for informado — nunca mistura clientes.
    """
    cliente_base = clean(cliente_base).upper()
    categoria = normalizar_categoria_sistema5(categoria) if categoria else ""

    # Sem cliente definido, não faz sentido listar fábricas
    if not cliente_base:
        return pd.DataFrame(columns=["fabrica", "ce_bri", "endereco_fabrica", "origem"])

    params_s5 = [cliente_base]
    where_s5_parts = ["UPPER(c.cliente_base) = UPPER(?)"]
    if categoria:
        where_s5_parts.append("c.categoria = ?")
        params_s5.append(categoria)
    where_s5 = "WHERE " + " AND ".join(where_s5_parts)

    df_s5 = pd.read_sql_query(f"""
    SELECT DISTINCT
        f.fabrica,
        f.ce_bri,
        f.endereco_fabrica,
        'Sistema 5' AS origem
    FROM sistema5_fabricas f
    INNER JOIN sistema5_clientes c ON c.id = f.cliente_id
    {where_s5}
    AND f.fabrica IS NOT NULL AND f.fabrica != ''
    ORDER BY f.fabrica
    """, conn, params=tuple(params_s5))

    df_reg = pd.read_sql_query("""
    SELECT DISTINCT
        fabrica,
        ce_bri,
        endereco_fabrica,
        'Registros' AS origem
    FROM registros
    WHERE UPPER(cliente_base) = UPPER(?)
    AND fabrica IS NOT NULL AND fabrica != ''
    ORDER BY fabrica
    """, conn, params=(cliente_base,))

    combinado = pd.concat([df_reg, df_s5], ignore_index=True)
    combinado = combinado.drop_duplicates(subset=["fabrica"], keep="first")
    return combinado.sort_values("fabrica").reset_index(drop=True)


@st.cache_data(ttl=30, show_spinner=False)
def buscar_fabricante_por_fabrica_s5(cliente_base, fabrica, ce_bri=""):
    """
    Para RECERTIFICACAO: busca a marca mais frequente nos itens
    ja salvos para esta fabrica (certificados ou sistema5_itens).
    """
    cliente_base = clean(cliente_base).upper()
    fabrica = clean(fabrica).upper()
    ce_bri = clean(ce_bri).upper()

    if ce_bri:
        resultado = pd.read_sql_query("""
        SELECT marca, COUNT(*) AS freq
        FROM itens
        WHERE certificado_id IN (
            SELECT id FROM certificados WHERE UPPER(ce_bri) = UPPER(?)
        )
        AND marca IS NOT NULL AND marca != ''
        GROUP BY UPPER(marca)
        ORDER BY freq DESC
        LIMIT 1
        """, conn, params=(ce_bri,))
        if not resultado.empty:
            return clean(resultado.iloc[0]["marca"])

    params = []
    where_parts = []
    if cliente_base:
        where_parts.append("UPPER(cliente_base) = UPPER(?)")
        params.append(cliente_base)
    if fabrica:
        where_parts.append("UPPER(fabrica) = UPPER(?)")
        params.append(fabrica)

    if not where_parts:
        return ""

    resultado = pd.read_sql_query(f"""
    SELECT marca, COUNT(*) AS freq
    FROM sistema5_itens
    WHERE {" AND ".join(where_parts)}
    AND marca IS NOT NULL AND marca != ''
    GROUP BY UPPER(marca)
    ORDER BY freq DESC
    LIMIT 1
    """, conn, params=tuple(params))

    return clean(resultado.iloc[0]["marca"]) if not resultado.empty else ""


def ip_processo_ja_existe_sistema5(ip_processo, cliente_base="", fabrica=""):
    """
    Retorna os processos já salvos com o mesmo IP.
    Filtra por cliente_base e fabrica se informados.
    """
    ip = normalizar_ip_processo(clean(ip_processo).upper())

    where = ["UPPER(a.ip_processo) = ?"]
    params = [ip]

    if clean(cliente_base):
        where.append("UPPER(c.cliente_base) = UPPER(?)")
        params.append(clean(cliente_base))

    if clean(fabrica):
        where.append("UPPER(f.fabrica) = UPPER(?)")
        params.append(clean(fabrica))

    sql = f"""
    SELECT
        a.id AS arquivo_id,
        c.categoria,
        c.cliente_base,
        f.fabrica,
        a.tipo_processo,
        a.ip_processo,
        a.data_processo,
        a.arquivo_nome,
        COUNT(i.id) AS qtd_itens
    FROM sistema5_arquivos a
    INNER JOIN sistema5_clientes c ON c.id = a.cliente_id
    INNER JOIN sistema5_fabricas f ON f.id = a.fabrica_id
    LEFT JOIN sistema5_itens i ON i.arquivo_id = a.id
    WHERE {" AND ".join(where)}
    GROUP BY a.id
    ORDER BY a.id DESC
    """

    return pd.read_sql_query(sql, conn, params=tuple(params))


def deletar_processo_sistema5(arquivo_id):
    """
    Remove um processo do Sistema 5 e todos os seus itens.
    Retorna o número de itens deletados.
    """
    arquivo_id = int(arquivo_id)

    qtd = cursor.execute(
        "SELECT COUNT(*) FROM sistema5_itens WHERE arquivo_id = ?",
        (arquivo_id,)
    ).fetchone()[0]

    cursor.execute("DELETE FROM sistema5_itens WHERE arquivo_id = ?", (arquivo_id,))
    cursor.execute("DELETE FROM sistema5_arquivos WHERE id = ?", (arquivo_id,))
    commit_seguro()

    return qtd


def mover_processo_para_fabrica(arquivo_id, nova_fabrica_id):
    """
    Move um processo (arquivo) de uma fábrica para outra dentro do mesmo cliente.
    Atualiza sistema5_arquivos.fabrica_id e todos os itens em sistema5_itens
    (fabrica, ce_bri, endereco_fabrica) num único commit.
    """
    arquivo_id    = int(arquivo_id)
    nova_fabrica_id = int(nova_fabrica_id)

    destino = cursor.execute(
        "SELECT fabrica, ce_bri, endereco_fabrica, cliente_id FROM sistema5_fabricas WHERE id = ?",
        (nova_fabrica_id,)
    ).fetchone()
    if not destino:
        raise ValueError("Fábrica de destino não encontrada no banco.")

    nova_fab_nome, nova_ce_bri, nova_endereco, nova_cliente_id = destino
    nova_ce_bri  = nova_ce_bri  or ""
    nova_endereco = nova_endereco or ""

    cursor.execute(
        "UPDATE sistema5_arquivos SET fabrica_id = ?, cliente_id = ? WHERE id = ?",
        (nova_fabrica_id, nova_cliente_id, arquivo_id)
    )
    cursor.execute(
        """UPDATE sistema5_itens
           SET fabrica = ?, ce_bri = ?, endereco_fabrica = ?
           WHERE arquivo_id = ?""",
        (nova_fab_nome, nova_ce_bri, nova_endereco, arquivo_id)
    )
    commit_seguro()
    return nova_fab_nome


def buscar_processos_por_ip(ip_texto):
    """Busca processos no Sistema 5 cujo IP contenha o texto digitado."""
    termo = f"%{clean(ip_texto)}%"
    return pd.read_sql_query("""
    SELECT
        a.id AS arquivo_id,
        c.categoria,
        c.cliente_base,
        f.fabrica,
        f.ce_bri,
        a.ip_processo,
        a.tipo_processo,
        a.data_processo
    FROM sistema5_arquivos a
    INNER JOIN sistema5_clientes c ON c.id = a.cliente_id
    INNER JOIN sistema5_fabricas f ON f.id = a.fabrica_id
    WHERE UPPER(a.ip_processo) LIKE UPPER(?)
    ORDER BY c.cliente_base, f.fabrica, a.ip_processo
    """, conn, params=(termo,))


def verificar_itens_sem_medidas(cliente_base=""):
    """
    Retorna itens do Sistema 5 cujo NOME não contém padrão de medidas.
    Padrões reconhecidos: 25X15, 25 X 15, 30CM, 15MM, 2M, etc.
    """
    _pat = re.compile(
        r'\b\d+\s*[Xx×]\s*\d+'   # 25X15, 25 x 15, 25×15
        r'|\b\d+\s*CM\b'          # 30CM, 30 CM
        r'|\b\d+\s*MM\b'          # 15MM
        r'|\b\d+\s*M\b',          # 2M
        re.IGNORECASE
    )
    params = []
    where  = "WHERE nome IS NOT NULL AND TRIM(nome) != ''"
    if clean(cliente_base):
        where += " AND UPPER(cliente_base) = UPPER(?)"
        params.append(clean(cliente_base))

    df = pd.read_sql_query(f"""
    SELECT
        cliente_base AS CLIENTE,
        fabrica      AS FABRICA,
        ip_processo  AS IP,
        tipo_processo AS TIPO,
        data_processo AS DATA,
        marca  AS MARCA,
        modelo AS MODELO,
        nome   AS NOME
    FROM sistema5_itens
    {where}
    ORDER BY cliente_base, fabrica, ip_processo, modelo
    """, conn, params=params if params else None)

    if df.empty:
        return df

    return df[~df["NOME"].apply(lambda n: bool(_pat.search(str(n))))].copy()


def listar_todas_fabricas_para_mover():
    """Lista todas as fábricas do Sistema 5 para usar como destino no mover em lote."""
    return pd.read_sql_query("""
    SELECT f.id, f.fabrica, f.ce_bri, c.cliente_base, c.categoria
    FROM sistema5_fabricas f
    INNER JOIN sistema5_clientes c ON c.id = f.cliente_id
    ORDER BY c.cliente_base, f.fabrica
    """, conn)


def salvar_inclusao_sistema5(categoria, cliente_base, fabrica, ce_bri, endereco_fabrica, tipo_processo, ip_processo, data_processo, arquivo_nome, df_itens):
    categoria = normalizar_categoria_sistema5(categoria)
    cliente_base = clean(cliente_base).upper()
    fabrica = clean(fabrica).upper()
    ce_bri = clean(ce_bri).upper()
    endereco_fabrica = clean(endereco_fabrica)
    tipo_processo = clean(tipo_processo).upper()
    ip_processo = normalizar_ip_processo(clean(ip_processo).upper())
    # Segurança extra (defesa em profundidade): garante que data_processo sempre
    # chega no banco em formato ISO (AAAA-MM-DD), mesmo que algum caminho não
    # normalize antes de chamar esta função. Formato errado quebra a ordenação
    # de "processo mais recente" em todo o sistema.
    data_processo, _ = normalizar_data_processo_manual(data_processo)

    cliente_id = get_or_create_cliente_sistema5(categoria, cliente_base)
    fabrica_id = get_or_create_fabrica_sistema5(cliente_id, fabrica, ce_bri, endereco_fabrica)

    cursor.execute("""
    INSERT INTO sistema5_arquivos (
        cliente_id, fabrica_id, tipo_processo, ip_processo, data_processo, arquivo_nome, data_upload
    )
    VALUES (?, ?, ?, ?, ?, ?, ?)
    """, (
        cliente_id,
        fabrica_id,
        tipo_processo,
        ip_processo,
        data_processo,
        arquivo_nome,
        datetime.now().strftime("%d/%m/%Y %H:%M:%S")
    ))
    commit_seguro()
    arquivo_id = cursor.lastrowid

    total = 0

    for _, r in df_itens.iterrows():
        cursor.execute("""
        INSERT INTO sistema5_itens (
            arquivo_id, cliente_base, categoria, fabrica, ce_bri, endereco_fabrica,
            tipo_processo, ip_processo, data_processo, familia, item, marca, modelo, nome, codigo, modelo_ref_key,
            arquivo_nome, data_upload
        )
        VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
        """, (
            arquivo_id,
            cliente_base,
            categoria,
            fabrica,
            ce_bri,
            endereco_fabrica,
            tipo_processo,
            ip_processo,
            data_processo,
            clean(r.get("FAMILIA", "")),
            clean(r.get("ITEM", "")),
            clean(r.get("MARCA", "")),
            clean(r.get("MODELO", "")),
            clean(r.get("NOME", "")),
            clean(r.get("CODIGO", "")),
            referencia_inicial_chave_modelo(clean(r.get("MODELO", ""))),
            arquivo_nome,
            datetime.now().strftime("%d/%m/%Y %H:%M:%S")
        ))
        total += 1

    commit_seguro()
    return total


@st.cache_data(ttl=60, show_spinner=False)
@st.cache_data(ttl=60, show_spinner=False)
def buscar_item_sistema5_confirmacao(valor_ref, _debug_ref=None):
    """_debug_ref: se igual a valor_ref (após limpeza), imprime diagnóstico
    detalhado de cada etapa da busca — usado para investigar por que uma
    referência específica está trazendo o processo errado."""
    ref = clean(valor_ref)
    if not ref:
        return None

    ref_key = referencia_chave_exata(ref)
    ref_sem_zero = ref.lstrip("0") or ref
    _mostrar_debug = _debug_ref and clean(_debug_ref) == ref

    if _mostrar_debug:
        st.info(f"🔎 **DIAGNÓSTICO para referência '{ref}'**")
        st.caption(f"ref_key calculada: {ref_key!r}")

        # Diagnóstico extra: mostra TODOS os registros que contêm essa referência
        # no MODELO, independente do modelo_ref_key salvo — revela se algum
        # processo tem modelo_ref_key diferente do esperado (o que faria ele
        # nunca ser encontrado pela busca exata do Passo 1).
        _todos_com_ref = pd.read_sql_query("""
        SELECT id AS ID_SISTEMA5, modelo AS MODELO, modelo_ref_key AS MODELO_REF_KEY,
               data_processo AS DATA_PROCESSO, arquivo_nome AS ARQUIVO_ORIGEM
        FROM sistema5_itens
        WHERE UPPER(modelo) LIKE UPPER(?)
        ORDER BY id DESC
        """, conn, params=(f"%{ref}%",))
        st.warning(f"🔎 **Diagnóstico extra** — TODOS os registros com '{ref}' no modelo (independente do modelo_ref_key salvo): {len(_todos_com_ref)} encontrado(s)")
        if not _todos_com_ref.empty:
            st.dataframe(_todos_com_ref, use_container_width=True)
            _chaves_distintas = _todos_com_ref["MODELO_REF_KEY"].unique()
            if len(_chaves_distintas) > 1:
                st.error(f"❌ ACHADO: existem {len(_chaves_distintas)} valores DIFERENTES de modelo_ref_key para a mesma referência: {list(_chaves_distintas)}. "
                        f"Isso faz a busca exata (Passo 1) só encontrar os que batem com '{ref_key}', ignorando os outros!")

    # 1. Busca por modelo_ref_key exato
    resultado = pd.read_sql_query("""
    SELECT
        marca AS MARCA,
        modelo AS MODELO,
        nome AS NOME,
        codigo AS CODIGO,
        ip_processo AS IP_PROCESSO,
        ce_bri AS CE_BRI,
        NULL AS REGISTRO,
        endereco_fabrica AS ENDERECO,
        fabrica AS FABRICA,
        tipo_processo AS TIPO_PROCESSO,
        data_processo AS DATA_PROCESSO,
        familia AS FAMILIA,
        arquivo_nome AS ARQUIVO_ORIGEM,
        modelo_ref_key AS MODELO_REF_KEY,
        data_upload AS DATA_UPLOAD_RAW,
        id AS ID_SISTEMA5
    FROM sistema5_itens
    WHERE modelo_ref_key = ?
      AND UPPER(COALESCE(familia,'')) NOT LIKE '%DESCONSIDER%'
    ORDER BY COALESCE(
        data_processo,
        substr(data_upload,7,4) || '-' || substr(data_upload,4,2) || '-' || substr(data_upload,1,2),
        '1900-01-01'
    ) DESC, id DESC
    LIMIT 10
    """, conn, params=(ref_key,))

    if _mostrar_debug:
        st.caption(f"Passo 1 — busca por modelo_ref_key = '{ref_key}': **{len(resultado)}** linha(s) encontrada(s)")
        if not resultado.empty:
            st.dataframe(resultado[["ID_SISTEMA5", "MODELO", "MODELO_REF_KEY", "DATA_PROCESSO", "DATA_UPLOAD_RAW", "ARQUIVO_ORIGEM"]], use_container_width=True)

    # CORRIGIDO: antes, se o Passo 1 achasse QUALQUER linha (mesmo 1 só), o sistema
    # parava e nunca tentava a busca mais ampla (LIKE) — então, se algum processo
    # tivesse modelo_ref_key salvo diferente (de uma versão antiga do sistema, ou
    # gravado errado), ele nunca entrava na disputa e o mais recente podia ficar de
    # fora silenciosamente. Agora SEMPRE roda os dois caminhos e junta os resultados,
    # garantindo que nenhum processo fique escondido por causa de modelo_ref_key
    # desatualizado.
    candidatos_like = pd.read_sql_query("""
    SELECT
        marca AS MARCA,
        modelo AS MODELO,
        nome AS NOME,
        codigo AS CODIGO,
        ip_processo AS IP_PROCESSO,
        ce_bri AS CE_BRI,
        NULL AS REGISTRO,
        endereco_fabrica AS ENDERECO,
        fabrica AS FABRICA,
        tipo_processo AS TIPO_PROCESSO,
        data_processo AS DATA_PROCESSO,
        familia AS FAMILIA,
        arquivo_nome AS ARQUIVO_ORIGEM,
        modelo_ref_key AS MODELO_REF_KEY,
        data_upload AS DATA_UPLOAD_RAW,
        id AS ID_SISTEMA5
    FROM sistema5_itens
    WHERE UPPER(modelo) LIKE UPPER(?)
      AND UPPER(COALESCE(familia,'')) NOT LIKE '%DESCONSIDER%'
    ORDER BY COALESCE(
        data_processo,
        substr(data_upload,7,4) || '-' || substr(data_upload,4,2) || '-' || substr(data_upload,1,2),
        '1900-01-01'
    ) DESC, id DESC
    LIMIT 100
    """, conn, params=(f"{ref}%",))

    if _mostrar_debug:
        st.caption(f"Passo 1b — busca ampla por LIKE '{ref}%' (roda sempre agora): **{len(candidatos_like)}** linha(s) encontrada(s)")
        if not candidatos_like.empty:
            st.dataframe(candidatos_like[["ID_SISTEMA5", "MODELO", "MODELO_REF_KEY", "DATA_PROCESSO"]], use_container_width=True)

    # Junta os dois conjuntos (exato + LIKE) e remove duplicatas por ID
    if not resultado.empty or not candidatos_like.empty:
        resultado = pd.concat([resultado, candidatos_like], ignore_index=True)
        resultado = resultado.drop_duplicates(subset=["ID_SISTEMA5"])
        # Reordena pelo critério de data (o concat pode ter embaralhado a ordem)
        resultado["_ord"] = resultado.apply(
            lambda r: r["DATA_PROCESSO"] if pd.notna(r["DATA_PROCESSO"]) and r["DATA_PROCESSO"]
            else (
                f"{r['DATA_UPLOAD_RAW'][6:10]}-{r['DATA_UPLOAD_RAW'][3:5]}-{r['DATA_UPLOAD_RAW'][0:2]}"
                if pd.notna(r["DATA_UPLOAD_RAW"]) and r["DATA_UPLOAD_RAW"] and len(str(r["DATA_UPLOAD_RAW"])) >= 10
                else "1900-01-01"
            ),
            axis=1
        )
        resultado = resultado.sort_values("_ord", ascending=False).drop(columns=["_ord"])

    # Valida que o modelo encontrado realmente corresponde à referência
    if not resultado.empty:
        resultado = filtrar_resultado_por_referencia_exata(resultado, ref)
        if _mostrar_debug:
            st.caption(f"Passo 2 — após juntar os dois caminhos + filtrar_resultado_por_referencia_exata: **{len(resultado)}** linha(s) restante(s)")
            if not resultado.empty:
                st.dataframe(resultado[["ID_SISTEMA5", "MODELO", "DATA_PROCESSO"]], use_container_width=True)
            else:
                st.warning("⚠️ O filtro de referência exata REMOVEU todos os candidatos! (isso indica que o MODELO salvo não bate com o padrão esperado da referência)")

    if resultado.empty:
        if _mostrar_debug:
            st.error("❌ Nenhum resultado em nenhuma das duas buscas — retornaria None.")
        return None

    item = resultado.iloc[0].to_dict()

    if _mostrar_debug:
        st.success(f"✅ ESCOLHIDO: MODELO='{item['MODELO']}' | DATA_PROCESSO={item['DATA_PROCESSO']} | ID={item['ID_SISTEMA5']}")

    ce_bri = clean(item.get("CE_BRI"))
    familia = clean(item.get("FAMILIA"))

    ip_bri_oficial = None
    if ce_bri and familia:
        oficial = pd.read_sql_query("""
        SELECT ip_bri
        FROM certificados
        WHERE UPPER(ce_bri) = UPPER(?)
        AND familia = ?
        ORDER BY rev DESC, id DESC
        LIMIT 1
        """, conn, params=(ce_bri, familia))
        if not oficial.empty:
            ip_bri_oficial = clean(oficial.iloc[0]["ip_bri"])

    if ip_bri_oficial:
        item["IP_BRI"] = ip_bri_oficial
    else:
        ip_bri_manual = buscar_ip_bri_manual_familia(ce_bri, familia) if "buscar_ip_bri_manual_familia" in globals() else ""
        item["IP_BRI"] = ip_bri_manual if ip_bri_manual else familia

    if ce_bri and familia:
        registro_complementar = pd.read_sql_query("""
        SELECT registro
        FROM registros
        WHERE UPPER(ce_bri) = UPPER(?)
        AND familia = ?
        ORDER BY id DESC
        LIMIT 1
        """, conn, params=(ce_bri, familia))
        if not registro_complementar.empty:
            item["REGISTRO"] = registro_complementar.iloc[0]["registro"]

    return item


@st.cache_data(ttl=20, show_spinner=False)
def listar_sistema5_resumo():
    return pd.read_sql_query("""
    SELECT
        c.categoria,
        c.cliente_base,
        f.fabrica,
        f.ce_bri,
        f.endereco_fabrica,
        a.tipo_processo,
        a.ip_processo,
        a.data_processo,
        a.arquivo_nome,
        COUNT(i.id) AS qtd_itens
    FROM sistema5_clientes c
    LEFT JOIN sistema5_fabricas f ON f.cliente_id = c.id
    LEFT JOIN sistema5_arquivos a ON a.fabrica_id = f.id
    LEFT JOIN sistema5_itens i ON i.arquivo_id = a.id
    GROUP BY c.categoria, c.cliente_base, f.fabrica, f.ce_bri, f.endereco_fabrica,
             a.tipo_processo, a.ip_processo, a.data_processo, a.arquivo_nome
    ORDER BY c.categoria, c.cliente_base, f.fabrica, a.data_processo DESC
    """, conn)


@st.cache_data(ttl=30, show_spinner=False)
def buscar_dados_fabrica_existente(cliente_base="", fabrica="", ce_bri=""):
    cliente_base = clean(cliente_base).upper()
    fabrica = clean(fabrica)
    ce_bri = clean(ce_bri).upper()

    where = []
    params = []

    if cliente_base:
        where.append("UPPER(cliente_base) = UPPER(?)")
        params.append(cliente_base)

    if fabrica:
        where.append("UPPER(fabrica) = UPPER(?)")
        params.append(fabrica)

    if ce_bri:
        where.append("UPPER(ce_bri) = UPPER(?)")
        params.append(ce_bri)

    if not where:
        return None

    sql = f"""
    SELECT
        cliente_base,
        fabrica,
        ce_bri,
        endereco_fabrica
    FROM registros
    WHERE {" AND ".join(where)}
    ORDER BY id DESC
    LIMIT 1
    """

    resultado = pd.read_sql_query(sql, conn, params=tuple(params))

    if resultado.empty:
        return None

    return resultado.iloc[0].to_dict()


@st.cache_data(ttl=30, show_spinner=False)
def listar_fabricas_existentes_para_sistema5(cliente_base=""):
    cliente_base = clean(cliente_base).upper()

    if cliente_base:
        return pd.read_sql_query("""
        SELECT DISTINCT
            fabrica,
            ce_bri,
            endereco_fabrica
        FROM registros
        WHERE UPPER(cliente_base) = UPPER(?)
        AND fabrica IS NOT NULL
        AND fabrica != ''
        ORDER BY fabrica
        """, conn, params=(cliente_base,))

    return pd.read_sql_query("""
    SELECT DISTINCT
        cliente_base,
        fabrica,
        ce_bri,
        endereco_fabrica
    FROM registros
    WHERE fabrica IS NOT NULL
    AND fabrica != ''
    ORDER BY cliente_base, fabrica
    """, conn)


def extrair_codigo_fabrica_nome(nome):
    texto = clean(nome).upper()

    # Remove a parte do IP do nome antes de procurar a fábrica
    # Ex: "INCLUSÃO 01-03-26 F01 IP-0707-26" → remove "IP-0707-26"
    texto_sem_ip = re.sub(r"\bIP[-\s]*[\d\-]+", "", texto)

    # Procura padrão de fábrica: F01, F02, F016, etc — exige espaço ou início antes do F
    match = re.search(r"(?<![A-Z\d])F\s*0*(\d{1,3})(?!\d)", texto_sem_ip)
    if match:
        numero = int(match.group(1))
        return f"F{numero:02d}"
    return ""


def normalizar_familia_aba_s5(nome_aba):
    texto = clean(nome_aba).upper()
    numeros = re.findall(r"\d+", texto)
    if numeros:
        return str(int(numeros[0]))
    return texto


@st.cache_data(ttl=30, show_spinner=False)
def buscar_fabrica_por_codigo_s5(cliente_base, codigo_fabrica):
    cliente_base = clean(cliente_base).upper()
    codigo_fabrica = clean(codigo_fabrica).upper()

    if not cliente_base or not codigo_fabrica:
        return None

    numero = re.sub(r"\D", "", codigo_fabrica)
    if numero:
        numero_int = str(int(numero))
        codigo_padrao = f"F{int(numero):02d}"
    else:
        numero_int = ""
        codigo_padrao = codigo_fabrica

    resultado = pd.read_sql_query("""
    SELECT
        cliente_base,
        fabrica,
        ce_bri,
        endereco_fabrica
    FROM registros
    WHERE UPPER(cliente_base) = UPPER(?)
    AND (
        UPPER(fabrica) LIKE ?
        OR UPPER(fabrica) LIKE ?
        OR UPPER(fabrica) LIKE ?
    )
    ORDER BY id DESC
    LIMIT 1
    """, conn, params=(
        cliente_base,
        f"%{codigo_padrao}%",
        f"%F{numero_int}%" if numero_int else f"%{codigo_fabrica}%",
        f"%FÁBRICA {numero_int}%" if numero_int else f"%{codigo_fabrica}%"
    ))

    if resultado.empty:
        return None

    return resultado.iloc[0].to_dict()


@st.cache_data(ttl=30, show_spinner=False)
def listar_processos_sistema5_cached():
    return pd.read_sql_query("""
    SELECT
        a.id AS arquivo_id,
        c.categoria,
        c.cliente_base,
        f.fabrica,
        f.ce_bri,
        f.endereco_fabrica,
        a.tipo_processo,
        a.ip_processo,
        a.data_processo,
        a.arquivo_nome,
        COUNT(i.id) AS qtd_itens
    FROM sistema5_arquivos a
    INNER JOIN sistema5_clientes c ON c.id = a.cliente_id
    INNER JOIN sistema5_fabricas f ON f.id = a.fabrica_id
    LEFT JOIN sistema5_itens i ON i.arquivo_id = a.id
    GROUP BY
        a.id, c.categoria, c.cliente_base, f.fabrica,
        f.ce_bri, f.endereco_fabrica, a.tipo_processo,
        a.ip_processo, a.data_processo, a.arquivo_nome
    ORDER BY c.categoria, c.cliente_base, f.fabrica, a.data_processo DESC, a.id DESC
    """, conn)


@st.cache_data(ttl=20, show_spinner=False)
def listar_processos_sistema5():
    return pd.read_sql_query("""
    SELECT
        a.id AS arquivo_id,
        c.categoria,
        c.cliente_base,
        f.fabrica,
        f.ce_bri,
        f.endereco_fabrica,
        a.tipo_processo,
        a.ip_processo,
        a.data_processo,
        a.arquivo_nome,
        COUNT(i.id) AS qtd_itens
    FROM sistema5_arquivos a
    INNER JOIN sistema5_clientes c
    ON c.id = a.cliente_id
    INNER JOIN sistema5_fabricas f
    ON f.id = a.fabrica_id
    LEFT JOIN sistema5_itens i
    ON i.arquivo_id = a.id
    GROUP BY
        a.id,
        c.categoria,
        c.cliente_base,
        f.fabrica,
        f.ce_bri,
        f.endereco_fabrica,
        a.tipo_processo,
        a.ip_processo,
        a.data_processo,
        a.arquivo_nome
    ORDER BY
        c.categoria,
        c.cliente_base,
        f.fabrica,
        a.data_processo DESC,
        a.id DESC
    """, conn)


@st.cache_data(ttl=30, show_spinner=False)
def buscar_itens_processo_sistema5(arquivo_id, termo=""):
    termo = clean(termo)

    if termo:
        return pd.read_sql_query("""
        SELECT
            familia,
            item,
            marca,
            modelo,
            nome,
            codigo,
            ce_bri,
            endereco_fabrica,
            tipo_processo,
            ip_processo,
            data_processo,
            arquivo_nome
        FROM sistema5_itens
        WHERE arquivo_id = ?
        AND (
            UPPER(modelo) LIKE UPPER(?)
            OR UPPER(marca) LIKE UPPER(?)
            OR UPPER(nome) LIKE UPPER(?)
            OR codigo LIKE ?
        )
        ORDER BY CAST(familia AS INTEGER), CAST(item AS INTEGER), modelo
        """, conn, params=(
            arquivo_id,
            f"{termo}%",
            f"%{termo}%",
            f"%{termo}%",
            f"%{termo}%"
        ))

    return pd.read_sql_query("""
    SELECT
        familia,
        item,
        marca,
        modelo,
        nome,
        codigo,
        ce_bri,
        endereco_fabrica,
        tipo_processo,
        ip_processo,
        data_processo,
        arquivo_nome
    FROM sistema5_itens
    WHERE arquivo_id = ?
    ORDER BY CAST(familia AS INTEGER), CAST(item AS INTEGER), modelo
    """, conn, params=(arquivo_id,))


# Limpeza automática de itens antigos do Sistema 5 marcados como DESCONSIDERADO
try:
    limpar_itens_desconsiderados_sistema5()
except Exception:
    pass


@st.cache_data(ttl=30, show_spinner=False)
def pesquisar_global_sistema5(termo):
    termo = clean(termo)

    if not termo:
        return pd.DataFrame()

    return pd.read_sql_query("""
    SELECT
        cliente_base,
        categoria,
        fabrica,
        ce_bri,
        endereco_fabrica,
        tipo_processo,
        ip_processo,
        data_processo,
        familia,
        item,
        marca,
        modelo,
        nome,
        codigo,
        arquivo_nome,
        data_upload
    FROM sistema5_itens
    WHERE
        UPPER(modelo) LIKE UPPER(?)
        OR codigo LIKE ?
        OR UPPER(marca) LIKE UPPER(?)
        OR UPPER(nome) LIKE UPPER(?)
        OR UPPER(ip_processo) LIKE UPPER(?)
        OR UPPER(ce_bri) LIKE UPPER(?)
        OR UPPER(fabrica) LIKE UPPER(?)
        OR UPPER(cliente_base) LIKE UPPER(?)
        OR UPPER(arquivo_nome) LIKE UPPER(?)
    ORDER BY
        COALESCE(
            data_processo,
            substr(data_upload,7,4) || '-' || substr(data_upload,4,2) || '-' || substr(data_upload,1,2),
            '1900-01-01'
        ) DESC,
        id DESC
    """, conn, params=(
        f"{termo}%",
        f"%{termo}%",
        f"%{termo}%",
        f"%{termo}%",
        f"%{termo}%",
        f"%{termo}%",
        f"%{termo}%",
        f"%{termo}%",
        f"%{termo}%"
    ))



# ==========================================
# PAINEL DE COBERTURA
# ==========================================

def cobertura_resumo_geral():
    dados = {}

    try:
        dados["clientes_registros"] = cursor.execute("""
        SELECT COUNT(DISTINCT cliente_base)
        FROM registros
        WHERE cliente_base IS NOT NULL AND cliente_base != ''
        """).fetchone()[0]
    except Exception:
        dados["clientes_registros"] = 0

    try:
        dados["fabricas_registros"] = cursor.execute("""
        SELECT COUNT(DISTINCT cliente_base || '|' || fabrica)
        FROM registros
        WHERE fabrica IS NOT NULL AND fabrica != ''
        """).fetchone()[0]
    except Exception:
        dados["fabricas_registros"] = 0

    try:
        dados["familias_registros"] = cursor.execute("""
        SELECT COUNT(DISTINCT ce_bri || '|' || familia)
        FROM registros
        WHERE ce_bri IS NOT NULL AND ce_bri != ''
        AND familia IS NOT NULL AND familia != ''
        """).fetchone()[0]
    except Exception:
        dados["familias_registros"] = 0

    try:
        dados["familias_sistema5"] = cursor.execute("""
        SELECT COUNT(DISTINCT ce_bri || '|' || familia)
        FROM sistema5_itens
        WHERE ce_bri IS NOT NULL AND ce_bri != ''
        AND familia IS NOT NULL AND familia != ''
        """).fetchone()[0]
    except Exception:
        dados["familias_sistema5"] = 0

    try:
        dados["itens_sistema5"] = cursor.execute("""
        SELECT COUNT(*)
        FROM sistema5_itens
        """).fetchone()[0]
    except Exception:
        dados["itens_sistema5"] = 0

    try:
        dados["pendentes_ip_bri"] = len(listar_ip_bri_pendentes_sistema5())
    except Exception:
        dados["pendentes_ip_bri"] = 0

    return dados


def cobertura_fabricas():
    return pd.read_sql_query("""
    WITH chaves AS (
        SELECT DISTINCT cliente_base, fabrica, ce_bri
        FROM registros
        WHERE fabrica IS NOT NULL AND fabrica != ''

        UNION

        SELECT DISTINCT cliente_base, fabrica, ce_bri
        FROM sistema5_itens
        WHERE fabrica IS NOT NULL AND fabrica != ''
    ),
    familias_reg AS (
        SELECT
            cliente_base,
            fabrica,
            ce_bri,
            COUNT(DISTINCT familia) AS familias_registros
        FROM registros
        WHERE fabrica IS NOT NULL AND fabrica != ''
        GROUP BY cliente_base, fabrica, ce_bri
    ),
    familias_s5 AS (
        SELECT
            cliente_base,
            fabrica,
            ce_bri,
            COUNT(DISTINCT familia) AS familias_sistema5,
            COUNT(*) AS itens_sistema5
        FROM sistema5_itens
        WHERE fabrica IS NOT NULL AND fabrica != ''
        GROUP BY cliente_base, fabrica, ce_bri
    ),
    pendentes AS (
        SELECT
            s.cliente_base,
            s.fabrica,
            s.ce_bri,
            COUNT(DISTINCT s.familia) AS familias_pendentes_ip_bri
        FROM sistema5_itens s
        LEFT JOIN certificados c
        ON UPPER(c.ce_bri) = UPPER(s.ce_bri)
        AND c.familia = s.familia
        LEFT JOIN ip_bri_familias m
        ON UPPER(m.ce_bri) = UPPER(s.ce_bri)
        AND m.familia = s.familia
        WHERE s.ce_bri IS NOT NULL
        AND s.ce_bri != ''
        AND s.familia IS NOT NULL
        AND s.familia != ''
        AND c.ip_bri IS NULL
        AND m.ip_bri IS NULL
        GROUP BY s.cliente_base, s.fabrica, s.ce_bri
    )
    SELECT
        chaves.cliente_base,
        chaves.fabrica,
        chaves.ce_bri,
        COALESCE(r.familias_registros, 0) AS familias_registros,
        COALESCE(s.familias_sistema5, 0) AS familias_sistema5,
        COALESCE(s.itens_sistema5, 0) AS itens_sistema5,
        COALESCE(p.familias_pendentes_ip_bri, 0) AS familias_pendentes_ip_bri
    FROM chaves
    LEFT JOIN familias_reg r
    ON UPPER(r.cliente_base) = UPPER(chaves.cliente_base)
    AND UPPER(r.fabrica) = UPPER(chaves.fabrica)
    AND UPPER(r.ce_bri) = UPPER(chaves.ce_bri)
    LEFT JOIN familias_s5 s
    ON UPPER(s.cliente_base) = UPPER(chaves.cliente_base)
    AND UPPER(s.fabrica) = UPPER(chaves.fabrica)
    AND UPPER(s.ce_bri) = UPPER(chaves.ce_bri)
    LEFT JOIN pendentes p
    ON UPPER(p.cliente_base) = UPPER(chaves.cliente_base)
    AND UPPER(p.fabrica) = UPPER(chaves.fabrica)
    AND UPPER(p.ce_bri) = UPPER(chaves.ce_bri)
    ORDER BY chaves.cliente_base, chaves.fabrica
    """, conn)

def cobertura_familias_detalhada(cliente_base="", fabrica="", ce_bri=""):
    where = []
    params = []

    if clean(cliente_base):
        where.append("UPPER(base.cliente_base) = UPPER(?)")
        params.append(clean(cliente_base))

    if clean(fabrica):
        where.append("UPPER(base.fabrica) = UPPER(?)")
        params.append(clean(fabrica))

    if clean(ce_bri):
        where.append("UPPER(base.ce_bri) = UPPER(?)")
        params.append(clean(ce_bri))

    filtro = ""
    if where:
        filtro = "WHERE " + " AND ".join(where)

    sql = f"""
    WITH base AS (
        SELECT
            cliente_base,
            fabrica,
            ce_bri,
            familia
        FROM registros
        WHERE familia IS NOT NULL AND familia != ''

        UNION

        SELECT
            cliente_base,
            fabrica,
            ce_bri,
            familia
        FROM sistema5_itens
        WHERE familia IS NOT NULL AND familia != ''
    ),
    itens_s5 AS (
        SELECT
            cliente_base,
            fabrica,
            ce_bri,
            familia,
            COUNT(*) AS qtd_itens_sistema5,
            GROUP_CONCAT(DISTINCT ip_processo) AS ips_processos
        FROM sistema5_itens
        GROUP BY cliente_base, fabrica, ce_bri, familia
    ),
    regs AS (
        SELECT
            cliente_base,
            fabrica,
            ce_bri,
            familia,
            MAX(registro) AS registro
        FROM registros
        GROUP BY cliente_base, fabrica, ce_bri, familia
    ),
    certs AS (
        SELECT
            ce_bri,
            familia,
            MAX(ip_bri) AS ip_bri_oficial,
            MAX(rev) AS rev
        FROM certificados
        GROUP BY ce_bri, familia
    ),
    manuais AS (
        SELECT
            ce_bri,
            familia,
            MAX(ip_bri) AS ip_bri_manual
        FROM ip_bri_familias
        GROUP BY ce_bri, familia
    )
    SELECT
        base.cliente_base,
        base.fabrica,
        base.ce_bri,
        base.familia,
        COALESCE(itens_s5.qtd_itens_sistema5, 0) AS qtd_itens_sistema5,
        COALESCE(regs.registro, '') AS registro,
        COALESCE(certs.ip_bri_oficial, '') AS ip_bri_oficial,
        COALESCE(manuais.ip_bri_manual, '') AS ip_bri_manual,
        CASE
            WHEN certs.ip_bri_oficial IS NOT NULL AND certs.ip_bri_oficial != '' THEN 'OFICIAL'
            WHEN manuais.ip_bri_manual IS NOT NULL AND manuais.ip_bri_manual != '' THEN 'MANUAL'
            ELSE 'PENDENTE'
        END AS status_ip_bri,
        COALESCE(certs.rev, '') AS rev,
        COALESCE(itens_s5.ips_processos, '') AS ips_processos
    FROM base
    LEFT JOIN itens_s5
    ON UPPER(itens_s5.cliente_base) = UPPER(base.cliente_base)
    AND UPPER(itens_s5.fabrica) = UPPER(base.fabrica)
    AND UPPER(itens_s5.ce_bri) = UPPER(base.ce_bri)
    AND itens_s5.familia = base.familia
    LEFT JOIN regs
    ON UPPER(regs.cliente_base) = UPPER(base.cliente_base)
    AND UPPER(regs.fabrica) = UPPER(base.fabrica)
    AND UPPER(regs.ce_bri) = UPPER(base.ce_bri)
    AND regs.familia = base.familia
    LEFT JOIN certs
    ON UPPER(certs.ce_bri) = UPPER(base.ce_bri)
    AND certs.familia = base.familia
    LEFT JOIN manuais
    ON UPPER(manuais.ce_bri) = UPPER(base.ce_bri)
    AND manuais.familia = base.familia
    {filtro}
    ORDER BY base.cliente_base, base.fabrica, CAST(base.familia AS INTEGER)
    """

    return pd.read_sql_query(sql, conn, params=tuple(params))


def cobertura_itens_por_familia(cliente_base, fabrica, ce_bri, familia):
    return pd.read_sql_query("""
    SELECT
        familia,
        item,
        marca,
        modelo,
        nome,
        codigo,
        tipo_processo,
        ip_processo,
        data_processo,
        arquivo_nome
    FROM sistema5_itens
    WHERE UPPER(cliente_base) = UPPER(?)
    AND UPPER(fabrica) = UPPER(?)
    AND UPPER(ce_bri) = UPPER(?)
    AND familia = ?
    ORDER BY CAST(item AS INTEGER), modelo
    """, conn, params=(
        clean(cliente_base),
        clean(fabrica),
        clean(ce_bri),
        normalizar_familia(familia)
    ))


def diagnosticar_referencia_confirmacao(ref):
    ref = clean(ref)

    if not ref:
        return {
            "status": "vazio",
            "item_confirmacao": None,
            "candidatos_sistema5": pd.DataFrame()
        }

    item = buscar_item_confirmacao(ref)

    candidatos = pd.read_sql_query("""
    SELECT
        cliente_base,
        categoria,
        fabrica,
        ce_bri,
        endereco_fabrica,
        tipo_processo,
        ip_processo,
        data_processo,
        familia,
        item,
        marca,
        modelo,
        nome,
        codigo,
        arquivo_nome,
        data_upload
    FROM sistema5_itens
    WHERE
        UPPER(modelo) LIKE UPPER(?)
        OR UPPER(modelo) LIKE UPPER(?)
        OR UPPER(modelo) LIKE UPPER(?)
    ORDER BY COALESCE(
        data_processo,
        substr(data_upload,7,4) || '-' || substr(data_upload,4,2) || '-' || substr(data_upload,1,2),
        '1900-01-01'
    ) DESC, id DESC
    LIMIT 50
    """, conn, params=(
        f"{ref}%",
        f"%{ref}%",
        f"%{normalizar_hifens_ref(ref)}%"
    ))

    if item:
        status = "encontrado_pela_confirmacao"
    elif not candidatos.empty:
        status = "existe_no_sistema5_mas_filtro_confirmacao_nao_aceitou"
    else:
        status = "nao_encontrado"

    return {
        "status": status,
        "item_confirmacao": item,
        "candidatos_sistema5": candidatos
    }



# ==========================================
# AUDITORIA / CORREÇÃO DE CÓDIGOS DE BARRAS
# ==========================================

@st.cache_data(ttl=20, show_spinner=False)
def listar_codigos_suspeitos_sistema5():
    """
    Lista códigos do Sistema 5 que podem ter perdido zero à esquerda.
    Principal suspeita: códigos com menos de 13 dígitos.
    """
    return pd.read_sql_query("""
    SELECT
        id,
        cliente_base,
        fabrica,
        ce_bri,
        familia,
        item,
        marca,
        modelo,
        nome,
        codigo,
        LENGTH(codigo) AS tamanho_codigo,
        ip_processo,
        arquivo_nome,
        data_upload
    FROM sistema5_itens
    WHERE codigo IS NOT NULL
    AND codigo != ''
    AND (
        LENGTH(codigo) < 13
        OR codigo LIKE '%.0'
    )
    ORDER BY cliente_base, fabrica, CAST(familia AS INTEGER), modelo
    """, conn)


@st.cache_data(ttl=20, show_spinner=False)
def listar_codigos_suspeitos_certificados():
    """
    Lista códigos dos certificados oficiais que podem ter perdido zero à esquerda.
    """
    return pd.read_sql_query("""
    SELECT
        i.id,
        c.ip_bri,
        c.ce_bri,
        c.familia,
        i.ordem,
        i.marca,
        i.modelo,
        i.nome,
        i.codigo,
        LENGTH(i.codigo) AS tamanho_codigo,
        c.arquivo_pdf,
        c.data_cadastro
    FROM itens i
    INNER JOIN certificados c
    ON c.id = i.certificado_id
    WHERE i.codigo IS NOT NULL
    AND i.codigo != ''
    AND (
        LENGTH(i.codigo) < 13
        OR i.codigo LIKE '%.0'
    )
    ORDER BY c.ip_bri, i.ordem
    """, conn)


def atualizar_codigo_sistema5_por_id(item_id, novo_codigo):
    item_id = int(item_id)
    codigo = codigo_barras_para_texto(novo_codigo)

    if not codigo:
        return False, "Código inválido. Informe um código com 8 a 14 dígitos."

    cursor.execute("""
    UPDATE sistema5_itens
    SET codigo = ?
    WHERE id = ?
    """, (codigo, item_id))

    commit_seguro()
    return True, f"Código atualizado no Sistema 5: {codigo}"


def atualizar_codigo_certificado_por_id(item_id, novo_codigo):
    item_id = int(item_id)
    codigo = codigo_barras_para_texto(novo_codigo)

    if not codigo:
        return False, "Código inválido. Informe um código com 8 a 14 dígitos."

    cursor.execute("""
    UPDATE itens
    SET codigo = ?
    WHERE id = ?
    """, (codigo, item_id))

    commit_seguro()
    return True, f"Código atualizado no banco de certificados: {codigo}"


def aplicar_zero_esquerda_sistema5_por_id(item_id, tamanho_final=13):
    item_id = int(item_id)
    tamanho_final = int(tamanho_final)

    atual = cursor.execute("""
    SELECT codigo
    FROM sistema5_itens
    WHERE id = ?
    """, (item_id,)).fetchone()

    if not atual:
        return False, "Item não encontrado."

    codigo_atual = re.sub(r"\D", "", clean(atual[0]))

    if not codigo_atual:
        return False, "Código atual inválido."

    novo_codigo = codigo_atual.zfill(tamanho_final)

    cursor.execute("""
    UPDATE sistema5_itens
    SET codigo = ?
    WHERE id = ?
    """, (novo_codigo, item_id))

    commit_seguro()
    return True, f"Código corrigido: {codigo_atual} → {novo_codigo}"


def aplicar_zero_esquerda_certificado_por_id(item_id, tamanho_final=13):
    item_id = int(item_id)
    tamanho_final = int(tamanho_final)

    atual = cursor.execute("""
    SELECT codigo
    FROM itens
    WHERE id = ?
    """, (item_id,)).fetchone()

    if not atual:
        return False, "Item não encontrado."

    codigo_atual = re.sub(r"\D", "", clean(atual[0]))

    if not codigo_atual:
        return False, "Código atual inválido."

    novo_codigo = codigo_atual.zfill(tamanho_final)

    cursor.execute("""
    UPDATE itens
    SET codigo = ?
    WHERE id = ?
    """, (novo_codigo, item_id))

    commit_seguro()
    return True, f"Código corrigido: {codigo_atual} → {novo_codigo}"

# ==========================================
# TABS
# ==========================================

aviso_backup_diario()

# ==========================================
# NAVEGAÇÃO SIDEBAR
# ==========================================
def _nav_btn(label, pagina, icon=""):
    ativo = st.session_state.get("pagina_ativa") == pagina
    prefix = "▶ " if ativo else "   "
    if st.sidebar.button(f"{prefix}{icon} {label}".strip(), key=f"nav_{pagina}", use_container_width=True):
        st.session_state["pagina_ativa"] = pagina
        st.rerun()

st.sidebar.markdown("---")
st.sidebar.markdown("<div style='color:rgba(255,255,255,0.35);font-size:10px;font-weight:700;letter-spacing:0.1em;text-transform:uppercase;padding:0 4px;margin-bottom:4px'>Principal</div>", unsafe_allow_html=True)
_nav_btn("Início",          "inicio",       "🏠")
_nav_btn("PDF → XML",       "pdf_xml",      "📄")
_nav_btn("XML Rápido",      "xml_rapido",   "⚡")
_nav_btn("Certificados",    "certificados", "🗄️")
_nav_btn("Registros",       "registros",    "📋")

st.sidebar.markdown("<div style='color:rgba(255,255,255,0.35);font-size:10px;font-weight:700;letter-spacing:0.1em;text-transform:uppercase;padding:0 4px;margin:8px 0 4px'>Operações</div>", unsafe_allow_html=True)
_nav_btn("Confirmação",     "confirmacao",  "✅")
_nav_btn("Pesquisa",        "pesquisa",     "🔍")
_nav_btn("Sistema 5",       "sistema5",     "📦")
_nav_btn("Gerar Etiquetas",  "gerar_etiquetas",  "🎨")
_nav_btn("Gerar Descrições", "gerar_descricoes", "✍️")

st.sidebar.markdown("<div style='color:rgba(255,255,255,0.35);font-size:10px;font-weight:700;letter-spacing:0.1em;text-transform:uppercase;padding:0 4px;margin:8px 0 4px'>Análise</div>", unsafe_allow_html=True)
_nav_btn("IP-BRI Pendentes","pendentes",    "⏳")
_nav_btn("Cobertura",       "cobertura",    "📊")
_nav_btn("Correção Códigos","correcao",     "🔧")
_nav_btn("Etiquetas",       "etiquetas",    "🏷️")

st.sidebar.markdown("<div style='color:rgba(255,255,255,0.35);font-size:10px;font-weight:700;letter-spacing:0.1em;text-transform:uppercase;padding:0 4px;margin:8px 0 4px'>Ajuda</div>", unsafe_allow_html=True)
_nav_btn("Funcionalidades", "funcionalidades", "📘")
st.sidebar.markdown("---")


try:
    st.sidebar.header("Backup Geral")
    st.sidebar.caption("Backup centralizado: baixar e restaurar tudo pelo mesmo local.")

    if st.sidebar.button("⬇️ GERAR E BAIXAR BACKUP ZIP", key="btn_gerar_backup_sidebar"):
        st.session_state["backup_zip_bytes"] = gerar_backup_geral_zip()
    if st.session_state.get("backup_zip_bytes"):
        st.sidebar.download_button(
            "💾 Baixar backup gerado",
            st.session_state["backup_zip_bytes"],
            "backup_geral_c_xml_br_engine.zip",
            "application/zip",
            key="backup_geral_sidebar"
        )

    st.sidebar.divider()
    st.sidebar.subheader("Restaurar backup")

    backup_geral_upload = st.sidebar.file_uploader(
        "Enviar certificados.db",
        type=["db"],
        key="backup_geral_upload_sidebar"
    )

    if backup_geral_upload:
        st.sidebar.warning("⚠️ Isso irá **substituir completamente** o banco atual.")
        confirmar_sidebar = st.sidebar.checkbox(
            "Confirmo que quero substituir o banco",
            key="confirmar_restore_sidebar"
        )
        if confirmar_sidebar:
            with open(DB_PATH, "wb") as f:
                f.write(backup_geral_upload.read())

            # Reconecta imediatamente ao banco novo — sem depender do rerun
            try:
                conn.close()
            except Exception:
                pass
            conn = _abrir_conexao(DB_PATH)
            cursor = conn.cursor()

            st.session_state["db_version"] = st.session_state.get("db_version", 0) + 1
            st.cache_data.clear()
            st.cache_resource.clear()
            st.sidebar.success("✅ Backup restaurado! Recarregando...")
            st.rerun()

except Exception as e:
    st.sidebar.warning(f"Backup geral indisponível: {e}")

# Navegação por session_state
if "pagina_ativa" not in st.session_state:
    st.session_state["pagina_ativa"] = "inicio"

_pagina = st.session_state["pagina_ativa"]

class _FakePage:
    def __init__(self, nome): self.nome = nome
    def __enter__(self): return self
    def __exit__(self, *a): pass

def _is_active(nome):
    return _pagina == nome

aba0 = _FakePage("inicio")
aba1 = _FakePage("pdf_xml")
aba1b = _FakePage("xml_rapido")
aba2 = _FakePage("certificados")
aba3 = _FakePage("registros")
aba4 = _FakePage("confirmacao")
aba5 = _FakePage("pesquisa")
aba6 = _FakePage("sistema5")
aba7 = _FakePage("pendentes")
aba8 = _FakePage("cobertura")
aba9 = _FakePage("correcao")

# ==========================================
# ABA 0 - INÍCIO / COMO USAR
# ==========================================

if _is_active("inicio"):
    # Busca métricas
    try:
        _m = metricas_banco()
        total_certs_home = _m.get("total_certs", 0)
        total_itens_home = _m.get("total_itens", 0)
        total_regs_home  = _m.get("total_regs", 0)
        total_s5_home    = _m.get("total_s5", 0)
        banco_vazio = total_certs_home == 0
    except Exception:
        total_certs_home = total_itens_home = total_regs_home = total_s5_home = 0
        banco_vazio = True

    status_html = (
        '<div class="status-badge empty">⚠️ Banco vazio — siga o fluxo abaixo para começar</div>'
        if banco_vazio else
        '<div class="status-badge ok">✅ Banco com dados. Pronto para uso</div>'
    )

    import streamlit.components.v1 as components
    components.html(f"""
<!DOCTYPE html>
<html>
<head>
<meta charset="utf-8">
<link href="https://fonts.googleapis.com/css2?family=Inter:wght@400;500;600;700&display=swap" rel="stylesheet">
<style>
* {{ box-sizing: border-box; margin: 0; padding: 0; }}
body {{ font-family: 'Inter', sans-serif; background: #F8FAFC; color: #1E293B; padding: 0; }}

/* Header */
.header {{ margin-bottom: 28px; }}
.header h1 {{ font-size: 28px; font-weight: 700; color: #0F172A; letter-spacing: -0.5px; margin-bottom: 4px; display: flex; align-items: center; gap: 10px; }}
.header p {{ font-size: 13px; color: #94A3B8; font-weight: 400; }}
.logo-dot {{ width: 10px; height: 10px; background: #E94560; border-radius: 50%; display: inline-block; }}

/* Métricas */
.metrics {{ display: grid; grid-template-columns: repeat(4, 1fr); gap: 14px; margin-bottom: 14px; }}
.metric-card {{ background: #fff; border: 1px solid #E2E8F0; border-radius: 12px; padding: 20px 22px; box-shadow: 0 1px 4px rgba(0,0,0,0.04); transition: box-shadow 0.15s, transform 0.15s; }}
.metric-card:hover {{ box-shadow: 0 6px 18px rgba(0,0,0,0.08); transform: translateY(-2px); }}
.metric-label {{ font-size: 11px; font-weight: 600; color: #94A3B8; text-transform: uppercase; letter-spacing: 0.07em; margin-bottom: 8px; }}
.metric-value {{ font-size: 34px; font-weight: 700; color: #0F172A; letter-spacing: -1px; line-height: 1; }}
.metric-card.red .metric-value {{ color: #E94560; }}

/* Status badge */
.status-badge {{ display: inline-flex; align-items: center; gap: 8px; font-size: 13px; font-weight: 500; padding: 10px 16px; border-radius: 8px; margin-bottom: 28px; }}
.status-badge.ok {{ background: rgba(16,185,129,0.08); color: #059669; border: 1px solid rgba(16,185,129,0.2); }}
.status-badge.empty {{ background: rgba(245,158,11,0.08); color: #D97706; border: 1px solid rgba(245,158,11,0.2); }}

/* Seção */
.section-title {{ font-size: 13px; font-weight: 600; color: #64748B; text-transform: uppercase; letter-spacing: 0.08em; margin-bottom: 14px; display: flex; align-items: center; gap: 8px; }}
.section-title::after {{ content: ''; flex: 1; height: 1px; background: #E2E8F0; }}

/* Fluxo */
.flow {{ display: grid; grid-template-columns: 1fr 32px 1fr 32px 1fr 32px 1fr; gap: 0; align-items: start; margin-bottom: 28px; }}
.flow-card {{ background: #fff; border: 1px solid #E2E8F0; border-radius: 12px; padding: 18px; box-shadow: 0 1px 4px rgba(0,0,0,0.04); }}
.flow-step {{ font-size: 10px; font-weight: 700; color: #E94560; text-transform: uppercase; letter-spacing: 0.1em; margin-bottom: 10px; }}
.flow-icon {{ width: 36px; height: 36px; background: rgba(233,69,96,0.08); border-radius: 8px; display: flex; align-items: center; justify-content: center; font-size: 18px; margin-bottom: 10px; }}
.flow-name {{ font-size: 14px; font-weight: 600; color: #0F172A; margin-bottom: 6px; }}
.flow-desc {{ font-size: 12px; color: #64748B; line-height: 1.5; margin-bottom: 10px; }}
.flow-link {{ font-size: 12px; color: #E94560; font-weight: 500; display: flex; align-items: center; gap: 4px; }}
.flow-arrow {{ display: flex; align-items: center; justify-content: center; height: 100%; padding-top: 50px; color: #CBD5E1; font-size: 20px; }}

/* Referência rápida */
.ref-grid {{ display: grid; grid-template-columns: repeat(3, 1fr); gap: 12px; }}
.ref-card {{ background: #fff; border: 1px solid #E2E8F0; border-radius: 10px; padding: 16px; }}
.ref-card h4 {{ font-size: 13px; font-weight: 600; color: #0F172A; margin-bottom: 8px; display: flex; align-items: center; gap: 6px; }}
.ref-card ul {{ padding-left: 14px; }}
.ref-card ul li {{ font-size: 12px; color: #64748B; line-height: 1.6; }}
</style>
</head>
<body>

<div class="header">
  <h1><span class="logo-dot"></span> C XML BR Engine</h1>
  <p>Sistema de gestão de certificados IP-BRI, geração de XML e preenchimento automático de confirmações.</p>
</div>

<div class="section-title">Status do banco</div>

<div class="metrics">
  <div class="metric-card">
    <div class="metric-label">Certificados</div>
    <div class="metric-value">{total_certs_home:,}</div>
  </div>
  <div class="metric-card">
    <div class="metric-label">Itens</div>
    <div class="metric-value">{total_itens_home:,}</div>
  </div>
  <div class="metric-card">
    <div class="metric-label">Registros</div>
    <div class="metric-value">{total_regs_home:,}</div>
  </div>
  <div class="metric-card {'red' if total_s5_home > 0 else ''}">
    <div class="metric-label">Itens Sistema 5</div>
    <div class="metric-value">{total_s5_home:,}</div>
  </div>
</div>

{status_html}

<div class="section-title">Referência rápida</div>

<div class="ref-grid">
  <div class="ref-card">
    <h4>📄 PDF → XML</h4>
    <ul>
      <li>Lê certificados e extrai itens</li>
      <li>Gera XML ISO-8859-1 automaticamente</li>
      <li>Salva no banco com histórico</li>
      <li>Ignora certificados com REV menor</li>
    </ul>
  </div>
  <div class="ref-card">
    <h4>🗄️ Banco de Certificados</h4>
    <ul>
      <li>Visualiza e busca certificados</li>
      <li>Pesquisa por referência ou código</li>
      <li>Histórico de alterações por IP-BRI</li>
      <li>Exporta banco em CSV</li>
    </ul>
  </div>
  <div class="ref-card">
    <h4>📋 Registros</h4>
    <ul>
      <li>Importa Excel de fábricas e CE-BRI</li>
      <li>Cada aba = uma fábrica</li>
      <li>Edita endereços sem reimportar</li>
      <li>Sincroniza com Sistema 5</li>
    </ul>
  </div>
  <div class="ref-card">
    <h4>✅ Confirmação</h4>
    <ul>
      <li>Células verdes = referências</li>
      <li>Células amarelas = campos</li>
      <li>Busca Sistema 5 primeiro</li>
      <li>Salva direto no PC</li>
    </ul>
  </div>
  <div class="ref-card">
    <h4>🔍 Pesquisa Avançada</h4>
    <ul>
      <li>IP-BRI → registro e fábrica</li>
      <li>Registro → IP-BRI</li>
      <li>Rápida verificação de vínculos</li>
    </ul>
  </div>
  <div class="ref-card">
    <h4>📦 Sistema 5</h4>
    <ul>
      <li>Processos sem certificado oficial</li>
      <li>Inclusões, manutenções, recertificações</li>
      <li>Integrado ao preenchimento</li>
      <li>Upload ZIP com vários processos</li>
    </ul>
  </div>
</div>

</body>
</html>
""", height=820, scrolling=False)

    st.divider()

    st.subheader("💾 Backup")
    st.markdown("""
O banco de dados fica salvo no servidor do Streamlit. **Se o app reiniciar ou ficar inativo por muito tempo, os dados podem ser perdidos.**

**Boas práticas:**
- Baixe o backup geral pelo menu lateral sempre que terminar uma sessão de cadastro
- Guarde o arquivo `.db` em um local seguro
- Para restaurar, use o menu lateral → "Restaurar backup" e envie o arquivo `.db`
""")
    st.warning("⚠️ O Streamlit Cloud pode reiniciar o app a qualquer momento. Faça backup com frequência.")

if _is_active("pdf_xml"):
    st.title("PDF → XML")

    uploaded_file = st.file_uploader(
        "Envie o PDF",
        type="pdf",
        key="xml_pdf"
    )

    if uploaded_file:
        with tempfile.TemporaryDirectory() as tmpdir:
            pdf_path = Path(tmpdir) / uploaded_file.name
            pdf_path.write_bytes(uploaded_file.read())

            rows = parse_pdf(pdf_path)

            if not rows:
                st.error("Nenhum item encontrado")
                st.stop()

            dados_certificado = extrair_dados_certificado(pdf_path)

            df = pd.DataFrame(
                rows,
                columns=[
                    "ORDEM",
                    "MARCA",
                    "MODELO",
                    "NOME",
                    "CODIGO"
                ]
            )

            st.success("Itens extraídos")

            st.dataframe(
                df,
                use_container_width=True
            )

            st.info(
                f"""
IP-BRI: {dados_certificado['ip_bri']}

CE-BRI: {dados_certificado['ce_bri']}

FAMÍLIA: {dados_certificado.get('familia')}

REV: {dados_certificado['rev']}
"""
            )

            registro_vinculado = buscar_registro_por_certificado(
                dados_certificado.get("ce_bri"),
                dados_certificado.get("familia")
            )

            if registro_vinculado.empty:
                st.warning("Nenhum registro vinculado encontrado para este CE-BRI + FAMÍLIA.")
            else:
                st.success("Registro vinculado encontrado ✅")
                st.dataframe(normalizar_df_para_exibicao(registro_vinculado), use_container_width=True)

            duplicados = verificar_codigos_duplicados(df)

            if not duplicados.empty:
                st.warning("Foram encontrados códigos duplicados neste PDF.")
                st.dataframe(normalizar_df_para_exibicao(duplicados), use_container_width=True)

            if not dados_certificado["ip_bri"]:
                st.warning("⚠️ IP-BRI não encontrado neste PDF. O XML foi gerado, mas o certificado **não foi salvo no banco**.")

            status, mensagem = salvar_ou_atualizar_certificado(
                dados_certificado,
                rows,
                uploaded_file.name
            )

            if status == "novo":
                st.success(mensagem)
                st.cache_data.clear()
            elif status == "atualizado":
                st.warning(mensagem)
                st.cache_data.clear()
            elif status == "ignorado":
                st.info(mensagem)
            else:
                st.error(mensagem)

            xml_comma = gerar_xml(df, ",")
            xml_dot = gerar_xml(df, ".")
            xml_sem_caractere = gerar_xml(df, "")

            zip_path = Path(tmpdir) / "resultado_xml.zip"

            with zipfile.ZipFile(zip_path, "w") as zipf:
                zipf.writestr(
                    "xml_virgula.xml",
                    xml_comma.encode(
                        "ISO-8859-1",
                        errors="replace"
                    )
                )

                zipf.writestr(
                    "xml_ponto.xml",
                    xml_dot.encode(
                        "ISO-8859-1",
                        errors="replace"
                    )
                )

                zipf.writestr(
                    "xml_sem_caractere.xml",
                    xml_sem_caractere.encode(
                        "ISO-8859-1",
                        errors="replace"
                    )
                )

            with open(zip_path, "rb") as f:
                st.download_button(
                    "Baixar XMLs",
                    f,
                    "resultado_xml.zip"
                )

# ==========================================
# ABA XML RÁPIDO — só gera XML, não salva nada
# ==========================================

if _is_active("xml_rapido"):
    st.title("⚡ XML Rápido")
    st.caption("Gera XML a partir de certificados INNAC. Nada é salvo no banco. Aceita um ou vários PDFs de uma vez.")

    _xr_files = st.file_uploader(
        "Envie um ou mais PDFs de certificados INNAC",
        type="pdf",
        accept_multiple_files=True,
        key="xr_pdf"
    )

    if _xr_files:
        with tempfile.TemporaryDirectory() as _xr_tmpdir:
            # Guarda os ZIPs de cada PDF para o botão "Baixar Todos"
            _xr_todos_bytes = {}   # {nome_zip: bytes}
            _xr_resumo = []        # para a tabela de resumo

            for _xr_file in _xr_files:
                _xr_pdf_stem = Path(_xr_file.name).stem   # nome sem extensão
                _xr_nome_zip = f"{_xr_pdf_stem}.zip"

                _xr_path = Path(_xr_tmpdir) / _xr_file.name
                _xr_path.write_bytes(_xr_file.read())

                with st.spinner(f"Processando {_xr_file.name}…"):
                    _xr_header = _extrair_header_innac(_xr_path)
                    _xr_rows   = _parse_pdf_innac(_xr_path)

                # --- Resultado por PDF em expander ---
                _xr_ok = bool(_xr_rows)
                _xr_qtd_decl = _xr_header.get('qtd_produtos') or 0
                _xr_qtd_ext  = len(_xr_rows)
                _xr_exp_label = (
                    f"{'✅' if _xr_ok and _xr_qtd_ext == _xr_qtd_decl else '⚠️'} "
                    f"{_xr_file.name}  —  "
                    f"{_xr_header['ip_bri'] or '?'}  |  "
                    f"{_xr_qtd_ext} item(s)"
                )

                with st.expander(_xr_exp_label, expanded=(len(_xr_files) == 1)):
                    if not _xr_rows:
                        st.error("Nenhum item encontrado. Verifique se é um certificado INNAC válido.")
                        _xr_resumo.append({
                            "Arquivo": _xr_file.name,
                            "IP-BRI": _xr_header['ip_bri'] or "—",
                            "Itens extraídos": 0,
                            "Declarados": _xr_qtd_decl or "—",
                            "Status": "❌ Erro"
                        })
                        continue

                    # Métricas
                    _xr_c1, _xr_c2, _xr_c3, _xr_c4 = st.columns(4)
                    _xr_c1.metric("IP-BRI",       _xr_header['ip_bri']       or "—")
                    _xr_c2.metric("CE-BRI",       _xr_header['ce_bri']       or "—")
                    _xr_c3.metric("Data Emissão", _xr_header['data_emissao'] or "—")
                    _xr_c4.metric("Revisão",      str(_xr_header['rev']) if _xr_header['rev'] is not None else "—")

                    if _xr_qtd_decl and _xr_qtd_ext != _xr_qtd_decl:
                        st.warning(
                            f"O PDF declara **{_xr_qtd_decl} produtos**, mas foram extraídos **{_xr_qtd_ext}**."
                        )

                    _xr_df = pd.DataFrame(
                        _xr_rows,
                        columns=["ORDEM", "MARCA", "MODELO", "NOME", "CODIGO"]
                    )

                    # Duplicatas
                    _xr_dup = _xr_df[_xr_df.duplicated(subset=["CODIGO"], keep=False)]
                    if not _xr_dup.empty:
                        st.warning("Códigos de barras duplicados:")
                        st.dataframe(normalizar_df_para_exibicao(_xr_dup), use_container_width=True)

                    # Aviso descrições longas
                    _xr_longas = _xr_df[_xr_df["NOME"].str.len() > 200]
                    if not _xr_longas.empty:
                        st.info(
                            f"{len(_xr_longas)} descrição(ões) >200 chars — serão abreviadas no XML "
                            "(REV.EXT., DET.BORD., F.IND., F.REST., CST.INV., CST.IND., FIX.COMP. …)."
                        )

                    st.dataframe(
                        _xr_df[["ORDEM", "MARCA", "MODELO", "NOME", "CODIGO"]],
                        use_container_width=True
                    )

                    # Gera os 3 XMLs
                    _xr_xml_v = _gerar_xml_innac(_xr_df, ",")
                    _xr_xml_p = _gerar_xml_innac(_xr_df, ".")
                    _xr_xml_s = _gerar_xml_innac(_xr_df, "")

                    # Cria ZIP com nome do PDF
                    _xr_zip_buf = Path(_xr_tmpdir) / _xr_nome_zip
                    with zipfile.ZipFile(_xr_zip_buf, "w") as _xr_zipf:
                        _xr_zipf.writestr("xml_virgula.xml",       _xr_xml_v.encode("ISO-8859-1", errors="replace"))
                        _xr_zipf.writestr("xml_ponto.xml",         _xr_xml_p.encode("ISO-8859-1", errors="replace"))
                        _xr_zipf.writestr("xml_sem_caractere.xml", _xr_xml_s.encode("ISO-8859-1", errors="replace"))

                    _xr_zip_bytes = _xr_zip_buf.read_bytes()
                    _xr_todos_bytes[_xr_nome_zip] = _xr_zip_bytes

                    # Botão de download individual
                    st.download_button(
                        label=f"⬇️ Baixar {_xr_nome_zip}",
                        data=_xr_zip_bytes,
                        file_name=_xr_nome_zip,
                        mime="application/zip",
                        key=f"dl_{_xr_pdf_stem}"
                    )

                    _xr_resumo.append({
                        "Arquivo": _xr_file.name,
                        "IP-BRI": _xr_header['ip_bri'] or "—",
                        "Itens extraídos": _xr_qtd_ext,
                        "Declarados": _xr_qtd_decl or "—",
                        "Status": "✅ OK" if _xr_qtd_ext == _xr_qtd_decl else "⚠️ Divergência"
                    })

            # --- Tabela de resumo (só aparece com 2+ PDFs) ---
            if len(_xr_files) > 1 and _xr_resumo:
                st.divider()
                st.subheader("Resumo do lote")
                st.dataframe(pd.DataFrame(_xr_resumo), use_container_width=True, hide_index=True)

                # Botão "Baixar Todos" — ZIP contendo todos os ZIPs individuais
                if _xr_todos_bytes:
                    _xr_master_path = Path(_xr_tmpdir) / "todos_xml_innac.zip"
                    with zipfile.ZipFile(_xr_master_path, "w") as _xr_mzf:
                        for _xr_zname, _xr_zbytes in _xr_todos_bytes.items():
                            _xr_mzf.writestr(_xr_zname, _xr_zbytes)
                    st.download_button(
                        label="⬇️ Baixar Todos (ZIP com todos os ZIPs)",
                        data=_xr_master_path.read_bytes(),
                        file_name="todos_xml_innac.zip",
                        mime="application/zip",
                        key="dl_todos_xr"
                    )

# ==========================================
# ABA 2 - BANCO
# ==========================================

if _is_active("certificados"):
    st.title("Banco de Certificados")

    # --- Dashboard de métricas ---
    try:
        _m = metricas_banco()
        total_certs       = _m.get("total_certs", 0)
        total_itens_banco = _m.get("total_itens", 0)
        total_marcas      = _m.get("total_marcas", 0)
        ultimo_cert_str   = _m.get("ultimo_cert", "—")

        col_m1, col_m2, col_m3, col_m4 = st.columns(4)
        col_m1.metric("📄 Certificados", total_certs)
        col_m2.metric("📦 Itens", total_itens_banco)
        col_m3.metric("🏷️ Marcas distintas", total_marcas)
        col_m4.metric("🆕 Último cadastro", ultimo_cert_str)
        st.divider()
    except Exception:
        pass

    st.subheader("Backup do Banco")

    backup_upload = st.file_uploader(
        "Importar backup certificados.db",
        type=["db"],
        key="backup_db"
    )

    if backup_upload:
        st.warning("⚠️ Isso irá **substituir completamente** o banco atual. Esta ação não pode ser desfeita.")
        confirmar_backup_aba2 = st.checkbox("Confirmo que quero substituir o banco atual", key="confirmar_backup_aba2")
        if confirmar_backup_aba2:
            with open(DB_PATH, "wb") as f:
                f.write(backup_upload.read())
            try:
                conn.close()
            except Exception:
                pass
            conn = _abrir_conexao(DB_PATH)
            cursor = conn.cursor()
            st.session_state["db_version"] = st.session_state.get("db_version", 0) + 1
            st.cache_data.clear()
            st.cache_resource.clear()
            st.success("✅ Backup importado! Recarregando...")
            st.rerun()

    if Path(DB_PATH).exists():
        with open(DB_PATH, "rb") as f:
            _db_bytes = f.read()
        st.download_button(
            "Baixar backup atualizado",
            _db_bytes,
            "certificados.db",
            "application/octet-stream"
        )

    st.divider()

    st.subheader("Exportar banco completo")

    todos = exportar_todos_itens()

    if todos.empty:
        st.info("Nenhum item cadastrado ainda.")
    else:
        st.download_button(
            "Baixar banco completo",
            todos.to_csv(
                index=False,
                sep=";"
            ).encode(
                "ISO-8859-1",
                errors="replace"
            ),
            "banco_completo.csv",
            "text/csv"
        )

    st.divider()

    st.subheader("Enviar certificado para o banco")

    banco_file = st.file_uploader(
        "Envie um certificado PDF",
        type="pdf",
        key="banco_pdf"
    )

    if banco_file:
        with tempfile.TemporaryDirectory() as tmpdir:
            pdf_path = Path(tmpdir) / banco_file.name
            pdf_path.write_bytes(banco_file.read())

            dados_certificado = extrair_dados_certificado(pdf_path)
            rows = parse_pdf(pdf_path)

            col1, col2, col3, col4, col5, col6 = st.columns(6)

            col1.metric("IP-BRI", dados_certificado["ip_bri"] or "Não encontrado")
            col2.metric("CE-BRI", dados_certificado["ce_bri"] or "Não encontrado")
            col3.metric("FAM", dados_certificado.get("familia") or "Não encontrada")
            col4.metric("REV", dados_certificado["rev"])
            col5.metric("Produto", dados_certificado["produto"] or "Não encontrado")
            col6.metric("Emissão", dados_certificado["data_emissao"] or "Não encontrado")

            registro_vinculado = buscar_registro_por_certificado(
                dados_certificado.get("ce_bri"),
                dados_certificado.get("familia")
            )

            if registro_vinculado.empty:
                st.warning("Nenhum registro vinculado encontrado para este CE-BRI + FAMÍLIA.")
            else:
                st.success("Registro vinculado encontrado ✅")
                st.dataframe(normalizar_df_para_exibicao(registro_vinculado), use_container_width=True)

            if rows:
                df_preview = pd.DataFrame(
                    rows,
                    columns=[
                        "ORDEM",
                        "MARCA",
                        "MODELO",
                        "NOME",
                        "CODIGO"
                    ]
                )

                st.dataframe(normalizar_df_para_exibicao(df_preview), use_container_width=True)

                duplicados = verificar_codigos_duplicados(df_preview)

                if not duplicados.empty:
                    st.warning("Foram encontrados códigos duplicados.")
                    st.dataframe(normalizar_df_para_exibicao(duplicados), use_container_width=True)

                if st.button("Salvar / Atualizar banco"):
                    status, mensagem = salvar_ou_atualizar_certificado(
                        dados_certificado,
                        rows,
                        banco_file.name
                    )

                    if status in ["novo", "atualizado"]:
                        st.success(mensagem)
                    elif status == "ignorado":
                        st.info(mensagem)
                    else:
                        st.error(mensagem)

    st.divider()

    st.subheader("Pesquisar item no banco")

    tipo_busca = st.selectbox(
        "Pesquisar por",
        ["Referência / Modelo", "Código de Barras"],
        key="tipo_busca"
    )

    termo_busca = st.text_input(
        "Digite a referência/modelo ou código",
        key="busca_item"
    )

    if termo_busca:
        campo = "i.modelo" if tipo_busca == "Referência / Modelo" else "i.codigo"

        resultado_item = pd.read_sql_query(f"""
        SELECT
            i.marca AS marca,
            i.modelo AS modelo,
            i.nome AS nome,
            i.codigo AS codigo,
            c.ip_bri AS ip_bri,
            c.ce_bri AS ce_bri,
            r.registro AS registro,
            c.familia AS fam,
            r.fabrica AS fabrica,
            c.rev AS rev,
            c.produto AS produto,
            c.data_emissao AS data_emissao,
            i.ordem AS ordem
        FROM itens i
        INNER JOIN certificados c
        ON c.id = i.certificado_id
        LEFT JOIN registros r
        ON UPPER(r.ce_bri) = UPPER(c.ce_bri)
        AND r.familia = c.familia
        WHERE {campo} LIKE ?
        ORDER BY i.marca, i.modelo, c.ip_bri, i.ordem
        """, conn, params=(f"%{termo_busca}%",))

        if resultado_item.empty:
            st.warning("Nenhum item encontrado.")
        else:
            st.success("Item encontrado ✅")
            st.dataframe(normalizar_df_para_exibicao(resultado_item), use_container_width=True)

            st.download_button(
                "Baixar resultado da pesquisa em CSV",
                resultado_item.to_csv(index=False, sep=";").encode(
                    "ISO-8859-1",
                    errors="replace"
                ),
                "resultado_pesquisa_item.csv",
                "text/csv"
            )

    st.divider()

    st.subheader("Itens Repetidos")

    if st.button("Ver itens repetidos"):
        itens_repetidos = pd.read_sql_query("""
        SELECT
            i.marca AS marca,
            i.modelo AS modelo,
            i.nome AS nome,
            i.codigo AS codigo,
            c.ip_bri AS ip_bri,
            c.ce_bri AS ce_bri,
            r.registro AS registro,
            c.familia AS fam,
            r.fabrica AS fabrica,
            c.rev AS rev,
            c.produto AS produto,
            c.data_emissao AS data_emissao,
            i.ordem AS ordem
        FROM itens i
        INNER JOIN certificados c
        ON c.id = i.certificado_id
        LEFT JOIN registros r
        ON UPPER(r.ce_bri) = UPPER(c.ce_bri)
        AND r.familia = c.familia
        WHERE i.codigo IN (
            SELECT codigo
            FROM itens
            WHERE codigo IS NOT NULL
            AND codigo != ''
            GROUP BY codigo
            HAVING COUNT(*) > 1
        )
        ORDER BY i.codigo, i.marca, i.modelo
        """, conn)

        if itens_repetidos.empty:
            st.success("Nenhum item repetido encontrado ✅")
        else:
            st.warning("Itens repetidos encontrados.")
            st.dataframe(normalizar_df_para_exibicao(itens_repetidos), use_container_width=True)

            st.download_button(
                "Baixar itens repetidos em CSV",
                itens_repetidos.to_csv(index=False, sep=";").encode(
                    "ISO-8859-1",
                    errors="replace"
                ),
                "itens_repetidos.csv",
                "text/csv"
            )

    st.divider()

    st.subheader("Filtro geral por Marca")

    marcas_banco = pd.read_sql_query("""
    SELECT DISTINCT marca
    FROM itens
    WHERE marca IS NOT NULL
    AND marca != ''
    ORDER BY marca
    """, conn)

    if marcas_banco.empty:
        st.info("Nenhuma marca cadastrada.")
    else:
        marca_filtro = st.selectbox(
            "Selecione uma marca",
            marcas_banco["marca"].tolist(),
            key="marca_filtro"
        )

        resultado_marca = pd.read_sql_query("""
        SELECT
            i.marca AS marca,
            i.modelo AS modelo,
            i.nome AS nome,
            i.codigo AS codigo,
            c.ip_bri AS ip_bri,
            c.ce_bri AS ce_bri,
            r.registro AS registro,
            c.familia AS fam,
            r.fabrica AS fabrica,
            c.rev AS rev,
            c.produto AS produto,
            c.data_emissao AS data_emissao,
            i.ordem AS ordem
        FROM itens i
        INNER JOIN certificados c
        ON c.id = i.certificado_id
        LEFT JOIN registros r
        ON UPPER(r.ce_bri) = UPPER(c.ce_bri)
        AND r.familia = c.familia
        WHERE i.marca = ?
        ORDER BY i.modelo, c.ip_bri
        """, conn, params=(marca_filtro,))

        st.dataframe(normalizar_df_para_exibicao(resultado_marca), use_container_width=True)

        st.download_button(
            "Baixar CSV da marca",
            resultado_marca.to_csv(index=False, sep=";").encode(
                "ISO-8859-1",
                errors="replace"
            ),
            f"{marca_filtro}.csv",
            "text/csv"
        )

    st.divider()

    st.subheader("Histórico por IP-BRI")

    ip_historico = st.text_input(
        "Digite o IP-BRI",
        key="historico_ip"
    )

    if ip_historico:
        historico = pd.read_sql_query("""
        SELECT
            ip_bri,
            rev_antiga,
            rev_nova,
            modelo,
            codigo,
            tipo_alteracao,
            campo_alterado,
            valor_antigo,
            valor_novo,
            arquivo_pdf,
            data_hora
        FROM historico_alteracoes
        WHERE ip_bri LIKE ?
        ORDER BY id DESC
        """, conn, params=(f"%{ip_historico}%",))

        if historico.empty:
            st.warning("Nenhum histórico encontrado.")
        else:
            grupos = historico.groupby(
                [
                    "rev_antiga",
                    "rev_nova",
                    "arquivo_pdf",
                    "data_hora"
                ],
                dropna=False
            )

            for (rev_antiga, rev_nova, arquivo_pdf, data_hora), grupo in grupos:
                with st.expander(
                    f"REV {rev_antiga} → REV {rev_nova} | {data_hora} | {arquivo_pdf}"
                ):
                    col1, col2, col3 = st.columns(3)

                    col1.metric(
                        "Itens novos",
                        len(grupo[grupo["tipo_alteracao"] == "ITEM_NOVO"])
                    )

                    col2.metric(
                        "Itens removidos",
                        len(grupo[grupo["tipo_alteracao"] == "ITEM_REMOVIDO"])
                    )

                    col3.metric(
                        "Campos alterados",
                        len(grupo[grupo["tipo_alteracao"] == "CAMPO_ALTERADO"])
                    )

                    for _, row in grupo.iterrows():
                        tipo = row["tipo_alteracao"]

                        if tipo == "CERTIFICADO_NOVO":
                            st.info(f"📄 Certificado novo: {row['valor_novo']}")

                        elif tipo == "CERTIFICADO_ATUALIZADO":
                            st.success(f"🔄 Atualizado: {row['valor_antigo']} → {row['valor_novo']}")

                        elif tipo == "REV_IGNORADA":
                            st.warning(f"⚠️ Revisão ignorada: {row['valor_antigo']} | {row['valor_novo']}")

                        elif tipo == "ITEM_REMOVIDO":
                            st.error(
                                f"🗑️ ITEM REMOVIDO\n\n"
                                f"Modelo: {row['modelo']}\n\n"
                                f"Código: {row['codigo']}\n\n"
                                f"Antes: {row['valor_antigo']}"
                            )

                        elif tipo == "ITEM_NOVO":
                            st.success(
                                f"➕ ITEM NOVO\n\n"
                                f"Modelo: {row['modelo']}\n\n"
                                f"Código: {row['codigo']}\n\n"
                                f"Novo: {row['valor_novo']}"
                            )

                        elif tipo == "CAMPO_ALTERADO":
                            st.warning(
                                f"✏️ CAMPO ALTERADO\n\n"
                                f"Modelo: {row['modelo']}\n\n"
                                f"Código: {row['codigo']}\n\n"
                                f"Campo: {row['campo_alterado']}\n\n"
                                f"Antes: {row['valor_antigo']}\n\n"
                                f"Depois: {row['valor_novo']}"
                            )

            st.download_button(
                "Baixar histórico CSV",
                historico.to_csv(index=False, sep=";").encode(
                    "ISO-8859-1",
                    errors="replace"
                ),
                "historico.csv",
                "text/csv"
            )

    st.divider()

    st.subheader("🗑️ Remover Certificado do Banco")
    st.caption("Use com cuidado. Remove permanentemente o certificado e todos os seus itens do banco.")

    ip_deletar = st.text_input(
        "Digite o IP-BRI exato para remover",
        placeholder="Ex: IP-BRI-0533/2023-15",
        key="ip_deletar"
    )

    if ip_deletar:
        cert_para_deletar = cursor.execute(
            "SELECT id, ip_bri, rev, produto FROM certificados WHERE ip_bri = ?",
            (ip_deletar.strip(),)
        ).fetchone()

        if not cert_para_deletar:
            st.warning("Certificado não encontrado no banco.")
        else:
            qtd_itens_cert = cursor.execute(
                "SELECT COUNT(*) FROM itens WHERE certificado_id = ?",
                (cert_para_deletar[0],)
            ).fetchone()[0]

            st.error(
                f"**{cert_para_deletar[1]}** | REV {cert_para_deletar[2]} | "
                f"{cert_para_deletar[3] or '—'} | {qtd_itens_cert} itens vinculados"
            )

            confirmar_delete = st.checkbox(
                f"Confirmo que quero remover permanentemente este certificado e seus {qtd_itens_cert} itens",
                key="confirmar_delete_cert"
            )

            if confirmar_delete:
                if st.button("🗑️ Remover do banco", key="btn_deletar_cert"):
                    try:
                        cursor.execute("DELETE FROM itens WHERE certificado_id = ?", (cert_para_deletar[0],))
                        cursor.execute("DELETE FROM certificados WHERE id = ?", (cert_para_deletar[0],))
                        commit_seguro()
                        st.success(f"Certificado {cert_para_deletar[1]} e seus itens foram removidos com sucesso ✅")
                    except Exception as e:
                        st.error(f"Erro ao remover: {e}")

# ==========================================
# ABA 3 - REGISTROS
# ==========================================

if _is_active("registros"):
    st.title("Registros")

    st.subheader("Enviar Excel de Registros")

    cliente_base = st.text_input(
        "Nome da Base / Cliente",
        value="",
        placeholder="Ex: BOLSA, MOHNISH, EMPRESA XYZ...",
        key="cliente_base_registros"
    )

    if not clean(cliente_base):
        st.warning("⚠️ Informe o nome da base/cliente antes de enviar o Excel.")
    else:
        st.info(f"Os registros serão salvos para a base: **{cliente_base.upper()}**")

    _modo_salvar_registros = st.radio(
        "Como salvar os registros deste Excel?",
        options=["substituir", "mesclar"],
        format_func=lambda v: {
            "substituir": "🗑️ Substituir tudo desta base (apaga tudo e recadastra do zero)",
            "mesclar": "🔄 Mesclar — manter fábricas já cadastradas, só adicionar registros novos",
        }[v],
        index=0,
        key="modo_salvar_registros_excel",
        help=(
            "Substituir: apaga TODOS os registros dessa base e recadastra com o "
            "que vier no Excel (comportamento de sempre). "
            "Mesclar: fábricas que ainda não existem são criadas normalmente; "
            "fábricas que JÁ existem mantêm o endereço que já está no sistema "
            "(ignora o endereço do upload pra elas) e só ganham os registros "
            "(família/CE-BRI) que ainda não estavam cadastrados — nada é apagado."
        )
    )

    with st.expander("ℹ️ Como funciona esta aba"):
        st.markdown("""
- Envie o Excel de registros do cliente
- Cada **aba do Excel** é tratada como uma fábrica separada
- Você pode renomear cada fábrica e informar o endereço antes de salvar
- Os registros vinculam o **CE-BRI** ao número de **registro** e à **fábrica**, permitindo que o Preenchimento de Confirmação preencha esses campos automaticamente
- No modo **Mesclar**, o nome da fábrica (depois de renomear, se você renomear) precisa bater exatamente com o nome já cadastrado pra ser reconhecida como existente — senão vira uma fábrica nova
""")

    registro_excel = st.file_uploader(
        "Envie o Excel de Registros",
        type=["xlsx", "xls"],
        key="excel_registros",
        disabled=not bool(clean(cliente_base))
    )

    if not clean(cliente_base) and registro_excel:
        st.error("Informe o nome da base antes de processar o Excel.")
        registro_excel = None

    if registro_excel:
        try:
            excel = pd.ExcelFile(registro_excel)
            abas = excel.sheet_names

            st.success(f"{len(abas)} abas encontradas ✅")

            st.subheader("Nomear Fábricas")
            st.caption("Se quiser, altere o nome de cada fábrica antes de salvar. Exemplo: F01 - MOHNISH / CE-BRI-XXXX")

            nomes_fabricas = {}
            enderecos_fabricas_por_aba = {}

            _fabricas_ja_cadastradas_cliente = set()
            if clean(cliente_base):
                _fabricas_ja_cadastradas_cliente = {
                    r[0] for r in cursor.execute(
                        "SELECT DISTINCT fabrica FROM registros WHERE UPPER(cliente_base) = UPPER(?)",
                        (cliente_base,)
                    ).fetchall()
                }

            for aba in abas:
                nome_sugerido = str(aba)

                nomes_fabricas[aba] = st.text_input(
                    f"Nome da fábrica para a aba: {aba}",
                    value=nome_sugerido,
                    key=f"nome_fabrica_{cliente_base}_{aba}"
                )

                _nome_atual_fabrica = clean(nomes_fabricas[aba])
                _fabrica_existe_no_modo_mesclar = (
                    _modo_salvar_registros == "mesclar"
                    and _nome_atual_fabrica in _fabricas_ja_cadastradas_cliente
                )

                if _fabrica_existe_no_modo_mesclar:
                    _end_existente_row = cursor.execute(
                        """SELECT endereco_fabrica FROM registros
                           WHERE UPPER(cliente_base) = UPPER(?) AND fabrica = ?
                           AND endereco_fabrica IS NOT NULL AND endereco_fabrica != '' LIMIT 1""",
                        (cliente_base, _nome_atual_fabrica)
                    ).fetchone()
                    _end_existente_txt = _end_existente_row[0] if _end_existente_row else "(sem endereço salvo)"
                    st.success(
                        f"✅ Fábrica **{_nome_atual_fabrica}** já cadastrada — endereço atual "
                        f"será mantido: _{_end_existente_txt}_ (só os registros novos serão adicionados)"
                    )
                    enderecos_fabricas_por_aba[aba] = ""  # ignorado no modo mesclar pra fábrica existente
                else:
                    if _modo_salvar_registros == "mesclar":
                        st.info(f"🆕 Fábrica **{_nome_atual_fabrica}** ainda não cadastrada — será criada.")
                    enderecos_fabricas_por_aba[aba] = st.text_area(
                        f"Endereço da fábrica para a aba: {aba}",
                        value="",
                        key=f"endereco_fabrica_{cliente_base}_{aba}",
                        height=80
                    )

            todas_fabricas = []

            for aba in abas:
                df_filtrado = ler_registros_aba(
                    registro_excel,
                    aba
                )

                if df_filtrado is None or df_filtrado.empty:
                    continue

                nome_final_fabrica = clean(nomes_fabricas.get(aba, aba)) or str(aba)
                endereco_final_fabrica = clean(enderecos_fabricas_por_aba.get(aba, ""))

                # Mantém o CE-BRI extraído da aba original, mas salva a fábrica com o nome escolhido.
                df_filtrado["FABRICA"] = nome_final_fabrica
                df_filtrado["ENDERECO_FABRICA"] = endereco_final_fabrica

                st.subheader(f"Prévia: {cliente_base} → {nome_final_fabrica}")

                st.dataframe(
                    df_filtrado,
                    use_container_width=True
                )

                todas_fabricas.append(df_filtrado)

            if todas_fabricas:
                banco_registros = pd.concat(
                    todas_fabricas,
                    ignore_index=True
                )

                st.divider()

                st.subheader("Banco Geral de Registros que será salvo")

                st.dataframe(
                    banco_registros,
                    use_container_width=True
                )

                _label_botao_salvar = (
                    "Salvar registros desta base"
                    if _modo_salvar_registros == "substituir"
                    else "🔄 Mesclar registros desta base"
                )
                if st.button(_label_botao_salvar, key="salvar_registros_base"):
                    enderecos_para_salvar = {}
                    if "ENDERECO_FABRICA" in banco_registros.columns:
                        for _, linha_endereco in banco_registros.drop_duplicates(subset=["FABRICA"]).iterrows():
                            enderecos_para_salvar[str(linha_endereco.get("FABRICA", ""))] = clean(linha_endereco.get("ENDERECO_FABRICA", ""))

                    if _modo_salvar_registros == "substituir":
                        total_salvo = salvar_registros_no_banco(
                            banco_registros,
                            registro_excel.name,
                            cliente_base,
                            enderecos_para_salvar
                        )
                        st.success(f"{total_salvo} registros salvos no banco para a base {cliente_base} ✅")
                    else:
                        _resumo_mesclagem = salvar_registros_mesclado(
                            banco_registros,
                            registro_excel.name,
                            cliente_base,
                            enderecos_para_salvar
                        )
                        _msg_resumo = (
                            f"✅ {_resumo_mesclagem['novos_registros']} registro(s) novo(s) adicionado(s), "
                            f"{_resumo_mesclagem['mantidos']} já existente(s) mantido(s) sem alteração."
                        )
                        if _resumo_mesclagem["fabricas_novas"]:
                            _msg_resumo += f"\n\n🆕 Fábrica(s) nova(s) criada(s): {', '.join(_resumo_mesclagem['fabricas_novas'])}"
                        if _resumo_mesclagem["fabricas_existentes"]:
                            _msg_resumo += f"\n\n✅ Fábrica(s) já cadastrada(s) (endereço mantido): {', '.join(_resumo_mesclagem['fabricas_existentes'])}"
                        st.success(_msg_resumo)
                    st.cache_data.clear()

                st.download_button(
                    "Baixar registros tratados CSV",
                    banco_registros.to_csv(
                        index=False,
                        sep=";"
                    ).encode(
                        "ISO-8859-1",
                        errors="replace"
                    ),
                    "registros_tratados.csv",
                    "text/csv"
                )

            else:
                st.warning("Nenhum registro válido encontrado no Excel.")

        except Exception as e:
            st.error(f"Erro ao ler Excel: {e}")

    st.divider()

    st.subheader("✏️ Editar registro de um item específico")
    st.caption("Selecione o cliente e a família pra ajustar só o número de registro, sem reenviar o Excel inteiro.")

    _clientes_disponiveis = [r[0] for r in cursor.execute(
        "SELECT DISTINCT cliente_base FROM registros WHERE cliente_base IS NOT NULL AND cliente_base != '' ORDER BY cliente_base"
    ).fetchall()]

    if not _clientes_disponiveis:
        st.info("Nenhum registro cadastrado ainda.")
    else:
        _cliente_edicao = st.selectbox("Cliente", _clientes_disponiveis, key="editar_registro_cliente")

        _familias_disponiveis = [r[0] for r in cursor.execute(
            """SELECT DISTINCT familia FROM registros
               WHERE UPPER(cliente_base) = UPPER(?) AND familia IS NOT NULL AND familia != ''
               ORDER BY familia""",
            (_cliente_edicao,)
        ).fetchall()]

        if not _familias_disponiveis:
            st.info("Nenhuma família encontrada para esse cliente.")
        else:
            _familia_edicao = st.selectbox("Família", _familias_disponiveis, key="editar_registro_familia")

            _itens_para_editar = cursor.execute(
                """SELECT id, cliente_base, fabrica, ce_bri, familia, registro, endereco_fabrica
                   FROM registros
                   WHERE UPPER(cliente_base) = UPPER(?) AND familia = ?
                   ORDER BY fabrica, ce_bri""",
                (_cliente_edicao, _familia_edicao)
            ).fetchall()

            if not _itens_para_editar:
                st.info("Nenhum registro encontrado para essa combinação.")
            else:
                if len(_itens_para_editar) > 1:
                    st.caption(
                        f"⚠️ {len(_itens_para_editar)} registro(s) encontrados para essa família "
                        f"(fábricas/CE-BRI diferentes). Ajuste cada um separadamente."
                    )

                for _item in _itens_para_editar:
                    _id_item, _cli, _fab, _ce, _fam, _reg_atual, _end = _item
                    with st.container(border=True):
                        col_e1, col_e2 = st.columns(2)
                        with col_e1:
                            st.text_input("Cliente", value=_cli or "", disabled=True, key=f"ro_cliente_{_id_item}")
                            st.text_input("Família", value=_fam or "", disabled=True, key=f"ro_familia_{_id_item}")
                            st.text_input("Fábrica / CE-BRI", value=f"{_fab or ''} | {_ce or ''}", disabled=True, key=f"ro_fabrica_{_id_item}")
                        with col_e2:
                            st.text_area("Endereço", value=_end or "", disabled=True, key=f"ro_endereco_{_id_item}", height=100)

                        _novo_registro = st.text_input(
                            "Registro (único campo editável)",
                            value=_reg_atual or "",
                            key=f"editar_registro_valor_{_id_item}"
                        )

                        if st.button("💾 Salvar novo registro", key=f"salvar_registro_{_id_item}"):
                            cursor.execute(
                                "UPDATE registros SET registro = ? WHERE id = ?",
                                (clean(_novo_registro), _id_item)
                            )
                            commit_seguro()
                            st.success(f"Registro atualizado para: {clean(_novo_registro)} ✅")
                            st.cache_data.clear()
                            st.rerun()

    st.divider()

    st.subheader("Backup dos Registros / Banco")

    backup_registro_upload = st.file_uploader(
        "Importar backup certificados.db",
        type=["db"],
        key="backup_db_registros"
    )

    if backup_registro_upload:
        st.warning("⚠️ Isso irá **substituir completamente** o banco atual. Esta ação não pode ser desfeita.")
        confirmar_backup_aba3 = st.checkbox("Confirmo que quero substituir o banco atual", key="confirmar_backup_aba3")
        if confirmar_backup_aba3:
            with open(DB_PATH, "wb") as f:
                f.write(backup_registro_upload.read())
            try:
                conn.close()
            except Exception:
                pass
            conn = _abrir_conexao(DB_PATH)
            cursor = conn.cursor()
            st.session_state["db_version"] = st.session_state.get("db_version", 0) + 1
            st.cache_data.clear()
            st.cache_resource.clear()
            st.success("✅ Backup importado! Recarregando...")
            st.rerun()

    if Path(DB_PATH).exists():
        with open(DB_PATH, "rb") as f:
            _db_bytes_reg = f.read()
        st.download_button(
            "Baixar backup atualizado",
            _db_bytes_reg,
            "certificados.db",
            "application/octet-stream",
            key="download_backup_registros"
        )

    registros_salvos = exportar_registros_banco()

    if not registros_salvos.empty:
        st.download_button(
            "Baixar registros salvos CSV",
            registros_salvos.to_csv(index=False, sep=";").encode(
                "ISO-8859-1",
                errors="replace"
            ),
            "registros_salvos.csv",
            "text/csv",
            key="download_registros_salvos"
        )

    st.divider()

    st.subheader("Visualizar Bases")

    bases_existentes = pd.read_sql_query("""
    SELECT DISTINCT cliente_base
    FROM registros
    WHERE cliente_base IS NOT NULL
    AND cliente_base != ''
    ORDER BY cliente_base
    """, conn)

    if bases_existentes.empty:
        st.info("Nenhuma base salva ainda.")
    else:
        base_selecionada = st.selectbox(
            "Selecione a base",
            bases_existentes["cliente_base"].tolist(),
            key="visualizar_base"
        )

        registros_base = pd.read_sql_query("""
        SELECT
            cliente_base,
            fabrica,
            ce_bri,
            familia,
            registro,
            endereco_fabrica
        FROM registros
        WHERE cliente_base = ?
        ORDER BY fabrica, CAST(familia AS INTEGER)
        """, conn, params=(base_selecionada,))

        fabricas = registros_base["fabrica"].unique().tolist()

        for fabrica in fabricas:
            bloco = registros_base[
                registros_base["fabrica"] == fabrica
            ]

            with st.expander(f"{base_selecionada} → {fabrica}", expanded=False):
                st.dataframe(
                    bloco[[
                        "familia",
                        "registro",
                        "ce_bri",
                        "endereco_fabrica"
                    ]],
                    use_container_width=True
                )

                st.download_button(
                    f"Baixar CSV - {fabrica}",
                    bloco.to_csv(index=False, sep=";").encode(
                        "ISO-8859-1",
                        errors="replace"
                    ),
                    f"{base_selecionada}_{str(fabrica).replace('/', '-')}.csv",
                    "text/csv",
                    key=f"download_{base_selecionada}_{fabrica}"
                )


    st.divider()

    st.subheader("Editar endereço e CE-BRI de fábrica já salva")
    st.caption("As alterações são aplicadas em todos os registros da fábrica e também no Sistema 5 automaticamente.")

    bases_para_editar = pd.read_sql_query("""
    SELECT DISTINCT cliente_base
    FROM registros
    WHERE cliente_base IS NOT NULL
    AND cliente_base != ''
    ORDER BY cliente_base
    """, conn)

    if bases_para_editar.empty:
        st.info("Nenhuma base disponível para edição.")
    else:
        base_editar = st.selectbox(
            "Base / Cliente",
            bases_para_editar["cliente_base"].tolist(),
            key="editar_endereco_base"
        )

        fabricas_para_editar = pd.read_sql_query("""
        SELECT
            fabrica,
            MAX(ce_bri) AS ce_bri,
            MAX(endereco_fabrica) AS endereco_fabrica
        FROM registros
        WHERE cliente_base = ?
        GROUP BY fabrica
        ORDER BY fabrica
        """, conn, params=(base_editar,))

        if fabricas_para_editar.empty:
            st.info("Nenhuma fábrica encontrada nessa base.")
        else:
            opcoes_fabricas = fabricas_para_editar["fabrica"].tolist()

            fabrica_editar = st.selectbox(
                "Fábrica",
                opcoes_fabricas,
                key="editar_endereco_fabrica"
            )

            linha_fabrica = fabricas_para_editar[
                fabricas_para_editar["fabrica"] == fabrica_editar
            ].iloc[0]

            endereco_atual = linha_fabrica.get("endereco_fabrica", "") or ""
            ce_bri_atual   = linha_fabrica.get("ce_bri", "") or ""

            if pd.isna(endereco_atual): endereco_atual = ""
            if pd.isna(ce_bri_atual):   ce_bri_atual   = ""

            col_ed1, col_ed2 = st.columns(2)

            with col_ed1:
                st.text_input(
                    "CE-BRI da fábrica",
                    value=str(ce_bri_atual),
                    disabled=True,
                    help="O CE-BRI vem da aba do Excel de registros e não pode ser alterado.",
                    key="exibir_ce_bri_readonly"
                )

            with col_ed2:
                st.caption("Será atualizado em Registros e Sistema 5")

            novo_endereco = st.text_area(
                "Endereço da fábrica",
                value="" if not endereco_atual or pd.isna(endereco_atual) else str(endereco_atual),
                height=120,
                key="editar_endereco_texto"
            )

            if st.button("💾 Salvar endereço", key="salvar_endereco_fabrica_existente"):
                alteracoes = 0

                # 1. Atualiza APENAS o endereço nos Registros — CE-BRI não é alterado
                cursor.execute("""
                UPDATE registros
                SET endereco_fabrica = ?
                WHERE UPPER(cliente_base) = UPPER(?)
                AND fabrica = ?
                """, (
                    clean(novo_endereco),
                    clean(base_editar).upper(),
                    fabrica_editar
                ))
                alteracoes += cursor.rowcount
                commit_seguro()

                # 2. Atualiza APENAS o endereço no sistema5_fabricas
                cursor.execute("""
                UPDATE sistema5_fabricas
                SET endereco_fabrica = ?
                WHERE id IN (
                    SELECT f.id
                    FROM sistema5_fabricas f
                    INNER JOIN sistema5_clientes c ON c.id = f.cliente_id
                    WHERE UPPER(c.cliente_base) = UPPER(?)
                    AND UPPER(f.fabrica) = UPPER(?)
                )
                """, (
                    clean(novo_endereco),
                    clean(base_editar).upper(),
                    fabrica_editar
                ))
                s5_fabricas_alt = cursor.rowcount
                commit_seguro()

                # 3. Atualiza APENAS o endereço no sistema5_itens
                cursor.execute("""
                UPDATE sistema5_itens
                SET endereco_fabrica = ?
                WHERE UPPER(fabrica) = UPPER(?)
                AND UPPER(cliente_base) = UPPER(?)
                """, (
                    clean(novo_endereco),
                    fabrica_editar,
                    clean(base_editar).upper()
                ))
                s5_itens_alt = cursor.rowcount
                commit_seguro()

                st.cache_data.clear()

                st.success(
                    f"✅ Endereço atualizado para **{fabrica_editar}**:\n\n"
                    f"- Registros: {alteracoes}\n"
                    f"- Sistema 5 Fábricas: {s5_fabricas_alt}\n"
                    f"- Itens salvos: {s5_itens_alt}"
                )

    st.divider()

    st.subheader("🔄 Sincronizar endereços do Sistema 5 com Registros")
    st.caption("Varre todos os processos do Sistema 5 e atualiza o endereço com base nos Registros. Também restaura CE-BRIs que possam ter sido alterados indevidamente.")

    if st.button("🔄 Sincronizar todos os endereços agora", key="sincronizar_enderecos_s5", type="primary"):
        try:
            # Busca todos os registros com endereço preenchido
            registros_ref = pd.read_sql_query("""
            SELECT
                UPPER(cliente_base) AS cliente_base,
                UPPER(fabrica) AS fabrica,
                MAX(ce_bri) AS ce_bri,
                MAX(endereco_fabrica) AS endereco_fabrica
            FROM registros
            WHERE endereco_fabrica IS NOT NULL
            AND endereco_fabrica != ''
            GROUP BY UPPER(cliente_base), UPPER(fabrica)
            """, conn)

            if registros_ref.empty:
                st.warning("Nenhum registro com endereço encontrado para sincronizar.")
            else:
                total_arquivos = 0
                total_itens = 0
                total_fabricas = 0

                barra = st.progress(0, text="Sincronizando...")
                total = len(registros_ref)

                for i, row in registros_ref.iterrows():
                    cliente = clean(row["cliente_base"]).upper()
                    fabrica = clean(row["fabrica"]).upper()
                    ce_bri  = clean(row["ce_bri"]).upper() if row["ce_bri"] else ""
                    endereco = clean(row["endereco_fabrica"]) if row["endereco_fabrica"] else ""

                    if not endereco:
                        continue

                    # Restaura CE-BRI original e atualiza endereço no sistema5_fabricas
                    cursor.execute("""
                    UPDATE sistema5_fabricas
                    SET endereco_fabrica = ?, ce_bri = ?
                    WHERE UPPER(fabrica) = ?
                    AND id IN (
                        SELECT f.id FROM sistema5_fabricas f
                        INNER JOIN sistema5_clientes c ON c.id = f.cliente_id
                        WHERE UPPER(c.cliente_base) = ?
                    )
                    """, (endereco, ce_bri, fabrica, cliente))
                    total_fabricas += cursor.rowcount

                    # Restaura CE-BRI original e atualiza endereço no sistema5_itens
                    cursor.execute("""
                    UPDATE sistema5_itens
                    SET endereco_fabrica = ?, ce_bri = ?
                    WHERE UPPER(fabrica) = ?
                    AND UPPER(cliente_base) = ?
                    """, (endereco, ce_bri, fabrica, cliente))
                    total_itens += cursor.rowcount

                    barra.progress(int((i + 1) / total * 100), text=f"Sincronizando {fabrica}...")

                commit_seguro()
                st.cache_data.clear()
                barra.empty()

                st.success(
                    f"✅ Sincronização concluída!\n\n"
                    f"- Fábricas Sistema 5: {total_fabricas} atualizadas\n"
                    f"- Itens sincronizados: {total_itens}"
                )

        except Exception as e:
            st.error(f"Erro na sincronização: {e}")

    st.divider()

    st.subheader("🔧 Corrigir CE-BRIs inconsistentes")
    st.caption("Varre todos os registros e corrige CE-BRIs que não batem com o CE-BRI principal da fábrica (extraído do nome da aba do Excel).")

    if st.button("🔧 Corrigir todos os CE-BRIs agora", key="corrigir_cebri_registros", type="primary"):
        try:
            import re as _re

            # Busca todas as fábricas com CE-BRI no nome
            fabricas = pd.read_sql_query("""
            SELECT DISTINCT cliente_base, fabrica
            FROM registros
            WHERE fabrica LIKE '% - CE-BRI-%'
            """, conn)

            total_registros = 0
            total_s5_fabricas = 0
            total_s5_itens = 0
            correcoes = []

            for _, row in fabricas.iterrows():
                cliente = row["cliente_base"]
                fabrica = row["fabrica"]

                # Extrai CE-BRI correto do nome da fábrica
                m = _re.search(r'CE-BRI-[\w-]+', fabrica)
                if not m:
                    continue
                ce_bri_correto = m.group(0).upper().strip()

                # Corrige na tabela registros (onde CE-BRI diferente do correto)
                cursor.execute("""
                UPDATE registros SET ce_bri = ?
                WHERE UPPER(cliente_base) = UPPER(?)
                AND UPPER(fabrica) = UPPER(?)
                AND (UPPER(ce_bri) != ? OR ce_bri IS NULL OR ce_bri = '')
                """, (ce_bri_correto, cliente, fabrica, ce_bri_correto))
                n = cursor.rowcount
                total_registros += n

                # Corrige no sistema5_fabricas
                cursor.execute("""
                UPDATE sistema5_fabricas SET ce_bri = ?
                WHERE UPPER(fabrica) = UPPER(?)
                AND id IN (
                    SELECT f.id FROM sistema5_fabricas f
                    INNER JOIN sistema5_clientes c ON c.id = f.cliente_id
                    WHERE UPPER(c.cliente_base) = UPPER(?)
                )
                AND (UPPER(COALESCE(ce_bri,'')) != ? OR ce_bri IS NULL)
                """, (ce_bri_correto, fabrica, cliente, ce_bri_correto))
                n2 = cursor.rowcount
                total_s5_fabricas += n2

                # Corrige no sistema5_itens
                cursor.execute("""
                UPDATE sistema5_itens SET ce_bri = ?
                WHERE UPPER(fabrica) = UPPER(?)
                AND UPPER(cliente_base) = UPPER(?)
                AND (UPPER(COALESCE(ce_bri,'')) != ? OR ce_bri IS NULL)
                """, (ce_bri_correto, fabrica, cliente, ce_bri_correto))
                n3 = cursor.rowcount
                total_s5_itens += n3

                if n > 0 or n2 > 0 or n3 > 0:
                    correcoes.append(f"{cliente} / {fabrica}: {n}R + {n2}F + {n3}I")

            commit_seguro()
            st.cache_data.clear()

            total = total_registros + total_s5_fabricas + total_s5_itens
            if total > 0:
                st.success(
                    f"✅ CE-BRIs corrigidos! Total: {total} registros\n\n"
                    f"- Tabela Registros: {total_registros}\n"
                    f"- Sistema 5 Fábricas: {total_s5_fabricas}\n"
                    f"- Sistema 5 Itens: {total_s5_itens}"
                )
            else:
                st.info(f"✅ Verificadas {len(fabricas)} fábricas — nenhum CE-BRI inconsistente encontrado.")

        except Exception as e:
            st.error(f"Erro ao corrigir CE-BRIs: {e}")

    st.divider()

    st.subheader("🗑️ Excluir registros de um cliente")
    st.caption("Remove permanentemente todos os registros de uma base/cliente do banco. Não afeta certificados nem Sistema 5.")

    bases_para_excluir = pd.read_sql_query("""
    SELECT DISTINCT cliente_base,
        COUNT(*) AS total_registros,
        COUNT(DISTINCT fabrica) AS total_fabricas
    FROM registros
    WHERE cliente_base IS NOT NULL AND cliente_base != ''
    GROUP BY cliente_base
    ORDER BY cliente_base
    """, conn)

    if bases_para_excluir.empty:
        st.info("Nenhuma base cadastrada para excluir.")
    else:
        base_excluir = st.selectbox(
            "Selecione o cliente para excluir",
            bases_para_excluir["cliente_base"].tolist(),
            key="excluir_registro_base"
        )

        linha_excluir = bases_para_excluir[bases_para_excluir["cliente_base"] == base_excluir].iloc[0]

        st.error(
            f"**{base_excluir}** — "
            f"{int(linha_excluir['total_registros'])} registro(s) em "
            f"{int(linha_excluir['total_fabricas'])} fábrica(s)"
        )

        confirmar_excluir_base = st.checkbox(
            f"Confirmo que quero excluir permanentemente todos os registros de {base_excluir}",
            key="confirmar_excluir_base"
        )

        if confirmar_excluir_base:
            if st.button("🗑️ Excluir registros deste cliente", key="btn_excluir_base"):
                try:
                    cursor.execute(
                        "DELETE FROM registros WHERE UPPER(cliente_base) = UPPER(?)",
                        (base_excluir,)
                    )
                    commit_seguro()
                    st.success(f"Todos os registros de **{base_excluir}** foram excluídos ✅")
                    st.rerun()
                except Exception as e:
                    st.error(f"Erro ao excluir: {e}")

if _is_active("confirmacao"):
    st.title("Preenchimento de Confirmação")

    with st.expander("🔎 Testar referência antes de preencher"):
        ref_teste_confirmacao = st.text_input(
            "Digite a referência para testar",
            placeholder="Ex: BAL-17-FAN-PR",
            key="teste_ref_confirmacao"
        )

        if st.button("Testar referência", key="btn_teste_ref_confirmacao"):
            diag_ref = diagnosticar_referencia_confirmacao(ref_teste_confirmacao)

            st.write(f"Status: **{diag_ref.get('status')}**")

            if diag_ref.get("item_confirmacao"):
                st.success("A confirmação encontrou este item ✅")
                st.json(diag_ref.get("item_confirmacao"))
            else:
                st.warning("A confirmação não encontrou este item.")

            candidatos_diag = diag_ref.get("candidatos_sistema5")

            if candidatos_diag is not None and not candidatos_diag.empty:
                st.info("Candidatos encontrados no Sistema 5:")
                st.dataframe(normalizar_df_para_exibicao(candidatos_diag), use_container_width=True)
            else:
                st.error("Nenhum candidato encontrado no Sistema 5 pela busca bruta.")

    with st.expander("ℹ️ Como funciona esta aba"):
        st.markdown("""
- Envie o Excel de confirmação com células coloridas:
  - 🟢 **Verde** → referência do produto (MODELO). O sistema busca este valor no banco
  - 🟡 **Amarelo** → campo a ser preenchido (MARCA, NOME, CODIGO, REGISTRO, ENDERECO...)
  - 🔵 **Azul** → cabeçalho da coluna (define o que preencher nas amarelas abaixo)
- Não é necessário ter uma coluna chamada "REF" — qualquer célula verde é considerada referência
- A busca é feita primeiro no **Sistema 5**, depois nos **Certificados**
- **Pré-requisito:** o banco precisa ter certificados e registros cadastrados para o preenchimento funcionar
""")

    # Aviso de banco vazio — aba 4 não funciona sem dados
    try:
        _m = metricas_banco()
        _total_cert_aba4 = _m.get("total_certs", 0)
        _total_s5_aba4   = _m.get("total_s5", 0)
        _total_reg_aba4  = _m.get("total_regs", 0)

        if _total_cert_aba4 == 0 and _total_s5_aba4 == 0:
            st.error(
                "⛔ O banco não tem nenhum certificado cadastrado. "
                "O preenchimento não vai encontrar nenhuma referência.\n\n"
                "**O que fazer:** vá à aba **PDF → XML** e envie os certificados primeiro."
            )
        elif _total_reg_aba4 == 0:
            st.warning(
                "⚠️ O banco tem certificados, mas **nenhum registro de fábrica** cadastrado. "
                "Os campos REGISTRO e ENDERECO não serão preenchidos.\n\n"
                "**O que fazer:** vá à aba **Registros** e importe o Excel de registros do cliente."
            )
        else:
            st.success(
                f"Banco pronto: {_total_cert_aba4} certificado(s) + {_total_s5_aba4} item(ns) do Sistema 5 + {_total_reg_aba4} registro(s) de fábrica ✅"
            )
    except Exception:
        pass

    # Toggle de filtro rigoroso
    filtro_rigoroso = st.toggle(
        "🔍 Filtro rigoroso de referência",
        value=False,
        key="filtro_referencia_rigoroso",
        help="Quando ativado, referência '0122' não preenche '0122-18'. Use quando houver modelos com sufixo numérico no mesmo certificado."
    )
    if filtro_rigoroso:
        st.info("⚠️ Filtro rigoroso ativado — referências serão buscadas de forma exata. Modelos com sufixo numérico (ex: 0122-18) não serão aceitos para a referência 0122.")

    with st.expander("🔎 Diagnosticar uma referência específica (investigar por que trouxe o processo errado)"):
        _ref_diag = st.text_input("Referência pra diagnosticar (ex: SK-1892)", key="ref_diagnostico_conf")
        if st.button("Rodar diagnóstico", key="btn_diag_conf") and clean(_ref_diag):
            buscar_item_sistema5_confirmacao.clear()  # força reexecução (ignora cache)
            _resultado_diag = buscar_item_confirmacao(clean(_ref_diag), _debug_ref=clean(_ref_diag))
            if _resultado_diag:
                st.markdown("**Resumo do item que seria usado:**")
                st.json({k: v for k, v in _resultado_diag.items() if v is not None})

    with st.expander("🛠️ Corrigir datas de processos já cadastrados (rodar uma vez)"):
        st.caption(
            "Se algum processo foi salvo com a data digitada manualmente em formato "
            "brasileiro (ex: 09/04/2026) antes da correção automática existir, isso "
            "quebra a ordenação de 'processo mais recente'. Essa ferramenta varre TODOS "
            "os processos já cadastrados no Sistema 5 e corrige o formato — sem apagar "
            "nem alterar nenhum outro dado."
        )
        if st.button("🔍 Verificar e corrigir datas", key="btn_corrigir_datas_s5"):
            with st.spinner("Verificando todos os processos cadastrados..."):
                _relatorio = corrigir_datas_processo_existentes()
            st.success(f"✅ {_relatorio['corrigidos']} data(s) corrigida(s).")
            st.caption(f"{_relatorio['ja_estavam_ok']} já estavam no formato certo (não precisaram de ajuste).")
            if _relatorio["nao_reconhecidos"]:
                st.warning(
                    f"⚠️ {len(_relatorio['nao_reconhecidos'])} data(s) não reconhecida(s) "
                    f"(formato estranho, não é nem ISO nem brasileiro comum) — precisam ser "
                    f"corrigidas manualmente:"
                )
                st.dataframe(pd.DataFrame(_relatorio["nao_reconhecidos"]), use_container_width=True)
            buscar_item_sistema5_confirmacao.clear()  # limpa cache pra refletir a correção já

    arquivo_confirmacao = st.file_uploader(
        "Envie o Excel de confirmação",
        type=["xlsx", "xlsm"],
        key="excel_confirmacao"
    )

    if arquivo_confirmacao:
        try:
            # ============================================================
            # ETAPA 1: Detectar referências com múltiplos candidatos
            # ============================================================
            from openpyxl import load_workbook as _lwb

            try:
                arquivo_confirmacao.seek(0)
                _wb_scan = _lwb(arquivo_confirmacao, data_only=True)
            except Exception:
                arquivo_confirmacao.seek(0)
                _wb_scan = _lwb(_limpar_excel_imagens(arquivo_confirmacao), data_only=True)

            # Coleta todas as referências verdes do arquivo
            _refs_unicas = set()
            for _ws in _wb_scan.worksheets:
                for _row in _ws.iter_rows():
                    for _cell in _row:
                        if eh_verde(_cell) and clean(_cell.value):
                            _val = ler_valor_celula_preservando_zeros(_cell)
                            if _val:
                                _refs_unicas.add(_val)

            # Verifica quais têm múltiplos candidatos COM MARCAS DIFERENTES
            _conflitos = {}
            for _ref in _refs_unicas:
                _candidatos = buscar_candidatos_confirmacao(_ref)
                if len(_candidatos) > 1:
                    _marcas_distintas = {
                        str(c.get("MARCA", "")).strip().upper()
                        for c in _candidatos
                        if str(c.get("MARCA", "")).strip()
                    }
                    if len(_marcas_distintas) > 1:
                        _conflitos[_ref] = _candidatos

            # ============================================================
            # ETAPA 2: Se há conflitos, mostrar tela de revisão
            # ============================================================
            if _conflitos:
                _escolhas_key = f"escolhas_confirmacao_{arquivo_confirmacao.name}"

                if _escolhas_key not in st.session_state:
                    st.session_state[_escolhas_key] = {}

                st.warning(f"⚠️ {len(_conflitos)} referência(s) encontrada(s) em múltiplas marcas. Escolha a marca correta antes de gerar o arquivo.")

                _todas_resolvidas = True
                for _ref, _cands in _conflitos.items():
                    st.markdown(f"**Referência: `{_ref}`** — {len(_cands)} marca(s) diferentes encontradas")
                    _opcoes = [
                        f"{c.get('MARCA','?')}  |  {c.get('MODELO','?')}  |  {c.get('IP_BRI') or c.get('IP_PROCESSO') or '—'}"
                        for c in _cands
                    ]
                    _escolha = st.radio(
                        f"Qual a marca correta para `{_ref}`?",
                        _opcoes,
                        key=f"radio_conf_{_ref}",
                        index=0
                    )
                    _idx_escolha = _opcoes.index(_escolha)
                    st.session_state[_escolhas_key][_ref] = _cands[_idx_escolha]
                    st.divider()

                if not st.button("✅ Confirmar escolhas e gerar arquivo", type="primary", key="confirmar_escolhas_conf"):
                    st.stop()

                _escolhas_pre = st.session_state.get(_escolhas_key, {})
            else:
                _escolhas_pre = {}

            # ============================================================
            # ETAPA 3: Gerar arquivo com escolhas confirmadas
            # ============================================================
            arquivo_confirmacao.seek(0)
            saida_excel, total_preenchidos, nao_encontrados = preencher_excel_confirmacao(
                arquivo_confirmacao,
                escolhas_pre=_escolhas_pre
            )

            st.success(f"Preenchimento concluído ✅ {total_preenchidos} células preenchidas.")

            if nao_encontrados:
                st.warning("Algumas referências verdes não foram encontradas no banco.")
                st.dataframe(
                    pd.DataFrame({"referencia_nao_encontrada": sorted(set(nao_encontrados))}),
                    use_container_width=True
                )

            # Nome sugerido baseado no arquivo original
            nome_original = arquivo_confirmacao.name.rsplit(".", 1)[0]
            nome_sugerido = f"{nome_original}_preenchida.xlsx"

            col_dl1, col_dl2 = st.columns(2)

            with col_dl1:
                # Download padrão pelo navegador
                st.download_button(
                    "⬇️ Baixar pelo navegador",
                    saida_excel,
                    nome_sugerido,
                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    key="download_confirmacao_browser"
                )

            with col_dl2:
                # Salvar direto no PC via janela nativa do Windows
                if st.button("💾 Salvar direto no PC", key="salvar_confirmacao_pc", type="primary"):
                    try:
                        import tkinter as tk
                        from tkinter import filedialog

                        root = tk.Tk()
                        root.withdraw()
                        root.wm_attributes("-topmost", True)

                        caminho = filedialog.asksaveasfilename(
                            parent=root,
                            title="Salvar confirmação preenchida",
                            defaultextension=".xlsx",
                            initialfile=nome_sugerido,
                            filetypes=[("Excel", "*.xlsx"), ("Todos os arquivos", "*.*")]
                        )
                        root.destroy()

                        if caminho:
                            with open(caminho, "wb") as f:
                                f.write(saida_excel)
                            st.success(f"✅ Arquivo salvo em:\n{caminho}")
                        else:
                            st.info("Salvamento cancelado.")

                    except Exception as e_save:
                        st.error(f"Erro ao salvar: {e_save}")

        except Exception as e:
            st.error(f"Erro ao preencher Excel: {e}")


# ==========================================
# ABA 5 - PESQUISA AVANÇADA
# ==========================================

if _is_active("pesquisa"):
    st.title("Pesquisa Avançada")

    with st.expander("ℹ️ Como funciona esta aba"):
        st.markdown("""
- Pesquise em **duas direções**:
  - Dado um **IP-BRI** → descubra o registro, fábrica e endereço vinculados
  - Dado um **REGISTRO** → descubra o IP-BRI e o fabricante
- Útil para conferir rapidamente se um certificado está vinculado a um cliente antes do preenchimento
""")

    tipo_pesquisa_avancada = st.radio(
        "O que você quer pesquisar?",
        [
            "Quero saber o REGISTRO de um IP-BRI",
            "Quero saber o IP-BRI de um REGISTRO"
        ],
        key="tipo_pesquisa_avancada"
    )

    if tipo_pesquisa_avancada == "Quero saber o REGISTRO de um IP-BRI":
        ip_pesquisa = st.text_input(
            "Digite o IP-BRI",
            placeholder="Ex: IP-BRI-0533/2023-15 ou 0533/2023-15",
            key="pesquisa_avancada_ip"
        )

        if ip_pesquisa:
            resultado_ip = pesquisar_por_ip_bri(ip_pesquisa)

            if resultado_ip.empty:
                st.warning("Nenhum registro encontrado para este IP-BRI.")
            else:
                st.success(f"{len(resultado_ip)} resultado(s) encontrado(s) ✅")
                st.dataframe(normalizar_df_para_exibicao(resultado_ip), use_container_width=True)

                st.download_button(
                    "Baixar resultado em CSV",
                    resultado_ip.to_csv(index=False, sep=";").encode(
                        "ISO-8859-1",
                        errors="replace"
                    ),
                    "pesquisa_por_ip_bri.csv",
                    "text/csv",
                    key="download_pesquisa_ip_bri"
                )

    else:
        registro_pesquisa = st.text_input(
            "Digite o REGISTRO",
            placeholder="Ex: 003889/2023",
            key="pesquisa_avancada_registro"
        )

        if registro_pesquisa:
            resultado_registro = pesquisar_por_registro(registro_pesquisa)

            if resultado_registro.empty:
                st.warning("Nenhum IP-BRI encontrado para este registro.")
            else:
                st.success(f"{len(resultado_registro)} resultado(s) encontrado(s) ✅")
                st.dataframe(normalizar_df_para_exibicao(resultado_registro), use_container_width=True)

                st.download_button(
                    "Baixar resultado em CSV",
                    resultado_registro.to_csv(index=False, sep=";").encode(
                        "ISO-8859-1",
                        errors="replace"
                    ),
                    "pesquisa_por_registro.csv",
                    "text/csv",
                    key="download_pesquisa_registro"
                )

# ==========================================
# ABA 6 - SISTEMA 5
# ==========================================

if _is_active("sistema5"):
    st.title("Sistema 5")

    if st.button("Limpar itens DESCONSIDERADOS já cadastrados", key="limpar_desconsiderados_s5"):
        removidos_desconsiderados = limpar_itens_desconsiderados_sistema5()
        st.success(f"{removidos_desconsiderados} item(ns) desconsiderado(s) removido(s) do Sistema 5 ✅")

    st.divider()

    # ---- Download lista de fábricas ----
    _df_fab_dl = pd.read_sql_query("""
        SELECT
            c.cliente_base   AS CLIENTE,
            f.ce_bri         AS "CE-BRI",
            f.endereco_fabrica AS ENDEREÇO
        FROM sistema5_fabricas f
        INNER JOIN sistema5_clientes c ON c.id = f.cliente_id
        ORDER BY c.cliente_base, f.ce_bri
    """, conn)

    if not _df_fab_dl.empty:
        _buf_fab = BytesIO()
        with pd.ExcelWriter(_buf_fab, engine='openpyxl') as _writer_fab:
            _df_fab_dl.to_excel(_writer_fab, index=False, sheet_name="Fábricas")
        st.download_button(
            "📥 Baixar lista de fábricas (Excel)",
            _buf_fab.getvalue(),
            "fabricas_sistema5.xlsx",
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            key="download_fabricas_s5"
        )

    st.divider()
    st.subheader("🔎 Pesquisa Global Sistema 5")

    termo_pesquisa_s5 = st.text_input(
        "Pesquisar item no Sistema 5",
        placeholder="Digite modelo, código, marca, descrição, IP, CE-BRI, fábrica ou cliente",
        key="pesquisa_global_sistema5"
    )

    if st.button("Pesquisar no Sistema 5", key="botao_pesquisa_global_s5"):
        if not clean(termo_pesquisa_s5):
            st.warning("Digite algo para pesquisar.")
        else:
            resultado_global_s5 = pesquisar_global_sistema5(termo_pesquisa_s5)

            if resultado_global_s5.empty:
                st.warning("Nenhum item encontrado no Sistema 5.")
            else:
                st.success(f"{len(resultado_global_s5)} item(ns) encontrado(s) no Sistema 5 ✅")
                st.dataframe(normalizar_df_para_exibicao(resultado_global_s5), use_container_width=True)

                st.download_button(
                    "Baixar resultado Sistema 5 CSV",
                    resultado_global_s5.to_csv(index=False, sep=";").encode(
                        "ISO-8859-1",
                        errors="replace"
                    ),
                    "pesquisa_global_sistema5.csv",
                    "text/csv",
                    key="download_pesquisa_global_s5"
                )


    with st.expander("ℹ️ Como funciona esta aba"):
        st.markdown("""
- Módulo para processos que ainda não têm certificado IP-BRI oficial emitido
- Importa Excel de **inclusões, manutenções ou recertificações** — o tipo é detectado pelo nome do arquivo
- **O arquivo precisa ter IP no nome** (ex: `INCLUSÃO 06-01-26 F01 IP-0094-26.xlsx`)
- Cada aba do Excel é tratada como uma **família** pelo número presente no nome da aba
- Os itens salvos aqui também são consultados pelo **Preenchimento de Confirmação**
- Organizado por: categoria → cliente → fábrica → processo
""")

    categoria_s5 = st.radio(
        "Categoria",
        ["SISTEMA 5 NOVO PROJETO", "SISTEMA 5 PROPRIOS", "SISTEMA 5 FOCUS"],
        horizontal=True,
        key="s5_categoria"
    )

    # --- Seleção de cliente ---
    if categoria_s5 == "SISTEMA 5 NOVO PROJETO":
        cliente_s5 = st.text_input(
            "Cliente / Solicitante",
            value="BOLSA",
            key="s5_cliente_np"
        )
    elif categoria_s5 == "SISTEMA 5 FOCUS":
        cliente_s5 = st.text_input(
            "Cliente / Solicitante",
            value="FOCUS",
            key="s5_cliente_focus"
        )
    else:
        # SISTEMA 5 PROPRIOS — usa dropdown com clientes já cadastrados + opção de novo
        clientes_proprios_banco = listar_clientes_sistema5_banco(categoria_s5)
        opcoes_clientes_proprios = clientes_proprios_banco["cliente_base"].tolist() if not clientes_proprios_banco.empty else []
        opcoes_clientes_proprios = ["➕ Novo cliente..."] + opcoes_clientes_proprios

        escolha_cliente_proprios = st.selectbox(
            "Cliente",
            opcoes_clientes_proprios,
            key="s5_cliente_proprios_select"
        )

        if escolha_cliente_proprios == "➕ Novo cliente...":
            cliente_s5 = st.text_input(
                "Nome do novo cliente",
                placeholder="Ex: EMPRESA XYZ",
                key="s5_cliente_proprios_novo"
            ).upper()
            if cliente_s5:
                st.info(f"Novo cliente: **{cliente_s5}** — será criado ao salvar o primeiro processo.")
        else:
            cliente_s5 = escolha_cliente_proprios
            st.success(f"Cliente selecionado: **{cliente_s5}**")

    st.subheader("Fábrica")

    # Busca fábricas das duas fontes: registros + sistema5_fabricas
    fabricas_disponiveis_s5 = listar_fabricas_sistema5_banco(cliente_s5, categoria_s5)

    usar_fabrica_existente = False
    fabrica_padrao  = ""
    ce_padrao       = ""
    endereco_padrao = ""
    idx_seguro      = 0
    dados_fabrica_sugeridos = None

    if not fabricas_disponiveis_s5.empty:
        usar_fabrica_existente = st.checkbox(
            "Usar fábrica já cadastrada",
            value=True,
            key="s5_usar_fabrica_existente"
        )
    else:
        if clean(cliente_s5):
            st.info(f"ℹ️ Nenhuma fábrica cadastrada para **{cliente_s5}** ainda. Preencha os campos abaixo manualmente — a fábrica será criada ao salvar o primeiro processo.")

    if usar_fabrica_existente and not fabricas_disponiveis_s5.empty:
        opcoes_fabrica_s5 = []

        for _, linha_fab in fabricas_disponiveis_s5.iterrows():
            fab    = clean(linha_fab.get("fabrica", ""))
            ce     = clean(linha_fab.get("ce_bri", ""))
            origem = clean(linha_fab.get("origem", ""))
            opcoes_fabrica_s5.append(f"{fab} | {ce} ({origem})")

        idx_escolhido_fab = st.selectbox(
            "Selecione uma fábrica",
            range(len(opcoes_fabrica_s5)),
            format_func=lambda i: opcoes_fabrica_s5[i],
            key=f"s5_fabrica_existente_select_{cliente_s5}"
        )

        # Garante que o índice nunca ultrapasse o tamanho da lista atual
        idx_seguro = min(idx_escolhido_fab, len(fabricas_disponiveis_s5) - 1)

        # Pega os dados direto do DataFrame pelo índice — sem parse de string
        linha_selecionada = fabricas_disponiveis_s5.iloc[idx_seguro]
        fabrica_padrao  = clean(linha_selecionada.get("fabrica", ""))
        ce_padrao       = clean(linha_selecionada.get("ce_bri", ""))
        endereco_padrao = clean(linha_selecionada.get("endereco_fabrica", ""))

        dados_fabrica_sugeridos = {
            "fabrica":          fabrica_padrao,
            "ce_bri":           ce_padrao,
            "endereco_fabrica": endereco_padrao,
        }

        st.success("Dados da fábrica carregados ✅")

    # key inclui cliente + índice da fábrica → sempre reseta ao trocar qualquer um dos dois
    _key_fab = f"{cliente_s5}_{idx_seguro}"

    col_f1, col_f2 = st.columns(2)

    with col_f1:
        fabrica_s5 = st.text_input(
            "Fábrica",
            value=fabrica_padrao,
            placeholder="Ex: F01, F02, FÁBRICA 01...",
            key=f"s5_fabrica_{_key_fab}"
        )

        ce_bri_s5 = st.text_input(
            "CE-BRI da fábrica",
            value=ce_padrao,
            placeholder="Ex: CE-BRI-INNAC-02484-01A",
            key=f"s5_ce_bri_{_key_fab}"
        )

    with col_f2:
        endereco_s5 = st.text_area(
            "Endereço da fábrica",
            value="" if not endereco_padrao or pd.isna(endereco_padrao) else str(endereco_padrao),
            placeholder="Cole aqui o endereço completo da fábrica",
            height=120,
            key=f"s5_endereco_{_key_fab}"
        )

    # Fabricante automático — mostra se a fábrica selecionada tem histórico
    if clean(fabrica_s5):
        fabricante_sugerido = buscar_fabricante_por_fabrica_s5(cliente_s5, fabrica_s5, ce_bri_s5)
        if fabricante_sugerido:
            st.info(f"🏭 Fabricante identificado para esta fábrica: **{fabricante_sugerido}** "
                    f"— será usado automaticamente em processos de RECERTIFICAÇÃO se o Excel não informar a marca.")
            st.session_state["s5_fabricante_sugerido"] = fabricante_sugerido
        else:
            st.session_state["s5_fabricante_sugerido"] = ""

    if clean(ce_bri_s5) and not clean(endereco_s5):
        dados_por_ce = buscar_dados_fabrica_existente(
            cliente_s5,
            "",
            ce_bri_s5
        )

        if dados_por_ce and clean(dados_por_ce.get("endereco_fabrica", "")):
            st.info("Existe endereço salvo para este CE-BRI nos Registros. Você pode copiar/preencher acima.")
            st.code(dados_por_ce.get("endereco_fabrica", ""))

    st.divider()
    st.subheader("Enviar Excel(s) de processo")

    modo_upload = st.radio(
        "Modo de envio",
        ["📄 Excels individuais", "🗜️ ZIP com vários processos"],
        horizontal=True,
        key="s5_modo_upload"
    )

    if modo_upload == "🗜️ ZIP com vários processos":
        zip_upload = st.file_uploader(
            "Envie o ZIP contendo os Excels (com IP no nome)",
            type=["zip"],
            key="s5_zip_upload"
        )

        arquivos_s5 = []

        if zip_upload:
            try:
                import zipfile as zipfile_mod
                from io import BytesIO

                with zipfile_mod.ZipFile(zip_upload, "r") as zf:
                    nomes = [
                        n for n in zf.namelist()
                        if n.lower().endswith((".xlsx", ".xls"))
                        and not n.startswith("__MACOSX")
                        and not "/." in n
                    ]

                    if not nomes:
                        st.error("Nenhum arquivo Excel encontrado dentro do ZIP.")
                    else:
                        st.info(f"📦 {len(nomes)} Excel(s) encontrado(s) no ZIP.")
                        for nome in nomes:
                            dados = zf.read(nome)
                            nome_base = nome.split("/")[-1]  # remove subpastas
                            buf = BytesIO(dados)
                            buf.name = nome_base
                            buf.size = len(dados)
                            arquivos_s5.append(buf)

            except Exception as e:
                st.error(f"Erro ao abrir ZIP: {e}")
    else:
        arquivos_s5 = st.file_uploader(
            "Envie um ou mais Excels com IP no nome",
            type=["xlsx", "xls"],
            accept_multiple_files=True,
            key="s5_excel_multi"
        )

    if arquivos_s5:
        # Validação de tamanho total — Streamlit Cloud rejeita lotes muito grandes
        LIMITE_MB_POR_LOTE = 150
        tamanho_total_mb = sum(a.size for a in arquivos_s5) / (1024 * 1024)

        fila = []
        erros = []

        if tamanho_total_mb > LIMITE_MB_POR_LOTE:
            qtd_arquivos = len(arquivos_s5)
            tamanho_medio = tamanho_total_mb / qtd_arquivos
            lote_sugerido = max(1, int(LIMITE_MB_POR_LOTE / tamanho_medio))
            st.error(
                f"⛔ Tamanho total dos arquivos: **{tamanho_total_mb:.1f}MB** "
                f"— excede o limite de {LIMITE_MB_POR_LOTE}MB por envio.\n\n"
                f"Envie em lotes de até **{lote_sugerido} arquivo(s)** por vez."
            )
        else:
            if tamanho_total_mb > 80:
                st.warning(
                    f"⚠️ Lote grande: {tamanho_total_mb:.1f}MB — "
                    f"o processamento pode demorar um pouco mais."
                )

            st.markdown(f"**{len(arquivos_s5)} arquivo(s) — {tamanho_total_mb:.1f}MB total.** Revise abaixo e clique em **Salvar todos** no final.")

            # ---- Pré-processamento: lê todos os arquivos e monta fila ----
            for idx_arquivo, arquivo_s5 in enumerate(arquivos_s5):
                ip_extraido          = extrair_ip_processo(arquivo_s5.name)
                data_extraida        = extrair_data_processo(arquivo_s5.name)
                tipo_detectado       = detectar_tipo_processo(arquivo_s5.name)
                fabrica_detectada_nm = extrair_codigo_fabrica_nome(arquivo_s5.name)
                dados_fab_arq        = buscar_fabrica_por_codigo_s5(cliente_s5, fabrica_detectada_nm) if fabrica_detectada_nm else None

                if not ip_extraido:
                    erros.append({"arquivo": arquivo_s5.name, "motivo": "IP não encontrado no nome do arquivo"})
                    continue

                # Fábrica selecionada manualmente SEMPRE tem prioridade.
                # A detecção automática pelo nome do arquivo só preenche CE-BRI e endereço
                # se o usuário não tiver informado — nunca sobrescreve a fábrica escolhida.
                fab_final = fabrica_s5
                ce_final  = ce_bri_s5
                end_final = endereco_s5

                # Se o usuário não preencheu fábrica manualmente, tenta puxar dos registros
                if not clean(fab_final) and dados_fab_arq:
                    fab_final = dados_fab_arq.get("fabrica", "")
                    ce_final  = dados_fab_arq.get("ce_bri", "")
                    end_final = dados_fab_arq.get("endereco_fabrica", "")
                elif dados_fab_arq:
                    # Fábrica manual definida — completa só CE-BRI e endereço se estiverem vazios
                    if not clean(ce_final):
                        ce_final  = dados_fab_arq.get("ce_bri", ce_final)
                    if not clean(end_final):
                        end_final = dados_fab_arq.get("endereco_fabrica", end_final)

                dup = ip_processo_ja_existe_sistema5(ip_extraido, cliente_s5, fab_final)

                try:
                    df_arq = ler_excel_inclusao_sistema5(arquivo_s5)
                except Exception as e:
                    erros.append({"arquivo": arquivo_s5.name, "motivo": str(e)})
                    continue

                if df_arq.empty:
                    erros.append({"arquivo": arquivo_s5.name, "motivo": "Nenhum item encontrado no Excel"})
                    continue

                fabricante_sessao = st.session_state.get("s5_fabricante_sugerido", "")
                if tipo_detectado == "RECERTIFICACAO" and fabricante_sessao:
                    mask = df_arq["MARCA"].isna() | (df_arq["MARCA"].str.strip() == "")
                    df_arq.loc[mask, "MARCA"] = fabricante_sessao

                fila.append({
                    "idx":       idx_arquivo,
                    "arquivo":   arquivo_s5,
                    "nome":      arquivo_s5.name,
                    "ip":        ip_extraido,
                    "tipo":      tipo_detectado,
                    "data":      data_extraida or "",
                    "fabrica":   fab_final,
                    "ce_bri":    ce_final,
                    "endereco":  end_final,
                    "df":        df_arq,
                    "duplicado": dup,
                    "auto_fab":  dados_fab_arq is not None,
                })

        if erros:
            st.error(f"⛔ {len(erros)} arquivo(s) com problema — serão ignorados no salvamento:")
            for err in erros:
                st.caption(f"• **{err['arquivo']}** — {err['motivo']}")

        if fila:
            st.divider()
            st.markdown(f"### Revisão — {len(fila)} processo(s) prontos")

            with st.expander("⚙️ Ajustes globais (aplicar a todos)", expanded=False):
                col_g1, col_g2 = st.columns(2)
                with col_g1:
                    tipos_global = ["— manter individual —", "INCLUSAO", "MANUTENCAO", "RECERTIFICACAO", "INICIAL", "OUTROS"]
                    tipo_global = st.selectbox("Forçar tipo para todos", tipos_global, key="s5_tipo_global")
                with col_g2:
                    data_global = st.text_input("Forçar data para todos (opcional)", placeholder="AAAA-MM-DD", key="s5_data_global")

            indices_selecionados = []

            for item in fila:
                tem_dup = not item["duplicado"].empty
                status_icon = "⚠️" if tem_dup else "✅"
                label_exp = (
                    f"{status_icon} {item['nome']} | "
                    f"IP: {item['ip']} | "
                    f"Tipo: {item['tipo']} | "
                    f"Fábrica: {item['fabrica']} | "
                    f"{len(item['df'])} itens"
                )

                with st.expander(label_exp, expanded=tem_dup):
                    col_r1, col_r2, col_r3 = st.columns(3)
                    col_r1.metric("IP", item["ip"])
                    col_r2.metric("Itens", len(item["df"]))
                    col_r3.metric("Fábrica", item["fabrica"])

                    if item["auto_fab"]:
                        st.success("Fábrica vinculada automaticamente pelos Registros ✅")

                    col_t1, col_t2 = st.columns(2)
                    with col_t1:
                        tipos = ["INCLUSAO", "MANUTENCAO", "RECERTIFICACAO", "INICIAL", "OUTROS"]
                        tipo_arq = st.selectbox(
                            "Tipo",
                            tipos,
                            index=tipos.index(item["tipo"]) if item["tipo"] in tipos else 0,
                            key=f"s5_tipo_rev_{item['idx']}_{item['nome']}"
                        )
                        item["tipo"] = tipo_arq

                    with col_t2:
                        data_arq = st.text_input(
                            "Data",
                            value=item["data"],
                            key=f"s5_data_rev_{item['idx']}_{item['nome']}",
                            help="Formato AAAA-MM-DD (ex: 2026-04-09). Outros formatos comuns "
                                 "(DD-MM-AAAA, DD/MM/AAAA) são convertidos automaticamente."
                        )
                        data_normalizada, _aviso_data = normalizar_data_processo_manual(data_arq)
                        item["data"] = data_normalizada
                        if _aviso_data:
                            st.warning(_aviso_data)
                        elif data_normalizada != clean(data_arq) and clean(data_arq):
                            st.caption(f"✏️ Convertida automaticamente para {data_normalizada}")

                    if tem_dup:
                        st.warning("⚠️ Já existe processo com este IP para esta fábrica:")
                        st.dataframe(normalizar_df_para_exibicao(item["duplicado"]), use_container_width=True)
                        item["substituir"] = st.checkbox(
                            "Substituir processo existente ao salvar",
                            value=False,
                            key=f"s5_sub_{item['idx']}_{item['nome']}"
                        )
                    else:
                        item["substituir"] = False

                    _df_preview = item["df"].head(10) if item["df"] is not None and not item["df"].empty else None
                    st.dataframe(normalizar_df_para_exibicao(_df_preview), use_container_width=True)
                    if len(item["df"]) > 10:
                        st.caption(f"... e mais {len(item['df']) - 10} item(ns)")

                    incluir = st.checkbox(
                        "Incluir este processo no salvamento",
                        value=not tem_dup,
                        key=f"s5_incluir_{item['idx']}_{item['nome']}"
                    )
                    if incluir:
                        indices_selecionados.append(item["idx"])

            # Aplicar ajustes globais
            for item in fila:
                if tipo_global != "— manter individual —":
                    item["tipo"] = tipo_global
                if clean(data_global):
                    data_global_norm, _aviso_data_global = normalizar_data_processo_manual(data_global)
                    item["data"] = data_global_norm
                    if _aviso_data_global:
                        st.warning(_aviso_data_global)

            st.divider()
            fila_final = [i for i in fila if i["idx"] in indices_selecionados]
            st.markdown(
                f"**{len(fila_final)} processo(s) selecionado(s)** "
                f"| {sum(len(i['df']) for i in fila_final)} itens no total"
            )

            if not clean(cliente_s5):
                st.error("⛔ Informe o cliente antes de salvar.")
            elif not fila_final:
                st.warning("Nenhum processo selecionado.")
            else:
                if st.button(
                    f"💾 Salvar {len(fila_final)} processo(s) no Sistema 5",
                    key="s5_salvar_todos",
                    type="primary"
                ):
                    salvos   = 0
                    falhos   = []
                    progress = st.progress(0, text="Salvando...")

                    for i_prog, item in enumerate(fila_final):
                        try:
                            if item["substituir"] and not item["duplicado"].empty:
                                for _, proc_dup in item["duplicado"].iterrows():
                                    deletar_processo_sistema5(proc_dup["arquivo_id"])

                            total_itens = salvar_inclusao_sistema5(
                                categoria_s5,
                                cliente_s5,
                                item["fabrica"],
                                item["ce_bri"],
                                item["endereco"],
                                item["tipo"],
                                item["ip"],
                                item["data"],
                                item["nome"],
                                item["df"]
                            )
                            salvos += 1
                            progress.progress(
                                (i_prog + 1) / len(fila_final),
                                text=f"Salvando {item['ip']}... ({i_prog+1}/{len(fila_final)})"
                            )
                        except Exception as e:
                            falhos.append({"arquivo": item["nome"], "erro": str(e)})

                    progress.empty()
                    st.cache_data.clear()

                    if salvos:
                        st.success(f"✅ {salvos} processo(s) salvos com sucesso!")
                    if falhos:
                        st.error(f"⛔ {len(falhos)} processo(s) com erro:")
                        for f in falhos:
                            st.caption(f"• **{f['arquivo']}** — {f['erro']}")
                    if salvos:
                        st.rerun()

    st.divider()

    with st.expander("🔍 Verificador de descrições sem medidas"):
        st.caption(
            "Varre todos os itens e mostra os que não têm medidas no NOME. "
            "A ordem correta é: medidas primeiro (ex: 25X15X10CM), depois o mecanismo."
        )

        _clientes_med = listar_clientes_sistema5_banco("SISTEMA 5 PROPRIOS")
        _opcoes_clientes_med = ["Todos os clientes"] + (
            _clientes_med["cliente_base"].tolist() if not _clientes_med.empty else []
        )
        _filtro_cliente_med = st.selectbox(
            "Filtrar por cliente",
            _opcoes_clientes_med,
            key="verif_med_cliente"
        )
        _cliente_filtro_med = (
            "" if _filtro_cliente_med == "Todos os clientes" else _filtro_cliente_med
        )

        if st.button("🔍 Verificar agora", key="btn_verificar_medidas"):
            with st.spinner("Verificando descrições..."):
                _df_sem_med = verificar_itens_sem_medidas(_cliente_filtro_med)
            st.session_state["resultado_sem_medidas"] = _df_sem_med

        _df_result_med = st.session_state.get("resultado_sem_medidas")
        if _df_result_med is not None:
            if _df_result_med.empty:
                st.success("Todos os itens têm medidas na descrição ✅")
            else:
                st.warning(f"{len(_df_result_med)} item(ns) sem medidas no NOME.")
                st.dataframe(normalizar_df_para_exibicao(_df_result_med), use_container_width=True)
                st.download_button(
                    "⬇️ Baixar lista CSV",
                    _df_result_med.to_csv(index=False, sep=";").encode("ISO-8859-1", errors="replace"),
                    "itens_sem_medidas.csv",
                    "text/csv",
                    key="download_sem_medidas"
                )

    with st.expander("🔁 Mover processos em lote para outra fábrica"):
        if "s5_mover_selecao" not in st.session_state:
            st.session_state["s5_mover_selecao"] = []

        st.caption("Digite o IP de um processo, adicione à fila, repita quantas vezes quiser e mova tudo de uma vez.")

        _col_ip1, _col_ip2 = st.columns([4, 1])
        with _col_ip1:
            _ip_busca_mv = st.text_input(
                "Número do processo (IP)",
                placeholder="Ex: IP-0094-26",
                key="s5_mover_ip_input"
            )
        with _col_ip2:
            st.write("")
            _btn_buscar_mv = st.button("Buscar", key="s5_mover_btn_buscar")

        if _btn_buscar_mv and clean(_ip_busca_mv):
            _res_busca_mv = buscar_processos_por_ip(_ip_busca_mv)
            st.session_state["s5_mover_resultados"] = (
                _res_busca_mv.to_dict("records") if not _res_busca_mv.empty else []
            )

        _resultados_mv = st.session_state.get("s5_mover_resultados", [])
        if _resultados_mv:
            st.markdown(f"**{len(_resultados_mv)} resultado(s):**")
            for _r in _resultados_mv:
                _aid = int(_r["arquivo_id"])
                _ja_sel = any(s["arquivo_id"] == _aid for s in st.session_state["s5_mover_selecao"])
                _cr1, _cr2 = st.columns([5, 1])
                _cr1.caption(
                    f"`{_r['ip_processo']}` · {_r['tipo_processo'] or ''} · "
                    f"{_r['fabrica']} · {_r['cliente_base']}"
                )
                if _ja_sel:
                    _cr2.success("✓ Na fila")
                else:
                    if _cr2.button("+ Adicionar", key=f"s5_add_mv_{_aid}"):
                        st.session_state["s5_mover_selecao"].append({
                            "arquivo_id": _aid,
                            "ip_processo": _r["ip_processo"],
                            "fabrica_atual": _r["fabrica"],
                            "cliente_base":  _r["cliente_base"],
                        })
                        st.rerun()

        _selecao_mv = st.session_state["s5_mover_selecao"]
        if _selecao_mv:
            st.divider()
            st.markdown(f"**Fila — {len(_selecao_mv)} processo(s) para mover:**")
            for _s in list(_selecao_mv):
                _cs1, _cs2 = st.columns([5, 1])
                _cs1.write(f"`{_s['ip_processo']}` · {_s['fabrica_atual']} · {_s['cliente_base']}")
                if _cs2.button("✕", key=f"s5_rem_mv_{_s['arquivo_id']}"):
                    st.session_state["s5_mover_selecao"] = [
                        x for x in _selecao_mv if x["arquivo_id"] != _s["arquivo_id"]
                    ]
                    st.rerun()

            if st.button("Limpar fila", key="s5_limpar_selecao_mv"):
                st.session_state["s5_mover_selecao"] = []
                st.session_state.pop("s5_mover_resultados", None)
                st.rerun()

            st.divider()
            _todas_fabs_mv = listar_todas_fabricas_para_mover()
            if _todas_fabs_mv.empty:
                st.warning("Nenhuma fábrica cadastrada no Sistema 5.")
            else:
                _opcoes_fab_mv = [
                    (int(r["id"]), f"{r['fabrica']} | {r.get('ce_bri','') or ''} ({r['cliente_base']})")
                    for _, r in _todas_fabs_mv.iterrows()
                ]
                _sel_fab_mv = st.selectbox(
                    "Mover todos para qual fábrica?",
                    options=range(len(_opcoes_fab_mv)),
                    format_func=lambda i: _opcoes_fab_mv[i][1],
                    key="s5_destino_fab_mv"
                )
                _destino_fab_id_mv   = _opcoes_fab_mv[_sel_fab_mv][0]
                _destino_fab_nome_mv = _opcoes_fab_mv[_sel_fab_mv][1].split(" | ")[0]

                _confirmar_lote_mv = st.checkbox(
                    f"Confirmo: mover {len(_selecao_mv)} processo(s) para **{_destino_fab_nome_mv}**",
                    key="s5_confirmar_lote_mv"
                )
                if _confirmar_lote_mv:
                    if st.button(
                        f"🔁 Mover {len(_selecao_mv)} processo(s) para {_destino_fab_nome_mv}",
                        key="s5_btn_mover_lote",
                        type="primary"
                    ):
                        _erros_mv = []
                        _ok_mv    = 0
                        for _s in _selecao_mv:
                            try:
                                mover_processo_para_fabrica(_s["arquivo_id"], _destino_fab_id_mv)
                                _ok_mv += 1
                            except Exception as _emv:
                                _erros_mv.append(f"{_s['ip_processo']}: {_emv}")

                        if _ok_mv:
                            st.success(f"✅ {_ok_mv} processo(s) movido(s) para **{_destino_fab_nome_mv}**")
                        for _emsg in _erros_mv:
                            st.error(_emsg)

                        st.session_state["s5_mover_selecao"] = []
                        st.session_state.pop("s5_mover_resultados", None)
                        st.cache_data.clear()
                        st.rerun()

    st.divider()

    # ---- Exclusão em lote por IP ----
    with st.expander("🗑️ Excluir processos em lote por IP", expanded=False):
        if "s5_excluir_fila" not in st.session_state:
            st.session_state["s5_excluir_fila"] = []

        _col_ex1, _col_ex2 = st.columns([3, 1])
        with _col_ex1:
            _ip_excluir = st.text_input(
                "Número do IP",
                placeholder="Ex: IP-0001-25",
                key="s5_excluir_ip_input"
            )
        with _col_ex2:
            st.markdown("<br>", unsafe_allow_html=True)
            _btn_buscar_ex = st.button("🔍 Buscar", key="s5_excluir_buscar")

        if _btn_buscar_ex and _ip_excluir.strip():
            _res_ex = buscar_processos_por_ip(_ip_excluir.strip())
            if _res_ex.empty:
                st.warning(f"Nenhum processo encontrado para '{_ip_excluir}'.")
            else:
                _fila_ids = {p["arquivo_id"] for p in st.session_state["s5_excluir_fila"]}
                _adicionados = 0
                for _, _row_ex in _res_ex.iterrows():
                    _aid = int(_row_ex["arquivo_id"])
                    if _aid not in _fila_ids:
                        st.session_state["s5_excluir_fila"].append({
                            "arquivo_id":   _aid,
                            "ip_processo":  _row_ex.get("ip_processo", ""),
                            "fabrica":      _row_ex.get("fabrica", ""),
                            "tipo":         _row_ex.get("tipo_processo", ""),
                            "data":         _row_ex.get("data_processo", ""),
                            "cliente_base": _row_ex.get("cliente_base", ""),
                        })
                        _adicionados += 1
                if _adicionados:
                    st.success(f"✅ {_adicionados} processo(s) adicionado(s) à fila.")
                else:
                    st.info("Processo(s) já estava(m) na fila.")

        # Exibe fila atual
        if st.session_state["s5_excluir_fila"]:
            st.markdown(f"**Fila de exclusão — {len(st.session_state['s5_excluir_fila'])} processo(s):**")
            for _idx_ex, _item_ex in enumerate(st.session_state["s5_excluir_fila"]):
                _c1, _c2 = st.columns([5, 1])
                with _c1:
                    st.caption(
                        f"• {_item_ex['ip_processo']}  |  "
                        f"{_item_ex['cliente_base']}  |  "
                        f"{_item_ex['fabrica']}  |  "
                        f"{_item_ex['tipo']}  |  "
                        f"{_item_ex['data']}"
                    )
                with _c2:
                    if st.button("✕", key=f"s5_ex_rem_{_idx_ex}_{_item_ex['arquivo_id']}"):
                        st.session_state["s5_excluir_fila"].pop(_idx_ex)
                        st.rerun()

            st.divider()
            _confirmar_excluir_lote = st.checkbox(
                f"Confirmo que quero excluir permanentemente os {len(st.session_state['s5_excluir_fila'])} processo(s) acima",
                key="s5_excluir_confirmar"
            )
            if _confirmar_excluir_lote:
                if st.button(
                    f"🗑️ Excluir {len(st.session_state['s5_excluir_fila'])} processo(s)",
                    key="s5_excluir_btn_confirmar",
                    type="primary"
                ):
                    _ids_excluir = [p["arquivo_id"] for p in st.session_state["s5_excluir_fila"]]
                    _ph = ",".join("?" * len(_ids_excluir))
                    _total_itens_ex = cursor.execute(
                        f"SELECT COUNT(*) FROM sistema5_itens WHERE arquivo_id IN ({_ph})",
                        _ids_excluir
                    ).fetchone()[0]
                    cursor.execute(
                        f"DELETE FROM sistema5_itens WHERE arquivo_id IN ({_ph})",
                        _ids_excluir
                    )
                    cursor.execute(
                        f"DELETE FROM sistema5_arquivos WHERE id IN ({_ph})",
                        _ids_excluir
                    )
                    commit_seguro()
                    st.success(
                        f"✅ {len(_ids_excluir)} processo(s) excluído(s) — {_total_itens_ex} itens removidos."
                    )
                    st.session_state["s5_excluir_fila"] = []
                    st.cache_data.clear()
                    st.rerun()
        else:
            st.info("Nenhum processo na fila. Busque por IP para adicionar.")

    st.divider()
    st.subheader("Estrutura salva no Sistema 5")

    processos_s5 = listar_processos_sistema5_cached()

    if processos_s5.empty:
        st.info("Nenhum processo salvo ainda.")
    else:
        categorias = processos_s5["categoria"].dropna().unique().tolist()

        for categoria in categorias:
            with st.expander(f"📁 {categoria}", expanded=True):
                bloco_categoria = processos_s5[processos_s5["categoria"] == categoria]
                clientes = bloco_categoria["cliente_base"].dropna().unique().tolist()

                for cliente in clientes:
                    bloco_cliente = bloco_categoria[bloco_categoria["cliente_base"] == cliente]
                    total_proc_cliente = len(bloco_cliente)
                    total_itens_cliente = int(bloco_cliente["qtd_itens"].fillna(0).sum())
                    fabricas_cliente = bloco_cliente["fabrica"].dropna().unique().tolist()

                    with st.expander(
                        f"👤 {cliente} — {len(fabricas_cliente)} fábrica(s) | {total_proc_cliente} processo(s) | {total_itens_cliente} itens",
                        expanded=False
                    ):
                        for fabrica in fabricas_cliente:
                            bloco_fabrica = bloco_cliente[bloco_cliente["fabrica"] == fabrica]
                            total_proc_fab = len(bloco_fabrica)
                            total_itens_fab = int(bloco_fabrica["qtd_itens"].fillna(0).sum())

                            with st.expander(
                                f"🏭 {fabrica} — {total_proc_fab} processo(s) | {total_itens_fab} itens",
                                expanded=False
                            ):
                            
                                for _, proc in bloco_fabrica.iterrows():
                                    arquivo_id = int(proc["arquivo_id"])
                                    titulo_processo = (
                                        f"📄 {proc.get('tipo_processo', '')} | "
                                        f"{proc.get('ip_processo', '')} | "
                                        f"{proc.get('data_processo', '')} | "
                                        f"{proc.get('arquivo_nome', '')} | "
                                        f"{int(proc.get('qtd_itens', 0) or 0)} itens"
                                    )

                                    with st.expander(titulo_processo, expanded=False):
                                        colp1, colp2, colp3 = st.columns(3)

                                        colp1.metric("IP", proc.get("ip_processo", ""))
                                        colp2.metric("Tipo", proc.get("tipo_processo", ""))
                                        colp3.metric("Itens", int(proc.get("qtd_itens", 0) or 0))

                                        st.caption(f"CE-BRI: {proc.get('ce_bri', '')}")
                                        st.caption(f"Endereço: {proc.get('endereco_fabrica', '')}")

                                        termo_item_processo = st.text_input(
                                            "Pesquisar item dentro deste processo",
                                            placeholder="Digite modelo, marca, descrição ou código",
                                            key=f"pesquisar_item_processo_{arquivo_id}"
                                        )

                                        itens_do_processo = buscar_itens_processo_sistema5(
                                            arquivo_id,
                                            termo_item_processo
                                        )

                                        if itens_do_processo.empty:
                                            st.warning("Nenhum item encontrado dentro deste processo.")
                                        else:
                                            st.success(f"{len(itens_do_processo)} item(ns) encontrado(s) neste processo ✅")
                                            st.dataframe(normalizar_df_para_exibicao(itens_do_processo), use_container_width=True)

                                            st.download_button(
                                                "Baixar itens deste processo CSV",
                                                itens_do_processo.to_csv(index=False, sep=";").encode(
                                                    "ISO-8859-1",
                                                    errors="replace"
                                                ),
                                                f"sistema5_processo_{arquivo_id}.csv",
                                                "text/csv",
                                                key=f"download_itens_processo_{arquivo_id}"
                                            )

                                        st.divider()
                                        confirmar_del_proc = st.checkbox(
                                            f"Confirmo que quero remover permanentemente este processo e seus {int(proc.get('qtd_itens', 0) or 0)} itens",
                                            key=f"confirmar_del_proc_{arquivo_id}"
                                        )
                                        if confirmar_del_proc:
                                            if st.button(
                                                "🗑️ Deletar processo",
                                                key=f"btn_del_proc_{arquivo_id}"
                                            ):
                                                try:
                                                    qtd_del = deletar_processo_sistema5(arquivo_id)
                                                    st.success(f"Processo {proc.get('ip_processo', '')} removido. {qtd_del} itens deletados ✅")
                                                    st.cache_data.clear()
                                                    st.rerun()
                                                except Exception as e:
                                                    st.error(f"Erro ao deletar: {e}")

                                # --- Excluir todos os processos desta fábrica ---
                                st.divider()
                                _key_fab_del = f"{cliente}_{fabrica}".replace(" ", "_")
                                confirmar_del_fab = st.checkbox(
                                    f"Confirmo que quero excluir todos os {total_proc_fab} processo(s) e {total_itens_fab} itens de {fabrica}",
                                    key=f"confirmar_del_fab_{_key_fab_del}"
                                )
                                if confirmar_del_fab:
                                    if st.button(
                                        f"🗑️ Excluir todos os processos de {fabrica}",
                                        key=f"btn_del_fab_{_key_fab_del}"
                                    ):
                                        try:
                                            ids_fab = [int(aid) for aid in bloco_fabrica["arquivo_id"].tolist()]

                                            # Delete em batch — um único commit no final
                                            placeholders = ",".join("?" * len(ids_fab))
                                            total_del_itens = cursor.execute(
                                                f"SELECT COUNT(*) FROM sistema5_itens WHERE arquivo_id IN ({placeholders})",
                                                ids_fab
                                            ).fetchone()[0]
                                            cursor.execute(
                                                f"DELETE FROM sistema5_itens WHERE arquivo_id IN ({placeholders})",
                                                ids_fab
                                            )
                                            cursor.execute(
                                                f"DELETE FROM sistema5_arquivos WHERE id IN ({placeholders})",
                                                ids_fab
                                            )
                                            commit_seguro()

                                            st.success(
                                                f"✅ {len(ids_fab)} processo(s) e {total_del_itens} itens de "
                                                f"**{fabrica}** excluídos com sucesso."
                                            )
                                            st.cache_data.clear()
                                            st.rerun()
                                        except Exception as e:
                                            st.error(f"Erro ao excluir: {e}")
                                            st.error(f"Erro ao excluir: {e}")

        st.download_button(
            "Baixar resumo Sistema 5 CSV",
            processos_s5.to_csv(index=False, sep=";").encode("ISO-8859-1", errors="replace"),
            "sistema5_resumo.csv",
            "text/csv",
            key="download_sistema5_resumo"
        )



# ==========================================
# ABA 7 - IP-BRI PENDENTES
# ==========================================

if _is_active("pendentes"):
    st.title("IP-BRI Pendentes")

    st.info(
        "Use esta aba para cadastrar manualmente o IP-BRI de famílias agrupadas do Sistema 5, "
        "mas ainda não possuem certificado oficial emitido/cadastrado."
    )

    st.subheader("Famílias pendentes encontradas no Sistema 5")

    pendentes_ip_bri = listar_ip_bri_pendentes_sistema5()

    if pendentes_ip_bri.empty:
        st.success("Nenhuma família pendente de IP-BRI encontrada ✅")
    else:
        st.warning(f"{len(pendentes_ip_bri)} família(s) pendente(s) de IP-BRI.")
        st.dataframe(normalizar_df_para_exibicao(pendentes_ip_bri), use_container_width=True)

        st.download_button(
            "Baixar pendentes CSV",
            pendentes_ip_bri.to_csv(index=False, sep=";").encode("ISO-8859-1", errors="replace"),
            "ip_bri_pendentes.csv",
            "text/csv",
            key="download_ip_bri_pendentes"
        )

        st.divider()
        st.subheader("Cadastrar IP-BRI de uma família pendente")

        opcoes_pendentes = []
        for _, linha in pendentes_ip_bri.iterrows():
            label = (
                f"{linha.get('cliente_base', '')} | "
                f"{linha.get('fabrica', '')} | "
                f"{linha.get('ce_bri', '')} | "
                f"FAM {linha.get('familia', '')} | "
                f"{int(linha.get('qtd_processos', 0) or 0)} processo(s)"
            )
            opcoes_pendentes.append(label)

        escolha_pendente = st.selectbox(
            "Selecione a família pendente",
            opcoes_pendentes,
            key="select_ip_bri_pendente"
        )

        idx_escolhido = opcoes_pendentes.index(escolha_pendente)
        linha_escolhida = pendentes_ip_bri.iloc[idx_escolhido]

        col_ip1, col_ip2 = st.columns(2)

        with col_ip1:
            st.text_input("CE-BRI", value=clean(linha_escolhida.get("ce_bri", "")), disabled=True, key="ip_pendente_ce_bri")
            st.text_input("Família", value=clean(linha_escolhida.get("familia", "")), disabled=True, key="ip_pendente_familia")

            registro_copiavel = clean(linha_escolhida.get("registro", ""))

            if registro_copiavel:
                st.code(registro_copiavel)

            st.code(clean(linha_escolhida.get("ips_processos", "")))

            st.write(f"{int(linha_escolhida.get('qtd_processos', 0) or 0)} processo(s) | {int(linha_escolhida.get('qtd_itens', 0) or 0)} item(ns)")

        with col_ip2:
            novo_ip_bri_manual = st.text_input("IP-BRI a cadastrar", placeholder="Ex: IP-BRI-1550-26 ou 1550-26", key="novo_ip_bri_manual")
            obs_ip_bri_manual = st.text_input("Observação", placeholder="Opcional", key="obs_ip_bri_manual")

        if st.button("Salvar IP-BRI desta família", key="salvar_ip_bri_pendente"):
            ok, mensagem = salvar_ip_bri_familia(
                linha_escolhida.get("ce_bri", ""),
                linha_escolhida.get("familia", ""),
                novo_ip_bri_manual,
                obs_ip_bri_manual
            )

            if ok:
                st.success(mensagem)
            else:
                st.error(mensagem)

    st.divider()
    st.subheader("Cadastro manual direto")

    col_manual1, col_manual2, col_manual3 = st.columns(3)

    with col_manual1:
        ce_manual = st.text_input("CE-BRI", placeholder="CE-BRI-INNAC-02484-01A", key="manual_ce_bri_ip")

    with col_manual2:
        fam_manual = st.text_input("Família", placeholder="Ex: 12", key="manual_familia_ip")

    with col_manual3:
        ip_manual = st.text_input("IP-BRI", placeholder="Ex: IP-BRI-1550-26", key="manual_ip_bri")

    obs_manual = st.text_input("Observação manual", placeholder="Opcional", key="manual_obs_ip_bri")

    if st.button("Salvar cadastro manual", key="salvar_ip_bri_manual_direto"):
        ok, mensagem = salvar_ip_bri_familia(ce_manual, fam_manual, ip_manual, obs_manual)

        if ok:
            st.success(mensagem)
        else:
            st.error(mensagem)

    st.divider()
    st.subheader("IP-BRI manuais já cadastrados")

    manuais = listar_ip_bri_manuais()

    if manuais.empty:
        st.info("Nenhum IP-BRI manual cadastrado ainda.")
    else:
        st.dataframe(normalizar_df_para_exibicao(manuais), use_container_width=True)

        st.download_button(
            "Baixar IP-BRI manuais CSV",
            manuais.to_csv(index=False, sep=";").encode("ISO-8859-1", errors="replace"),
            "ip_bri_manuais.csv",
            "text/csv",
            key="download_ip_bri_manuais"
        )

        st.divider()
        st.subheader("Editar ou excluir IP-BRI manual")

        opcoes_manuais = []

        for _, linha in manuais.iterrows():
            label = (
                f"ID {linha.get('id', '')} | "
                f"{linha.get('ce_bri', '')} | "
                f"FAM {linha.get('familia', '')} | "
                f"{linha.get('ip_bri', '')}"
            )
            opcoes_manuais.append(label)

        escolha_manual = st.selectbox(
            "Selecione o cadastro para editar/excluir",
            opcoes_manuais,
            key="select_editar_ip_bri_manual"
        )

        idx_manual = opcoes_manuais.index(escolha_manual)
        linha_manual = manuais.iloc[idx_manual]

        registro_manual_id = int(linha_manual.get("id"))

        col_edit1, col_edit2, col_edit3 = st.columns(3)

        with col_edit1:
            st.caption("CE-BRI vinculado")
            st.code(clean(linha_manual.get("ce_bri", "")))

        with col_edit2:
            st.caption("Família vinculada")
            st.code(clean(linha_manual.get("familia", "")))

        with col_edit3:
            novo_ip_editado = st.text_input(
                "Novo IP-BRI",
                value=clean(linha_manual.get("ip_bri", "")),
                key=f"edit_ip_bri_novo_{registro_manual_id}"
            )

        nova_obs_editada = st.text_input(
            "Observação",
            value=clean(linha_manual.get("observacao", "")),
            key=f"edit_ip_obs_nova_{registro_manual_id}"
        )

        col_btn1, col_btn2 = st.columns(2)

        with col_btn1:
            if st.button("Salvar alteração do IP-BRI", key="btn_salvar_edit_ip_manual"):
                ok, mensagem = atualizar_ip_bri_manual_por_id(
                    registro_manual_id,
                    novo_ip_editado,
                    nova_obs_editada
                )

                if ok:
                    st.success(mensagem)
                else:
                    st.error(mensagem)

        with col_btn2:
            confirmar_exclusao_ip = st.checkbox(
                "Confirmo que quero excluir este cadastro",
                key="confirmar_excluir_ip_manual"
            )

            if st.button("Excluir cadastro selecionado", key="btn_excluir_ip_manual"):
                if not confirmar_exclusao_ip:
                    st.warning("Marque a confirmação antes de excluir.")
                else:
                    ok, mensagem = excluir_ip_bri_manual_por_id(registro_manual_id)

                    if ok:
                        st.success(mensagem)
                    else:
                        st.error(mensagem)



# ==========================================
# ABA 8 - PAINEL DE COBERTURA
# ==========================================

if _is_active("cobertura"):
    st.title("Painel de Cobertura")

    st.info(
        "Visão geral por cliente, fábrica, CE-BRI, família e quantidade de itens. "
        "Use para saber quantas fábricas existem, quantas famílias cada fábrica possui "
        "e quantos itens existem em cada família."
    )

    resumo = cobertura_resumo_geral()

    c1, c2, c3, c4, c5 = st.columns(5)
    c1.metric("Clientes", resumo.get("clientes_registros", 0))
    c2.metric("Fábricas", resumo.get("fabricas_registros", 0))
    c3.metric("Famílias Registros", resumo.get("familias_registros", 0))
    c4.metric("Famílias Sistema 5", resumo.get("familias_sistema5", 0))
    c5.metric("Pendentes IP-BRI", resumo.get("pendentes_ip_bri", 0))

    st.divider()

    st.subheader("Resumo por fábrica")

    df_fabricas = cobertura_fabricas()

    if df_fabricas.empty:
        st.info("Nenhuma fábrica encontrada.")
    else:
        st.dataframe(normalizar_df_para_exibicao(df_fabricas), use_container_width=True)

        st.download_button(
            "Baixar resumo por fábrica CSV",
            df_fabricas.to_csv(index=False, sep=";").encode("ISO-8859-1", errors="replace"),
            "painel_cobertura_fabricas.csv",
            "text/csv",
            key="download_cobertura_fabricas"
        )

    st.divider()

    st.subheader("Detalhamento por família")

    if df_fabricas.empty:
        st.info("Cadastre registros ou processos no Sistema 5 para visualizar o detalhamento.")
    else:
        clientes_opcoes = ["TODOS"] + sorted(df_fabricas["cliente_base"].dropna().astype(str).unique().tolist())
        cliente_filtro = st.selectbox("Cliente", clientes_opcoes, key="cobertura_cliente")

        df_base_filtro = df_fabricas.copy()
        if cliente_filtro != "TODOS":
            df_base_filtro = df_base_filtro[df_base_filtro["cliente_base"].astype(str) == cliente_filtro]

        fabricas_opcoes = ["TODAS"] + sorted(df_base_filtro["fabrica"].dropna().astype(str).unique().tolist())
        fabrica_filtro = st.selectbox("Fábrica", fabricas_opcoes, key="cobertura_fabrica")

        ce_opcoes_base = df_base_filtro.copy()
        if fabrica_filtro != "TODAS":
            ce_opcoes_base = ce_opcoes_base[ce_opcoes_base["fabrica"].astype(str) == fabrica_filtro]

        ce_opcoes = ["TODOS"] + sorted(ce_opcoes_base["ce_bri"].dropna().astype(str).unique().tolist())
        ce_filtro = st.selectbox("CE-BRI", ce_opcoes, key="cobertura_ce_bri")

        df_familias = cobertura_familias_detalhada(
            "" if cliente_filtro == "TODOS" else cliente_filtro,
            "" if fabrica_filtro == "TODAS" else fabrica_filtro,
            "" if ce_filtro == "TODOS" else ce_filtro
        )

        if df_familias.empty:
            st.warning("Nenhuma família encontrada para os filtros selecionados.")
        else:
            st.dataframe(normalizar_df_para_exibicao(df_familias), use_container_width=True)

            st.download_button(
                "Baixar detalhamento por família CSV",
                df_familias.to_csv(index=False, sep=";").encode("ISO-8859-1", errors="replace"),
                "painel_cobertura_familias.csv",
                "text/csv",
                key="download_cobertura_familias"
            )

            st.divider()
            st.subheader("Itens de uma família específica")

            opcoes_familia = []
            for _, linha in df_familias.iterrows():
                label = (
                    f"{linha.get('cliente_base', '')} | "
                    f"{linha.get('fabrica', '')} | "
                    f"{linha.get('ce_bri', '')} | "
                    f"FAM {linha.get('familia', '')} | "
                    f"{linha.get('qtd_itens_sistema5', 0)} itens"
                )
                opcoes_familia.append(label)

            escolha_fam = st.selectbox(
                "Selecione uma família para ver os itens",
                opcoes_familia,
                key="select_cobertura_familia_itens"
            )

            idx_fam = opcoes_familia.index(escolha_fam)
            linha_fam = df_familias.iloc[idx_fam]

            df_itens_fam = cobertura_itens_por_familia(
                linha_fam.get("cliente_base", ""),
                linha_fam.get("fabrica", ""),
                linha_fam.get("ce_bri", ""),
                linha_fam.get("familia", "")
            )

            if df_itens_fam.empty:
                st.info("Essa família não possui itens no Sistema 5.")
            else:
                st.success(f"{len(df_itens_fam)} item(ns) encontrado(s) nesta família.")
                st.dataframe(normalizar_df_para_exibicao(df_itens_fam), use_container_width=True)

                st.download_button(
                    "Baixar itens da família CSV",
                    df_itens_fam.to_csv(index=False, sep=";").encode("ISO-8859-1", errors="replace"),
                    "painel_cobertura_itens_familia.csv",
                    "text/csv",
                    key="download_cobertura_itens_familia"
                )



# ==========================================
# ABA 9 - CORREÇÃO DE CÓDIGOS DE BARRAS
# ==========================================

if _is_active("correcao"):
    st.title("Correção de Códigos de Barras")

    st.warning(
        "Use esta aba para auditar códigos que podem ter perdido zero à esquerda. "
        "O sistema NÃO corrige tudo sozinho sem confirmação, para evitar alterar códigos válidos por engano."
    )

    st.subheader("Códigos suspeitos no Sistema 5")

    suspeitos_s5 = listar_codigos_suspeitos_sistema5()

    if suspeitos_s5.empty:
        st.success("Nenhum código suspeito encontrado no Sistema 5 ✅")
    else:
        st.warning(f"{len(suspeitos_s5)} código(s) suspeito(s) no Sistema 5.")
        st.dataframe(normalizar_df_para_exibicao(suspeitos_s5), use_container_width=True)

        st.download_button(
            "Baixar suspeitos Sistema 5 CSV",
            suspeitos_s5.to_csv(index=False, sep=";").encode("ISO-8859-1", errors="replace"),
            "codigos_suspeitos_sistema5.csv",
            "text/csv",
            key="download_suspeitos_s5"
        )

        st.divider()
        st.subheader("Corrigir item do Sistema 5")

        opcoes_s5 = []
        for _, linha in suspeitos_s5.iterrows():
            label = (
                f"ID {linha.get('id')} | "
                f"{linha.get('modelo', '')} | "
                f"{linha.get('codigo', '')} | "
                f"{linha.get('arquivo_nome', '')}"
            )
            opcoes_s5.append(label)

        escolha_s5 = st.selectbox(
            "Selecione o item suspeito do Sistema 5",
            opcoes_s5,
            key="select_codigo_suspeito_s5"
        )

        idx_s5 = opcoes_s5.index(escolha_s5)
        linha_s5 = suspeitos_s5.iloc[idx_s5]

        st.code(
            f"Modelo: {linha_s5.get('modelo', '')}\n"
            f"Código atual: {linha_s5.get('codigo', '')}\n"
            f"Família: {linha_s5.get('familia', '')}\n"
            f"Arquivo: {linha_s5.get('arquivo_nome', '')}"
        )

        novo_codigo_s5 = st.text_input(
            "Novo código correto",
            value=str(linha_s5.get("codigo", "")),
            key="novo_codigo_s5"
        )

        col_s5_a, col_s5_b = st.columns(2)

        with col_s5_a:
            if st.button("Salvar código manual no Sistema 5", key="salvar_codigo_manual_s5"):
                ok, mensagem = atualizar_codigo_sistema5_por_id(
                    linha_s5.get("id"),
                    novo_codigo_s5
                )
                if ok:
                    st.success(mensagem)
                else:
                    st.error(mensagem)

        with col_s5_b:
            tamanho_final_s5 = st.selectbox(
                "Completar com zero à esquerda até",
                [13, 14, 12],
                key="tamanho_zero_s5"
            )

            if st.button("Aplicar zero à esquerda no Sistema 5", key="aplicar_zero_s5"):
                ok, mensagem = aplicar_zero_esquerda_sistema5_por_id(
                    linha_s5.get("id"),
                    tamanho_final_s5
                )
                if ok:
                    st.success(mensagem)
                else:
                    st.error(mensagem)

    st.divider()

    st.subheader("Códigos suspeitos nos Certificados Oficiais")

    suspeitos_cert = listar_codigos_suspeitos_certificados()

    if suspeitos_cert.empty:
        st.success("Nenhum código suspeito encontrado nos certificados oficiais ✅")
    else:
        st.warning(f"{len(suspeitos_cert)} código(s) suspeito(s) nos certificados oficiais.")
        st.dataframe(normalizar_df_para_exibicao(suspeitos_cert), use_container_width=True)

        st.download_button(
            "Baixar suspeitos Certificados CSV",
            suspeitos_cert.to_csv(index=False, sep=";").encode("ISO-8859-1", errors="replace"),
            "codigos_suspeitos_certificados.csv",
            "text/csv",
            key="download_suspeitos_certificados"
        )

        st.divider()
        st.subheader("Corrigir item dos Certificados")

        opcoes_cert = []
        for _, linha in suspeitos_cert.iterrows():
            label = (
                f"ID {linha.get('id')} | "
                f"{linha.get('ip_bri', '')} | "
                f"{linha.get('modelo', '')} | "
                f"{linha.get('codigo', '')}"
            )
            opcoes_cert.append(label)

        escolha_cert = st.selectbox(
            "Selecione o item suspeito dos certificados",
            opcoes_cert,
            key="select_codigo_suspeito_cert"
        )

        idx_cert = opcoes_cert.index(escolha_cert)
        linha_cert = suspeitos_cert.iloc[idx_cert]

        st.code(
            f"IP-BRI: {linha_cert.get('ip_bri', '')}\n"
            f"Modelo: {linha_cert.get('modelo', '')}\n"
            f"Código atual: {linha_cert.get('codigo', '')}\n"
            f"Família: {linha_cert.get('familia', '')}"
        )

        novo_codigo_cert = st.text_input(
            "Novo código correto",
            value=str(linha_cert.get("codigo", "")),
            key="novo_codigo_cert"
        )

        col_cert_a, col_cert_b = st.columns(2)

        with col_cert_a:
            if st.button("Salvar código manual nos Certificados", key="salvar_codigo_manual_cert"):
                ok, mensagem = atualizar_codigo_certificado_por_id(
                    linha_cert.get("id"),
                    novo_codigo_cert
                )
                if ok:
                    st.success(mensagem)
                else:
                    st.error(mensagem)

        with col_cert_b:
            tamanho_final_cert = st.selectbox(
                "Completar com zero à esquerda até",
                [13, 14, 12],
                key="tamanho_zero_cert"
            )

            if st.button("Aplicar zero à esquerda nos Certificados", key="aplicar_zero_cert"):
                ok, mensagem = aplicar_zero_esquerda_certificado_por_id(
                    linha_cert.get("id"),
                    tamanho_final_cert
                )
                if ok:
                    st.success(mensagem)
                else:
                    st.error(mensagem)


# ==========================================
# ABA CONFERÊNCIA DE ETIQUETAS
# ==========================================

if _is_active("etiquetas"):
    st.title("🏷️ Conferência de Etiquetas")
    st.caption("Confere automaticamente se as etiquetas do Word batem com os dados do Excel ORDER LIST.")

    col_et1, col_et2 = st.columns(2)
    with col_et1:
        excel_etiqueta = st.file_uploader(
            "📊 Excel ORDER LIST",
            type=["xlsx", "xls"],
            key="excel_order_list"
        )
    with col_et2:
        word_etiqueta = st.file_uploader(
            "📄 Word com etiquetas",
            type=["docx"],
            key="word_etiquetas"
        )

    if excel_etiqueta and word_etiqueta:
        debug_manual = st.checkbox("🐛 Modo debug (mostrar detalhes técnicos de extração)", key="debug_etiquetas")
        try:
            import re as _re
            try:
                from docx import Document as _Document
            except ImportError:
                st.error(
                    "⛔ Biblioteca **python-docx** não está instalada.\n\n"
                    "Execute no terminal e reinicie o app:\n"
                    "```\npip install python-docx\n```"
                )
                st.stop()
            from lxml import etree as _ET
            from openpyxl import load_workbook as _lwb

            _ns = {'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}

            def _remover_desenhos_do_xlsx(_bytes_xlsx):
                """Remove os desenhos/imagens incorporados (xl/drawings/* e as
                imagens que eles referenciam) de dentro do .xlsx, ANTES de abrir
                com openpyxl. Bug conhecido: algumas imagens coladas em células
                (fotos de produto, comum em planilhas de packing list) têm um
                efeito/preenchimento no XML do desenho que o openpyxl não aceita
                (ValueError 'Min value is -1000000' ao processar
                SpreadsheetDrawing) — e isso quebra a leitura da planilha
                INTEIRA, mesmo só precisando dos valores das células. Não
                depende de qual extensão de imagem é (aqui eram .png/.jpeg, não
                os formatos que o _limpar_excel_imagens já tratava). Não afeta
                cor de preenchimento de CÉLULA (fica em xl/styles.xml, parte
                totalmente separada do arquivo — a detecção do cabeçalho por
                cor continua funcionando normalmente)."""
                import zipfile as _zfX, re as _reX
                from io import BytesIO as _BioX

                try:
                    with _zfX.ZipFile(_BioX(_bytes_xlsx)) as _zin:
                        _nomes = _zin.namelist()
                        if not any(n.startswith("xl/drawings/") for n in _nomes):
                            return _bytes_xlsx  # nada a remover, evita reescrever à toa

                        _saida = _BioX()
                        with _zfX.ZipFile(_saida, "w", _zfX.ZIP_DEFLATED) as _zout:
                            for _item in _zin.infolist():
                                _nome = _item.filename
                                if _nome.startswith("xl/drawings/") or _nome.startswith("xl/media/"):
                                    continue
                                _dados = _zin.read(_nome)
                                if _nome.startswith("xl/worksheets/sheet") and _nome.endswith(".xml"):
                                    _dados = _reX.sub(rb'<drawing[^>]*/>', b'', _dados)
                                elif _nome.startswith("xl/worksheets/_rels/") and _nome.endswith(".rels"):
                                    _dados = _reX.sub(
                                        rb'<Relationship[^>]*Type="[^"]*?/drawing"[^>]*/>', b'', _dados
                                    )
                                elif _nome == "[Content_Types].xml":
                                    _dados = _reX.sub(rb'<Override[^>]*PartName="/xl/drawings/[^"]*"[^>]*/>', b'', _dados)
                                    _dados = _reX.sub(rb'<Override[^>]*PartName="/xl/media/[^"]*"[^>]*/>', b'', _dados)
                                _zout.writestr(_item, _dados)
                        return _saida.getvalue()
                except Exception:
                    return _bytes_xlsx  # qualquer problema aqui -> segue com o original

            # ---- Parser do Excel ----
            def _abrir_xls_via_xlrd(file, erro_original):
                """Lê .xls legado direto com xlrd (sem LibreOffice) e devolve um
                workbook openpyxl em memória, compatível com o resto do código.
                Não lê cor de célula (xlrd não expõe isso de forma simples), então
                a detecção de cabeçalho cai pro método de texto 'REFERENCIA'."""
                try:
                    import xlrd as _xlrd
                except ImportError:
                    return None
                try:
                    from openpyxl import Workbook as _WB
                    file.seek(0)
                    _dados = file.read()
                    file.seek(0)
                    _wbx = _xlrd.open_workbook(file_contents=_dados)
                    _wsx = _wbx.sheet_by_index(0)
                    _wb_novo = _WB()
                    _ws_novo = _wb_novo.active
                    for _r in range(_wsx.nrows):
                        for _c in range(_wsx.ncols):
                            _val = _wsx.cell_value(_r, _c)
                            if _val != "":
                                _ws_novo.cell(row=_r+1, column=_c+1, value=_val)
                    return _wb_novo
                except Exception:
                    return None

            def _abrir_excel_com_fallback_libreoffice(file, erro_original):
                """Se o Excel não abrir com openpyxl (comum em .xls legado/binário),
                tenta primeiro ler com xlrd (rápido, sem LibreOffice); se não der,
                converte pra .xlsx via LibreOffice e tenta de novo."""
                # 1) Tenta xlrd (não depende do LibreOffice, evita o código 81)
                _wb_xlrd = _abrir_xls_via_xlrd(file, erro_original)
                if _wb_xlrd is not None:
                    if debug_manual:
                        st.caption("🔎 Excel .xls lido direto com xlrd (sem LibreOffice).")
                    return _wb_xlrd

                # 2) Fallback: converte via LibreOffice (com env isolado + nolockcheck)
                if not _soffice_path or _soffice_path == "soffice":
                    raise erro_original

                import tempfile as _tmpX
                import subprocess as _spX
                import os as _osX

                _dir_temp = _tmpX.mkdtemp(prefix="xls_conv_")
                try:
                    _nome_original = getattr(file, "name", "arquivo.xls")
                    _ext_original = _nome_original.rsplit(".", 1)[-1].lower() if "." in _nome_original else "xls"
                    _caminho_entrada = _osX.path.join(_dir_temp, f"entrada.{_ext_original}")

                    file.seek(0)
                    with open(_caminho_entrada, "wb") as _f:
                        _f.write(file.read())
                    file.seek(0)

                    _profile_dir = _tmpX.mkdtemp(prefix="lo_profile_xls_")
                    _profile_uri = _montar_uri_perfil(_profile_dir)

                    _env_xls = dict(_osX.environ)
                    _env_xls["HOME"] = _dir_temp
                    _env_xls["TMPDIR"] = _dir_temp
                    _env_xls["TMP"] = _dir_temp
                    _env_xls["TEMP"] = _dir_temp

                    _spX.run(
                        [_soffice_path, "--headless", "--norestore", "--invisible", "--nologo",
                         "--nofirststartwizard", "--nodefault", "--nolockcheck",
                         f"-env:UserInstallation={_profile_uri}",
                         "--convert-to", "xlsx",
                         _caminho_entrada, "--outdir", _dir_temp],
                        capture_output=True, timeout=180, env=_env_xls
                    )

                    _caminho_saida = _osX.path.join(_dir_temp, "entrada.xlsx")
                    if not _osX.path.exists(_caminho_saida):
                        raise erro_original

                    if debug_manual:
                        st.caption("🔎 Excel convertido de .xls (legado) para .xlsx via LibreOffice antes de abrir.")

                    return _lwb(_caminho_saida, data_only=True)
                finally:
                    import shutil as _shX
                    _shX.rmtree(_dir_temp, ignore_errors=True)

            def _parse_excel_order(file):
                # Remove desenhos/imagens ANTES de tentar abrir — proativo, não
                # só em resposta a exceção, porque planilhas de packing list
                # costumam ter fotos de produto coladas nas células, e é isso
                # que quebra o openpyxl (ver _remover_desenhos_do_xlsx). Não
                # precisamos das imagens aqui, só dos valores das células.
                try:
                    file.seek(0)
                    _bytes_sem_desenhos = _remover_desenhos_do_xlsx(file.read())
                    file.seek(0)
                except Exception:
                    _bytes_sem_desenhos = None

                try:
                    if _bytes_sem_desenhos is not None:
                        import io as _io_order
                        wb = _lwb(_io_order.BytesIO(_bytes_sem_desenhos), data_only=True)
                    else:
                        wb = _lwb(file, data_only=True)
                except Exception as _e_wb:
                    try:
                        file.seek(0)
                    except Exception:
                        pass
                    wb = _abrir_excel_com_fallback_libreoffice(file, _e_wb)
                ws = wb.active
                rows = list(ws.iter_rows(values_only=True))
                if not rows:
                    return {}

                # Procura a linha real do cabeçalho. Método principal: cor de
                # preenchimento azul que a fábrica usa nessas planilhas de fatura
                # pra marcar o cabeçalho (dados vêm em amarelo #FFFF00).
                # Aceita os dois azuis mais comuns: #00CCFF (neon) e #00B0F0 (celeste).
                # Fallback: procura o texto "REFERENCIA" caso a cor não seja encontrada.
                _CORES_CABECALHO = ("00CCFF", "00B0F0")
                _idx_cabecalho = None

                try:
                    for _i, _linha_celulas in enumerate(ws.iter_rows(min_row=1, max_row=min(30, ws.max_row))):
                        for _celula in _linha_celulas:
                            _fill = _celula.fill
                            _rgb = getattr(getattr(_fill, "fgColor", None), "rgb", None) if _fill else None
                            if _rgb and any(c in str(_rgb).upper() for c in _CORES_CABECALHO):
                                _idx_cabecalho = _i
                                break
                        if _idx_cabecalho is not None:
                            break
                except Exception:
                    _idx_cabecalho = None

                if debug_manual and _idx_cabecalho is not None:
                    st.caption(f"🔎 Cabeçalho do Excel encontrado pela cor azul na linha {_idx_cabecalho + 1}.")

                # Fallback por texto, caso a cor não seja encontrada
                if _idx_cabecalho is None:
                    for _i, _row in enumerate(rows[:30]):
                        _valores_linha = [str(v).strip().upper() if v else "" for v in _row]
                        if any(v == "REFERENCIA" or v.startswith("REFERENCIA") for v in _valores_linha):
                            _idx_cabecalho = _i
                            if debug_manual:
                                st.caption(f"🔎 Cabeçalho do Excel encontrado por texto (REFERENCIA) na linha {_i + 1} — cor azul não encontrada.")
                            break

                if _idx_cabecalho is None:
                    if debug_manual:
                        st.warning("⚠️ Não encontrei o cabeçalho nem pela cor azul nem pelo texto 'REFERENCIA' nas primeiras 30 linhas do Excel.")
                    return {}

                headers = [str(h).strip().upper() if h else "" for h in rows[_idx_cabecalho]]

                def col(row, name):
                    try:
                        i = next(i for i, h in enumerate(headers) if name in h)
                        v = row[i]
                        return str(v).strip() if v else ""
                    except StopIteration:
                        return ""

                itens = {}
                for row in rows[_idx_cabecalho + 1:]:
                    ref = col(row, "REFERENCIA")
                    if not ref or ref.upper() == "REFERENCIA":
                        continue
                    itens[ref] = {
                        'referencia': ref,
                        'marca':      col(row, "MARCA"),
                        'codigo':     col(row, "CÓDIGO DE BARRAS").replace(" ", ""),
                        'nome':       col(row, "NOME"),
                        'fabrica':    col(row, "FABRICA"),
                        'familia':    col(row, "FAMILIA"),
                        'registro':   col(row, "REGISTRO"),
                    }
                return itens

            # ---- Localizar executável do LibreOffice (soffice) ----
            _soffice_path = _localizar_soffice()
            if debug_manual:
                st.caption(f"🔎 Executável do LibreOffice localizado em: {_soffice_path}")
                if _soffice_path == "soffice":
                    st.warning("⚠️ Não encontrei o soffice nos caminhos padrão. Verifique se o LibreOffice está instalado.")

            # ---- Localizar executável do Tesseract OCR ----
            def _localizar_tesseract():
                import shutil as _sh1
                import os as _os1
                try:
                    _pasta_app = _os1.path.dirname(_os1.path.abspath(__file__))
                except Exception:
                    _pasta_app = ""

                candidatos = [
                    _sh1.which("tesseract"),
                    _sh1.which("tesseract.exe"),
                    # Pasta ao lado do app.py -> torna o pacote portátil
                    _os1.path.join(_pasta_app, "Tesseract-OCR", "tesseract.exe"),
                    _os1.path.expandvars(r"%LOCALAPPDATA%\Programs\Tesseract-OCR\tesseract.exe"),
                    r"C:\Program Files\Tesseract-OCR\tesseract.exe",
                    r"C:\Program Files (x86)\Tesseract-OCR\tesseract.exe",
                ]
                for c in candidatos:
                    if c and _os1.path.exists(c):
                        return c
                return None

            _tesseract_path = _localizar_tesseract()
            _ocr_disponivel = False
            if _tesseract_path:
                try:
                    import pytesseract as _pytess
                    _pytess.pytesseract.tesseract_cmd = _tesseract_path
                    _ocr_disponivel = True
                except ImportError:
                    _ocr_disponivel = False
            if debug_manual:
                if _tesseract_path:
                    st.caption(f"🔎 Tesseract OCR localizado em: {_tesseract_path} | pytesseract disponível: {_ocr_disponivel}")
                else:
                    st.warning("⚠️ Não encontrei o tesseract.exe. OCR do registro ficará desativado (regex em texto ainda será tentado).")

            # ---- Parser do Word (etiquetas) ----
            word_etiqueta.seek(0)
            _docx_bytes_global = word_etiqueta.read()
            word_etiqueta.seek(0)

            # ---- OTIMIZAÇÃO 1: converter TODOS os EMFs em UMA única chamada do LibreOffice ----
            # (em vez de abrir um processo novo do soffice pra cada etiqueta, o que é o maior custo)
            def _converter_emfs_em_lote(_docx_bytes, _soffice_exe):
                import zipfile as _zfL
                import tempfile as _tmpL
                import os as _osL
                import subprocess as _spL
                from io import BytesIO as _BioL

                _cache = {}
                if not _soffice_exe or _soffice_exe == "soffice":
                    return _cache

                _lote_dir = _tmpL.mkdtemp(prefix="emf_lote_")
                _mapa_arquivo_para_tgt = {}
                _diag_lote = {"erro": None, "retcode": None, "stderr": "", "soffice": _soffice_exe, "n_emf": 0, "n_png": 0}

                try:
                    with _zfL.ZipFile(_BioL(_docx_bytes)) as _zz:
                        _emf_tgts = [n[5:] for n in _zz.namelist() if n.startswith("word/media/") and n.lower().endswith(".emf")]
                        if not _emf_tgts:
                            return _cache

                        for _i, _tgt in enumerate(_emf_tgts):
                            _dados = _zz.read(f"word/{_tgt}")
                            _nome_local = f"item_{_i}.emf"
                            _caminho_local = _osL.path.join(_lote_dir, _nome_local)
                            with open(_caminho_local, "wb") as _f:
                                _f.write(_dados)
                            _mapa_arquivo_para_tgt[_nome_local] = _tgt

                    _diag_lote["n_emf"] = len(_mapa_arquivo_para_tgt)
                    _profile_dir = _tmpL.mkdtemp(prefix="lo_profile_lote_")
                    _profile_uri = _montar_uri_perfil(_profile_dir)

                    _arquivos_emf = [_osL.path.join(_lote_dir, n) for n in _mapa_arquivo_para_tgt]

                    # Ambiente isolado + --nolockcheck: evita o erro código 81 do LibreOffice
                    # (mesma correção aplicada na conversão do selo).
                    _env_lote = dict(_osL.environ)
                    _env_lote["HOME"] = _lote_dir
                    _env_lote["TMPDIR"] = _lote_dir
                    _env_lote["TMP"] = _lote_dir
                    _env_lote["TEMP"] = _lote_dir

                    _res_lote = _spL.run(
                        [_soffice_exe, "--headless", "--norestore", "--invisible", "--nologo",
                         "--nofirststartwizard", "--nodefault", "--nolockcheck",
                         f"-env:UserInstallation={_profile_uri}",
                         "--convert-to", "png",
                         *_arquivos_emf, "--outdir", _lote_dir],
                        capture_output=True, timeout=180, env=_env_lote
                    )
                    _diag_lote["retcode"] = _res_lote.returncode
                    _diag_lote["stderr"] = (_res_lote.stderr.decode('utf-8','ignore')[:300] if _res_lote.stderr else "")

                    for _nome_local, _tgt in _mapa_arquivo_para_tgt.items():
                        _png_path = _osL.path.join(_lote_dir, _nome_local.replace(".emf", ".png"))
                        if _osL.path.exists(_png_path):
                            with open(_png_path, "rb") as _f:
                                _cache[_tgt] = _f.read()
                    _diag_lote["n_png"] = len(_cache)

                    import shutil as _shL
                    _shL.rmtree(_profile_dir, ignore_errors=True)
                except Exception as _e_lote:
                    _diag_lote["erro"] = f"{type(_e_lote).__name__}: {_e_lote}"
                finally:
                    import shutil as _shL2
                    _shL2.rmtree(_lote_dir, ignore_errors=True)

                _cache["__diag__"] = _diag_lote
                return _cache

            def _converter_emfs_individual(_docx_bytes, _soffice_exe, _progress=None):
                """Fallback: converte as EMFs em blocos pequenos (5 por vez), cada bloco
                com processo e perfil isolados. Mais lento que o lote único, mas contorna
                o código 81 do LibreOffice (conflito de instância) que acontece em alguns PCs."""
                import zipfile as _zfI, tempfile as _tmpI, os as _osI
                import subprocess as _spI, glob as _globI, uuid as _uuidI
                from io import BytesIO as _BioI

                _cache = {}
                if not _soffice_exe or _soffice_exe == "soffice":
                    return _cache

                with _zfI.ZipFile(_BioI(_docx_bytes)) as _zz:
                    _emf_tgts = [n[5:] for n in _zz.namelist()
                                 if n.startswith("word/media/") and n.lower().endswith(".emf")]
                    _dados_por_tgt = {_t: _zz.read(f"word/{_t}") for _t in _emf_tgts}

                _total = len(_emf_tgts)
                _BLOCO = 5
                _feitos = 0
                for _ini in range(0, _total, _BLOCO):
                    _bloco = _emf_tgts[_ini:_ini + _BLOCO]
                    _base = _osI.path.join(_tmpI.gettempdir(),
                                           f"emf_ind_{_osI.getpid()}_{_uuidI.uuid4().hex}")
                    _osI.makedirs(_base, exist_ok=True)
                    _mapa = {}
                    for _j, _tgt in enumerate(_bloco):
                        _nl = f"bc_{_j}.emf"
                        with open(_osI.path.join(_base, _nl), "wb") as _f:
                            _f.write(_dados_por_tgt[_tgt])
                        _mapa[_nl] = _tgt
                    _perfil = _montar_uri_perfil(_osI.path.join(_base, "prof"))
                    _env = dict(_osI.environ)
                    _env["HOME"] = _base; _env["TMPDIR"] = _base
                    _env["TMP"] = _base; _env["TEMP"] = _base
                    try:
                        _spI.run(
                            [_soffice_exe, "--headless", "--norestore", "--invisible", "--nologo",
                             "--nofirststartwizard", "--nodefault", "--nolockcheck",
                             f"-env:UserInstallation={_perfil}",
                             "--convert-to", "png",
                             *[_osI.path.join(_base, n) for n in _mapa],
                             "--outdir", _base],
                            capture_output=True, timeout=90, env=_env
                        )
                        for _nl, _tgt in _mapa.items():
                            _pp = _osI.path.join(_base, _nl.replace(".emf", ".png"))
                            if _osI.path.exists(_pp):
                                with open(_pp, "rb") as _f:
                                    _cache[_tgt] = _f.read()
                    except Exception:
                        pass
                    finally:
                        import shutil as _shI
                        _shI.rmtree(_base, ignore_errors=True)
                    _feitos += len(_bloco)
                    if _progress:
                        _progress(_feitos, _total)
                return _cache

            def _converter_emfs_via_listener(_docx_bytes, _soffice_exe, _log_conf=None):
                """MÉTODO PRINCIPAL: converte todos os códigos de barras (EMF) do
                Word usando o listener persistente do LibreOffice — sem abrir
                uma instância nova pra cada um, o que evita o código 81. Retorna
                o mesmo formato que os métodos antigos: {arquivo_media: bytes_png}.
                _log_conf: lista opcional que recebe o log detalhado passo a passo
                de cada conversão (pra auditoria completa do que aconteceu)."""
                import zipfile as _zfU, tempfile as _tmpU, os as _osU
                from io import BytesIO as _BioU

                _cache = {}
                _suporte = _verificar_suporte_uno(_soffice_exe)
                if _log_conf is not None:
                    _log_conf.append(f"Suporte UNO: {_suporte}")
                if not _suporte["suportado"]:
                    return None  # sinaliza pra quem chamou tentar outro método

                if not _iniciar_listener_libreoffice(_soffice_exe):
                    if _log_conf is not None:
                        _log_conf.append("Falha ao iniciar o listener persistente.")
                    return None

                with _zfU.ZipFile(_BioU(_docx_bytes)) as _zz:
                    _emf_tgts = [n[5:] for n in _zz.namelist()
                                 if n.startswith("word/media/") and n.lower().endswith(".emf")]
                    _dados_por_tgt = {_t: _zz.read(f"word/{_t}") for _t in _emf_tgts}

                if _log_conf is not None:
                    _log_conf.append(f"Total de códigos de barras (EMF) a converter: {len(_dados_por_tgt)}")

                _base = _tmpU.mkdtemp(prefix="emf_uno_")
                try:
                    _caminhos = {}
                    for _tgt, _dados in _dados_por_tgt.items():
                        _nome_seguro = _tgt.replace("/", "_")
                        _emf_path = _osU.path.join(_base, _nome_seguro)
                        _png_path = _osU.path.join(_base, _nome_seguro + ".png")
                        with open(_emf_path, "wb") as _f:
                            _f.write(_dados)
                        _caminhos[_tgt] = (_emf_path, _png_path)

                    # OTIMIZAÇÃO: converte em BLOCOS (não tudo de uma vez, não um
                    # a um) numa mesma ponte UNO por bloco — reduz drasticamente
                    # as reconexões (era isso que dominava o tempo em lotes
                    # grandes: uma lista com ~200 códigos de barras pagava ~200
                    # reconexões, uma por arquivo) SEM depender de um processo
                    # gigante único: se um EMF travar, só o bloco dele (até
                    # _TAMANHO_BLOCO arquivos) espera o timeout, os demais blocos
                    # nem são afetados — mesma tolerância a falha de antes, só
                    # que ~25x menos reconexões.
                    _TAMANHO_BLOCO = 25
                    _tgts_ordenados = list(_caminhos.keys())
                    _faltando = dict(_dados_por_tgt)  # tgt -> dados, vai sendo reduzido
                    for _ini in range(0, len(_tgts_ordenados), _TAMANHO_BLOCO):
                        _bloco_tgts = _tgts_ordenados[_ini:_ini + _TAMANHO_BLOCO]
                        _jobs_bloco = [(_caminhos[_t][0], _caminhos[_t][1], "draw_png_Export") for _t in _bloco_tgts]
                        if _log_conf is not None:
                            _log_conf.append(f"Bloco {_ini // _TAMANHO_BLOCO + 1}: {len(_jobs_bloco)} arquivo(s)")
                        _resultado_bloco = _converter_lote_via_uno_listener(_jobs_bloco, _soffice_exe, _log=_log_conf)

                        if _resultado_bloco is not None:
                            for _tgt in _bloco_tgts:
                                _emf_path, _png_path = _caminhos[_tgt]
                                _ok, _erro = _resultado_bloco.get(_emf_path, (False, "sem resultado no bloco"))
                                if _ok and _osU.path.exists(_png_path):
                                    with open(_png_path, "rb") as _f:
                                        _cache[_tgt] = _f.read()
                                    _faltando.pop(_tgt, None)
                                elif _log_conf is not None:
                                    _log_conf.append(f"⚠️ {_tgt} não saiu no bloco ({_erro}) — tentando individualmente")
                        elif _log_conf is not None:
                            _log_conf.append(f"Bloco {_ini // _TAMANHO_BLOCO + 1} falhou por completo — caindo pro método arquivo-a-arquivo pra esse bloco.")

                    # Método de reserva: qualquer arquivo que não saiu no lote
                    # (ou o lote inteiro falhou) é convertido individualmente,
                    # exatamente como funcionava antes desta otimização.
                    for _tgt, _dados in _faltando.items():
                        _emf_path, _png_path = _caminhos[_tgt]
                        if _log_conf is not None:
                            _log_conf.append(f"--- Convertendo (individual) {_tgt} ---")
                        _ok, _erro = _converter_via_uno_listener(_emf_path, _png_path, "draw_png_Export", _soffice_exe, _log=_log_conf)
                        if _ok and _osU.path.exists(_png_path):
                            with open(_png_path, "rb") as _f:
                                _cache[_tgt] = _f.read()
                        elif _log_conf is not None:
                            _log_conf.append(f"❌ {_tgt} NÃO convertido: {_erro}")
                finally:
                    import shutil as _shU
                    _shU.rmtree(_base, ignore_errors=True)

                return _cache

            # Limpa qualquer LibreOffice pendurado antes de começar (garante início
            # limpo, evita código 81 causado por resquício de execução anterior) e
            # limpa de novo ao final (não deixa nada pra próxima conferência/geração).
            # MÉTODO PRINCIPAL: usa o listener persistente (uma instância sempre
            # aberta, sem abrir/fechar repetidamente — evita o código 81 na raiz).
            _log_conferencia = []
            with st.spinner("Convertendo códigos de barras..."):
                _emf_png_cache = _converter_emfs_via_listener(_docx_bytes_global, _soffice_path, _log_conf=_log_conferencia)

            if _emf_png_cache is not None:
                _n_convertidos = len(_emf_png_cache)
                _diag = None
                if debug_manual:
                    st.caption(f"🔎 Convertido via listener persistente: {_n_convertidos} código(s) de barras.")
            else:
                # Listener não disponível nesta instalação — cai pro método antigo
                if debug_manual:
                    st.caption("🔎 Listener persistente indisponível — usando método antigo (abre uma instância por conversão).")
                with _TarefaLibreOfficeLimpa():
                    with st.spinner("Convertendo códigos de barras (lote único)..."):
                        _emf_png_cache = _converter_emfs_em_lote(_docx_bytes_global, _soffice_path)

                    # Extrai o diagnóstico (e remove do cache pra não atrapalhar)
                    _diag = _emf_png_cache.pop("__diag__", None)
                    _n_convertidos = len(_emf_png_cache)

                    # FALLBACK: se o lote único falhou (0 convertidos ou código 81),
                    # tenta a conversão individual em blocos pequenos.
                    if _diag and _diag.get("n_emf", 0) > 0 and _n_convertidos == 0 and _soffice_path and _soffice_path != "soffice":
                        _barra = st.progress(0.0, text="O lote único falhou. Convertendo um a um (mais lento, porém confiável)...")
                        def _upd(feito, total):
                            _barra.progress(min(feito/total, 1.0), text=f"Convertendo códigos de barras: {feito}/{total}")
                        _emf_png_cache = _converter_emfs_individual(_docx_bytes_global, _soffice_path, _progress=_upd)
                        _barra.empty()
                        _n_convertidos = len(_emf_png_cache)
                        if _n_convertidos > 0:
                            st.success(f"✅ Conversão individual funcionou: {_n_convertidos} códigos de barras convertidos.")

            # (mensagem de debug já mostrada acima, em cada ramo — listener ou fallback)

            # Se converteu ZERO (ou o soffice não foi achado), avisa com o motivo real
            if _diag and _diag.get("n_emf", 0) > 0 and _n_convertidos == 0:
                _msg = "⚠️ **Nenhum código de barras pôde ser convertido.** "
                if not _soffice_path or _soffice_path == "soffice":
                    _msg += "O LibreOffice não foi encontrado no PC. Verifique se está instalado em C:\\LibreOfficePortable."
                elif _diag.get("erro"):
                    _msg += f"Erro: {_diag['erro']}"
                elif _diag.get("retcode") == 81:
                    _msg += (f"O LibreOffice retornou código 81 (conflito de instância). "
                             f"Feche TODAS as janelas do LibreOffice e recarregue a página. "
                             f"Executável usado: {_diag.get('soffice')}")
                else:
                    _msg += (f"LibreOffice usado: {_diag.get('soffice')} | "
                             f"código de saída={_diag.get('retcode')} | "
                             f"EMFs={_diag.get('n_emf')} PNGs gerados={_diag.get('n_png')} | "
                             f"stderr='{_diag.get('stderr','')}'")
                st.error(_msg)
            elif _diag and _diag.get("n_emf", 0) > 0 and _n_convertidos < _diag.get("n_emf", 0):
                st.warning(f"⚠️ Convertidos {_n_convertidos} de {_diag.get('n_emf')} códigos de barras. "
                           f"Alguns podem não ter sido lidos (código de saída={_diag.get('retcode')}).")

            # Log detalhado, passo a passo, de tudo que aconteceu na conversão dos
            # códigos de barras — pra auditar exatamente o que o sistema fez
            # (se achou o listener já aberto, se precisou reiniciar, cada erro).
            if _log_conferencia:
                with st.expander(f"📋 Log detalhado da conversão ({len(_log_conferencia)} linhas)", expanded=(_n_convertidos == 0)):
                    _texto_log = "\n".join(_log_conferencia)
                    st.code(_texto_log, language=None)
                    st.download_button(
                        "⬇️ Baixar log completo (.txt)",
                        _texto_log.encode("utf-8"),
                        "log_conferencia_etiquetas.txt",
                        "text/plain",
                        key="download_log_conferencia"
                    )

            # ---- OTIMIZAÇÃO 2: cache de OCR por conteúdo de imagem ----
            # (o mesmo selo de registro se repete em várias etiquetas -> só faz OCR uma vez)
            _ocr_cache = {}

            # ---- Corretor ortográfico (erro de escrita/digitação) ----
            _spell_disponivel = False
            try:
                from spellchecker import SpellChecker as _SpellChecker
                _corretor = _SpellChecker(language='pt')
                _spell_disponivel = True
            except ImportError:
                _corretor = None
            if debug_manual:
                st.caption(f"🔎 Corretor ortográfico (pyspellchecker) disponível: {_spell_disponivel}")

            # Termos técnicos/siglas/marcas conhecidas que não estão em dicionário comum
            # de português, mas são válidos nesse contexto -> não devem ser marcados como erro.
            _TERMOS_TECNICOS_CONHECIDOS = {
                "pvc", "sac", "cnpj", "cep", "inmetro", "innac", "ocp", "tuv",
                "exworks", "ipbri", "cebri", "bls", "toys", "gafi", "coaf",
                "eua", "ean", "rmb", "cbm", "ctns", "ctn", "pcs", "qty",
                "fax", "tel", "skype", "email",
                # Vocabulário regulatório/técnico válido, mas raro no dicionário genérico
                "dispersante", "dispersantes", "espessante", "espessantes",
                "hipersensível", "hipersensíveis", "hipersensivel", "hipersensiveis",
                "substituível", "substituíveis", "substituivel", "substituiveis",
            }

            _ortografia_cache = {}

            def _extrair_trechos_para_ortografia(texto_completo, referencia):
                """Isola só os trechos com texto de verdade (ATENÇÃO, INDICAÇÃO, ADVERTÊNCIA
                e a linha de referência/nome do produto) e ignora dados do importador/
                solicitante/CNPJ/SAC, que são informações do cliente e não devem ser
                verificadas como erro de escrita."""
                trechos = []

                m_atencao = _re.search(
                    r'ATEN[ÇC][ÃA]O\s*:?(.*?)(?=INDICA[ÇC][ÃA]O\s*:|ADVERT[ÊE]NCIA\s*:|$)',
                    texto_completo, _re.IGNORECASE | _re.DOTALL
                )
                if m_atencao:
                    trechos.append(m_atencao.group(1))

                m_indicacao = _re.search(
                    r'INDICA[ÇC][ÃA]O\s*:?(.*?)(?=ADVERT[ÊE]NCIA\s*:|$)',
                    texto_completo, _re.IGNORECASE | _re.DOTALL
                )
                if m_indicacao:
                    trechos.append(m_indicacao.group(1))

                m_advertencia = _re.search(
                    r'ADVERT[ÊE]NCIA\s*:?(.*?)(?=GUARDAR A EMBALAGEM|$)',
                    texto_completo, _re.IGNORECASE | _re.DOTALL
                )
                if m_advertencia:
                    trechos.append(m_advertencia.group(1))

                # Linha de referência/nome do produto (ex: "ZS-0316 - BRINQUEDO CONJUNTO...")
                if referencia:
                    _idx = texto_completo.upper().find(referencia.upper())
                    if _idx != -1:
                        _resto = texto_completo[_idx:_idx + 250]
                        _m_marca = _re.search(r'MARCA\s*:', _resto, _re.IGNORECASE)
                        if _m_marca:
                            _resto = _resto[:_m_marca.start()]
                        trechos.append(_resto)

                return " ".join(trechos)

            def _verificar_ortografia_etiqueta(texto_etiqueta, palavras_conhecidas_item):
                """Retorna lista de palavras que parecem erro de escrita (não é erro de português,
                é tipo letra faltando/trocada: 'rinquedo' em vez de 'brinquedo')."""
                if not _spell_disponivel or not texto_etiqueta:
                    return []

                import hashlib as _hashlib3
                _chave_cache = _hashlib3.md5(
                    (texto_etiqueta + "|" + ",".join(sorted(palavras_conhecidas_item))).encode("utf-8")
                ).hexdigest()
                if _chave_cache in _ortografia_cache:
                    return _ortografia_cache[_chave_cache]

                _palavras_texto = _re.findall(r'[A-ZÀ-Üa-zà-ü]{4,}', texto_etiqueta)
                _palavras_unicas = {}
                for _p in _palavras_texto:
                    _p_lower = _p.lower()
                    if _p_lower not in _palavras_unicas:
                        _palavras_unicas[_p_lower] = _p  # guarda a forma original pra exibir

                _candidatas = set(_palavras_unicas.keys()) - _TERMOS_TECNICOS_CONHECIDOS - palavras_conhecidas_item
                _desconhecidas = _corretor.unknown(_candidatas) if _candidatas else set()

                _resultado = sorted(_palavras_unicas[p] for p in _desconhecidas)
                _ortografia_cache[_chave_cache] = _resultado
                return _resultado

            def _parse_word_etiqueta(table, _docx_bytes_global=_docx_bytes_global):
                cell = table.rows[0].cells[0]
                root = _ET.fromstring(_ET.tostring(cell._element, encoding='unicode'))
                xml_str = _ET.tostring(cell._element, encoding='unicode')

                # Junta o texto respeitando a estrutura de parágrafos do Word:
                # - DENTRO de um parágrafo, sem espaço artificial entre runs -> evita
                #   quebrar palavra ao meio (o Word às vezes divide "APONTAR" em vários <w:t>)
                # - ENTRE parágrafos diferentes, insere espaço -> evita colar o fim de uma
                #   linha com o início da próxima, ex: "...DE PLÁSTICO" + "MARCA: BLS TOYS"
                #   virando "PLÁSTICOMARCA"
                _paragrafos = root.findall('.//w:p', _ns)
                _partes_texto = []
                for _p in _paragrafos:
                    _textos_paragrafo = [t.text for t in _p.findall('.//w:t', _ns) if t.text]
                    _partes_texto.append(''.join(_textos_paragrafo))
                texto = _re.sub(r'\s+', ' ', ' '.join(_partes_texto)).strip()


                # Referência (em negrito)
                referencia = None
                for run in root.findall('.//w:r', _ns):
                    rpr = run.find('w:rPr', _ns)
                    if rpr is not None and rpr.find('w:b', _ns) is not None:
                        t = run.find('w:t', _ns)
                        if t is not None and t.text:
                            bt = t.text.strip()
                            # Remove apóstrofo/aspas inicial que o Excel/Word às vezes coloca
                            # na frente de um texto pra forçá-lo como texto (ex: 'ZS-0207)
                            bt = bt.lstrip("'\"`´")
                            if _re.match(r'[A-Z0-9][\w-]+ - [A-ZÁÉÍÓÚÃÕÇ]', bt):
                                referencia = bt
                                break

                # Plano B: se a referência não veio em negrito (algumas etiquetas não
                # marcam a referência em negrito de forma consistente), usa a ESTRUTURA
                # de parágrafos em vez de tentar adivinhar o FORMATO do código: no
                # layout padrão, o parágrafo da referência é sempre o último parágrafo
                # não-vazio antes do parágrafo que começa com "MARCA:" (pode haver
                # parágrafos vazios entre eles, por causa do espaçamento da imagem do
                # chorão). Isso é muito mais confiável que uma regex baseada em contagem
                # de dígitos: códigos reais variam demais ("BG-A1-9", "BG-5", "BG-T1",
                # "BG-NVC5A", "BG-23-1"...), e a regex antiga deixava passar boa parte
                # das etiquetas sem achar referência nenhuma.
                if not referencia:
                    _partes_normalizadas = [_re.sub(r'\s+', ' ', p).strip() for p in _partes_texto]
                    _idx_marca = None
                    for _pi, _ptxt in enumerate(_partes_normalizadas):
                        if _re.match(r'MARCA\s*:', _ptxt, _re.IGNORECASE):
                            _idx_marca = _pi
                            break
                    if _idx_marca is not None:
                        for _pi in range(_idx_marca - 1, -1, -1):
                            if _partes_normalizadas[_pi]:
                                referencia = _partes_normalizadas[_pi]
                                break

                # Limpeza final: tira qualquer apóstrofo/aspa residual das pontas
                if referencia:
                    referencia = referencia.strip().strip("'\"`´").strip()

                # Marca
                m = _re.search(r'MARCA:\s+([A-Z][A-Z ]+?)(?:\s{2,}|\s*ATENÇÃO|\s*$)', texto)
                marca = m.group(1).strip() if m else None

                # Lote
                m = _re.search(r'Lote:\s*(\d{6})', texto)
                lote = m.group(1) if m else None

                # Data de Fabricação (ex: "Julho/2026")
                m = _re.search(
                    r'Data de Fabrica[çc][ãa]o:\s*([A-Za-zÀ-ü]+)\s*/\s*(\d{4})',
                    texto, _re.IGNORECASE
                )
                if m:
                    data_fabricacao_mes = m.group(1).strip()
                    data_fabricacao_ano = m.group(2).strip()
                    data_fabricacao = f"{data_fabricacao_mes}/{data_fabricacao_ano}"
                else:
                    data_fabricacao_mes = None
                    data_fabricacao_ano = None
                    data_fabricacao = None

                # Número de registro (selo "REGISTRO ... 004 871/2025 ... Compulsório")
                m = _re.search(
                    r'REGISTRO\D{0,20}?(?<!\d)(\d{3}\s?\d{3})(?!\d)\s*/\s*(\d{4})',
                    texto, _re.IGNORECASE
                )
                _reg_num = _re.sub(r'\s+', '', m.group(1)) if m else None
                registro = f"{_reg_num}/{m.group(2)}" if m else None

                # Indicativo
                m = _re.search(r'PARTIR DE\s+(\d+)\s+(ANOS?|MESES?)', texto, _re.IGNORECASE)
                indicativo = f"{m.group(1)} {m.group(2).upper()}" if m else None

                # Restritivo
                m = _re.search(r'MENORES DE\s+0?(\d+)\s*(ANOS?|MESES?)?', texto, _re.IGNORECASE)
                restritivo = f"{m.group(1)} {(m.group(2) or 'ANOS').upper()}" if m else None

                # ---- Processar imagens desta etiqueta ----
                tem_chorao = False
                tem_pilha = False
                tem_bateria_png = False
                codigos_barras_lidos = []

                try:
                    import zipfile as _zf2
                    import subprocess as _sp2
                    from PIL import Image as _Img2
                    from io import BytesIO as _Bio2
                    import tempfile as _tmp2
                    import os as _os2

                    with _zf2.ZipFile(_Bio2(_docx_bytes_global)) as _zzip:
                        _rels2 = _zzip.read("word/_rels/document.xml.rels").decode('utf-8')
                        _rmap = dict(_re.findall(r'Id="(rId\d+)"[^>]*Target="([^"]+)"', _rels2))
                        _embs = _re.findall(r'r:embed="(rId\d+)"', xml_str)

                        for _rid2 in _embs:
                            _tgt = _rmap.get(_rid2, '')
                            if not _tgt.startswith('media/'):
                                continue

                            _img_bytes = _zzip.read(f"word/{_tgt}")
                            _ext = _tgt.split('.')[-1].lower()

                            if debug_manual:
                                _tam_kb = len(_img_bytes) / 1024
                                _dims_txt = ""
                                if _ext in ('jpeg', 'jpg', 'png', 'bmp', 'gif'):
                                    try:
                                        _dw, _dh = _Img2.open(_Bio2(_img_bytes)).size
                                        _dims_txt = f" | dimensões={_dw}x{_dh}px"
                                    except Exception:
                                        _dims_txt = " | (não foi possível abrir com PIL)"
                                st.caption(
                                    f"🔎 [{referencia}] imagem encontrada: {_tgt} "
                                    f"| ext={_ext} | {_tam_kb:.1f}KB{_dims_txt}"
                                )

                            if _ext in ('jpeg', 'jpg', 'png'):
                                try:
                                    _img = _Img2.open(_Bio2(_img_bytes))
                                    w, h = _img.size

                                    # Chorão: JPEG pequeno ~123x129
                                    if 100 <= w <= 150 and 100 <= h <= 160:
                                        tem_chorao = True

                                    # Bateria/pilha PNG
                                    if _ext == 'png' and w < 200 and h < 200:
                                        tem_bateria_png = True
                                        tem_pilha = True

                                    # Selo "Segurança/REGISTRO/Compulsório" -> imagem maior,
                                    # tenta OCR pra extrair o número do registro
                                    if registro is None and w > 120 and h > 80 and _ocr_disponivel:
                                        import hashlib as _hashlib2
                                        _cache_key_ocr = _hashlib2.md5(_img_bytes).hexdigest()
                                        _cache_hit_ocr = _cache_key_ocr in _ocr_cache

                                    if registro is None and w > 120 and h > 80 and _ocr_disponivel and _cache_hit_ocr:
                                        registro = _ocr_cache[_cache_key_ocr]
                                        if debug_manual:
                                            st.caption(
                                                f"🔎 [{referencia}] OCR {_tgt}: resultado reaproveitado do cache -> "
                                                f"{registro or '(não encontrado)'}"
                                            )

                                    if registro is None and w > 120 and h > 80 and _ocr_disponivel and not _cache_hit_ocr:
                                        try:
                                            from PIL import ImageEnhance as _IE2
                                            from PIL import ImageFilter as _IF2
                                            from PIL import ImageChops as _ICh2
                                            from collections import Counter as _Counter2

                                            def _preparar_variantes(_base_img, _regiao):
                                                """Gera as variantes de imagem SOB DEMANDA (lazy). Antes, as 6
                                                versões (resize 4x + realce + threshold + máscaras de cor RGB)
                                                eram todas montadas de uma vez antes do OCR começar — trabalho
                                                pesado de PIL desperdiçado sempre que o early-exit já ia parar
                                                na 1ª ou 2ª variante. Agora cada variante só é calculada quando
                                                o loop de OCR realmente chega nela."""
                                                _cinza_base = _base_img.convert("L")
                                                _bw, _bh = _cinza_base.size
                                                # Escala adaptativa: garante pelo menos 400px de largura (mínimo
                                                # necessário pro Tesseract ler dígitos pequenos com precisão).
                                                # Cap em 6x pra não explodir memória em imagens já grandes.
                                                _escala = max(4, min(6, max(int(400 / max(_bw, 1)), int(300 / max(_bh, 1)))))
                                                _cinza = _cinza_base.resize((_bw * _escala, _bh * _escala), _Img2.LANCZOS)
                                                _cinza = _IE2.Contrast(_cinza).enhance(2.2)
                                                yield (_cinza, f'{_regiao}/contraste')
                                                # Threshold 140 — padrão
                                                yield (_cinza.point(lambda p: 255 if p > 140 else 0), f'{_regiao}/bin-140')
                                                # Threshold 110 — melhor pra texto escuro/cinza claro
                                                yield (_cinza.point(lambda p: 255 if p > 110 else 0), f'{_regiao}/bin-110')
                                                # Threshold 165 — melhor pra texto branco sobre fundo claro
                                                yield (_cinza.point(lambda p: 255 if p > 165 else 0), f'{_regiao}/bin-165')
                                                yield (_cinza.point(lambda p: 0 if p > 140 else 255), f'{_regiao}/binário-inv')
                                                yield (_cinza.filter(_IF2.UnsharpMask(radius=2, percent=200, threshold=2)), f'{_regiao}/nítida')
                                                # Denoised: blur suave antes do threshold reduz ruído de compressão JPEG
                                                _blur = _cinza.filter(_IF2.GaussianBlur(radius=1))
                                                yield (_blur.point(lambda p: 255 if p > 128 else 0), f'{_regiao}/denoised-bin')

                                                # Máscara por cor: isola texto realmente branco (R,G,B altos ao
                                                # mesmo tempo) ou realmente escuro (R,G,B baixos ao mesmo tempo).
                                                # Selos como o do INMETRO têm o número em branco sobre fundo
                                                # laranja/amarelo com marca d'água repetida ("Segurança Infantil")
                                                # na mesma cor do fundo -> converter pra cinza mistura tudo e o
                                                # threshold de luminância pega ruído da marca d'água junto com o
                                                # número. Olhando os 3 canais de cor separadamente, a marca d'água
                                                # nunca chega a R=G=B extremos ao mesmo tempo, só o texto real.
                                                _rgb = _base_img.convert("RGB").resize((_bw * _escala, _bh * _escala), _Img2.LANCZOS)
                                                _r, _g, _b = _rgb.split()
                                                _LIMIAR_CLARO = 185
                                                _LIMIAR_ESCURO = 90
                                                _r_claro = _r.point(lambda p: 255 if p > _LIMIAR_CLARO else 0)
                                                _g_claro = _g.point(lambda p: 255 if p > _LIMIAR_CLARO else 0)
                                                _b_claro = _b.point(lambda p: 255 if p > _LIMIAR_CLARO else 0)
                                                yield (_ICh2.darker(_ICh2.darker(_r_claro, _g_claro), _b_claro), f'{_regiao}/máscara-branca')
                                                _r_escuro = _r.point(lambda p: 255 if p < _LIMIAR_ESCURO else 0)
                                                _g_escuro = _g.point(lambda p: 255 if p < _LIMIAR_ESCURO else 0)
                                                _b_escuro = _b.point(lambda p: 255 if p < _LIMIAR_ESCURO else 0)
                                                yield (_ICh2.darker(_ICh2.darker(_r_escuro, _g_escuro), _b_escuro), f'{_regiao}/máscara-escura')

                                            # Múltiplos recortes: o INMETRO Safety Seal tem o número em
                                            # posições variadas dependendo da versão/escala. Tentamos:
                                            # 1) Faixa do meio (mais comum): entre 25%–75% da altura
                                            # 2) Faixa inferior (registros no rodapé): entre 55%–100%
                                            # 3) Faixa superior (registros no topo): entre 0%–50%
                                            # O early-exit no _tentar_ocr garante que só percorre os
                                            # recortes realmente necessários.
                                            _crop_centro = _img.crop((int(w * 0.05), int(h * 0.22), int(w * 0.95), int(h * 0.78)))
                                            _crop_baixo  = _img.crop((int(w * 0.05), int(h * 0.52), int(w * 0.95), h))
                                            _crop_cima   = _img.crop((int(w * 0.05), 0, int(w * 0.95), int(h * 0.52)))

                                            _achou_por_debug = []

                                            def _tentar_ocr(_variantes, _psms_texto, _psms_num):
                                                """Roda as tentativas e devolve lista de candidatos encontrados.
                                                OTIMIZAÇÃO: cada chamada ao tesseract sobe um processo novo (~100-
                                                300ms) — rodar TODAS as combinações de variante×psm×oem sempre
                                                (até 36 chamadas por imagem) é o maior custo de tempo da
                                                Conferência. Assim que 2 tentativas concordam no mesmo número já
                                                temos a mesma confiança que a votação por maioria daria no final,
                                                então paramos cedo. Casos difíceis (sem 2 concordâncias) continuam
                                                caindo no varrimento completo, mantendo a robustez de antes."""
                                                _candidatos = []

                                                def _ja_confiante():
                                                    if len(_candidatos) < 2:
                                                        return False
                                                    _, _n_top = _Counter2(_candidatos).most_common(1)[0]
                                                    return _n_top >= 2

                                                # Um único passe pela lista de variantes (agora um gerador lazy):
                                                # pra CADA variante, tenta os psms de texto e depois os de número
                                                # antes de passar pra próxima — isso é o que permite parar cedo
                                                # SEM nunca ter chegado a montar (resize/realce/máscara) as
                                                # variantes seguintes, que é o trabalho de PIL mais caro daqui.
                                                for _img_tentativa, _rotulo_img in _variantes:
                                                    for _psm in _psms_texto:
                                                        try:
                                                            _txt = _pytess.image_to_string(
                                                                _img_tentativa, lang='por', config=f'--psm {_psm} --oem 3'
                                                            )
                                                        except Exception:
                                                            _txt = ""
                                                        if debug_manual:
                                                            _achou_por_debug.append(
                                                                f"{_rotulo_img}/psm{_psm}: {_re.sub(chr(10), ' | ', _txt.strip())[:120]}"
                                                            )
                                                        # Padrão principal: NNN NNN / YYYY ou NNNNNN / YYYY
                                                        m_ = _re.search(
                                                            r'(?<!\d)(\d{3})\s{0,2}(\d{3})(?!\d)\s{0,3}[/]\s{0,3}(\d{4})(?!\d)',
                                                            _txt
                                                        )
                                                        if not m_:
                                                            # Padrão compacto sem separador visível: NNNNNN/YYYY
                                                            m_ = _re.search(r'(?<!\d)(\d{3})(\d{3})\s{0,3}/\s{0,3}(\d{4})(?!\d)', _txt)
                                                        if m_:
                                                            _candidatos.append(f"{m_.group(1)}{m_.group(2)}/{m_.group(3)}")
                                                            if _ja_confiante() and not debug_manual:
                                                                return _candidatos

                                                    for _psm in _psms_num:
                                                        for _oem in ('3', '0'):
                                                            try:
                                                                _txt_num = _pytess.image_to_string(
                                                                    _img_tentativa,
                                                                    config=f'--psm {_psm} --oem {_oem} -c tessedit_char_whitelist=0123456789/'
                                                                )
                                                            except Exception:
                                                                _txt_num = ""
                                                            if debug_manual:
                                                                _achou_por_debug.append(
                                                                    f"{_rotulo_img}/psm{_psm}/oem{_oem}/só-números: "
                                                                    f"{_re.sub(chr(10), ' | ', _txt_num.strip())[:120]}"
                                                                )
                                                            m_n = _re.search(r'(?<!\d)(\d{6})(?!\d)\s{0,3}/\s{0,3}(\d{4})(?!\d)', _txt_num)
                                                            if not m_n:
                                                                # Permite espaços entre grupos de 3 dígitos
                                                                m_n = _re.search(r'(?<!\d)(\d{3})\s{0,2}(\d{3})(?!\d)\s{0,3}/\s{0,3}(\d{4})(?!\d)', _txt_num)
                                                                if m_n:
                                                                    _candidatos.append(f"{m_n.group(1)}{m_n.group(2)}/{m_n.group(3)}")
                                                                    if _ja_confiante() and not debug_manual:
                                                                        return _candidatos
                                                                    continue
                                                            if m_n:
                                                                _candidatos.append(f"{m_n.group(1)}/{m_n.group(2)}")
                                                                if _ja_confiante() and not debug_manual:
                                                                    return _candidatos
                                                return _candidatos

                                            # Estratégia em cascata: começa pelo recorte mais provável e
                                            # só avança pro seguinte se não encontrou nada (early-exit).
                                            # PSM 6 = bloco uniforme, PSM 11 = texto esparso, PSM 3 = auto,
                                            # PSM 13 = linha crua (útil quando o número está isolado).
                                            _candidatos_registro = []
                                            for _crop_img, _crop_nome in (
                                                (_crop_centro, 'recorte-centro'),
                                                (_crop_baixo,  'recorte-baixo'),
                                                (_crop_cima,   'recorte-cima'),
                                            ):
                                                if _candidatos_registro:
                                                    break
                                                _variantes_crop = _preparar_variantes(_crop_img, _crop_nome)
                                                _candidatos_registro = _tentar_ocr(_variantes_crop, ('6', '11', '3', '13'), ('7', '8', '6'))

                                            # Fallback final: imagem inteira sem recorte
                                            if not _candidatos_registro:
                                                _variantes_completa = _preparar_variantes(_img, 'completa')
                                                _candidatos_registro = _tentar_ocr(_variantes_completa, ('6', '11', '3', '13'), ('7', '8', '6'))

                                            # Votação por maioria: usa o valor que mais se repetiu entre
                                            # todas as tentativas, em vez de confiar na primeira leitura
                                            if _candidatos_registro:
                                                _contagem = _Counter2(_candidatos_registro)
                                                registro, _n_votos = _contagem.most_common(1)[0]
                                                if debug_manual:
                                                    st.caption(
                                                        f"🔎 [{referencia}] Votação de registro: {dict(_contagem)} "
                                                        f"-> escolhido '{registro}' ({_n_votos} de {len(_candidatos_registro)} votos)"
                                                    )

                                            if debug_manual:
                                                for _linha_dbg in _achou_por_debug:
                                                    st.caption(f"🔎 [{referencia}] OCR {_tgt} — {_linha_dbg}")
                                        except Exception as _e_ocr:
                                            if debug_manual:
                                                st.warning(f"🔎 [{referencia}] erro no OCR de {_tgt}: {_e_ocr}")

                                        _ocr_cache[_cache_key_ocr] = registro

                                except Exception:
                                    pass

                            elif _ext == 'emf':
                                # Buscar PNG já convertido no lote (ver _converter_emfs_em_lote)
                                _debug = debug_manual
                                try:
                                    _png_bytes_cache = _emf_png_cache.get(_tgt)

                                    if _debug:
                                        if _png_bytes_cache:
                                            st.caption(f"🔎 [{referencia}] EMF {_tgt}: PNG obtido do cache em lote ✅")
                                        else:
                                            st.warning(f"⚠️ [{referencia}] EMF {_tgt}: não encontrado no cache em lote (conversão pode ter falhado)")

                                    if _png_bytes_cache:
                                        _img_orig = _Img2.open(_Bio2(_png_bytes_cache))

                                        # Múltiplas estratégias de leitura — código de barras
                                        # às vezes vem pequeno numa imagem grande (página inteira),
                                        # e o pyzbar falha se não estiver bem definido.
                                        from PIL import ImageEnhance as _IE
                                        from pyzbar.pyzbar import decode as _decode

                                        def _tentar_ler_barcode(_im):
                                            _w0, _h0 = _im.size
                                            _variantes_bc = []
                                            # 1) cinza + contraste (como era antes)
                                            _g = _IE.Contrast(_im.convert("L")).enhance(3.0)
                                            _variantes_bc.append(_g)
                                            # 2) imagem original
                                            _variantes_bc.append(_im)
                                            # 3) recorte da área com conteúdo (remove bordas brancas)
                                            try:
                                                from PIL import ImageOps as _IO
                                                _inv = _IO.invert(_im.convert("L"))
                                                _bbox = _inv.getbbox()
                                                if _bbox:
                                                    _crop = _im.crop(_bbox)
                                                    _variantes_bc.append(_crop)
                                                    # 4) recorte ampliado 2x
                                                    _cw, _ch = _crop.size
                                                    _variantes_bc.append(_crop.resize((_cw*2, _ch*2), _Img2.LANCZOS))
                                            except Exception:
                                                pass
                                            # 5) original ampliada (se for pequena)
                                            if _w0 < 400:
                                                _variantes_bc.append(_im.resize((_w0*2, _h0*2), _Img2.LANCZOS))

                                            for _v in _variantes_bc:
                                                _r = _decode(_v)
                                                if _r:
                                                    return _r
                                            return []

                                        _barcodes = _tentar_ler_barcode(_img_orig)
                                        if _debug:
                                            st.caption(f"🔎 [{referencia}] EMF {_tgt}: barcodes_encontrados={len(_barcodes)}")
                                        for _bc in _barcodes:
                                            _num = _bc.data.decode('utf-8').strip()
                                            if _num:
                                                codigos_barras_lidos.append(_num)
                                                tem_pilha = True  # EMF com barcode = imagem de pilha presente
                                except Exception as _e_emf:
                                    if _debug:
                                        st.warning(f"EMF debug [{_tgt}] erro geral: {_e_emf}")

                except Exception as _e_outer:
                    if debug_manual:
                        st.warning(f"🔎 [{referencia}] erro geral ao processar imagens: {_e_outer}")
                    # Fallback para detecção por XML
                    tem_chorao = 'Restri' in xml_str and '0-3' in xml_str
                    tem_pilha = bool(_re.search(r'PILHAS|BATERIAS SUBSTITU', texto, _re.IGNORECASE))

                return {
                    'referencia': referencia,
                    'marca': marca,
                    'lote': lote,
                    'registro': registro,
                    'indicativo': indicativo,
                    'restritivo': restritivo,
                    'tem_chorao': tem_chorao,
                    'tem_pilha': tem_pilha,
                    'codigos_barras': codigos_barras_lidos,
                    'texto_completo': texto,
                    'data_fabricacao': data_fabricacao,
                    'data_fabricacao_mes': data_fabricacao_mes,
                    'data_fabricacao_ano': data_fabricacao_ano,
                }

            # ---- Extrair indicativo/restritivo do NOME do Excel ----
            def _extrair_idade_nome(nome):
                nome_upper = nome.upper()
                ind = _re.search(r'INDICATIVO\s*\+(\d+)\s*(ANOS?|MESES?)', nome_upper)
                rest = _re.search(r'RESTRITIVO\s*-(\d+)\s*(ANOS?|MESES?)', nome_upper)
                # fallback sem palavra INDICATIVO/RESTRITIVO
                if not ind:
                    ind = _re.search(r'\+(\d+)\s*(ANOS?|MESES?)', nome_upper)
                if not rest:
                    rest = _re.search(r'-(\d+)\s*(ANOS?|MESES?)', nome_upper)
                ind_val = f"{ind.group(1)} {ind.group(2)}" if ind else None
                rest_val = f"{rest.group(1)} {rest.group(2)}" if rest else None
                return ind_val, rest_val

            def _tem_pilha_nome(nome):
                return bool(_re.search(r'PILHA|MOTOR .* PILHA|À PILHA', nome.upper()))

            # ---- Carregar dados ----
            with st.spinner("Processando arquivos..."):
                excel_etiqueta.seek(0)
                itens_excel = _parse_excel_order(excel_etiqueta)

                word_etiqueta.seek(0)
                doc_word = _Document(word_etiqueta)
                etiquetas_word = [_parse_word_etiqueta(t) for t in doc_word.tables if t.rows]

            st.info(f"📊 Excel: **{len(itens_excel)}** itens | 📄 Word: **{len(etiquetas_word)}** etiquetas")

            # ---- Detecção de etiquetas duplicadas ----
            _contagem_refs = {}
            for _e in etiquetas_word:
                _r = _e.get('referencia', '')
                if _r:
                    _contagem_refs[_r] = _contagem_refs.get(_r, 0) + 1
            _refs_duplicadas = {_r: _c for _r, _c in _contagem_refs.items() if _c > 1}
            if _refs_duplicadas:
                st.error(
                    f"🔁 **{len(_refs_duplicadas)} referência(s) duplicada(s) no Word** "
                    f"— a mesma etiqueta aparece mais de uma vez:"
                )
                for _rd, _cd in _refs_duplicadas.items():
                    st.caption(f"• `{_rd}` — aparece **{_cd}×**")

            st.divider()

            # ---- Gerador de Word corrigido ----
            def _gerar_word_corrigido(_doc_bytes, _resultados, _etiquetas_word, _itens_excel):
                """Retorna (bytes_corrigido, lista_de_correcoes_aplicadas)."""
                from io import BytesIO as _BioC
                _agora_c = datetime.now(ZoneInfo("America/Sao_Paulo"))
                _MESES_C = ['','Janeiro','Fevereiro','Março','Abril','Maio','Junho',
                            'Julho','Agosto','Setembro','Outubro','Novembro','Dezembro']
                _mes_vig = _MESES_C[_agora_c.month]
                _ano_vig = str(_agora_c.year)

                _doc_c = _Document(_BioC(_doc_bytes))
                _tables_c = [t for t in _doc_c.tables if t.rows]

                # ref -> índice de tabela
                _ref_idx = {}
                for _i, _etq in enumerate(_etiquetas_word):
                    _rf = _etq.get('referencia', '')
                    if _rf and _i < len(_tables_c):
                        _ref_idx[_rf] = _i

                def _setar_para(_para, _novo):
                    if _para.runs:
                        _para.runs[0].text = _novo
                        for _run in _para.runs[1:]:
                            _run.text = ''

                _log_corr = []

                for _res in _resultados:
                    _erros_c = _res.get('erros', [])
                    if not _erros_c:
                        continue
                    _ref_c = _res['ref']
                    _tidx = _ref_idx.get(_ref_c)
                    if _tidx is None:
                        continue
                    _tab = _tables_c[_tidx]

                    _fix_data      = False
                    _fix_reg       = None   # (errado_sem_espaco, certo)
                    _fix_ind       = None   # novo valor, ex: "3 ANOS"
                    _fix_rest      = None   # novo valor, ex: "3 ANOS"

                    for _err in _erros_c:
                        if 'Data de fabricação' in _err and 'esperado=' in _err:
                            _fix_data = True
                        elif 'Registro divergente' in _err:
                            _m = _re.search(r"etiqueta='([^']+)'.*?excel='([^']+)'", _err)
                            if _m:
                                _fix_reg = (
                                    _re.sub(r'\D', '', _m.group(1)),
                                    _m.group(2).strip()
                                )
                        elif 'Indicativo:' in _err:
                            _m = _re.search(r"nome='\+([^']+)'", _err)
                            if _m:
                                _fix_ind = _m.group(1).strip()
                        elif 'Restritivo:' in _err:
                            _m = _re.search(r"nome='-([^']+)'", _err)
                            if _m:
                                _fix_rest = _m.group(1).strip()

                    if not any([_fix_data, _fix_reg, _fix_ind, _fix_rest]):
                        continue

                    _seen_cells = set()
                    for _row in _tab.rows:
                        for _cell in _row.cells:
                            if id(_cell) in _seen_cells:
                                continue
                            _seen_cells.add(id(_cell))
                            for _para in _cell.paragraphs:
                                _txt_orig = _para.text
                                if not _txt_orig.strip():
                                    continue
                                _txt_novo = _txt_orig

                                if _fix_data:
                                    _txt_novo = _re.sub(
                                        r'([A-Za-zÀ-ü]+)\s*/\s*(\d{4})',
                                        f'{_mes_vig}/{_ano_vig}',
                                        _txt_novo, flags=_re.IGNORECASE
                                    )

                                if _fix_reg:
                                    _errado_num, _certo = _fix_reg
                                    if len(_errado_num) >= 6:
                                        _pat_reg = _re.escape(_errado_num[:3]) + r'\s?' + _re.escape(_errado_num[3:])
                                        _txt_novo = _re.sub(_pat_reg, _certo, _txt_novo, flags=_re.IGNORECASE)

                                if _fix_ind:
                                    _txt_novo = _re.sub(
                                        r'PARTIR DE\s+\d+\s+(?:ANOS?|MESES?)',
                                        f'PARTIR DE {_fix_ind}',
                                        _txt_novo, flags=_re.IGNORECASE
                                    )

                                if _fix_rest:
                                    _txt_novo = _re.sub(
                                        r'MENORES DE\s+0?\d+\s*(?:ANOS?|MESES?)?',
                                        f'MENORES DE {_fix_rest}',
                                        _txt_novo, flags=_re.IGNORECASE
                                    )

                                if _txt_novo != _txt_orig:
                                    _setar_para(_para, _txt_novo)
                                    _log_corr.append(
                                        f"[{_ref_c}] '{_txt_orig.strip()[:60]}' → '{_txt_novo.strip()[:60]}'"
                                    )

                _buf_c = _BioC()
                _doc_c.save(_buf_c)
                return _buf_c.getvalue(), _log_corr

            # ---- Conferência ----
            resultados = []

            for etiq in etiquetas_word:
                ref_etiq = etiq.get('referencia', '')
                if not ref_etiq:
                    continue

                # Busca no Excel pela referência da etiqueta
                item_excel = itens_excel.get(ref_etiq)

                # Tentativa de match parcial se não encontrou exato
                if not item_excel:
                    for ref_ex, dados_ex in itens_excel.items():
                        if ref_etiq.upper() in ref_ex.upper() or ref_ex.upper() in ref_etiq.upper():
                            item_excel = dados_ex
                            break

                if not item_excel:
                    resultados.append({
                        'ref': ref_etiq,
                        'status': '❌',
                        'erros': ['Referência não encontrada no Excel'],
                        'ok': [],
                        'avisos': [],
                        'prova_real': {
                            'cod_barras_excel': '—',
                            'cod_barras_etiqueta': ', '.join(etiq.get('codigos_barras', []) or []) or '—',
                            'registro_excel': '—',
                            'registro_etiqueta': etiq.get('registro') or '—',
                        },
                    })
                    continue

                erros = []
                ok = []
                avisos = []

                # 1. Marca
                if etiq.get('marca') and item_excel.get('marca'):
                    _marca_etiq_norm = _re.sub(r'\s+', ' ', etiq['marca'].upper()).strip()
                    _marca_excel_norm = _re.sub(r'\s+', ' ', item_excel['marca'].upper()).strip()
                    if _marca_etiq_norm == _marca_excel_norm:
                        ok.append('Marca ✅')
                    else:
                        erros.append(f"Marca: etiqueta='{etiq['marca']}' | excel='{item_excel['marca']}'")

                # 2. Indicativo de idade
                ind_excel, rest_excel = _extrair_idade_nome(item_excel.get('nome', ''))
                if etiq.get('indicativo') and ind_excel:
                    etiq_ind = etiq['indicativo'].upper().replace('ANOS', 'ANO').strip()
                    exc_ind = ind_excel.upper().replace('ANOS', 'ANO').strip()
                    if etiq_ind == exc_ind or etiq['indicativo'].split()[0] == ind_excel.split()[0]:
                        ok.append(f"Indicativo +{ind_excel} ✅")
                    else:
                        erros.append(f"Indicativo: etiqueta='+{etiq['indicativo']}' | nome='+{ind_excel}'")

                # 3. Restritivo
                if etiq.get('restritivo') and rest_excel:
                    if etiq['restritivo'].split()[0] == rest_excel.split()[0]:
                        ok.append(f"Restritivo -{rest_excel} ✅")
                    else:
                        erros.append(f"Restritivo: etiqueta='-{etiq['restritivo']}' | nome='-{rest_excel}'")

                # 4. Imagem chorão (obrigatória se tiver -3 anos)
                if rest_excel and rest_excel.startswith('3'):
                    if etiq.get('tem_chorao'):
                        ok.append('Imagem chorão ✅')
                    else:
                        erros.append('Imagem chorão ausente (produto -3 anos)')

                # 5. Imagem de pilha/bateria
                precisa_pilha = _tem_pilha_nome(item_excel.get('nome', ''))
                if precisa_pilha:
                    if etiq.get('tem_pilha'):
                        ok.append('Imagem pilha/bateria ✅')
                    else:
                        erros.append('Imagem de pilha/bateria ausente (produto com PILHA no nome)')

                # 6. Código de barras
                cod_excel = item_excel.get('codigo', '').replace(' ', '').strip()
                cods_lidos = etiq.get('codigos_barras', [])
                if cod_excel and cods_lidos:
                    if cod_excel in cods_lidos:
                        ok.append(f'Código de barras {cod_excel} ✅')
                    else:
                        erros.append(f"Código de barras divergente: etiqueta={cods_lidos} | excel={cod_excel}")
                elif cod_excel and not cods_lidos:
                    erros.append(f'Código de barras não detectado na etiqueta (esperado: {cod_excel})')

                # 7. Número de registro
                registro_excel = (item_excel.get('registro') or '').strip()
                registro_etiq = (etiq.get('registro') or '').strip()

                def _norm_registro(v):
                    return _re.sub(r'\D', '', v or '')

                def _diferenca_de_1_digito(a, b):
                    """True se as duas strings têm o mesmo tamanho e diferem em exatamente 1 caractere -> sinal de possível erro de OCR, não etiqueta errada."""
                    if len(a) != len(b) or len(a) == 0:
                        return False
                    return sum(1 for _x, _y in zip(a, b) if _x != _y) == 1

                if registro_excel and registro_etiq:
                    _norm_excel = _norm_registro(registro_excel)
                    _norm_etiq = _norm_registro(registro_etiq)
                    if _norm_excel == _norm_etiq:
                        ok.append(f'Registro {registro_etiq} ✅')
                    elif _diferenca_de_1_digito(_norm_excel, _norm_etiq):
                        avisos.append(
                            f"⚠️ Registro possivelmente lido errado pelo OCR (difere em 1 dígito) — "
                            f"confira a etiqueta física: etiqueta_lida='{registro_etiq}' | excel='{registro_excel}'"
                        )
                    else:
                        erros.append(f"Registro divergente: etiqueta='{registro_etiq}' | excel='{registro_excel}'")
                elif registro_excel and not registro_etiq:
                    avisos.append(
                        f"⚠️ Registro não lido pelo OCR (esperado: {registro_excel}) — "
                        f"confira o selo fisicamente. Pode ser artefato de compressão ou tamanho de imagem."
                    )

                # 8. Ortografia / erro de escrita (não confunde com erro de português/gramática)
                # Só verifica ATENÇÃO, INDICAÇÃO, ADVERTÊNCIA e a linha de referência/nome —
                # ignora dados do importador/solicitante/CNPJ/SAC (informação do cliente).
                _texto_etiqueta_completo = etiq.get('texto_completo', '')
                _trechos_relevantes = _extrair_trechos_para_ortografia(_texto_etiqueta_completo, ref_etiq)
                # Vocabulário conhecido do item = tudo que já vem validado no Excel
                # (nome + marca + referência). Assim palavras de marca (ex: "FOR KIDS")
                # nunca são marcadas como erro de escrita.
                _texto_vocab_item = " ".join([
                    (item_excel.get('nome', '') or ''),
                    (item_excel.get('marca', '') or ''),
                    (item_excel.get('referencia', '') or ''),
                    (ref_etiq or ''),
                ]).lower()
                _palavras_conhecidas_item = set(
                    _re.findall(r'[A-ZÀ-Üa-zà-ü]{3,}', _texto_vocab_item)
                )
                _erros_escrita = _verificar_ortografia_etiqueta(_trechos_relevantes, _palavras_conhecidas_item)
                if _erros_escrita:
                    avisos.append(
                        f"✏️ Possível erro de escrita na etiqueta (confira manualmente): {', '.join(_erros_escrita)}"
                    )

                # 9. Data de Fabricação = mês vigente (mês E ano atuais)
                _MESES_PT = {
                    'janeiro': 1, 'fevereiro': 2, 'março': 3, 'marco': 3, 'abril': 4,
                    'maio': 5, 'junho': 6, 'julho': 7, 'agosto': 8, 'setembro': 9,
                    'outubro': 10, 'novembro': 11, 'dezembro': 12,
                }
                _data_fab = etiq.get('data_fabricacao')
                _mes_etiq_txt = (etiq.get('data_fabricacao_mes') or '').strip().lower()
                _ano_etiq_txt = (etiq.get('data_fabricacao_ano') or '').strip()

                if _data_fab:
                    _agora = datetime.now(ZoneInfo("America/Sao_Paulo"))
                    _mes_atual = _agora.month
                    _ano_atual = _agora.year
                    _mes_etiq_num = _MESES_PT.get(_mes_etiq_txt)

                    if _mes_etiq_num == _mes_atual and _ano_etiq_txt == str(_ano_atual):
                        ok.append(f'Data de fabricação {_data_fab} ✅')
                    else:
                        _nomes_mes = ['', 'Janeiro','Fevereiro','Março','Abril','Maio','Junho',
                                      'Julho','Agosto','Setembro','Outubro','Novembro','Dezembro']
                        _mes_vigente_txt = f"{_nomes_mes[_mes_atual]}/{_ano_atual}"
                        erros.append(
                            f"📅 Data de fabricação fora do mês vigente: etiqueta='{_data_fab}' | esperado='{_mes_vigente_txt}'"
                        )
                # Se a etiqueta não trouxer data de fabricação, não bloqueia (não é erro,
                # só não dá pra conferir) — fica silencioso.

                # ---- Prova real: dados brutos lado a lado para conferência manual ----
                prova_real = {
                    'cod_barras_excel': cod_excel or '—',
                    'cod_barras_etiqueta': ', '.join(cods_lidos) if cods_lidos else '—',
                    'registro_excel': registro_excel or '—',
                    'registro_etiqueta': registro_etiq or '—',
                }

                resultados.append({
                    'ref': ref_etiq,
                    'status': '✅' if not erros and not avisos else ('⚠️' if not erros else '❌'),
                    'erros': erros,
                    'ok': ok,
                    'avisos': avisos,
                    'prova_real': prova_real,
                    'duplicada': ref_etiq in _refs_duplicadas,
                })

            # ---- Exibir resultados ----
            total_ok = sum(1 for r in resultados if r['status'] == '✅')
            total_err = sum(1 for r in resultados if r['status'] == '❌')
            total_aviso = sum(1 for r in resultados if r['status'] == '⚠️')

            col_r1, col_r2, col_r3, col_r4 = st.columns(4)
            col_r1.metric("Total conferidas", len(resultados))
            col_r2.metric("✅ OK", total_ok)
            col_r3.metric("⚠️ Confira manualmente", total_aviso)
            col_r4.metric("❌ Com problemas", total_err)

            # ---- Botão: Word com correções automáticas ----
            _tem_erros_corrigiveis = any(
                any(
                    ('Data de fabricação' in e and 'esperado=' in e)
                    or 'Registro divergente' in e
                    or 'Indicativo:' in e
                    or 'Restritivo:' in e
                    for e in r.get('erros', [])
                )
                for r in resultados
            )

            if _tem_erros_corrigiveis:
                st.info(
                    "✏️ Foram encontrados erros que podem ser corrigidos automaticamente "
                    "(data, registro, indicativo, restritivo). "
                    "Código de barras e imagens **não são corrigidos** — sinalizados apenas."
                )
                if st.button("📥 Gerar Word com correções automáticas", key="btn_gerar_word_corrigido", type="primary"):
                    try:
                        _bytes_corr, _log_corr = _gerar_word_corrigido(
                            _docx_bytes_global, resultados, etiquetas_word, itens_excel
                        )
                        st.download_button(
                            "⬇️ Baixar Word corrigido",
                            _bytes_corr,
                            "etiquetas_corrigidas.docx",
                            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                            key="dl_word_corrigido"
                        )
                        if _log_corr:
                            with st.expander(f"📋 {len(_log_corr)} substituição(ões) aplicada(s)", expanded=False):
                                for _lc in _log_corr:
                                    st.caption(_lc)
                        else:
                            st.warning("Nenhuma substituição foi aplicada (os campos podem estar em imagens, não em texto).")
                    except Exception as _ec:
                        st.error(f"Erro ao gerar Word corrigido: {_ec}")
                        import traceback as _tb
                        st.code(_tb.format_exc())

            st.divider()

            # Filtro
            filtro_conf = st.radio(
                "Mostrar:",
                ["Todas", "Apenas com erro", "Apenas OK"],
                horizontal=True,
                key="filtro_conferencia"
            )

            for r in resultados:
                if filtro_conf == "Apenas com erro" and r['status'] == '✅':
                    continue
                if filtro_conf == "Apenas OK" and r['status'] != '✅':
                    continue

                _dup_badge = "  🔁 DUPLICADA" if r.get('duplicada') else ""
                with st.expander(f"{r['status']} {r['ref'][:70]}{_dup_badge}"):
                    if r['erros']:
                        for e in r['erros']:
                            st.error(f"❌ {e}")
                    if r.get('avisos'):
                        for a in r['avisos']:
                            st.warning(a)
                    if r['ok']:
                        for o in r['ok']:
                            st.success(o)
                    if not r['erros'] and not r['ok'] and not r.get('avisos'):
                        st.info("Sem dados suficientes para conferir")

                    pr = r.get('prova_real')
                    if pr:
                        st.markdown("**🔍 Prova real (conferência manual)**")
                        pcol1, pcol2 = st.columns(2)
                        with pcol1:
                            st.caption("Código de barras — Excel")
                            st.code(pr['cod_barras_excel'])
                            st.caption("Registro — Excel")
                            st.code(pr['registro_excel'])
                        with pcol2:
                            st.caption("Código de barras — Etiqueta")
                            st.code(pr['cod_barras_etiqueta'])
                            st.caption("Registro — Etiqueta")
                            st.code(pr['registro_etiqueta'])

        except Exception as e:
            st.error(f"Erro ao processar: {e}")
            import traceback
            st.code(traceback.format_exc())


# ==========================================
# GERAR ETIQUETAS (módulo novo, em construção)
# ==========================================

if _is_active("gerar_etiquetas"):
    st.title("🎨 Gerar Etiquetas")
    st.caption("Módulo pra gerar etiquetas automaticamente a partir do Excel ORDER LIST. Fase 1: biblioteca de imagens (selos e assets genéricos).")

    with st.expander("ℹ️ Como vai funcionar (roadmap)"):
        st.markdown("""
Esse módulo está sendo construído por etapas:

1. ✅ **Fase 1 (esta tela):** biblioteca de selos "Segurança/REGISTRO" por Cliente+Fábrica+Família, e imagens genéricas reaproveitáveis (chorão, pilha/bateria).
2. ⏳ **Fase 2:** motor de geração — monta uma etiqueta (Word) por item do Excel, usando os selos cadastrados aqui + gerando o código de barras automaticamente + montando os textos de ATENÇÃO/INDICAÇÃO/ADVERTÊNCIA.
3. ⏳ **Fase 3:** geração em lote (todo o Excel de uma vez) + exportação também em PDF.

Cadastre os selos e os assets genéricos aqui primeiro — sem isso, a geração da etiqueta não tem imagem pra usar.
""")

    tab_gerar, tab_clientes, tab_selos, tab_textos, tab_genericos, tab_lista = st.tabs([
        "🏭 Gerar etiquetas",
        "👤 Clientes (importadores)",
        "📌 Modelos de selo",
        "📝 Textos padrão da etiqueta",
        "🖼️ Assets Genéricos (chorão/pilha)",
        "📋 Ver tudo cadastrado"
    ])

    with tab_gerar:
        st.subheader("Gerar etiquetas a partir do Excel")
        st.caption("Suba o mesmo Packing List (Excel) que você usa na conferência. O sistema lê os itens e monta as etiquetas.")

        # ETAPA 1: dados gerais da geração
        st.markdown("##### 1️⃣ Dados gerais")
        _clientes_disp = listar_clientes_etiqueta()
        if _clientes_disp.empty:
            st.warning("Nenhum cliente cadastrado. Cadastre um cliente na aba '👤 Clientes' antes de gerar.")
        else:
            colg1, colg2, colg3 = st.columns(3)
            with colg1:
                _cliente_gerar = st.selectbox("Cliente (importador)", _clientes_disp["apelido"].tolist(), key="gerar_cliente")
            with colg2:
                _origem_gerar = st.text_input("Origem", value="CHINA", key="gerar_origem")
            with colg3:
                _com_solicitante = st.radio("Tem solicitante?", ["Não", "Sim"], horizontal=True, key="gerar_tem_solic")

            _solicitante_cnpj = ""
            if _com_solicitante == "Sim":
                _solicitante_cnpj = st.text_input("CNPJ do solicitante", key="gerar_solic_cnpj",
                    help="Aparece quando um cliente empresta o registro para outro.")

            # De onde puxar o REGISTRO (busca por fábrica+família)
            st.markdown("**De onde puxar o número de registro?**")
            colr1, colr2 = st.columns(2)
            with colr1:
                _sistema_reg = st.radio("Sistema", ["Sistema 5 Novo Projeto (Bolsa)", "Sistema 5 Próprios", "Sistema 5 Focus"],
                    key="gerar_sistema_reg")
            with colr2:
                if _sistema_reg == "Sistema 5 Novo Projeto (Bolsa)":
                    _cliente_base_reg = "BOLSA"
                    st.info("Cliente base: **BOLSA**")
                elif _sistema_reg == "Sistema 5 Focus":
                    _cliente_base_reg = "FOCUS"
                    st.info("Cliente base: **FOCUS**")
                else:
                    _clientes_base_disp = listar_clientes_base_registros()
                    if _clientes_base_disp:
                        _cliente_base_reg = st.selectbox("Cliente (base de registros)", _clientes_base_disp, key="gerar_clientebase_reg")
                    else:
                        _cliente_base_reg = ""
                        st.warning("Nenhum cliente na base de Registros.")

            st.markdown("##### 2️⃣ Dados que valem pra todas as etiquetas deste lote")
            cold1, cold2 = st.columns(2)
            with cold1:
                _data_fab_gerar = st.text_input("Data de Fabricação", value="", placeholder="Ex: Julho/2026", key="gerar_data_fab")
            with cold2:
                _lote_gerar = st.text_input("Lote", value="", placeholder="Ex: 202607", key="gerar_lote")

            st.markdown("##### 3️⃣ Enviar o Excel (Packing List)")
            _excel_gerar = st.file_uploader("Excel do Packing List", type=["xls", "xlsx"], key="gerar_excel_upload")

            if _excel_gerar is not None:
                import pandas as _pd_g
                from openpyxl import load_workbook as _lwb_g

                # --- Ler o Excel e localizar cabeçalho (mesma lógica da conferência) ---
                @st.cache_data(show_spinner=False)
                def _carregar_itens_excel_geracao(_file_bytes, _nome):
                    import tempfile as _tmp_g, os as _os_g, subprocess as _sub_g, glob as _glob_g
                    _dir = _tmp_g.mkdtemp()
                    _path = _os_g.path.join(_dir, _nome)
                    with open(_path, "wb") as f:
                        f.write(_file_bytes)
                    # Converte .xls -> .xlsx se necessário
                    if _nome.lower().endswith(".xls"):
                        _sp = _localizar_soffice()
                        try:
                            _sub_g.run([_sp, "--headless", "--convert-to", "xlsx", "--outdir", _dir, _path,
                                        f"-env:UserInstallation=file://{_dir}/lo"], capture_output=True, timeout=90)
                            _xlsxs = _glob_g.glob(_os_g.path.join(_dir, "*.xlsx"))
                            if _xlsxs:
                                _path = _xlsxs[0]
                        except Exception:
                            pass
                    _wb = _lwb_g(_path, data_only=True)
                    _ws = _wb.active
                    # Acha cabeçalho por cor azul
                    _CORES_HDR = ("00CCFF", "00B0F0")
                    _idx = None
                    for _i, _linha in enumerate(_ws.iter_rows(min_row=1, max_row=min(30, _ws.max_row))):
                        for _cel in _linha:
                            _rgb = getattr(getattr(_cel.fill, 'fgColor', None), 'rgb', None) if _cel.fill else None
                            if _rgb and any(c in str(_rgb).upper() for c in _CORES_HDR):
                                _idx = _i; break
                        if _idx is not None: break
                    _rows = list(_ws.iter_rows(values_only=True))
                    if _idx is None:
                        for _i, _r in enumerate(_rows[:30]):
                            if any("REFERENCIA" in str(c).upper() for c in _r if c):
                                _idx = _i; break
                    if _idx is None:
                        return None
                    _headers = [str(h).strip() if h else "" for h in _rows[_idx]]
                    def _ci(nome):
                        for _j, _h in enumerate(_headers):
                            if nome in _h.upper(): return _j
                        return None
                    _cols = {k: _ci(k) for k in ['REFERENCIA','MARCA','CÓDIGO DE BARRAS','NOME','FABRICA','FAMILIA','REGISTRO']}
                    _itens = []
                    for _r in _rows[_idx+1:]:
                        _ref = _r[_cols['REFERENCIA']] if _cols['REFERENCIA'] is not None else None
                        if not _ref or str(_ref).strip().upper() in ("REFERENCIA", ""):
                            continue
                        def _get(col):
                            _ci_ = _cols.get(col)
                            if _ci_ is None or _ci_ >= len(_r): return ""
                            _v = _r[_ci_]
                            return str(_v).strip() if _v is not None else ""
                        _itens.append({
                            'referencia_full': _get('REFERENCIA'),
                            'marca': _get('MARCA'),
                            'codigo_barras': _get('CÓDIGO DE BARRAS'),
                            'nome': _get('NOME'),
                            'fabrica': _get('FABRICA'),
                            'familia': _get('FAMILIA'),
                            'registro': _get('REGISTRO'),
                        })
                    return _itens

                _excel_gerar.seek(0)
                _itens_lidos = _carregar_itens_excel_geracao(_excel_gerar.read(), _excel_gerar.name)

                if not _itens_lidos:
                    st.error("Não consegui identificar os itens no Excel. Confira se é o Packing List no padrão usado na conferência.")
                else:
                    st.success(f"✅ {len(_itens_lidos)} itens identificados no Excel.")

                    st.markdown("##### 4️⃣ Configurar cada item")
                    st.caption("Pra cada item: escolha o tipo de etiqueta, a variante de selo e a quantidade de peças. A idade é detectada do Excel (corrija se precisar).")

                    _TIPOS_ETIQUETA = {
                        "padrao": "Padrão",
                        "pilha": "C/ Pilha",
                        "maquiagem": "Maquiagem",
                        "massa": "Massa de Modelar",
                    }
                    _VARIANTES = list(_VARIANTES_SELO_VALIDAS.keys())
                    _df_textos_todos = listar_textos_etiqueta()
                    _NENHUM_TEXTO = "— nenhum —"

                    for _i, _item in enumerate(_itens_lidos):
                        _ref_curta, _nome_curto = separar_referencia_nome(_item['referencia_full'])
                        _idade_det = detectar_idade_do_nome(_item['nome'])
                        _tem_pilha_det = detectar_tem_pilha_do_nome(_item['nome'])
                        _tipo_sugerido = "pilha" if _tem_pilha_det else "padrao"

                        with st.expander(f"{_ref_curta} — {_nome_curto[:50]}", expanded=(_i < 3)):
                            cc1, cc2, cc3, cc4 = st.columns(4)
                            with cc1:
                                _tipo = st.selectbox("Tipo", list(_TIPOS_ETIQUETA.keys()),
                                    format_func=lambda t: _TIPOS_ETIQUETA[t],
                                    index=list(_TIPOS_ETIQUETA.keys()).index(_tipo_sugerido),
                                    key=f"tipo_{_i}")
                            with cc2:
                                _variante = st.selectbox("Selo", _VARIANTES,
                                    format_func=lambda v: _VARIANTES_SELO_VALIDAS[v], key=f"variante_{_i}")
                            with cc3:
                                _pecas = st.number_input("Peças", min_value=1, value=1, step=1, key=f"pecas_{_i}")
                            with cc4:
                                _idade_num = st.number_input("Idade (nº)", min_value=0,
                                    value=(_idade_det['numero'] if _idade_det else 3), step=1, key=f"idadenum_{_i}")
                                _idade_uni = st.selectbox("Unidade", ["ANOS", "MESES"],
                                    index=(0 if not _idade_det or _idade_det['unidade']=="ANOS" else 1), key=f"idadeuni_{_i}")
                            st.caption(f"Registro (do Excel): {_item['registro'] or '(vazio)'} | Código: {_item['codigo_barras'] or '(vazio)'}")

                            st.markdown("**Textos da etiqueta (ATENÇÃO/INDICAÇÃO/etc):**")
                            _categorias_deste_tipo = CATEGORIAS_POR_TIPO_ETIQUETA.get(_tipo, [])
                            _cols_txt = st.columns(len(_categorias_deste_tipo)) if _categorias_deste_tipo else []
                            for _cat, _col_txt in zip(_categorias_deste_tipo, _cols_txt):
                                with _col_txt:
                                    if not _df_textos_todos.empty:
                                        _opcoes_df = _df_textos_todos[
                                            (_df_textos_todos['tipo'] == _tipo) & (_df_textos_todos['categoria'] == _cat)
                                        ]
                                    else:
                                        _opcoes_df = _df_textos_todos
                                    _opcoes = [_NENHUM_TEXTO] + list(_opcoes_df['titulo'])
                                    st.selectbox(_cat.title(), _opcoes,
                                        index=(1 if len(_opcoes) > 1 else 0),
                                        key=f"txt_{_i}_{_cat}")
                            if not _categorias_deste_tipo:
                                st.caption("Esse tipo não usa textos de ATENÇÃO/INDICAÇÃO configuráveis.")

                    st.markdown("##### 5️⃣ Gerar")
                    if st.button("🏭 Gerar etiquetas (Word + PDF)", type="primary", key="btn_gerar_lote"):
                        _cli = buscar_cliente_etiqueta(_cliente_gerar)
                        if not _cli:
                            st.error("Cliente não encontrado.")
                        else:
                            # Monta a config lendo as escolhas de cada item do session_state
                            _itens_config = []
                            for _i, _item in enumerate(_itens_lidos):
                                _tipo_item = st.session_state.get(f"tipo_{_i}", "padrao")
                                _titulos_texto = {}
                                for _cat in CATEGORIAS_POR_TIPO_ETIQUETA.get(_tipo_item, []):
                                    _escolha = st.session_state.get(f"txt_{_i}_{_cat}")
                                    if _escolha and _escolha != "— nenhum —":
                                        _titulos_texto[_cat] = _escolha
                                _itens_config.append({
                                    'item': _item,
                                    'tipo': _tipo_item,
                                    'variante': st.session_state.get(f"variante_{_i}", _VARIANTES[0]),
                                    'pecas': st.session_state.get(f"pecas_{_i}", 1),
                                    'idade_num': st.session_state.get(f"idadenum_{_i}", 3),
                                    'idade_uni': st.session_state.get(f"idadeuni_{_i}", "ANOS"),
                                    'titulos_texto': _titulos_texto,
                                })
                            # Busca as imagens de chorão e pilha do banco (uma vez)
                            _chorao_bytes, _ = buscar_asset_generico("chorao")
                            _pilha_bytes, _ = buscar_asset_generico("pilha")

                            with st.spinner("Gerando etiquetas... (a conversão dos selos pode levar um tempo)"):
                                _resultado = _gerar_lote_etiquetas(
                                    _itens_config, _cli, _origem_gerar, _solicitante_cnpj,
                                    _data_fab_gerar, _lote_gerar, _cliente_base_reg,
                                    chorao_png_bytes=_chorao_bytes, pilha_png_bytes=_pilha_bytes
                                )
                            if _resultado and _resultado.get('docx'):
                                st.success("✅ Etiquetas geradas!")
                                st.download_button("⬇️ Baixar Word (.docx)", _resultado['docx'],
                                    "etiquetas.docx",
                                    "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                                    key="dl_etiq_docx")
                                st.caption("Pra ter em PDF: abra o Word e use 'Salvar como PDF'.")
                                if _resultado.get('avisos'):
                                    st.markdown("**Avisos:**")
                                    for _av in _resultado['avisos'][:40]:
                                        st.warning(_av)
                            else:
                                st.error("Não foi possível gerar as etiquetas.")

    with tab_clientes:
        st.subheader("Cadastro de clientes (importadores)")
        st.caption("Preencha manualmente conforme o cartão CNPJ. Esses dados vão no lado esquerdo da etiqueta (Importador, CNPJ, endereço, SAC, Origem). Tudo editável — se o cliente mudar algo, é só atualizar aqui.")

        _clientes_existentes = listar_clientes_etiqueta()
        _opcoes_edit = ["➕ Novo cliente"] + (_clientes_existentes["apelido"].tolist() if not _clientes_existentes.empty else [])
        _escolha_cliente = st.selectbox("Cliente", _opcoes_edit, key="sel_cliente_etiqueta_edit")

        _dados_edit = None
        if _escolha_cliente != "➕ Novo cliente":
            _dados_edit = buscar_cliente_etiqueta(_escolha_cliente)

        col_c1, col_c2 = st.columns(2)
        with col_c1:
            _apelido = st.text_input("Apelido/identificação (curto, pra você achar)", value=(_dados_edit["apelido"] if _dados_edit else ""), placeholder="Ex: ATEF", key="cli_apelido")
            _razao = st.text_input("Razão social (como no cartão CNPJ)", value=(_dados_edit["razao_social"] if _dados_edit else ""), placeholder="Ex: ATEF COMERCIO IMPORTACAO E EXPORTACAO DE BRINQUEDOS LTDA", key="cli_razao")
            _cnpj = st.text_input("CNPJ", value=(_dados_edit["cnpj"] if _dados_edit else ""), placeholder="Ex: 47.327.976/0002-18", key="cli_cnpj")
        with col_c2:
            _endereco = st.text_area("Endereço completo", value=(_dados_edit["endereco"] if _dados_edit else ""), placeholder="Rua, nº, sala, bairro, cidade - UF, CEP", height=100, key="cli_endereco")
            _sac = st.text_input("SAC (email)", value=(_dados_edit["sac"] if _dados_edit else ""), placeholder="Ex: atefcomercio@outlook.com", key="cli_sac")
            st.caption("ℹ️ A Origem (ex: CHINA) é perguntada na hora de gerar a etiqueta, não aqui — porque pode variar de um lote pra outro do mesmo cliente.")

        col_b1, col_b2 = st.columns([1, 1])
        with col_b1:
            if st.button("💾 Salvar cliente", key="salvar_cliente_etiqueta", type="primary"):
                if not (clean(_apelido) and clean(_razao) and clean(_cnpj)):
                    st.error("Preencha ao menos Apelido, Razão social e CNPJ.")
                else:
                    salvar_cliente_etiqueta(
                        _apelido, _razao, _cnpj, _endereco, _sac,
                        (_dados_edit["origem"] if _dados_edit else ""),
                        id_existente=(_dados_edit["id"] if _dados_edit else None)
                    )
                    st.success(f"Cliente '{_apelido}' salvo ✅")
                    st.cache_data.clear()
                    st.rerun()
        with col_b2:
            if _dados_edit and st.button("🗑️ Excluir este cliente", key="excluir_cliente_etiqueta"):
                excluir_cliente_etiqueta(_dados_edit["id"])
                st.success("Cliente excluído ✅")
                st.cache_data.clear()
                st.rerun()

        if not _clientes_existentes.empty:
            st.divider()
            st.markdown("**Clientes cadastrados:**")
            st.dataframe(normalizar_df_para_exibicao(_clientes_existentes), use_container_width=True)

    with tab_selos:
        st.subheader("Modelos de selo (Word editável)")
        st.caption("Cadastre UMA VEZ os 4 modelos de selo em Word. O número de registro dentro deles é editável — na geração da etiqueta, o sistema substitui automaticamente pelo registro correto do cliente (vindo da aba Registros). Não usa mais OCR.")

        st.info("Cada modelo é um arquivo .docx com o selo desenhado e o número de registro como texto. O sistema detecta o número sozinho e troca pelo certo na hora de gerar.")

        for _vk, _vlabel in _VARIANTES_SELO_VALIDAS.items():
            with st.container(border=True):
                _modelo_atual = buscar_modelo_selo(_vk)
                col_m1, col_m2 = st.columns([2, 1])
                with col_m1:
                    st.markdown(f"**{_vlabel}**")
                    if _modelo_atual:
                        st.success(f"✅ Cadastrado: {_modelo_atual['nome']} | registro detectado no modelo: {_modelo_atual['registro_exemplo'] or '(não detectado)'}")
                    else:
                        st.warning("Ainda não cadastrado.")
                with col_m2:
                    if _modelo_atual and st.button("🗑️ Remover", key=f"remover_modelo_{_vk}"):
                        excluir_modelo_selo(_vk)
                        st.success("Modelo removido ✅")
                        st.cache_data.clear()
                        st.rerun()

                _up_modelo = st.file_uploader(f"Enviar/substituir modelo {_vlabel}", type=["docx"], key=f"upload_modelo_{_vk}")
                if _up_modelo and st.button(f"💾 Salvar modelo {_vlabel}", key=f"salvar_modelo_{_vk}"):
                    _up_modelo.seek(0)
                    _reg_detectado = salvar_modelo_selo(_vk, _up_modelo.read(), _up_modelo.name)
                    if _reg_detectado:
                        st.success(f"Modelo salvo ✅ | Número de registro detectado no modelo: {_reg_detectado} (será substituído na geração)")
                    else:
                        st.warning("Modelo salvo, mas não achei um número de registro (formato NNN NNN/AAAA) dentro dele. Confira se o número está como texto editável no Word.")
                    st.cache_data.clear()
                    st.rerun()

        st.divider()
        st.markdown("**🧪 Testar substituição de registro**")
        st.caption("Veja como fica um selo com um registro específico, antes de usar na geração real.")
        col_t1, col_t2 = st.columns(2)
        with col_t1:
            _var_teste = st.selectbox("Variante", list(_VARIANTES_SELO_VALIDAS.keys()), format_func=lambda v: _VARIANTES_SELO_VALIDAS[v], key="var_teste_selo")
        with col_t2:
            _reg_teste = st.text_input("Registro (formato do banco, ex: 003565/2023)", key="reg_teste_selo")
        if st.button("Gerar selo de teste", key="btn_teste_selo"):
            if not clean(_reg_teste):
                st.error("Informe um número de registro.")
            else:
                _docx_gerado = gerar_selo_com_registro(_var_teste, _reg_teste)
                if _docx_gerado:
                    st.download_button(
                        "⬇️ Baixar selo de teste (.docx)",
                        _docx_gerado,
                        f"selo_teste_{_var_teste}.docx",
                        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                        key="download_selo_teste"
                    )
                    st.success("Selo gerado! Baixe e confira se o número ficou correto.")
                else:
                    st.error(f"O modelo '{_VARIANTES_SELO_VALIDAS[_var_teste]}' ainda não foi cadastrado.")

    with tab_textos:
        st.subheader("Textos padrão da etiqueta")
        st.caption("Cadastre aqui os blocos de texto que entram nas etiquetas (advertências, instruções de pilha, restritivos de idade, indicações etc). Na geração, você escolhe qual texto usar em cada item.")

        with st.expander("📥 Importar os textos extraídos dos seus modelos reais (MODELO_ETIQUETA.docx)", expanded=False):
            st.caption(
                f"Insere de uma vez só os {len(_TEXTOS_PADRAO_SEED)} textos revisados com você "
                "(ATENÇÃO/INDICAÇÃO/ADVERTÊNCIA/CUIDADOS DE USO/COMPOSIÇÃO), separados por tipo de etiqueta. "
                "Pode rodar mais de uma vez sem duplicar — só insere o que ainda não existe."
            )
            if st.button("Importar agora", key="btn_importar_textos_seed"):
                _inseridos, _ja_existiam = importar_textos_padrao_seed()
                st.success(f"✅ {_inseridos} texto(s) importado(s). {_ja_existiam} já existiam e foram ignorados.")
                st.cache_data.clear()
                st.rerun()

        st.info("💡 **Como marcar negrito:** coloque o trecho entre dois asteriscos de cada lado. Ex: `**ATENÇÃO:** não recomendável...` — na etiqueta gerada, 'ATENÇÃO:' fica em negrito e os asteriscos somem.")

        st.info("🎯 **Marcador de idade:** onde entra a idade/meses, escreva `{IDADE}`. Na geração, o sistema detecta do Excel e substitui automaticamente (ex: '03 (TRÊS) ANOS' ou '06 MESES'). Você poderá corrigir na hora se ele errar.\n\nEx: `ATENÇÃO: NÃO RECOMENDÁVEL PARA CRIANÇAS MENORES DE {IDADE} POR CONTER PARTES PEQUENAS...`")

        _CATEGORIAS_TEXTO = ["ATENÇÃO", "INDICAÇÃO", "ADVERTÊNCIA", "CUIDADOS DE USO", "COMPOSIÇÃO", "RESTRITIVO", "OUTROS"]
        _TIPOS_TEXTO = {"padrao": "Padrão", "pilha": "C/ Pilha", "maquiagem": "Maquiagem", "massa": "Massa de Modelar"}

        col_txt1, col_txt2 = st.columns([1, 2])
        with col_txt1:
            tipo_texto = st.selectbox("Tipo de etiqueta", list(_TIPOS_TEXTO.keys()), format_func=lambda t: _TIPOS_TEXTO[t], key="tipo_texto_etiqueta")
            cat_texto = st.selectbox("Categoria", _CATEGORIAS_TEXTO, key="cat_texto_etiqueta")
            titulo_texto = st.text_input("Título/apelido (pra você identificar)", placeholder="Ex: Advertência padrão pilha", key="titulo_texto_etiqueta")
        with col_txt2:
            conteudo_texto = st.text_area("Conteúdo do texto (use **texto** para negrito)", height=150, placeholder="Ex: **ATENÇÃO:** NÃO RECOMENDÁVEL PARA CRIANÇAS MENORES DE 03 ANOS...", key="conteudo_texto_etiqueta")

        # Pré-visualização ao vivo de como vai ficar (negrito aplicado, asteriscos removidos)
        if clean(conteudo_texto):
            st.markdown("**Pré-visualização (como vai aparecer na etiqueta):**")
            _preview_texto = conteudo_texto
            if "{IDADE}" in _preview_texto:
                _preview_texto = _preview_texto.replace("{IDADE}", "03 (TRÊS) ANOS")
                st.caption("(No exemplo abaixo, {IDADE} foi preenchido com '03 (TRÊS) ANOS' só pra ilustrar — na geração real, vem do Excel.)")
            st.markdown(
                f"<div style='border:1px solid #ddd;border-radius:6px;padding:10px;background:#fafafa'>{texto_para_preview_html(_preview_texto)}</div>",
                unsafe_allow_html=True
            )

        if st.button("💾 Salvar texto", key="salvar_texto_etiqueta_btn"):
            if not (clean(titulo_texto) and clean(conteudo_texto)):
                st.error("Preencha o título e o conteúdo do texto.")
            else:
                salvar_texto_etiqueta(cat_texto, titulo_texto, conteudo_texto, tipo=tipo_texto)
                st.success("Texto salvo ✅")
                st.cache_data.clear()
                st.rerun()

        st.divider()
        st.markdown("**Textos já cadastrados:**")
        df_textos = listar_textos_etiqueta()
        if df_textos.empty:
            st.info("Nenhum texto cadastrado ainda.")
        else:
            for _tipo_grupo, _df_grupo in df_textos.groupby('tipo'):
                st.markdown(f"##### {_TIPOS_TEXTO.get(_tipo_grupo, _tipo_grupo)}")
                for _, row in _df_grupo.iterrows():
                    with st.expander(f"[{row['categoria']}] {row['titulo']}"):
                        # Mostra já formatado (negrito aplicado), como vai sair na etiqueta
                        st.markdown(
                            f"<div style='border:1px solid #eee;border-radius:6px;padding:8px;background:#fafafa'>{texto_para_preview_html(row['conteudo'])}</div>",
                            unsafe_allow_html=True
                        )
                        st.caption(f"Texto salvo (com marcação): {row['conteudo'][:150]}{'...' if len(str(row['conteudo'])) > 150 else ''}")
                        if st.button("🗑️ Excluir", key=f"excluir_texto_{row['id']}"):
                            excluir_texto_etiqueta(row['id'])
                            st.success("Texto excluído ✅")
                            st.cache_data.clear()
                            st.rerun()

    with tab_genericos:
        st.subheader("Imagens genéricas reaproveitáveis")
        st.caption("Essas imagens são as mesmas em qualquer etiqueta (não mudam por cliente/fábrica) — cadastra uma vez só e o sistema reaproveita sempre que aplicável (chorão em produtos -3 anos, pilha em produtos com PILHA no nome).")

        col_gen1, col_gen2 = st.columns(2)

        with col_gen1:
            st.markdown("**Imagem do chorão** (ícone -3 anos)")
            chorao_atual_bytes, chorao_atual_nome = buscar_asset_generico("chorao")
            if chorao_atual_bytes:
                st.image(chorao_atual_bytes, caption=f"Atual: {chorao_atual_nome}", width=150)
            else:
                st.info("Nenhuma imagem de chorão cadastrada ainda.")

            novo_chorao = st.file_uploader("Enviar/substituir imagem do chorão", type=["jpg", "jpeg", "png"], key="upload_chorao_generico")
            if novo_chorao and st.button("💾 Salvar imagem do chorão", key="salvar_chorao_generico"):
                novo_chorao.seek(0)
                salvar_asset_generico("chorao", novo_chorao.read(), novo_chorao.name)
                st.success("Imagem do chorão salva ✅")
                st.cache_data.clear()
                st.rerun()

        with col_gen2:
            st.markdown("**Imagem de pilha/bateria**")
            pilha_atual_bytes, pilha_atual_nome = buscar_asset_generico("pilha")
            if pilha_atual_bytes:
                st.image(pilha_atual_bytes, caption=f"Atual: {pilha_atual_nome}", width=150)
            else:
                st.info("Nenhuma imagem de pilha/bateria cadastrada ainda.")

            nova_pilha = st.file_uploader("Enviar/substituir imagem de pilha/bateria", type=["jpg", "jpeg", "png"], key="upload_pilha_generico")
            if nova_pilha and st.button("💾 Salvar imagem de pilha/bateria", key="salvar_pilha_generico"):
                nova_pilha.seek(0)
                salvar_asset_generico("pilha", nova_pilha.read(), nova_pilha.name)
                st.success("Imagem de pilha/bateria salva ✅")
                st.cache_data.clear()
                st.rerun()

    with tab_lista:
        st.subheader("Selos cadastrados por Cliente/Fábrica/Família")
        df_selos = listar_selos_registro()

        if df_selos.empty:
            st.info("Nenhum selo cadastrado ainda.")
        else:
            st.dataframe(normalizar_df_para_exibicao(df_selos), use_container_width=True)

            opcoes_excluir_selo = [
                f"ID {row['id']} | {row['cliente_base']} | {row['fabrica']} | Família {row['familia']}"
                for _, row in df_selos.iterrows()
            ]
            escolha_excluir_selo = st.selectbox("Selo pra excluir", opcoes_excluir_selo, key="select_excluir_selo")
            if st.button("🗑️ Excluir selo selecionado", key="btn_excluir_selo"):
                id_para_excluir = int(escolha_excluir_selo.split("ID ")[1].split(" |")[0])
                excluir_selo_registro(id_para_excluir)
                st.success("Selo excluído ✅")
                st.cache_data.clear()
                st.rerun()

        st.divider()
        st.subheader("Assets genéricos cadastrados")
        df_genericos = listar_assets_genericos()
        if df_genericos.empty:
            st.info("Nenhum asset genérico cadastrado ainda.")
        else:
            st.dataframe(normalizar_df_para_exibicao(df_genericos), use_container_width=True)


# ==========================================
# PÁGINA: GERAR DESCRIÇÕES
# ==========================================
if _is_active("gerar_descricoes"):
    import re as _re_gd
    import base64 as _b64_gd

    st.title("✍️ Gerar Descrições")
    st.caption(
        "Gera descrições técnicas no padrão do banco de dados a partir de um arquivo "
        "com medidas e quantidades de peças. As imagens dos produtos embutidas no Excel "
        "são extraídas e exibidas para identificação visual de cada item."
    )

    # ---- Análise dos padrões existentes no banco ----
    # Formato real das descrições:
    # "{TIPO PRODUTO}, MEDIDAS {L}X{A}X{P}CM, MECANISMO {MEC}, C/ {N} PEÇAS"
    # As medidas ficam APÓS a palavra "MEDIDAS" e ANTES de ", MECANISMO".
    def _analisar_padroes_gd():
        _nomes_raw = pd.read_sql_query(
            """SELECT DISTINCT UPPER(TRIM(nome)) AS nome
               FROM sistema5_itens
               WHERE nome IS NOT NULL AND TRIM(nome) != ''
               ORDER BY nome""",
            conn
        )['nome'].tolist()

        # Captura dimensões depois de "MEDIDAS" (palavra-chave obrigatória)
        _pat_med = _re_gd.compile(
            r'MEDIDAS?\s*[:\-]?\s*'
            r'((?:\d+(?:[.,]\d+)?\s*[Xx×*]\s*)+\d+(?:[.,]\d+)?)'
            r'\s*(CM|MM|M\b)?',
            _re_gd.IGNORECASE
        )
        # Captura mecanismo depois de "MECANISMO"
        _pat_mec = _re_gd.compile(
            r'MECANISMOS?\s*[:\-]?\s*([^,\n]+)',
            _re_gd.IGNORECASE
        )
        # Captura quantidade de peças
        _pat_qtd = _re_gd.compile(
            r'C[/\s]\s*(\d+)\s*PE[CÇ]AS?',
            _re_gd.IGNORECASE
        )

        _com_medida, _sem_medida = [], []
        _tipos_freq, _mecs_freq  = {}, {}

        for _n in _nomes_raw:
            _n = _n.strip()
            _m_med = _pat_med.search(_n)
            if _m_med:
                _dims_str = _m_med.group(1)
                _unid_str = (_m_med.group(2) or 'CM').upper()
                _n_dims   = len(_re_gd.split(r'\s*[Xx×*]\s*', _dims_str))

                # Tipo de produto = tudo antes de "MEDIDAS" (sem vírgula/espaço final)
                _produto = _n[:_m_med.start()].strip().rstrip(',').strip()
                _produto = _re_gd.sub(r'\s+', ' ', _produto)

                # Mecanismo
                _m_mec  = _pat_mec.search(_n)
                _mec    = _m_mec.group(1).strip().rstrip(',').strip() if _m_mec else ''

                # Quantidade
                _m_qtd  = _pat_qtd.search(_n)
                _qtd_v  = int(_m_qtd.group(1)) if _m_qtd else None

                _com_medida.append({
                    'nome_original': _n, 'n_dims': _n_dims, 'unidade': _unid_str,
                    'produto': _produto, 'mecanismo': _mec, 'qtd': _qtd_v
                })
                if _produto:
                    _tipos_freq[_produto] = _tipos_freq.get(_produto, 0) + 1
                if _mec:
                    _mecs_freq[_mec] = _mecs_freq.get(_mec, 0) + 1
            else:
                _sem_medida.append(_n)

        _df_com = pd.DataFrame(_com_medida) if _com_medida else pd.DataFrame()
        _tipos_df = (
            pd.DataFrame(sorted(_tipos_freq.items(), key=lambda x: -x[1]),
                         columns=['Tipo de Produto', 'Ocorrências']).head(50)
            if _tipos_freq else pd.DataFrame()
        )
        _mecs_df = (
            pd.DataFrame(sorted(_mecs_freq.items(), key=lambda x: -x[1]),
                         columns=['Mecanismo', 'Ocorrências']).head(30)
            if _mecs_freq else pd.DataFrame()
        )
        return _df_com, _sem_medida, len(_nomes_raw), _tipos_df, _mecs_df

    if "gd_padroes" not in st.session_state:
        st.session_state["gd_padroes"] = None

    with st.expander("📊 Padrões aprendidos do banco", expanded=False):
        st.caption(
            "Formato detectado: **{TIPO PRODUTO}, MEDIDAS {L}X{A}X{P}CM, MECANISMO {MEC}, C/ {N} PEÇAS**"
        )
        if st.button("🔍 Analisar padrões agora", key="btn_anal_padroes"):
            with st.spinner("Analisando…"):
                st.session_state["gd_padroes"] = _analisar_padroes_gd()
        _anal = st.session_state["gd_padroes"]
        if _anal:
            _df_com_a, _sem_med_a, _total_a, _tipos_a, _mecs_a = _anal
            _col_a1, _col_a2, _col_a3, _col_a4 = st.columns(4)
            _col_a1.metric("Total de descrições", _total_a)
            _col_a2.metric("Com MEDIDAS", len(_df_com_a))
            _col_a3.metric("Sem MEDIDAS", len(_sem_med_a))
            _col_a4.metric("Com qtd de peças",
                int(_df_com_a['qtd'].notna().sum()) if not _df_com_a.empty else 0)
            if not _df_com_a.empty:
                _col_b1, _col_b2, _col_b3 = st.columns(3)
                with _col_b1:
                    st.markdown("**Nº de dimensões mais comum:**")
                    st.dataframe(_df_com_a['n_dims'].value_counts()
                                 .rename_axis('Nº dims').reset_index(name='Ocorrências'),
                                 use_container_width=True, hide_index=True)
                with _col_b2:
                    st.markdown("**Top tipos de produto:**")
                    st.dataframe(_tipos_a, use_container_width=True, hide_index=True)
                with _col_b3:
                    st.markdown("**Top mecanismos:**")
                    if not _mecs_a.empty:
                        st.dataframe(_mecs_a, use_container_width=True, hide_index=True)
                    else:
                        st.caption("Nenhum 'MECANISMO' encontrado nas descrições.")
        else:
            st.info("Clique em 'Analisar padrões agora' para ver os formatos do banco.")

        # ---- Debug: mostra exemplos reais do banco + testador de regex ----
        with st.expander("🔬 Debug: ver exemplos reais do banco de dados", expanded=False):
            _ex_rows = pd.read_sql_query(
                "SELECT DISTINCT UPPER(TRIM(nome)) AS nome FROM sistema5_itens "
                "WHERE nome IS NOT NULL AND TRIM(nome)!='' ORDER BY nome LIMIT 30",
                conn
            )['nome'].tolist()
            if _ex_rows:
                st.markdown("**Primeiros 30 registros da coluna `nome`:**")
                for _i, _ex_n in enumerate(_ex_rows, 1):
                    st.code(f"{_i:02d}. {_ex_n}", language=None)
            else:
                st.warning("Nenhum registro encontrado na tabela sistema5_itens.")

            st.divider()
            st.markdown("**Teste rápido de regex:** cole uma descrição abaixo e veja o que o sistema extrai")
            _ex_input = st.text_area("Cole uma descrição exemplo aqui:", key="gd_debug_ex", height=80)
            if _ex_input.strip():
                import re as _re_dbg
                _p_med = _re_dbg.compile(
                    r'MEDIDAS?\s*[:\-]?\s*((?:\d+(?:[.,]\d+)?\s*[Xx×*]\s*)+\d+(?:[.,]\d+)?)\s*(CM|MM|M\b)?',
                    _re_dbg.IGNORECASE
                )
                _p_mec = _re_dbg.compile(r'MECANISMOS?\s*[:\-]?\s*([^,\n]+)', _re_dbg.IGNORECASE)
                _p_qtd = _re_dbg.compile(r'C[/\s]\s*(\d+)\s*PE[CÇ]AS?', _re_dbg.IGNORECASE)
                _m_med_d = _p_med.search(_ex_input.strip().upper())
                _m_mec_d = _p_mec.search(_ex_input.strip().upper())
                _m_qtd_d = _p_qtd.search(_ex_input.strip().upper())
                st.markdown("**Resultado da extração:**")
                if _m_med_d:
                    _prod_d = _ex_input.strip().upper()[:_m_med_d.start()].strip().rstrip(',').strip()
                    st.success(f"✅ MEDIDAS encontrado: `{_m_med_d.group(1)}` (unidade: `{_m_med_d.group(2) or 'CM'}`)")
                    st.info(f"Tipo de produto extraído: `{_prod_d}`")
                else:
                    st.error("❌ Palavra 'MEDIDAS' não encontrada — o sistema não consegue extrair as dimensões desta descrição")
                    st.markdown("**Dica:** A descrição precisa conter a palavra `MEDIDAS` seguida das dimensões, ex: `MEDIDAS 30X20X10CM`")
                if _m_mec_d:
                    st.success(f"✅ MECANISMO encontrado: `{_m_mec_d.group(1).strip()}`")
                else:
                    st.warning("⚠️ Palavra 'MECANISMO' não encontrada")
                if _m_qtd_d:
                    st.success(f"✅ Quantidade: `{_m_qtd_d.group(1)} PEÇAS`")
                else:
                    st.warning("⚠️ 'C/ N PEÇAS' não encontrado")

    st.divider()

    # ---- Funções de suporte a imagens ----
    def _extrair_imagens_xlsx(_file_bytes):
        """
        Extrai imagens embutidas do xlsx e mapeia pelo índice da linha de dados (0-based).
        Retorna dict {data_row_0based: img_bytes}.
        Imagens em Excel são shapes "flutuantes" — têm âncora (row, col) mas não ficam
        literalmente dentro de células. Usamos _from.row do TwoCellAnchor (openpyxl)
        para saber em qual linha a imagem começa.
        """
        _mapa = {}
        try:
            from openpyxl import load_workbook as _lwb_img
            _wb_img = _lwb_img(BytesIO(_file_bytes))
            _ws_img = _wb_img.active

            # Descobrir linha do cabeçalho (igual ao _parse_excel_order)
            _idx_hdr = None
            for _i, _row_cells in enumerate(_ws_img.iter_rows(min_row=1, max_row=min(30, _ws_img.max_row))):
                for _cell in _row_cells:
                    _rgb_hdr = getattr(getattr(getattr(_cell, 'fill', None), 'fgColor', None), 'rgb', None)
                    if _rgb_hdr and any(c in str(_rgb_hdr).upper() for c in ("00CCFF", "00B0F0")):
                        _idx_hdr = _i
                        break
                    _val_hdr = str(_cell.value or '').strip().upper()
                    if _val_hdr in ('REFERENCIA', 'REFERÊNCIA', 'REF'):
                        _idx_hdr = _i
                if _idx_hdr is not None:
                    break
            if _idx_hdr is None:
                _idx_hdr = 0  # fallback: assume primeira linha como cabeçalho

            for _img_obj in (_ws_img._images or []):
                try:
                    _anchor = _img_obj.anchor
                    # TwoCellAnchor tem _from; OneCellAnchor também tem _from
                    if hasattr(_anchor, '_from'):
                        _row_xlsx = _anchor._from.row   # 0-indexed (linha 1 do Excel = índice 0)
                    elif hasattr(_anchor, 'row'):
                        _row_xlsx = _anchor.row
                    else:
                        continue

                    # Índice relativo à primeira linha de dados
                    _data_idx = _row_xlsx - (_idx_hdr + 1)
                    if _data_idx < 0:
                        continue

                    _img_bytes_raw = _img_obj._data()
                    if _img_bytes_raw:
                        # Guarda só a primeira imagem por linha (geralmente só há uma)
                        if _data_idx not in _mapa:
                            _mapa[_data_idx] = _img_bytes_raw
                except Exception:
                    continue
        except Exception:
            pass
        return _mapa

    def _img_para_uri(_img_bytes, _max_px=160):
        """Converte bytes de imagem para data URI base64 (thumbnail pequeno para a tabela)."""
        try:
            from PIL import Image as _PILg
            _pil = _PILg.open(BytesIO(_img_bytes)).convert("RGB")
            _pil.thumbnail((_max_px, _max_px), _PILg.LANCZOS)
            _buf_uri = BytesIO()
            _pil.save(_buf_uri, format='PNG', optimize=True)
            _enc = _b64_gd.b64encode(_buf_uri.getvalue()).decode()
            return f"data:image/png;base64,{_enc}"
        except Exception:
            return None

    def _ocr_texto_imagem(_img_bytes):
        """
        Tenta extrair texto da imagem do produto via Tesseract.
        Útil quando a imagem tem legenda ou nome do produto impresso.
        Retorna string com texto encontrado (pode ser vazio).
        """
        try:
            import pytesseract as _pytess_gd
            from PIL import Image as _PILg2, ImageEnhance as _IEg2
            _pil2 = _PILg2.open(BytesIO(_img_bytes)).convert("L")
            _w2, _h2 = _pil2.size
            _pil2 = _pil2.resize((_w2 * 3, _h2 * 3), _PILg2.LANCZOS)
            _pil2 = _IEg2.Contrast(_pil2).enhance(2.0)
            _txt2 = _pytess_gd.image_to_string(_pil2, lang='por', config='--psm 11 --oem 3')
            return _re_gd.sub(r'\s+', ' ', _txt2).strip()
        except Exception:
            return ''

    # ---- Chave API Claude (visível, com fallback de env var) ----
    import os as _os_gd
    _env_key_gd = _os_gd.environ.get("ANTHROPIC_API_KEY", "")
    _saved_key_gd = st.session_state.get("gd_api_key_val", _env_key_gd)

    _col_key1, _col_key2 = st.columns([3, 1])
    with _col_key1:
        _api_key_gd = st.text_input(
            "🔑 Chave API Anthropic (para identificar produtos via IA)",
            value=_saved_key_gd,
            type="password", key="gd_api_key_input",
            placeholder="sk-ant-api03-...",
            help="Obtenha em console.anthropic.com. Com a chave, o sistema analisa a foto, "
                 "identifica o tipo de produto, busca similares no banco e gera a descrição completa."
        )
    with _col_key2:
        st.write("")
        st.write("")
        if _api_key_gd:
            st.success("✅ IA ativa")
        else:
            st.error("❌ Sem chave")

    if _api_key_gd:
        st.session_state["gd_api_key_val"] = _api_key_gd
    else:
        _api_key_gd = _saved_key_gd

    if not _api_key_gd:
        st.info(
            "**Como funciona com a chave API:**\n"
            "1. Claude analisa a foto do produto\n"
            "2. Identifica o tipo (ex: VARINHA MÁGICA, BRINQUEDO ANIMAL)\n"
            "3. Busca produtos similares no banco de dados\n"
            "4. Se achar similar → troca só as medidas, preserva o resto\n"
            "5. Se não achar → gera descrição completa no padrão INMETRO\n\n"
            "Sem a chave, só funciona busca por referência (REF)."
        )

    st.divider()

    # ---- Upload do arquivo ----
    st.subheader("1. Arquivo de medidas e quantidades")
    st.caption(
        "Envie o Excel com células **azuis** como cabeçalho (REF, MEDIDAS, FOTO DO PRODUTO, DESCRIÇÃO). "
        "A coluna DESCRIÇÃO (células amarelas) será preenchida e o arquivo retornado completo."
    )

    _arq_gd = st.file_uploader(
        "Arquivo (.xlsx, .xls, .csv)",
        type=["xlsx", "xls", "csv"],
        key="upload_gerar_desc"
    )

    if _arq_gd:
        try:
            _arq_gd.seek(0)
            _bytes_gd_arq = _arq_gd.read()
            _arq_gd.seek(0)

            # ---- Leitura por cores (openpyxl): detecta cabeçalhos azuis ----
            # Azul: FF00B0F0 — células que identificam as colunas
            # Amarelo: FFFFFF00 — células DESCRIÇÃO a preencher
            from openpyxl import load_workbook as _lwb_gd

            # ---- Lê Excel detectando colunas por células AZUIS (FF00B0F0) ----
            _wb_gd = _lwb_gd(BytesIO(_bytes_gd_arq))
            _ws_gd = _wb_gd.active

            def _rgb_cel(_c):
                try: return str(_c.fill.fgColor.rgb).upper()
                except: return ''

            _COL_ALIASES = {
                'REF':    ['REF','REFERENCIA','REFERÊNCIA','MODELO'],
                'MED':    ['MEDIDAS','MED','DIM'],
                'FOTO':   ['FOTO','IMG','IMAG','PHOTO'],
                'DESC':   ['DESCRI','DESC'],
                'MARCA':  ['MARCA','BRAND'],
                'QTD':    ['QTD','PECAS','PEÇAS','PCS'],
                'PESO':   ['PESO','WEIGHT','KG'],
                'NOME':   ['NOME','TIPO','PRODUTO'],
            }
            _col_pos   = {}   # campo → coluna Excel (1-based)
            _hdr_row   = None

            for _ri, _row in enumerate(
                _ws_gd.iter_rows(min_row=1, max_row=min(10, _ws_gd.max_row)), start=1
            ):
                _found_blue = False
                for _cell in _row:
                    _rgb = _rgb_cel(_cell)
                    _is_blue = 'B0F0' in _rgb
                    _val     = str(_cell.value or '').upper().strip()
                    # Sem cor: tenta também detectar por texto (fallback)
                    for _campo, _als in _COL_ALIASES.items():
                        if _campo not in _col_pos:
                            for _al in _als:
                                if _al in _val:
                                    _col_pos[_campo] = _cell.column
                                    if _is_blue: _found_blue = True
                                    break
                if _col_pos and (_found_blue or _hdr_row is None):
                    _hdr_row = _ri
                    if _found_blue:
                        break

            _data_row0 = (_hdr_row or 1) + 1

            # Extrai dados apenas das colunas identificadas
            _dados_gd = []
            for _ri in range(_data_row0, _ws_gd.max_row + 1):
                def _gv(_campo):
                    _c = _col_pos.get(_campo)
                    return str(_ws_gd.cell(_ri, _c).value or '').strip().lstrip("'") if _c else ''
                _ref = _gv('REF')
                # Célula DESC: verifica se é amarela (a preencher) ou já preenchida
                _desc_cell   = _ws_gd.cell(_ri, _col_pos['DESC']) if 'DESC' in _col_pos else None
                _desc_exist  = str(_desc_cell.value or '').strip() if _desc_cell else ''
                _desc_amarelo = 'FFFF00' in _rgb_cel(_desc_cell) if _desc_cell else False
                _precisa     = (not _desc_exist) and (_ref or _desc_amarelo)
                if _ref or _precisa:
                    _dados_gd.append({
                        '_ri': _ri, '_precisa': _precisa,
                        'REF': _ref, 'MED': _gv('MED'), 'NOME': _gv('NOME'),
                        'MARCA': _gv('MARCA'), 'QTD': _gv('QTD'), 'PESO': _gv('PESO'),
                        'DESC_EXIST': _desc_exist,
                    })

            with st.spinner("Extraindo imagens embutidas…"):
                _imgs_gd = _extrair_imagens_xlsx(_bytes_gd_arq)
            _n_imgs = len(_imgs_gd)
            _n_prec = sum(1 for r in _dados_gd if r['_precisa'])

            st.success(
                f"✅ **{len(_dados_gd)}** produto(s) · **{_n_prec}** aguardando descrição"
                + (f" · **{_n_imgs}** foto(s)" if _n_imgs else " · sem fotos embutidas")
            )
            _tag_c = " | ".join(f"`{v}`={k}" for k, v in _col_pos.items())
            if _tag_c: st.caption("Colunas detectadas: " + _tag_c)

            with st.expander("👁️ Dados lidos (primeiros 8)", expanded=False):
                st.dataframe(pd.DataFrame(_dados_gd).drop(
                    columns=['_ri','_precisa'], errors='ignore').head(8),
                    use_container_width=True)

            # ---- Funções de geração ----
            _pat_med_parse = _re_gd.compile(
                r'((?:\d+(?:[.,]\d+)?\s*[Xx×*]\s*)+\d+(?:[.,]\d+)?)\s*(CM|MM|M\b)?',
                _re_gd.IGNORECASE
            )
            def _parse_med(_s):
                _m = _pat_med_parse.search(str(_s or ''))
                if not _m: return '', 'CM'
                return _re_gd.sub(r'\s','',_m.group(1)), (_m.group(2) or 'CM').upper()

            _pat_med_sub = _re_gd.compile(
                r'MEDIDAS?\s*[:\-]?\s*[\d.,*×xX\s]+(?:CM|MM|M\b)?',
                _re_gd.IGNORECASE
            )
            def _subst_med(_base, _dims, _unid):
                _n = 'MEDIDAS ' + _dims.upper() + ' ' + _unid.upper()
                if _pat_med_sub.search(_base):
                    return _pat_med_sub.sub(_n, _base, count=1)
                _i = _base.find(',')
                return (_base[:_i]+', '+_n+_base[_i:]) if _i >= 0 else _base+', '+_n

            def _buscar_ref(_ref):
                if not _ref.strip(): return None, None
                for _sql, _p, _orig in [
                    ("SELECT nome FROM sistema5_itens WHERE nome IS NOT NULL "
                     "AND UPPER(TRIM(modelo))=UPPER(TRIM(?)) GROUP BY UPPER(TRIM(nome)) "
                     "ORDER BY COUNT(*) DESC LIMIT 1",
                     [_ref], 'banco (ref exata)'),
                    ("SELECT nome FROM sistema5_itens WHERE nome IS NOT NULL "
                     "AND UPPER(modelo) LIKE UPPER(?) GROUP BY UPPER(TRIM(nome)) "
                     "ORDER BY COUNT(*) DESC LIMIT 1",
                     [f'%{_ref}%'], 'banco (ref parcial)'),
                ]:
                    _q = pd.read_sql_query(_sql, conn, params=_p)
                    if not _q.empty: return _q.iloc[0]['nome'], _orig
                return None, None

            def _buscar_tipo_db(_tipo):
                """Busca no DB pelo tipo identificado. Valida que o resultado bate com o tipo."""
                if not _tipo: return None, None
                _pals = [p for p in _tipo.upper().split() if len(p) > 3]
                if not _pals: return None, None
                # Palavra principal (1ª palavra significativa) deve estar no resultado
                _palavra_chave = _pals[0]
                for _n in range(min(len(_pals), 3), 0, -1):
                    _conds = ' AND '.join(
                        f"UPPER(nome) LIKE '%{p}%'" for p in _pals[:_n]
                    )
                    try:
                        _q = pd.read_sql_query(
                            f"SELECT nome, COUNT(*) AS cnt FROM sistema5_itens "
                            f"WHERE nome IS NOT NULL AND {_conds} "
                            f"GROUP BY UPPER(TRIM(nome)) ORDER BY cnt DESC LIMIT 3",
                            conn
                        )
                        if not _q.empty:
                            # Filtra: o resultado deve conter a palavra-chave principal
                            for _, _row_db in _q.iterrows():
                                _nome_db = str(_row_db['nome']).upper()
                                if _palavra_chave in _nome_db:
                                    return _row_db['nome'], f'banco+ia ({_tipo})'
                    except Exception:
                        pass
                return None, None

            def _buscar_exemplos_db(_tipo):
                """Exemplos do DB do mesmo tipo + gerais para estudo de padrão."""
                _exs: list = []
                if _tipo:
                    _pals = [p for p in _tipo.upper().split() if len(p) > 3]
                    if _pals:
                        # Pega até 15 exemplos específicos do tipo
                        for _np in range(min(len(_pals), 2), 0, -1):
                            _cond = ' AND '.join(f"UPPER(nome) LIKE '%{p}%'" for p in _pals[:_np])
                            try:
                                _q = pd.read_sql_query(
                                    f"SELECT DISTINCT UPPER(TRIM(nome)) AS n FROM sistema5_itens "
                                    f"WHERE nome IS NOT NULL AND LENGTH(TRIM(nome))>20 "
                                    f"AND ({_cond}) ORDER BY RANDOM() LIMIT 15",
                                    conn
                                )['n'].tolist()
                                _exs.extend(_q)
                                if len(_exs) >= 8: break
                            except Exception:
                                pass
                # Complementa com exemplos gerais (padrão gramatical)
                try:
                    _q2 = pd.read_sql_query(
                        "SELECT DISTINCT UPPER(TRIM(nome)) AS n FROM sistema5_itens "
                        "WHERE nome IS NOT NULL AND LENGTH(TRIM(nome))>30 "
                        "AND UPPER(nome) LIKE '%MEDIDAS%' AND UPPER(nome) LIKE '%ANOS%' "
                        "ORDER BY RANDOM() LIMIT 10", conn
                    )['n'].tolist()
                    _exs = list(dict.fromkeys(_exs + _q2))
                except Exception:
                    pass
                return _exs[:18]

            def _abreviar_desc(_desc, _max=200):
                """Aplica abreviações padrão INMETRO — mesmas regras da função limitar_nome."""
                if len(_desc) <= _max:
                    return _desc
                _d = _desc
                _d = _d.replace("PRODUZIDO", "PROD.")
                _d = _d.replace("INDICATIVO", "IND.")
                _d = _d.replace("RESTRITIVO", "REST.")
                _d = _re_gd.sub(r'\bANOS\b', 'A', _d, flags=_re_gd.IGNORECASE)
                _d = _re_gd.sub(r'\bMESES\b', 'MES.', _d, flags=_re_gd.IGNORECASE)
                _d = _re_gd.sub(r'\bINJE[ÇC][ÃA]O\b|\bINJECAO\b', 'INJ.', _d, flags=_re_gd.IGNORECASE)
                _d = _re_gd.sub(r'\bM[ÁA]XIMA\b', 'MAX.', _d, flags=_re_gd.IGNORECASE)
                _d = _re_gd.sub(r'\bPL[ÁA]STICO\b', 'PLAST.', _d, flags=_re_gd.IGNORECASE)
                _d = _d.replace("CONTROLE", "CONT.")
                _d = _d.replace("REMOTO", "REM.")
                _d = _d.replace("VELOCIDADE", "VEL.")
                _d = _d.replace("MEDIDAS", "MED.")
                return _d[:_max].rstrip(',; ') if len(_d) > _max else _d

            # ---- Regras de faixa etária — Portaria INMETRO 302/2021 ----
            # Norma: ABNT ISO/TR 8124-8 + ABNT NBR NM 300
            # INDICATIVO = faixa mínima recomendada (desenvolvimento cognitivo)
            # RESTRITIVO = proibição por segurança (peças pequenas, risco de asfixia, etc.)
            _FAIXA_RULES = [
                # (palavras-no-tipo,          indicativo,    restritivo)
                # ---- Bebê / primeiros anos ----
                (['CHOCALHO'],               '+18 MESES',  None),
                (['MORDEDOR'],               '+0 MESES',   None),
                (['PELÚCIA','PELUCIA'],      '+0 MESES',   None),
                (['BEBÊ','BEBE'],            '+0 MESES',   None),
                # ---- Apertar / animal squeeze ----
                (['APERTAR'],                '+18 MESES',  '-3 ANOS'),
                (['ANIMAL','APERTAR'],       '+18 MESES',  '-3 ANOS'),
                # ---- Bola ----
                (['BOLA','GIRAT'],           '+3 ANOS',    '-3 ANOS'),
                (['BOLA','INTERATIV'],       '+3 ANOS',    '-3 ANOS'),
                (['BOLA','SALTIT'],          '+18 MESES',  None),
                (['BOLA'],                   '+3 ANOS',    None),
                # ---- Pião ----
                (['PIÃO','GIRATORI'],        '+3 ANOS',    '-3 ANOS'),
                (['PIAO','GIRATORI'],        '+3 ANOS',    '-3 ANOS'),
                (['PIÃO'],                   '+3 ANOS',    '-3 ANOS'),
                (['PIAO'],                   '+3 ANOS',    '-3 ANOS'),
                (['BEYBLADE'],               '+6 ANOS',    '-3 ANOS'),
                # ---- Fidget / cubo ----
                (['FIDGET','SPINNER'],       '+6 ANOS',    '-3 ANOS'),
                (['CUBO','MÁGICO'],          '+6 ANOS',    '-3 ANOS'),
                (['CUBO','MAGICO'],          '+6 ANOS',    '-3 ANOS'),
                # ---- Varinha ----
                (['VARINHA','MÁGICA'],       '+3 ANOS',    '-3 ANOS'),
                (['VARINHA','MAGICA'],       '+3 ANOS',    '-3 ANOS'),
                (['VARINHA'],               '+3 ANOS',    '-3 ANOS'),
                # ---- Pistola / armas de brinquedo ----
                (['PISTOLA','ÁGUA'],         '+3 ANOS',    '-3 ANOS'),
                (['PISTOLA','AGUA'],         '+3 ANOS',    '-3 ANOS'),
                (['PISTOLA','BOLHA'],        '+3 ANOS',    '-3 ANOS'),
                (['PISTOLA'],                '+6 ANOS',    '-3 ANOS'),
                (['ESPADA'],                 '+3 ANOS',    '-3 ANOS'),
                # ---- Boneca / figura ----
                (['BONECA'],                 '+3 ANOS',    '-3 ANOS'),
                (['BONECOS'],                '+3 ANOS',    '-3 ANOS'),
                (['FIGURA','AÇÃO'],          '+3 ANOS',    '-3 ANOS'),
                (['FIGURA','ACAO'],          '+3 ANOS',    '-3 ANOS'),
                # ---- Carrinho / veículo ----
                (['CONTROLE','REMOTO'],      '+6 ANOS',    '-3 ANOS'),
                (['CARRINHO'],               '+3 ANOS',    '-3 ANOS'),
                (['VEÍCULO'],                '+3 ANOS',    '-3 ANOS'),
                (['VEICULO'],                '+3 ANOS',    '-3 ANOS'),
                # ---- Maquiagem / acessório ----
                (['MAQUIAGEM'],              '+3 ANOS',    '-3 ANOS'),
                (['ESMALTE'],                '+6 ANOS',    '-3 ANOS'),
                (['JÓIA','INFANTIL'],        '+3 ANOS',    '-3 ANOS'),
                (['JOIA','INFANTIL'],        '+3 ANOS',    '-3 ANOS'),
                # ---- Eletrônico / pilha ----
                (['ELETRÔNICO','PILHA'],     '+3 ANOS',    '-3 ANOS'),
                (['ELETRONICO','PILHA'],     '+3 ANOS',    '-3 ANOS'),
                # ---- Dardos / projétil ----
                (['DARDO'],                  '+6 ANOS',    '-6 ANOS'),
                (['NERF'],                   '+6 ANOS',    '-3 ANOS'),
                # ---- Jogo químico / ciência ----
                (['QUÍMICO'],                '+8 ANOS',    '-8 ANOS'),
                (['QUIMICO'],                '+8 ANOS',    '-8 ANOS'),
                (['CIÊNCIA'],                '+8 ANOS',    '-3 ANOS'),
                (['CIENCIA'],                '+8 ANOS',    '-3 ANOS'),
                # ---- Conjunto genérico ----
                (['CONJUNTO'],               '+3 ANOS',    '-3 ANOS'),
                (['BRINQUEDO'],              '+3 ANOS',    '-3 ANOS'),
            ]

            def _determinar_faixa(_tipo_produto, _desc_base=''):
                """Retorna (indicativo, restritivo) com base no tipo e na portaria 302/2021."""
                _t = (_tipo_produto or '').upper()
                _d = (_desc_base or '').upper()
                _texto = _t + ' ' + _d
                for _pals, _ind, _res in _FAIXA_RULES:
                    if all(p in _texto for p in _pals):
                        return _ind, _res
                # Fallback genérico: a maioria dos brinquedos tem peças pequenas
                return '+3 ANOS', '-3 ANOS'

            # ---- API Claude: 2 chamadas separadas ----
            def _claude_tipo(_img_bytes, _nome_orig, _api_key):
                """Passo 1: identifica o tipo EXATO do produto (curto, PT-BR, distingue subtipos)."""
                if not _api_key: return None
                try:
                    import anthropic as _ant
                    _client = _ant.Anthropic(api_key=_api_key)
                    _prompt = (
                        "Analise a imagem com atenção e identifique o tipo EXATO do produto.\n"
                        "Responda APENAS com o tipo em PORTUGUÊS BRASILEIRO, MAIÚSCULAS, máximo 5 palavras.\n\n"
                        "═══ DISTINÇÕES CRÍTICAS — NÃO CONFUNDA ═══\n\n"
                        "PIÃO vs BOLA GIRATÓRIA — leia com atenção:\n"
                        "  PIÃO = brinquedo que GIRA SOBRE UMA PONTA no chão/superfície\n"
                        "     O CORPO pode ser QUALQUER formato: cônico clássico, achatado (beyblade),\n"
                        "     em formato de hambúrguer, esférico com ponta, metálico giroscópio, etc.\n"
                        "     → Se vê PONTA NA BASE para girar = PIÃO (independente do formato do corpo)\n"
                        "     ex: pião de madeira, beyblade, pião hambúrguer, pião metálico eletrogalvanizado\n\n"
                        "  BOLA GIRATÓRIA = bola de borracha/silicone que você APERTA COM A MÃO\n"
                        "     e ela gira/vibra ENQUANTO VOCÊ SEGURA — não tem ponta, não gira no chão\n"
                        "     ex: Pressure Rotating Ball (bola de apertar com dinossauro dentro)\n\n"
                        "FIDGET / SPINNER:\n"
                        "  FIDGET SPINNER = disco PLANO com rolamentos, segura no centro e gira as aletas\n"
                        "  CUBO MAGICO = cubo que gira nos eixos (tipo Rubik)\n\n"
                        "VARINHA:\n"
                        "  VARINHA MAGICA = bastão longo, pode ter luz/som na ponta\n"
                        "  NÃO é um conjunto de fantasia com saia/fantasia a menos que apareça tudo na foto\n\n"
                        "OUTROS EXEMPLOS CORRETOS:\n"
                        "  BOLA GIRATÓRIA DINOSSAURO\n"
                        "  BOLA INTERATIVA COM LUZ\n"
                        "  PIÃO GIRATÓRIO (só se for cônico com ponta)\n"
                        "  VARINHA MAGICA LUMINOSA\n"
                        "  BRINQUEDO ANIMAL DE APERTAR\n"
                        "  CARRINHO CONTROLE REMOTO\n"
                        "  PISTOLA DE BOLHA DE SABAO\n"
                        "  PISTOLA DE AGUA\n"
                        "  BONECA\n"
                        "  CUBO MAGICO\n"
                        "  ESPADA BRINQUEDO\n"
                        "  IOIO\n\n"
                        "VOCABULÁRIO OBRIGATÓRIO PT-BR:\n"
                        "  PIÃO (nunca PEONZA/TROMPO/PERINOLA)\n"
                        "  BONECA (nunca DOLL)\n"
                        "  VARINHA (nunca WAND)\n"
                        "  CARRINHO (nunca CAR)\n\n"
                        f"Nome original do fabricante (pode estar em chinês/inglês): {_nome_orig}\n\n"
                        "Responda SOMENTE o tipo, sem mais nada."
                    )
                    _content: list = []
                    if _img_bytes:
                        import base64 as _b64v
                        _b6 = _b64v.b64encode(_img_bytes).decode()
                        _mt = "image/png" if _img_bytes[:4] == b'\x89PNG' else "image/jpeg"
                        _content.append({"type": "image",
                                          "source": {"type": "base64",
                                                     "media_type": _mt, "data": _b6}})
                    _content.append({"type": "text", "text": _prompt})
                    _r = _client.messages.create(
                        model="claude-sonnet-4-6", max_tokens=25,
                        messages=[{"role": "user", "content": _content}]
                    )
                    _tipo_raw = _r.content[0].text.strip().upper()
                    # Corrige palavras em espanhol/inglês que escaparam
                    _tipo_raw = _tipo_raw.replace('PEONZA', 'PIAO').replace('PERINOLA', 'PIAO')
                    _tipo_raw = _tipo_raw.replace('TROMPO', 'PIAO')
                    _tipo_raw = _tipo_raw.replace('ROTATING BALL', 'BOLA GIRATORIA')
                    _tipo_raw = _tipo_raw.replace('PRESSURE BALL', 'BOLA GIRATORIA')
                    _tipo_raw = _tipo_raw.replace('SPINNING TOP', 'PIAO GIRATORIO')
                    return _tipo_raw
                except Exception:
                    return None

            def _claude_adaptar_base(_img_bytes, _base_desc, _tipo, _dims, _unid, _api_key):
                """Adapta descrição do banco ao produto real — remove o que não aparece na foto."""
                if not _api_key: return None
                try:
                    import anthropic as _ant
                    _client = _ant.Anthropic(api_key=_api_key)
                    _dim_str = f"{_dims} {_unid}" if _dims else "VERIFICAR"
                    _ind_s, _res_s = _determinar_faixa(_tipo, _base_desc)
                    _prompt = (
                        "Você é especialista em certificação INMETRO Brasil (Portaria 302/2021).\n\n"
                        "Tenho uma descrição de produto SIMILAR no banco de dados:\n"
                        f"BASE: {_base_desc}\n\n"
                        f"Tipo identificado na imagem: {_tipo}\n"
                        f"Medidas reais: {_dim_str}\n\n"
                        "FAIXA ETÁRIA correta para este tipo (Portaria 302/2021):\n"
                        f"  INDICATIVO: {_ind_s}\n"
                        + (f"  RESTRITIVO: {_res_s}\n" if _res_s else
                           "  RESTRITIVO: não obrigatório\n") +
                        "\nTAREFA: Adapte a descrição BASE para o produto REAL visível nesta foto.\n\n"
                        "REGRAS ABSOLUTAS:\n"
                        "1. Olhe a imagem com MUITA ATENÇÃO — descreva SOMENTE o que está visível\n"
                        "2. Se a BASE tem 'ARMAS', 'LANÇA PIÃO', mas a foto mostra SÓ um pião → REMOVA as armas\n"
                        "3. Se a BASE tem 'SAIA', 'ASINHA', 'FANTASIA', 'TIARA', mas foto mostra SÓ a varinha → REMOVA\n"
                        "4. Se a BASE tem 'CONJUNTO C/ 2 PEÇAS' mas a foto mostra 1 → corrija a quantidade\n"
                        "5. Mantenha o padrão gramatical INMETRO da BASE\n"
                        f"6. Use as medidas: MEDIDAS {_dim_str}\n"
                        f"7. Use as idades: INDICATIVO {_ind_s}"
                        + (f", RESTRITIVO {_res_s}" if _res_s else "") + "\n"
                        "8. PORTUGUÊS BRASILEIRO — PIÃO (nunca PEONZA), TUDO MAIÚSCULAS\n"
                        "9. LIMITE: máximo 200 caracteres no total\n"
                        "   Use abreviações se necessário (mesmas do sistema INMETRO):\n"
                        "   INDICATIVO→IND.  RESTRITIVO→REST.  PRODUZIDO→PROD.\n"
                        "   ANOS→A  MESES→MES.  INJEÇÃO→INJ.  PLÁSTICO→PLAST.\n"
                        "   CONTROLE→CONT.  REMOTO→REM.  VELOCIDADE→VEL.  MEDIDAS→MED.\n"
                        "10. Retorne APENAS a descrição adaptada, sem explicações, sem aspas"
                    )
                    _content: list = []
                    if _img_bytes:
                        import base64 as _b64v
                        _b6 = _b64v.b64encode(_img_bytes).decode()
                        _mt = "image/png" if _img_bytes[:4] == b'\x89PNG' else "image/jpeg"
                        _content.append({"type": "image",
                                          "source": {"type": "base64",
                                                     "media_type": _mt, "data": _b6}})
                    _content.append({"type": "text", "text": _prompt})
                    _r = _client.messages.create(
                        model="claude-sonnet-4-6", max_tokens=350,
                        messages=[{"role": "user", "content": _content}]
                    )
                    _desc = _r.content[0].text.strip().upper()
                    _desc = _desc.replace('PEONZA', 'PIÃO').replace('PIAO ', 'PIÃO ')
                    return _abreviar_desc(_desc)
                except Exception:
                    return None

            def _claude_gerar(_img_bytes, _nome_orig, _tipo, _dims, _unid, _qtd, _exemplos, _api_key):
                """Passo 4: gera descrição INMETRO — estuda padrão do DB + usa regras da portaria 302/2021."""
                if not _api_key: return None
                try:
                    import anthropic as _ant
                    _client = _ant.Anthropic(api_key=_api_key)
                    _dim_str = f"{_dims} {_unid}" if _dims else "VERIFICAR"
                    _exs_txt = '\n'.join(f'  {i+1}. {e}' for i, e in enumerate(_exemplos))
                    _ind_sug, _res_sug = _determinar_faixa(_tipo)
                    _faixa_hint = (
                        f"Portaria INMETRO 302/2021 — sugestão para este tipo:\n"
                        f"  INDICATIVO sugerido: {_ind_sug}\n"
                        + (f"  RESTRITIVO sugerido: {_res_sug}\n" if _res_sug else
                           "  RESTRITIVO: não obrigatório para este tipo\n")
                        + "  (ajuste se a imagem indicar produto para outra faixa)\n"
                    )
                    _prompt = (
                        "Você é um especialista em certificação INMETRO do Brasil. "
                        "Sua função é gerar descrições técnicas de brinquedos para uso em certificação oficial.\n\n"
                        "═══ BANCO DE DADOS INMETRO — ESTUDE ESTE PADRÃO ═══\n"
                        "Analise os exemplos abaixo e aprenda a gramática/estrutura exata:\n"
                        + _exs_txt +
                        "\n\n═══ PRODUTO A DESCREVER ═══\n"
                        f"Tipo identificado na foto: {_tipo or 'BRINQUEDO'}\n"
                        f"Medidas do produto: {_dim_str}\n"
                        f"Quantidade de peças: {_qtd or '1'}\n\n"
                        "═══ FAIXA ETÁRIA (Portaria INMETRO 302/2021) ═══\n"
                        + _faixa_hint +
                        "\n═══ DISTINÇÕES OBRIGATÓRIAS ═══\n"
                        "Antes de escrever, identifique EXATAMENTE o que está na foto:\n"
                        "  PIÃO = brinquedo que GIRA SOBRE UMA PONTA no chão/superfície\n"
                        "     Corpo pode ser cônico, hambúrguer, esférico com ponta, metálico — qualquer formato\n"
                        "     → Se tem PONTA NA BASE para girar = PIÃO\n"
                        "  BOLA GIRATÓRIA = bola que você APERTA COM A MÃO e gira enquanto segura\n"
                        "     Não tem ponta, não gira no chão — ex: Pressure Rotating Ball\n"
                        "  FIDGET SPINNER = disco PLANO com rolamentos\n"
                        "  VARINHA = bastão longo — NÃO inclui saia/fantasia a menos que apareça\n\n"
                        "═══ REGRAS ABSOLUTAS ═══\n"
                        "1. PORTUGUÊS BRASILEIRO:\n"
                        "   PIÃO (NUNCA PEONZA/TROMPO), BONECA (nunca doll)\n"
                        "   CARRINHO (nunca car), VARINHA (nunca wand)\n"
                        "   BOLA GIRATÓRIA (nunca Rotating Ball)\n"
                        "2. TUDO EM MAIÚSCULAS\n"
                        "3. Descreva APENAS o produto principal visível na imagem\n"
                        "   PIÃO = tem PONTA NA BASE para girar no chão (qualquer formato de corpo)\n"
                        "   BOLA GIRATÓRIA = bola que se APERTA COM A MÃO, sem ponta\n"
                        "   Se é só varinha → sem saia/fantasia\n"
                        "   Se é só pião → sem lançador/armas\n"
                        f"4. Inclua: MEDIDAS {_dim_str}\n"
                        "5. Inclua: MECANISMO [A PILHA / A CORDA / MANUAL / ELETRONICO / PRESSAO...]\n"
                        "6. Inclua: C/ [N] PEÇAS (se mais de 1)\n"
                        f"7. Inclua: INDICATIVO {_ind_sug}  (ajuste se necessário)\n"
                        + (f"8. Inclua: RESTRITIVO {_res_sug}  (obrigatório — peças pequenas)\n"
                           if _res_sug else
                           "8. Omita RESTRITIVO (não necessário para este tipo)\n") +
                        "9. Material: DE PLÁSTICO / DE METAL / DE MADEIRA\n"
                        "10. Siga a estrutura dos exemplos EXATAMENTE\n"
                        "11. LIMITE: máximo 200 caracteres no total\n"
                        "    Use abreviações se necessário (padrão INMETRO):\n"
                        "    INDICATIVO→IND.  RESTRITIVO→REST.  PRODUZIDO→PROD.\n"
                        "    ANOS→A  MESES→MES.  INJEÇÃO→INJ.  PLÁSTICO→PLAST.\n"
                        "    CONTROLE→CONT.  REMOTO→REM.  VELOCIDADE→VEL.  MEDIDAS→MED.\n\n"
                        "Retorne APENAS a descrição técnica completa, sem explicações, sem aspas."
                    )
                    _content: list = []
                    if _img_bytes:
                        import base64 as _b64v
                        _b6 = _b64v.b64encode(_img_bytes).decode()
                        _mt = "image/png" if _img_bytes[:4] == b'\x89PNG' else "image/jpeg"
                        _content.append({"type": "image",
                                          "source": {"type": "base64",
                                                     "media_type": _mt, "data": _b6}})
                    _content.append({"type": "text", "text": _prompt})
                    _r = _client.messages.create(
                        model="claude-sonnet-4-6", max_tokens=400,
                        messages=[{"role": "user", "content": _content}]
                    )
                    _desc = _r.content[0].text.strip().upper()
                    # Pós-processamento: corrige palavras em espanhol/inglês escapadas
                    _desc = _desc.replace('PEONZA', 'PIÃO').replace('PERINOLA', 'PIÃO')
                    _desc = _desc.replace('TROMPO', 'PIÃO').replace('PIAO ', 'PIÃO ')
                    return _abreviar_desc(_desc)
                except Exception:
                    return None

            # ---- Geração automática: 4 etapas + adaptação visual + faixas etárias + limite 200 chars ----
            _file_key = f"gd_v8_{_arq_gd.name}_{len(_dados_gd)}"
            if _file_key not in st.session_state:
                _rows_out = []
                _prog = st.progress(0, text="Gerando…")
                for _di, _d in enumerate(_dados_gd):
                    _prog.progress((_di+1)/max(len(_dados_gd),1),
                                   text=f"Produto {_di+1}/{len(_dados_gd)}: {_d['REF']}…")
                    if not _d['_precisa']:
                        _rows_out.append({
                            '_ri': _d['_ri'], 'REF': _d['REF'], 'MARCA': _d['MARCA'],
                            'DESC': _d['DESC_EXIST'], 'ORIGEM': 'já preenchida',
                            'IMAGEM': None, 'TIPO_IA': ''
                        })
                        continue

                    _dims, _unid = _parse_med(_d['MED'])
                    _data_idx  = _d['_ri'] - _data_row0
                    _img_b     = _imgs_gd.get(_data_idx)
                    _img_uri   = _img_para_uri(_img_b) if _img_b else None
                    _tipo_prod = ''

                    # ETAPA 1 — busca exata/parcial por referência no DB
                    _base, _orig = _buscar_ref(_d['REF'])

                    if not _base:
                        # ETAPA 2 — Claude identifica o TIPO do produto (chamada curta, visão)
                        _tipo_prod = _claude_tipo(_img_b, _d['NOME'], _api_key_gd) or ''

                        # ETAPA 3 — busca DB pelo tipo identificado pela IA
                        if _tipo_prod:
                            _base, _orig = _buscar_tipo_db(_tipo_prod)

                    # ETAPA 3b — fallback sem API: busca pelo NOME original
                    # (funciona quando o nome tem palavras em português/inglês rastreáveis)
                    if not _base and not _api_key_gd and _d['NOME']:
                        _nome_words = [
                            w for w in str(_d['NOME']).upper().split()
                            if len(w) > 3 and w.isascii()
                        ]
                        for _nw in _nome_words[:3]:
                            try:
                                _q_fb = pd.read_sql_query(
                                    "SELECT nome, COUNT(*) AS cnt FROM sistema5_itens "
                                    "WHERE nome IS NOT NULL AND UPPER(nome) LIKE ? "
                                    "GROUP BY UPPER(TRIM(nome)) ORDER BY cnt DESC LIMIT 1",
                                    conn, params=[f'%{_nw}%']
                                )
                                if not _q_fb.empty:
                                    _base = _q_fb.iloc[0]['nome']
                                    _orig = f'banco (nome: {_nw})'
                                    break
                            except Exception:
                                pass

                    if _base:
                        # Detecta se a base tem componentes extras que não aparecem na imagem
                        _EXTRAS = {'ARMA', 'ARMAS', 'LANÇA', 'LANCA', 'SAIA', 'ASINHA',
                                   'FANTASIA', 'ROUPA', 'TIARA', 'COROA', 'ESPADA',
                                   'MASCARA', 'CAPA', 'CHAPEU'}
                        _base_up = _base.upper()
                        _tipo_up = (_tipo_prod or '').upper()
                        # Se a base tem palavras extras E o tipo identificado não inclui elas
                        _tem_extra = any(
                            ex in _base_up and ex not in _tipo_up
                            for ex in _EXTRAS
                        )
                        _orig_banco = _orig or 'banco'
                        if _tem_extra and _api_key_gd and _img_b:
                            # ETAPA 3c — adapta a base removendo o que não está na foto
                            _desc_adapt = _claude_adaptar_base(
                                _img_b, _base_up, _tipo_prod, _dims, _unid, _api_key_gd
                            )
                            if _desc_adapt:
                                _desc_final = _desc_adapt
                                _orig = _orig_banco + ' adaptado'
                            else:
                                _desc_final = _subst_med(_base_up, _dims, _unid) if _dims else _base_up
                        else:
                            # Base compatível: troca só as MEDIDAS
                            _desc_final = _subst_med(_base_up, _dims, _unid) if _dims else _base_up
                    else:
                        # ETAPA 4 — Claude gera do zero com exemplos relevantes
                        _exs = _buscar_exemplos_db(_tipo_prod)
                        _desc_gerada = _claude_gerar(
                            _img_b, _d['NOME'], _tipo_prod, _dims, _unid, _d['QTD'], _exs, _api_key_gd
                        )
                        if _desc_gerada:
                            _desc_final = _desc_gerada
                            if _dims and 'MEDIDAS' not in _desc_final:
                                _desc_final = _subst_med(_desc_final, _dims, _unid)
                            _orig = 'claude vision'
                        else:
                            _desc_final = (
                                '[REVISAR] PRODUTO DESCONHECIDO'
                                + (f', MEDIDAS {_dims} {_unid}' if _dims else '')
                            )
                            _orig = 'revisar'

                    _rows_out.append({
                        '_ri': _d['_ri'], 'REF': _d['REF'], 'MARCA': _d['MARCA'],
                        'DESC': _desc_final, 'ORIGEM': _orig or 'banco',
                        'IMAGEM': _img_uri, 'TIPO_IA': _tipo_prod
                    })

                _prog.empty()
                st.session_state['gd_resultado']  = _rows_out
                st.session_state['gd_bytes_orig'] = _bytes_gd_arq
                st.session_state[_file_key]        = True

            # ---- Resultado ----
            _res_gd = st.session_state.get('gd_resultado')
            if _res_gd:
                _df_res = pd.DataFrame(_res_gd)
                # banco = ref exata/parcial + banco+ia (tipo encontrado no DB via Claude)
                _n_banco  = int(_df_res['ORIGEM'].str.contains('banco', na=False).sum())
                _n_claude = int(_df_res['ORIGEM'].str.contains('claude vision', na=False).sum())
                _n_rev    = int(_df_res['ORIGEM'].str.contains('revisar', na=False).sum())
                _mc1,_mc2,_mc3,_mc4 = st.columns(4)
                _mc1.metric("Total", len(_df_res))
                _mc2.metric("Do banco", _n_banco,
                            help="Ref exata, ref parcial ou produto similar encontrado no DB")
                _mc3.metric("Via Claude AI", _n_claude,
                            help="IA identificou o tipo e gerou descrição sem base no banco")
                _mc4.metric("Revisar", _n_rev,
                            help="API não configurada ou produto não identificado")

                if _n_rev > 0 and not _api_key_gd:
                    st.warning(
                        f"⚠️ **{_n_rev}** produto(s) ficaram como [REVISAR] porque a chave API Claude "
                        "não está configurada. Com a chave, o sistema identifica o tipo do produto "
                        "na foto, busca similares no banco e gera a descrição automaticamente."
                    )
                elif _n_rev > 0:
                    st.warning(
                        f"⚠️ **{_n_rev}** produto(s) não identificados. Edite a descrição manualmente."
                    )

                # Garante coluna TIPO_IA existe
                if 'TIPO_IA' not in _df_res.columns:
                    _df_res['TIPO_IA'] = ''
                # Coluna com contagem de caracteres
                _df_res['CHARS'] = _df_res['DESC'].str.len()
                _n_long = int((_df_res['CHARS'] > 200).sum())
                if _n_long > 0:
                    st.warning(f"⚠️ {_n_long} descrição(ões) acima de 200 caracteres — edite antes de baixar.")
                st.info("💡 DESC é editável. CHARS mostra o tamanho — máximo 200. TIPO IA = o que Claude identificou na foto.")
                _df_edit = st.data_editor(
                    _df_res[['IMAGEM','REF','MARCA','TIPO_IA','DESC','CHARS','ORIGEM']],
                    use_container_width=True, num_rows="fixed", key="gd_editor",
                    column_config={
                        'IMAGEM':  st.column_config.ImageColumn("Foto", width="small"),
                        'TIPO_IA': st.column_config.TextColumn("Tipo IA", disabled=True,
                                                                help="Tipo identificado pelo Claude na foto"),
                        'CHARS':   st.column_config.NumberColumn("Chars", disabled=True,
                                                                  help="Nº de caracteres — máximo 200"),
                        'ORIGEM':  st.column_config.TextColumn(disabled=True),
                        'REF':     st.column_config.TextColumn("Referência", disabled=True),
                        'MARCA':   st.column_config.TextColumn(disabled=True),
                    }
                )

                # Gera o Excel original com a coluna DESCRIÇÃO preenchida
                _col_desc_pos = _col_pos.get('DESC')
                if _col_desc_pos:
                    _wb_out = _lwb_gd(BytesIO(st.session_state.get('gd_bytes_orig', _bytes_gd_arq)))
                    _ws_out = _wb_out.active
                    for _row_r in _df_edit.itertuples():
                        _ri_excel = _res_gd[_row_r.Index]['_ri']
                        _ws_out.cell(_ri_excel, _col_desc_pos).value = _row_r.DESC
                    _buf_orig = BytesIO()
                    _wb_out.save(_buf_orig)
                    st.download_button(
                        "⬇️ Baixar arquivo original com DESCRIÇÃO preenchida",
                        _buf_orig.getvalue(),
                        _arq_gd.name.replace('.xlsx','_descricoes.xlsx').replace('.xls','_descricoes.xls'),
                        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                        key="dl_gd_orig"
                    )
                else:
                    _buf_gd = BytesIO()
                    _df_edit[['REF','MARCA','DESC','ORIGEM']].to_excel(
                        _buf_gd, index=False, engine='openpyxl')
                    st.download_button(
                        "⬇️ Baixar descrições geradas (Excel)",
                        _buf_gd.getvalue(), "descricoes_geradas.xlsx",
                        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                        key="dl_gd_xlsx"
                    )
                st.caption(
                    "🔵 banco = base do banco com medidas trocadas  |  "
                    "🟢 claude vision = identificado pela IA  |  "
                    "🔴 revisar = sem API configurada"
                )

        except Exception as _e_gd_up:
            st.error(f"Erro ao processar arquivo: {_e_gd_up}")
            import traceback as _tb_gd_up
            st.code(_tb_gd_up.format_exc())


# ==========================================
# PÁGINA: FUNCIONALIDADES
# ==========================================
if _is_active("funcionalidades"):
    st.title("📘 Funcionalidades")
    st.caption("Explicação de cada regra e funcionalidade do sistema, aba por aba — como o sistema inteiro funciona, não só o que foi adicionado recentemente.")

    with st.expander("🏠 Início", expanded=False):
        st.markdown("""
Tela de status geral: mostra quantos certificados, itens, registros e itens do Sistema 5 existem no banco, e um resumo rápido de cada aba.

**Backup:** o banco de dados (`certificados.db`) fica salvo no servidor onde o Streamlit está rodando. Se o app reiniciar (o que pode acontecer sem aviso, dependendo de onde está hospedado), os dados podem ser perdidos. Por isso existe o backup geral no menu lateral — baixa um `.zip`/`.db` com tudo, e permite restaurar depois. Recomendado baixar sempre que terminar uma sessão de cadastro.
""")

    with st.expander("📄 PDF → XML", expanded=False):
        st.markdown("""
**O que faz:** lê um certificado IP-BRI em PDF, extrai os itens (tabelas dentro do PDF) e os dados gerais do certificado, gera XML pronto pra importação, e salva/atualiza o certificado no banco.

**Extração de dados do certificado** (via texto do PDF, com regex):
- `IP-BRI`: padrão `IP-BRI-\\d+/\\d+-\\d+`
- `CE-BRI`: padrão `CE-BRI-[A-Z0-9-]+`
- `FAMÍLIA`: os dígitos finais do IP-BRI (depois do último `-`)
- `REV`: procura "REV", "REVISÃO" ou variações seguido de número; se não achar em uma linha só, olha a linha seguinte
- `Data de Emissão`: padrão `DD/MM/AAAA`

**Extração dos itens:** lê todas as tabelas de todas as páginas do PDF (via `pdfplumber`). Só aceita linhas onde a coluna de ordem é um número de 3-4 dígitos (filtra cabeçalhos/lixo). Ordena os itens pela ordem extraída.

**Regra de REV ao salvar no banco** (evita sobrescrever com versão mais antiga por engano):
- Certificado novo (IP-BRI não existe no banco) → sempre salva
- REV nova ≤ REV do banco **e já existem itens salvos** → **ignora** (não altera nada, registra no histórico como "REV_IGNORADA")
- REV nova ≤ REV do banco **mas o certificado existe sem nenhum item vinculado** (situação de dado incompleto) → insere os itens mesmo assim, sem exigir REV maior (auto-correção de dado quebrado)
- REV nova > REV do banco → atualiza normalmente, compara os itens antigos com os novos (registra as diferenças no histórico) e substitui

**Verificações extras:** avisa se há códigos de barras duplicados dentro do mesmo PDF; avisa se não achou registro vinculado (por CE-BRI + Família) na tabela de registros; se não achar o IP-BRI no PDF, gera o XML mas **não salva** o certificado no banco (evita gravar dado incompleto/não identificável).

**Saída:** ZIP com 3 variações de XML — `xml_virgula.xml`, `xml_ponto.xml` e `xml_sem_caractere.xml` — diferindo só em como o campo `Modelo` termina (`,`, `.`, ou nada).
""")

    with st.expander("🗄️ Banco de Certificados", expanded=False):
        st.markdown("""
**O que faz:** visão geral e gestão direta do banco de certificados.

- **Enviar certificado direto pro banco:** mesmo fluxo de extração do PDF → XML, mas sem gerar o XML — só pra popular/atualizar o banco.
- **Pesquisar item:** busca por Referência/Modelo ou por Código de Barras, juntando dados de `itens` + `certificados` + `registros` (join por CE-BRI + Família).
- **Itens Repetidos:** lista itens cujo código de barras aparece mais de uma vez no banco inteiro (possível duplicidade cadastral).
- **Filtro por Marca:** filtra a base completa por marca.
- **Histórico por IP-BRI:** mostra o histórico de alterações registrado (toda vez que um certificado é criado, atualizado, ou tem REV ignorada, fica um registro em `historico_alteracoes`).
- **Remover certificado:** exclusão manual de um certificado e seus itens vinculados.
- **Backup/Restore:** importar ou baixar o `.db` inteiro; exportar o banco completo em CSV.
""")

    with st.expander("📋 Registros", expanded=False):
        st.markdown("""
**O que faz:** cadastra os registros de certificação (número de registro + CE-BRI + Família + Fábrica) por cliente/base, usados pelo Preenchimento de Confirmação.

**Upload de Excel:** cada **aba** do Excel enviado é tratada como uma **fábrica separada**. Antes de salvar, dá pra renomear a fábrica e informar o endereço de cada uma. Ao salvar, todos os registros anteriores daquele `cliente_base` são apagados e substituídos pelos novos (evita duplicidade ao reenviar o mesmo cliente).

**Editar registro de um item específico:** seleciona Cliente → Família, e edita só o campo Registro num cartão onde Cliente/Família/Fábrica-CE-BRI/Endereço aparecem travados (somente leitura) — pra corrigir um número pontual sem reenviar o Excel inteiro. Se a mesma combinação Cliente+Família tiver mais de uma Fábrica/CE-BRI, aparece um cartão por combinação.

**Editar endereço/CE-BRI de fábrica já salva:** edição direta de endereço e CE-BRI pra uma fábrica já cadastrada, sem precisar reenviar Excel.

**Sincronizar endereços do Sistema 5 com Registros:** propaga o endereço cadastrado na tabela `registros` pros itens e fábricas equivalentes dentro do Sistema 5 (evita ter que cadastrar o mesmo endereço duas vezes em dois módulos diferentes).

**Corrigir CE-BRIs inconsistentes:** varre todos os registros cuja fábrica tem um CE-BRI no próprio nome (formato `NOME - CE-BRI-XXXX`) e corrige qualquer CE-BRI salvo que não bata com esse — nas tabelas `registros`, `sistema5_fabricas` e `sistema5_itens` ao mesmo tempo, mantendo tudo consistente.

**Excluir registros de um cliente:** remove todos os registros de uma base/cliente específico.
""")

    with st.expander("✅ Confirmação", expanded=False):
        st.markdown("""
**O que faz:** preenche automaticamente um Excel de confirmação, usando o banco (Sistema 5 + Certificados) como fonte de dados.

**Regra de cores (definida pela cor de preenchimento da célula, não pelo nome da coluna):**
- 🟢 **Verde** → é uma referência (MODELO). Qualquer célula verde na linha é considerada referência a buscar — não precisa ter uma coluna chamada "REF".
- 🟡 **Amarelo** → é um campo a ser preenchido (MARCA, NOME, CODIGO, REGISTRO, ENDERECO, IP_BRI, CE_BRI, FAMILIA, ITEM, TIPO_PROCESSO, DATA_PROCESSO, ARQUIVO_ORIGEM, IP_PROCESSO).
- 🔵 **Azul** → cabeçalho da coluna (o texto dessa célula diz *qual* campo vai nas células amarelas abaixo).

As cores são detectadas por **faixa de RGB** (não precisa ser um verde/amarelo/azul exato — tolera variações de tom entre diferentes templates de Excel).

**Prioridade de busca (quando não há conflito):** primeiro procura no **Sistema 5** (pega o processo mais recente); se não achar lá, cai pro **Certificado oficial** (pega a revisão mais recente).

**Filtro rigoroso (toggle opcional):** quando desligado (padrão), a referência "0122" também aceita casar com "0122-18" (referências com sufixo). Quando ligado, exige correspondência exata — útil quando o mesmo certificado tem vários modelos com sufixo numérico que não podem ser confundidos entre si.

**Regra de desambiguação** (quando existe mais de um candidato pra mesma referência): o sistema extrai a **palavra-núcleo** do campo MODELO (a primeira palavra depois de "BRINQUEDO" que não seja um termo genérico como CONJUNTO/KIT/SET/MINI/GRANDE/PEQUENO/MÉDIO/SORTIDO). Se o núcleo é igual entre os candidatos (ex: `BRINQUEDO CONJUNTO BONECA` e `BRINQUEDO BONECA DE PLÁSTICO`, núcleo `BONECA` nos dois), trata como o mesmo item e escolhe automaticamente pela prioridade padrão acima, sem perguntar nada. Se o núcleo é diferente (`TRATOR` vs `CARRINHO`), pede pra você escolher manualmente antes de gerar o arquivo.

**Testar referência antes de preencher:** ferramenta de diagnóstico que mostra exatamente o que a busca encontraria pra uma referência específica, sem precisar subir o Excel inteiro.
""")

    with st.expander("🔍 Pesquisa Avançada", expanded=False):
        st.markdown("""
Busca bidirecional simples:
- **IP-BRI → Registro/Fábrica:** dado um IP-BRI, descobre o número de registro e a fábrica vinculados (join por CE-BRI + Família).
- **Registro → IP-BRI:** caminho inverso — dado um número de registro, descobre qual(is) IP-BRI usa(m) ele.

Útil pra conferir rapidamente se um certificado está vinculado a um cliente antes de rodar o preenchimento de confirmação.
""")

    with st.expander("📦 Sistema 5", expanded=False):
        st.markdown("""
**O que faz:** módulo pra processos que **ainda não têm certificado IP-BRI oficial emitido** — inclusões, manutenções e recertificações em andamento.

**Três categorias:**
- **SISTEMA 5 NOVO PROJETO** — cliente padrão "BOLSA" (a entidade certificadora); usado quando o processo é conduzido pela Bolsa.
- **SISTEMA 5 PRÓPRIOS** — o cliente escolhe/cadastra sua própria base (empresa própria conduzindo o processo com registro próprio, não pela Bolsa).
- **SISTEMA 5 FOCUS** — cliente padrão "FOCUS"; categoria separada com processos e registros exclusivos, independente do Novo Projeto e Próprios.

**Detecção automática do tipo de processo** pelo **nome do arquivo** enviado (não pelo conteúdo):
- Contém "MANUT" → Manutenção
- Contém "RECERT" → Recertificação
- Contém "INICIAL" → Inicial
- Contém "INCLUS" → Inclusão
- Nenhum desses → "Outros"

**Requisito de nome de arquivo:** precisa ter o IP no nome (ex: `INCLUSÃO 06-01-26 F01 IP-0094-26.xlsx`) — é assim que o sistema sabe qual IP-BRI (ainda pendente) está associado àquele processo.

**Estrutura de organização:** Categoria → Cliente → Fábrica → Processo. Cada **aba do Excel** enviado é tratada como uma **família**, identificada pelo número presente no nome da aba.

Os itens cadastrados aqui são consultados automaticamente pelo **Preenchimento de Confirmação** (prioridade antes do certificado oficial).
""")

    with st.expander("⏳ IP-BRI Pendentes", expanded=False):
        st.markdown("""
Lista famílias do Sistema 5 que já têm processos/itens cadastrados, mas **ainda não têm um IP-BRI oficial vinculado**. Permite cadastrar manualmente esse IP-BRI assim que ele for emitido, associando à combinação CE-BRI + Família certa — sem precisar esperar o certificado PDF chegar pra vincular tudo.

Também existe um cadastro manual direto (CE-BRI + Família + IP-BRI), pra quando você já sabe o vínculo sem precisar navegar pela lista de pendentes.
""")

    with st.expander("📊 Painel de Cobertura", expanded=False):
        st.markdown("""
Visão geral hierárquica: quantos clientes, fábricas, famílias (tanto em Registros quanto em Sistema 5) e quantas famílias/IP-BRIs pendentes existem.

Permite navegar em cascata: escolhe um **Cliente** → filtra as **Fábricas** dele → filtra os **CE-BRIs** dessa fábrica → vê o **detalhamento por família** → e por fim os **itens específicos** de uma família escolhida. Serve pra enxergar rapidamente onde tem lacuna de cadastro (ex: uma fábrica sem nenhuma família com IP-BRI ainda).
""")

    with st.expander("🔧 Correção de Códigos de Barras", expanded=False):
        st.markdown("""
**O que faz:** audita códigos de barras que **podem** ter perdido um zero à esquerda (erro comum quando o Excel trata o código como número em vez de texto).

**Critério de suspeita** (não corrige nada automaticamente, só sinaliza pra revisão manual):
- Código com **menos de 13 dígitos**, ou
- Código termina em **".0"** (sinal claro de que virou número float em algum momento e perdeu formatação)

Roda separadamente pro Sistema 5 e pros Certificados oficiais. Pra cada item suspeito, dá pra corrigir manualmente digitando o valor certo, ou aplicar "completar com zero à esquerda até X dígitos" (13, 14 ou 12) com um clique.

**Importante:** o sistema nunca corrige em massa sem confirmação — é sempre item por item, justamente para não arriscar alterar um código que só parece suspeito mas está correto.
""")

    with st.expander("🏷️ Conferência de Etiquetas", expanded=True):
        st.markdown("""
**O que faz:** confere automaticamente se as etiquetas do Word batem com os dados do Excel ORDER LIST, item por item.

**Entradas:** um Excel (ORDER LIST) e um Word (com uma tabela por etiqueta).

**Como acha os dados no Excel:** procura o cabeçalho da tabela por **cor de preenchimento azul** (`#00CCFF` neon ou `#00B0F0` celeste), que é como as planilhas de fatura marcam a linha de cabeçalho de verdade (ignorando linhas de carta/fatura antes dela, tipo nome da fábrica, endereço, cliente etc). Se não achar nenhuma célula com essas cores, cai num plano B: procura a linha que contém o texto "REFERENCIA".

**Checagens feitas por etiqueta (1 a 8):**
1. **Marca** — compara texto da etiqueta com o Excel
2. **Indicativo de idade** (ex: +3 ANOS) — compara com o que está no nome do item no Excel
3. **Restritivo de idade** (ex: -3 ANOS) — idem
4. **Imagem do chorão** — obrigatória se o produto é pra -3 anos; detecta pelo tamanho da imagem (~123×129px)
5. **Imagem de pilha/bateria** — obrigatória se o nome do produto tem "PILHA"; detecta por tamanho de imagem PNG
6. **Código de barras** — lê o código de dentro da imagem (formato EMF) da etiqueta e compara com o Excel
7. **Número de registro** — lê o selo "Segurança/REGISTRO" (que é imagem, não texto) via OCR e compara com o Excel
8. **Ortografia/erro de escrita** — verifica só os textos de ATENÇÃO, INDICAÇÃO, ADVERTÊNCIA e a linha de referência/nome; ignora dados do importador/CNPJ/SAC (informação do cliente, não é pra checar)

**Resultado de cada checagem:**
- ✅ verde = bateu
- ❌ vermelho = divergência real
- ⚠️ amarelo = "confira manualmente" — usado quando o sistema não tem certeza se é erro de verdade ou limitação da leitura automática (ver abaixo)

---
#### 🔍 Código de barras — como funciona por trás
- O código de barras na etiqueta vem como uma imagem em formato **EMF** (um formato vetorial do Windows).
- O sistema converte **todos os EMFs da etiqueta de uma vez só** (uma única chamada ao LibreOffice pra tudo, não uma por etiqueta) — isso foi feito assim porque abrir o LibreOffice repetidamente é o passo mais lento do processo.
- Depois de convertido pra PNG, usa a biblioteca `pyzbar` pra ler o código de barras de dentro da imagem.
- Requer o **LibreOffice** instalado (o sistema procura automaticamente em: pasta local `C:\\LibreOfficePortable`, `Program Files`, ou o caminho de rede configurado — nessa ordem, priorizando sempre o mais rápido).

---
#### 🔍 Número de registro (OCR) — como funciona por trás
- O selo "Segurança / REGISTRO / Compulsório" é uma **imagem**, não texto — então não dá pra ler com regex, precisa de **OCR** (reconhecimento de texto em imagem), feito com **Tesseract**.
- O sistema recorta só a região central da imagem (onde fica o número), amplia 4x, e testa várias versões: contraste normal, binarização (preto/branco), binarização invertida, e nitidez artificial (sharpen).
- Pra cada versão, tenta ler com múltiplos modos do Tesseract (incluindo um modo restrito a só dígitos e `/`, mais preciso pra esse tipo de número curto) e com os dois motores do Tesseract (padrão e legado).
- **Votação por maioria:** roda todas essas tentativas (primeiro só no recorte central; só cai pra imagem completa se o recorte não achar nada) e usa o valor que mais se repetiu entre elas, em vez de confiar na primeira leitura. No modo debug, aparece a distribuição de votos.
- **Cache por imagem:** se a mesma imagem de selo aparece em várias etiquetas (comum quando o mesmo registro vale pra vários produtos/cores), o OCR só roda **uma vez** — as outras aproveitam o resultado já calculado.
- **Limite físico:** a imagem do selo vem em resolução baixa (~438×198px) já embutida no Word. Nenhuma técnica de imagem "inventa" detalhe que não existe na imagem original — por isso ainda pode acontecer erro de leitura ocasional, principalmente confusão entre dígitos parecidos (8↔3, 0↔8, 1↔7).

---
#### ⚠️ Avisos amarelos — o que significam
O sistema distingue **erro real** (❌) de **possível limitação da leitura automática** (⚠️), pra você saber quando vale a pena olhar a etiqueta física antes de rejeitar:
- **Registro com ⚠️:** o valor lido pela etiqueta difere do Excel em **exatamente 1 dígito** (mesmo tamanho, só um caractere diferente) — padrão típico de confusão de OCR. Confira a "Prova Real" antes de considerar erro de verdade.
- **Ortografia com ⚠️:** encontrou uma palavra que não está no dicionário de português nem na lista de termos técnicos conhecidos nem no vocabulário do próprio item no Excel — pode ser erro de escrita real (ex: "APONT" por "APONTAR") ou um termo técnico novo que ainda não está na lista.

Divergências que **não** se encaixam nesses padrões (ex: registro completamente diferente, marca errada) continuam como ❌ vermelho.

---
#### 🔍 Prova Real — conferência manual
Em cada etiqueta conferida, aparece um bloco **"Prova Real"** mostrando lado a lado, sem interpretação nenhuma:
- Código de barras: valor do Excel × valor lido da etiqueta
- Registro: valor do Excel × valor lido da etiqueta

Serve pra você conferir os dados brutos com os próprios olhos, mesmo quando o sistema já deu ✅ ou ❌ — principalmente útil nos casos de ⚠️, onde a leitura automática tem menos certeza.

---
#### 🐛 Modo debug
Checkbox disponível na Conferência de Etiquetas. Quando ligado, mostra:
- Caminho onde o LibreOffice e o Tesseract foram encontrados
- Toda imagem embutida em cada etiqueta (nome do arquivo, extensão, tamanho, dimensões)
- Cada tentativa de OCR (variante de imagem + modo de leitura + resultado bruto lido)
- A distribuição de votos usada pra escolher o número de registro
- Erros técnicos que normalmente ficam escondidos (ex: falha de conversão, biblioteca faltando)

Deixa desligado no uso normal (deixa a tela mais limpa); liga só quando algo não bate e você quer entender o motivo.
""")

    with st.expander("⚙️ Regras técnicas gerais (bastidores)", expanded=False):
        st.markdown("""
- **LibreOffice e Tesseract são localizados automaticamente**, testando (em ordem): PATH do sistema, `Program Files`, cópia local em pasta específica, e por último caminho de rede (mais lento, usado só se não achar nada local).
- **Excel `.xls` legado:** se o arquivo não abrir com a biblioteca padrão (openpyxl, que só lê `.xlsx`), o sistema converte automaticamente pra `.xlsx` via LibreOffice antes de processar.
- **Cabeçalho do Excel (Conferência de Etiquetas e Geração):** detectado pela cor azul de preenchimento (`#00CCFF` neon ou `#00B0F0` celeste) nas primeiras 30 linhas; se não achar, cai pro método antigo (procurar o texto "REFERENCIA").
- **Cache de OCR e conversão em lote:** otimizações feitas pra reduzir o tempo de processamento sem perder precisão — a lógica de comparação em si não muda, só evita repetir trabalho em cima da mesma imagem.
- **Todas as tabelas do banco** (`certificados`, `itens`, `registros`, `sistema5_itens`, `sistema5_arquivos`, `sistema5_clientes`, `sistema5_fabricas`, `historico_alteracoes`, `ip_bri_familias`) ficam no mesmo arquivo SQLite (`certificados.db`), o que é o que possibilita os backups/restaurações completas pelo menu lateral.
""")
